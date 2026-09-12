# -*- coding: utf-8 -*-
"""内帐总账 · 共用样式与布局常量"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

THIN = Side(style='thin', color='9E9E9E')
MED  = Side(style='medium', color='404040')
BOX  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

F_TITLE = Font(name='微软雅黑', size=15, bold=True, color='1F3864')
F_H2    = Font(name='微软雅黑', size=10, bold=True, color='1F4E79')
F_HDR   = Font(name='微软雅黑', size=9,  bold=True, color='FFFFFF')
F_HDR2  = Font(name='微软雅黑', size=9,  bold=True, color='1F3864')
F_TXT   = Font(name='微软雅黑', size=9)
F_IN    = Font(name='微软雅黑', size=9, color='0000C0')
F_LINK  = Font(name='微软雅黑', size=9, color='006100')
F_TOT   = Font(name='微软雅黑', size=9, bold=True)
F_NOTE  = Font(name='微软雅黑', size=8, color='808080')
F_BIG   = Font(name='微软雅黑', size=16, bold=True, color='1F3864')

FILL_HDR  = PatternFill('solid', fgColor='2F5597')   # 表头深蓝
FILL_HDR2 = PatternFill('solid', fgColor='D6E4F0')   # 次级表头
FILL_IN   = PatternFill('solid', fgColor='FFF7E0')   # 录入区淡黄
FILL_AUTO = PatternFill('solid', fgColor='EFEFEF')   # 自动分录区灰
FILL_TOT  = PatternFill('solid', fgColor='DDEBF7')
FILL_CHK  = PatternFill('solid', fgColor='E2EFDA')
FILL_CARD = PatternFill('solid', fgColor='F2F6FC')
FILL_WARN = PatternFill('solid', fgColor='FCE4E4')

C  = Alignment(horizontal='center', vertical='center', wrap_text=True)
CL = Alignment(horizontal='left',   vertical='center', wrap_text=True)
CR = Alignment(horizontal='right',  vertical='center')

MONEY = '#,##0.00;[Red]-#,##0.00;"-"'
QTY   = '#,##0.###;[Red]-#,##0.###;"-"'
PCT   = '0.00%'
DATE  = 'yyyy-mm-dd'
YM    = 'yyyy-mm'

# ---------------- 容量 ----------------
ACT_R0, ACT_R1 = 6, 35        # 资金账户 30
DEP_R0, DEP_R1 = 6, 25        # 部门 20
EMP_R0, EMP_R1 = 6, 65        # 员工 60
GDS_R0, GDS_R1 = 6, 205       # 商品 200
CUS_R0, CUS_R1 = 6, 155       # 客户 150
SUP_R0, SUP_R1 = 6, 155       # 供应商 150
ACC_R0, ACC_R1 = 6, 125       # 会计科目 120
CAT_R0, CAT_R1 = 6, 85        # 收支科目 80
OC_R0,  OC_R1  = 6, 35        # 期初·资金 30
OR_R0,  OR_R1  = 6, 155       # 期初·应收 150
OP_R0,  OP_R1  = 6, 155       # 期初·应付 150
OG_R0,  OG_R1  = 6, 205       # 期初·库存 200
OA_R0,  OA_R1  = 6, 125       # 期初·科目 120
CASH_R0, CASH_R1 = 5, 504     # 资金流水 500
SAL_R0,  SAL_R1  = 5, 404     # 销售录入 400
BUY_R0,  BUY_R1  = 5, 404     # 采购录入 400
OTH_R0,  OTH_R1  = 5, 304     # 其他凭证 300（计提/折旧/摊销/结转损益）
FA_R0,   FA_R1   = 6, 105     # 固定资产 100
PAY_R0,  PAY_R1  = 6, 65      # 工资表 60

SH_HOME, SH_HELP = '首页', '操作说明'
SH_PARAM, SH_BASE, SH_OPEN = '科目参数', '基础资料', '期初数据'
SH_CASH, SH_SAL, SH_BUY = '资金流水', '销售录入', '采购录入'
SH_OTH = '其他凭证'
SH_TB, SH_PL, SH_BS, SH_CF = '科目余额表', '利润表', '资产负债表', '现金流量表'
SH_VOU, SH_INV = '记账凭证', '库存收发存'
SH_AR, SH_AP, SH_BAL = '应收跟进', '应付跟进', '账户余额'
SH_FA, SH_PAYROLL = '固定资产台账', '工资表'
SH_AUX, SH_ANA, SH_CHK = '辅助核算', '经营分析', '财务勾稽核查'
SH_SOA_C, SH_SOA_S = '客户对账单', '供应商对账单'

Q = lambda s: f"'{s}'"          # 表名加引号，跨工具安全

BAL_FIXED = [0]

def bal(f):
    """配平公式末尾的右括号（本模板所有嵌套 IF 都以右括号串收尾）"""
    d = f.count('(') - f.count(')')
    if d > 0:
        BAL_FIXED[0] += 1
        return f + ')' * d
    while d < 0 and f.endswith(')'):
        f = f[:-1]; d += 1; BAL_FIXED[0] += 1
    return f

def put(ws, coord, value=None, font=F_TXT, fill=None, align=C, fmt=None, border=BOX):
    c = ws[coord]
    if value is not None:
        if isinstance(value, str) and value.startswith('=') and value.count('(') != value.count(')'):
            value = bal(value)
        c.value = value
    c.font = font
    if fill is not None:
        c.fill = fill
    if align is not None:
        c.alignment = align
    if fmt is not None:
        c.number_format = fmt
    if border is not None:
        c.border = border
    return c

def hdr(ws, coord, text, span=None, font=F_HDR, fill=FILL_HDR):
    if span:
        ws.merge_cells(span)
    put(ws, coord, text, font=font, fill=fill)

def headers(ws, row, start_col, names, widths_map=None, fill=FILL_HDR, font=F_HDR):
    """一次写一整行表头，返回 {列名: 列字母}"""
    m = {}
    for i, n in enumerate(names):
        col = L(start_col + i)
        put(ws, f'{col}{row}', n, font=font, fill=fill)
        m[n] = col
    ws.row_dimensions[row].height = 30
    return m

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

def title(ws, text, last_col, sub=None, row=1):
    ws.merge_cells(f'A{row}:{last_col}{row}')
    put(ws, f'A{row}', text, font=F_TITLE, align=CL, border=None)
    ws.row_dimensions[row].height = 26
    if sub:
        ws.merge_cells(f'A{row+1}:{last_col}{row+1}')
        put(ws, f'A{row+1}', sub, font=F_NOTE, align=CL, border=None)
        ws.row_dimensions[row+1].height = 15

def page(ws, titles=None, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if titles:
        ws.print_title_rows = titles

def dv_list(ws, rng, source, title_='请从下拉中选择', msg='必须选择下拉中已维护的值，手工输入会导致取数错位'):
    dv = DataValidation(type='list', formula1=source, allow_blank=True,
                        showErrorMessage=True, errorStyle='stop',
                        errorTitle=title_, error=msg)
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv

def dv_num(ws, rng, op='greaterThanOrEqual', f1='0', title_='只能填数字',
           msg='本列必须是数字，填成文本会让整笔业务被静默丢弃'):
    dv = DataValidation(type='decimal', operator=op, formula1=f1, allow_blank=True,
                        showErrorMessage=True, errorStyle='stop', errorTitle=title_, error=msg)
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv

def dv_date(ws, rng, y_cell):
    dv = DataValidation(type='date', operator='between',
                        formula1=f'DATE({y_cell},1,1)', formula2=f'DATE({y_cell},12,31)',
                        allow_blank=True, showErrorMessage=True, errorStyle='stop',
                        errorTitle='日期超出会计年度',
                        error='日期必须落在首页设定的会计年度内，否则该笔业务不会进入任何报表')
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv
