# -*- coding: utf-8 -*-
"""A063 样式与小工具"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL

C_MAIN, C_BASE, C_BIZ, C_RPT, C_WARN, C_DASH = '1F4E79', '2E75B6', 'ED7D31', '70AD47', 'C00000', '7030A0'

FILL_IN   = PatternFill('solid', fgColor='FFFBEA')   # 手工填
FILL_AUTO = PatternFill('solid', fgColor='F2F2F2')   # 公式算
FILL_TOT  = PatternFill('solid', fgColor='E2EFDA')   # 合计
FILL_WARN = PatternFill('solid', fgColor='FCE4E4')   # 预警
FILL_KPI  = PatternFill('solid', fgColor='FFF2CC')   # 看板
FILL_SEC  = PatternFill('solid', fgColor='DDEBF7')   # 分区标题

F_TITLE = Font(name='微软雅黑', size=15, bold=True, color='FFFFFF')
F_HDR   = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
F_HDR2  = Font(name='微软雅黑', size=10, bold=True, color='1F4E79')
F_IN    = Font(name='微软雅黑', size=10)
F_AUTO  = Font(name='微软雅黑', size=10, color='7F7F7F')
F_TOT   = Font(name='微软雅黑', size=10, bold=True, color='1F4E79')
F_WARN  = Font(name='微软雅黑', size=10, bold=True, color='C00000')
F_NOTE  = Font(name='微软雅黑', size=9, color='8B5E00')
F_TXT   = Font(name='微软雅黑', size=10)
F_BIG   = Font(name='微软雅黑', size=14, bold=True, color='1F4E79')
F_LINK  = Font(name='微软雅黑', size=10, color='0563C1', underline='single')
F_SEC   = Font(name='微软雅黑', size=11, bold=True, color='1F4E79')

_T = Side('thin', color='BFBFBF')
BOX = Border(left=_T, right=_T, top=_T, bottom=_T)
C  = Alignment('center', 'center', wrap_text=True)
CL_= Alignment('left', 'center', wrap_text=True)
CR = Alignment('right', 'center')
CT = Alignment('left', 'top', wrap_text=True)

MONEY = '#,##0.00;[Red]\\-#,##0.00;\\-'
MONEY0= '#,##0;[Red]\\-#,##0;\\-'
NUM   = '#,##0;[Red]\\-#,##0;\\-'
PCT   = '0.0%'
DATEF = 'yyyy-mm-dd'
DATEQ = 'yyyy-mm-dd;;;@'          # 自动格：0 不画成 1900-01-00
TXT   = '@'
DAYS  = '0"天";[Red]\\-0"天";\\-'


def put(ws, addr, value=None, font=F_IN, fill=None, fmt=None, align=C, border=BOX):
    c = ws[addr]
    if value is not None:
        c.value = value
    c.font = font
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align is not None:
        c.alignment = align
    if border is not None:
        c.border = border
    return c


def title(ws, text, last_col, note=None, color=C_MAIN):
    put(ws, 'A1', text, font=F_TITLE, fill=PatternFill('solid', fgColor=color), align=CL_, border=None)
    ws.merge_cells(f'A1:{last_col}1')
    ws.row_dimensions[1].height = 30
    if note:
        put(ws, 'A2', note, font=F_NOTE, align=CT, border=None)
        ws.merge_cells(f'A2:{last_col}2')
        ws.row_dimensions[2].height = max(24, 13 * (note.count('\n') + 1) + 8)


def headers(ws, row, names, fill_color=C_MAIN, height=32, font=None, fill=None):
    f = fill or PatternFill('solid', fgColor=fill_color)
    for i, n in enumerate(names, 1):
        put(ws, f'{CL(i)}{row}', n, font=font or F_HDR, fill=f)
    ws.row_dimensions[row].height = height


def widths(ws, mp):
    for k, v in mp.items():
        ws.column_dimensions[k].width = v


def page(ws, titles=None, landscape=True):
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if titles:
        ws.print_title_rows = titles
