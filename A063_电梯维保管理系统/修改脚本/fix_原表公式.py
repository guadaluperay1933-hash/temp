# -*- coding: utf-8 -*-
"""在客户原表《电梯档案 + 各维保组应收》上补公式（保持原框架，只做小调整）

原表结构：
  电梯档案 —— 一台梯一行
  刘丕斌徐少斌夏凯 / 陆莹杨凯 / 刘坤田超 / 陈彪姚永华梅思阳 / 胡春平 / Sheet5(空白模板)
      —— 各维保组的合同应收，一个合同(或一期)一行，第 2 行是合计
  三种列布局：A 型（无合同备注）/ B 型（多一列合同备注）/ C 型（陆莹杨凯，K 列是付款约定、没有税率）

这次改动：
  ① 收款区：在「收现金额」后面插 1 列「收款不开票」，在「开票应收额」后面插 5 列
     「到账未开票 / 未开票应收 / 应收余额 / 回款情况 / 收款备注」，其余原列只是整体右移
  ② 公式全部铺到第 NR 行；修掉原表 4 处坑（见 FIXES）
  ③ 到期提醒顺带提示「合同未回」和免保状态
  ④ 每张组表末尾追加：维保状态、付款周期、分期进度、截至今天应收(按期)、逾期未收、往年欠款、
     项目累计欠款，以及 5 个隐藏辅助列
  ⑤ 电梯档案：W 维保状态下拉，X 免保到期日，Y 维保状态提醒，Z 年检/校验提醒
  ⑥ 新增《汇总》《填写说明》两张表

跑法：python3 fix_原表公式.py [源.xlsx] [输出.xlsx]
"""
import sys, os, copy
import openpyxl
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.dimensions import DimensionHolder, ColumnDimension
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, '参考', '原_2027版_电梯档案与各组应收.xlsx')
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'A063_电梯档案与维保应收_公式版.xlsx')

NR = 500          # 组表公式铺到第几行
TOT_END = 1001    # 合计行 / 汇总统计到第几行（沿用原表的 3:1001）
LIFT_NR = 800     # 电梯档案公式铺到第几行

ACC = '_ * #,##0.00_ ;_ * \\-#,##0.00_ ;_ * "-"??_ ;_ @_ '     # 原表用的会计格式
DATEF = 'yyyy"年"m"月"d"日";@'
AUTO_FONT_COLOR = '1F4E79'
FILL_NEW_IN = PatternFill('solid', fgColor='FFF2CC')     # 新增·手填列表头
FILL_NEW_AUTO = PatternFill('solid', fgColor='DDEBF7')   # 新增·公式列表头
FILL_HELP = PatternFill('solid', fgColor='EDEDED')

HTXT = {'合同未回': '有合同,已送出未回,未见合同,未签'}
STATES = ['正常维保', '安装免保', '技术免保', '质保期内', '暂停维保', '已解约']
CYCLES = ['一次性', '年付', '半年付', '季度付', '月付', '半月付']
NOSCHED_STATES = ['安装免保', '技术免保', '质保期内', '暂停维保', '已解约']

# 三种布局：逻辑列 → 原表列字母
LAYOUT_A = dict(g='G', h='H', i='I', j='J', k='K', l='L', m='M', n='N', o='O', p='P', q='Q', s='R',
                typ='S', kind='T', tax='U', taxamt='V', last='V', notes=['T'])
LAYOUT_B = dict(g='H', h='I', i='J', j='K', k='L', l='M', m='N', n='O', o='P', p='Q', q='R', s='S',
                typ='T', kind='U', tax='V', taxamt='W', last='W', notes=['G', 'U'])
LAYOUT_C = dict(g='G', h='H', i='I', j='J', pay='K', k='L', l='M', m='N', n='O', o='P', p='Q', q='R', s='S',
                typ='T', kind='U', tax=None, taxamt=None, last='U', notes=['K', 'U'])
TEAMS = [('刘丕斌徐少斌夏凯', LAYOUT_A), ('陆莹杨凯', LAYOUT_C), ('刘坤田超', LAYOUT_B),
         ('陈彪姚永华梅思阳', LAYOUT_B), ('胡春平', LAYOUT_A), ('Sheet5', LAYOUT_A)]

LOG = []


def parse_pay(txt):
    """从原表的备注文字里认付款周期（只认得很有把握的写法，认不出就空着＝按一次性算）"""
    t = str(txt or '')
    if '半月' in t:
        return '半月付'
    if '季度' in t or '每季' in t:
        return '季度付'
    if '半年' in t:
        return '半年付'
    if '每月' in t or '月付' in t or '按月' in t:
        return '月付'
    if '七日内付' in t or '一次性' in t or '年付' in t:
        return '一次性'
    return None


def old_dims(ws, upto):
    """原来每一列的列定义（宽度 + 整列默认格式），按列号展开"""
    d = {}
    for key, cd in ws.column_dimensions.items():
        lo = cd.min or CI(key)
        hi = cd.max or lo
        for c in range(lo, min(hi, upto) + 1):
            d[c] = cd
    return d


def clone_dim(ws, cd, idx, width=None, hidden=None):
    nd = copy.copy(cd) if cd is not None else ColumnDimension(ws, index=CL(idx))
    if cd is not None:
        nd._style = copy.copy(cd._style)
    nd.index = CL(idx)
    nd.min = nd.max = idx
    if width is not None:
        nd.width = width
    if hidden is not None:
        nd.hidden = hidden
    return nd


def style_like(dst, src, fmt=None, auto=False, wrap=None):
    if src is not None and src.has_style:
        dst._style = copy.copy(src._style)
    if fmt:
        dst.number_format = fmt
    if auto:
        f = copy.copy(dst.font)
        dst.font = Font(name=f.name or '宋体', sz=f.sz or 11, b=f.b, color=AUTO_FONT_COLOR)
    if wrap is not None:
        a = dst.alignment
        dst.alignment = Alignment(horizontal=a.horizontal, vertical=a.vertical or 'center', wrap_text=wrap)


wb = openpyxl.load_workbook(SRC, rich_text=True)   # 不开 rich_text 会把单元格里局部标红的字变成普通字
COLS = {}   # 每张组表改完之后的列字母，汇总表要用

# ══════════════════════════════════════════════════════════════
# 一、各维保组应收表
# ══════════════════════════════════════════════════════════════
for sn, L in TEAMS:
    ws = wb[sn]
    qi = CI(L['q'])
    last_old = CI(L['last'])
    maxr = ws.max_row

    # 右边有没有别的内容会被追加块盖掉
    for r in range(1, maxr + 1):
        for c in range(last_old + 1, last_old + 3):
            if ws.cell(r, c).value not in (None, ''):
                raise SystemExit(f'{sn}!{CL(c)}{r} 有内容，追加列会盖掉它：{ws.cell(r, c).value!r}')

    # 付款周期：插列之前先按原备注认出来
    pay0 = {}
    for r in range(3, maxr + 1):
        txt = ' '.join(str(ws[f'{c}{r}'].value or '') for c in L['notes'])
        v = parse_pay(txt)
        if v:
            pay0[r] = v

    # 原来 J 列是手填数字的行：跟公式算出来一样才换成公式，不一样就保留手填
    keepJ = set()
    for r in range(3, maxr + 1):
        v = ws[f"{L['j']}{r}"].value
        if isinstance(v, (int, float)):
            g, h, i = (ws[f"{L[x]}{r}"].value for x in 'ghi')
            try:
                calc = float(g) * float(h) * (float(i) if i not in (None, '') else 1)
            except (TypeError, ValueError):
                calc = None
            if calc is not None and abs(calc - v) < 0.005:
                LOG.append(f'{sn}!{L["j"]}{r} 手填 {v} = 单价×台量×年限，换成公式')
            else:
                keepJ.add(r)
                LOG.append(f'{sn}!{L["j"]}{r} 手填 {v} ≠ 单价×台量×年限({calc})，保留手填')

    d0 = old_dims(ws, last_old + 3)

    # ── 插列：收现金额后 1 列；开票应收额后 5 列 ──
    ws.insert_cols(qi + 1, 1)
    ws.insert_cols(qi + 3, 5)

    def nc(old):          # 原列字母 → 新列字母
        idx = CI(old)
        if idx <= qi:
            return old
        if idx == qi + 1:
            return CL(idx + 1)
        return CL(idx + 6)

    C = {k: nc(v) for k, v in L.items() if k in 'g h i j k l m n o p q s typ kind tax taxamt last pay'.split() and v}
    C.update(A='A', B='B', Cc='C', D='D', E='E', F='F')
    C['x0'] = CL(qi + 1)      # 收款不开票（手填）
    C['t'] = CL(qi + 3)       # 到账未开票
    C['u'] = CL(qi + 4)       # 未开票应收
    C['v'] = CL(qi + 5)       # 应收余额
    C['w'] = CL(qi + 6)       # 回款情况
    C['bz'] = CL(qi + 7)      # 收款备注（手填）
    a0 = CI(C['last']) + 1
    for k in ['zt', 'zq', 'pg', 'jd', 'yq', 'wn', 'lj', 'key', 'kd', 'nd', 'tn', 'kn']:
        C[k] = CL(a0)
        a0 += 1
    COLS[sn] = C

    # 列定义（宽度 + 整列默认格式）跟着右移；新列借用旁边原列的整列格式
    dh = DimensionHolder(worksheet=ws)
    for c, cd in d0.items():
        n = CI(nc(CL(c)))
        dh[CL(n)] = clone_dim(ws, cd, n, width=max(cd.width or 9, 9) if c == 1 else None)
    qd = d0.get(qi)
    NEWW = [('x0', 12, qd), ('t', 12, qd), ('u', 12, qd), ('v', 13, qd), ('w', 36, None), ('bz', 22, None),
            ('zt', 10, None), ('zq', 9, None), ('pg', 16, None), ('jd', 13, qd), ('yq', 12, qd), ('wn', 12, qd),
            ('lj', 13, qd), ('key', 20, None), ('kd', 11, None), ('nd', 7, None), ('tn', 7, None), ('kn', 7, None)]
    for k, wd, base in NEWW:
        n = CI(C[k])
        dh[CL(n)] = clone_dim(ws, base, n, width=wd, hidden=k in ('key', 'kd', 'nd', 'tn', 'kn'))
    ws.column_dimensions = dh

    # ── 表头 ──
    hsrc = ws[f"{C['s']}1"]           # 原「开票应收额」表头，拿它的样式
    if L is LAYOUT_C and ws[f"{C['pay']}1"].value in (None, ''):
        ws[f"{C['pay']}1"] = '付款约定'
        style_like(ws[f"{C['pay']}1"], hsrc)
    for k, sub in [('p', '(开了票/要开票的)'), ('q', '(不开票的现金)')]:
        c = ws[f'{C[k]}1']
        base = str(c.value or '').split('\n')[0].strip()
        c.value = base + '\n' + sub
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    NEWH = [('x0', '收款不开票\n(转账/收据)', True), ('t', '到账未开票\n(待补票)', False),
            ('u', '未开票应收', False), ('v', '应收余额\n(欠款)', False), ('w', '回款情况', False),
            ('bz', '收款备注', True), ('zt', '维保状态', True), ('zq', '付款周期', True),
            ('pg', '分期进度', False), ('jd', '截至今天应收\n(按付款周期)', False), ('yq', '逾期未收', False),
            ('wn', '往年欠款\n(同项目)', False), ('lj', '项目累计欠款\n(各年合计)', False),
            ('key', '项目键', False), ('kd', '开票已到账', False), ('nd', '年度', False),
            ('tn', '总期数', False), ('kn', '已到期数', False)]
    for k, txt, is_in in NEWH:
        c = ws[f'{C[k]}1']
        c.value = txt
        style_like(c, hsrc, wrap=True)
        c.fill = FILL_NEW_IN if is_in else FILL_NEW_AUTO
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 0, 42)

    # ── 第 2 行合计 ──
    tsrc = ws[f"{C['q']}2"]
    for k in ['h', 'j', 'k', 'l', 'n', 'p', 'q', 'x0', 's', 't', 'u', 'v', 'jd', 'yq'] + (['taxamt'] if C.get('taxamt') else []):
        c = ws[f'{C[k]}2']
        c.value = f'=SUM({C[k]}3:{C[k]}{TOT_END})'
        style_like(c, tsrc, fmt=('General' if k == 'h' else ACC))

    # ── 数据行公式 ──
    A_, B_, D_, E_ = 'A', 'B', 'D', 'E'
    for r in range(3, NR + 1):
        R = lambda k: f'${C[k]}{r}'
        BLK = f'TRIM($B{r}&$C{r})=""'
        UNRET = (f'OR(ISNUMBER(SEARCH("未",$A{r})),ISNUMBER(SEARCH("送出",$A{r})),'
                 f'AND(TRIM($A{r})="",NOT(ISNUMBER($D{r})),N({R("g")})=0))')
        FREE = f'OR({R("zt")}="安装免保",{R("zt")}="技术免保",{R("zt")}="质保期内")'
        NOSCH = 'OR(' + ','.join(f'{R("zt")}="{x}"' for x in NOSCHED_STATES) + ')'
        KEYR = f"${C['key']}$3:${C['key']}${TOT_END}"
        RENEW = f'COUNTIFS({KEYR},{R("key")},$E$3:$E${TOT_END},">"&$E{r})>0'
        src_in = ws[f"{C['q']}{r}"] if ws[f"{C['q']}{r}"].has_style else ws[f"{C['q']}3"]
        src_txt = ws[f'F{r}'] if ws[f'F{r}'].has_style else ws['F3']

        def put(k, val, fmt=ACC, auto=True, wrap=None, src=None):
            c = ws[f'{C[k]}{r}'] if k in C else ws[f'{k}{r}']
            c.value = val
            style_like(c, src if src is not None else src_in, fmt=fmt, auto=auto, wrap=wrap)

        # F 到期提醒（沿用原来的【已到期】【即将到期，剩余N天】【正常】写法）
        put('F', f'=IF({BLK},"",IF({UNRET},"【合同未回】","")'
                 f'&IF({R("zt")}="已解约","",IF(NOT(ISNUMBER($E{r})),IF({UNRET},"","【未填合同期】"),'
                 f'IF({RENEW},"【已续签】",IF($E{r}<TODAY(),"【已到期】",IF($E{r}-TODAY()<=提醒天数,'
                 f'"【即将到期，剩余"&($E{r}-TODAY())&"天】","【正常】")))))'
                 f'&IF(OR({R("zt")}="",{R("zt")}="正常维保"),"","【"&{R("zt")}&"】"))',
            fmt='General', src=src_txt, auto=False, wrap=True)
        # J 合同应收 / L 总应收
        if r not in keepJ:
            put('j', f'=IF(OR(N({R("g")})=0,N({R("h")})=0),"",'
                     f'ROUND(N({R("g")})*N({R("h")})*IF(N({R("i")})=0,1,N({R("i")})),2))', auto=False,
                src=ws[f"{C['j']}{r}"] if ws[f"{C['j']}{r}"].has_style else src_in)
        put('l', f'=IF({BLK},"",IF(AND(N({R("j")})=0,N({R("k")})=0),IF(N({R("n")})>0,ROUND(N({R("n")}),2),""),'
                 f'ROUND(N({R("j")})+N({R("k")}),2)))', auto=False,
            src=ws[f"{C['l']}{r}"] if ws[f"{C['l']}{r}"].has_style else src_in)
        # 收款区
        c = ws[f"{C['x0']}{r}"]
        style_like(c, src_in, fmt=ACC)
        put('s', f'=IF({BLK},"",IF(N({R("n")})=0,"",MAX(0,MIN(N({R("n")})-N({R("p")}),N({R("v")})))))')
        put('t', f'=IF({BLK},"",IF(N({R("p")})<=N({R("n")}),"",N({R("p")})-N({R("n")})))')
        put('v', f'=IF({BLK},"",IF(AND(N({R("l")})=0,N({R("p")})+N({R("q")})+N({R("x0")})=0),"",'
                 f'ROUND(N({R("l")})-N({R("p")})-N({R("q")})-N({R("x0")}),2)))')
        put('u', f'=IF(OR({BLK},{R("v")}=""),"",MAX(0,N({R("v")}))-N({R("s")}))')
        RCV = f'(N({R("p")})+N({R("q")})+N({R("x0")}))'
        YQ = R('yq')
        OVERINV = f'(N({R("n")})-MIN(N({R("n")}),N({R("p")}))-N({R("s")}))'
        put('w', f'=IF({BLK},"",IF(N({R("l")})=0,IF({RCV}=0,IF({FREE},"免保，不计费","—未填单价"),'
                 f'"未填单价，已收 "&TEXT({RCV},"#,##0.00")),'
                 f'IF(N({R("v")})>0.005,IF(AND(ISNUMBER({YQ}),N({YQ})<=0.005),"未到期待收 ","欠 ")'
                 f'&TEXT(N({R("v")}),"#,##0.00")'
                 f'&IF(AND(ISNUMBER({YQ}),N({YQ})>0.005,N({YQ})<N({R("v")})-0.005),"（其中已逾期 "&TEXT(N({YQ}),"#,##0.00")&"）","")'
                 f'&IF(N({R("s")})>0,"（开票未到账 "&TEXT(N({R("s")}),"#,##0.00")&"）","")'
                 f'&IF(N({R("u")})>0,"（未开票 "&TEXT(N({R("u")}),"#,##0.00")&"）",""),'
                 f'IF(N({R("v")})<-0.005,"多收 "&TEXT(-N({R("v")}),"#,##0.00"),"✔ 已结清")))'
                 f'&IF(N({R("t")})>0,"｜到账未开票 "&TEXT(N({R("t")}),"#,##0.00")&"，待补票","")'
                 f'&IF({OVERINV}>0.005,"｜多开票 "&TEXT({OVERINV},"#,##0.00")&"，核对发票","")'
                 f'&IF(N({R("q")})>0,"｜收现 "&TEXT(N({R("q")}),"#,##0.00"),"")'
                 f'&IF(N({R("x0")})>0,"｜收款不开票 "&TEXT(N({R("x0")}),"#,##0.00"),"")'
                 f'&IF(AND(N({R("j")})=0,N({R("k")})=0,N({R("l")})>0),"｜未填单价，总应收按开票金额暂估",""))',
            fmt='General', wrap=True)
        c = ws[f"{C['bz']}{r}"]
        style_like(c, src_txt, fmt='General', wrap=True)
        # 税额：原来写死 /1.06，改成跟着税率走
        if C.get('taxamt'):
            put('taxamt', f'=IF(OR(N({R("n")})=0,N({R("tax")})=0),"",ROUND(N({R("n")})/(1+N({R("tax")}))*N({R("tax")}),2))',
                auto=False, src=ws[f"{C['taxamt']}{r}"] if ws[f"{C['taxamt']}{r}"].has_style else src_in)
        # 追加块
        for k in ['zt', 'zq']:
            style_like(ws[f'{C[k]}{r}'], src_txt, fmt='General')
        NPY = (f'IF({R("zq")}="半年付",2,IF({R("zq")}="季度付",4,IF({R("zq")}="月付",12,'
               f'IF({R("zq")}="半月付",24,1))))')
        ONCE = f'OR({R("zq")}="",{R("zq")}="一次性")'
        put('tn', f'=IF(OR({BLK},NOT(ISNUMBER($D{r})),NOT(ISNUMBER($E{r})),{NOSCH}),"",IF($E{r}<$D{r},"",'
                  f'IF({ONCE},1,MAX(1,ROUND(($E{r}-$D{r}+1)/365*{NPY},0)))))', fmt='0')
        put('kn', f'=IF({R("tn")}="","",MIN({R("tn")},MAX(0,INT((TODAY()-付款宽限天数-$D{r})/(($E{r}-$D{r}+1)/{R("tn")}))+1)))', fmt='0')
        put('pg', f'=IF({R("tn")}="","",IF({ONCE},IF({R("kn")}>=1,"一次性·已到收款期","一次性·未到期"),'
                  f'{R("zq")}&" 第"&{R("kn")}&"/"&{R("tn")}&"期"))', fmt='General')
        put('jd', f'=IF(OR({R("tn")}="",N({R("l")})=0),"",ROUND(N({R("j")})*{R("kn")}/{R("tn")},2)+N({R("k")}))')
        put('yq', f'=IF({R("jd")}="","",MAX(0,ROUND({R("jd")}-{RCV},2)))')
        VR = f"${C['v']}$3:${C['v']}${TOT_END}"
        KR = f"${C['key']}$3:${C['key']}${TOT_END}"
        DR = f'$D$3:$D${TOT_END}'
        put('wn', f'=IF(OR({BLK},NOT(ISNUMBER($D{r}))),"",ROUND({R("lj")}-SUMIFS({VR},{KR},{R("key")},{DR},">="&$D{r}),2))')
        put('lj', f'=IF({BLK},"",SUMIFS({VR},{KR},{R("key")}))')
        put('key', f'=IF({BLK},"",TRIM($B{r})&"｜"&TRIM($C{r}))', fmt='General')
        put('kd', f'=IF({BLK},"",MIN(N({R("n")}),N({R("p")})))')
        put('nd', f'=IF(ISNUMBER($D{r}),YEAR($D{r}),"")', fmt='0')
        if r in pay0:
            ws[f"{C['zq']}{r}"] = pay0[r]

    # ── 下拉 ──
    dv = DataValidation(type='list', formula1=f'"{HTXT["合同未回"]}"', allow_blank=True,
                        showErrorMessage=False, showInputMessage=True,
                        promptTitle='合同回签', prompt='合同没回来就选「已送出未回 / 未见合同 / 未签」；也可以照旧手写，含“未”或“送出”都算没回。')
    ws.add_data_validation(dv); dv.add(f'A3:A{NR}')
    for k, lst, tip in [('zt', STATES, '安装免保 / 技术免保 / 质保期内会在「到期提醒」里一起提示。'),
                        ('zq', CYCLES, '决定「截至今天应收」按几期算：年付每年 1 期、半年付 2 期、季度付 4 期、月付 12 期、半月付 24 期，空着按一次性（整份合同一次付）。')]:
        dv = DataValidation(type='list', formula1='"' + ','.join(lst) + '"', allow_blank=True,
                            showErrorMessage=True, showInputMessage=True, promptTitle='填写提示', prompt=tip,
                            errorTitle='只能从下拉里选', error='公式按这几个字来认，请从下拉里选。')
        ws.add_data_validation(dv); dv.add(f'{C[k]}3:{C[k]}{NR}')

    # ── 条件格式 ──
    red = PatternFill('solid', fgColor='FFC7CE'); org = PatternFill('solid', fgColor='FFEB9C')
    ws.conditional_formatting.add(f'F3:F{NR}', FormulaRule(
        formula=['OR(ISNUMBER(SEARCH("已到期",$F3)),ISNUMBER(SEARCH("合同未回",$F3)))'],
        fill=red, font=Font(color='9C0006')))
    ws.conditional_formatting.add(f'F3:F{NR}', FormulaRule(
        formula=['ISNUMBER(SEARCH("即将到期",$F3))'], fill=org, font=Font(color='9C5700')))
    ws.conditional_formatting.add(f"{C['w']}3:{C['w']}{NR}", FormulaRule(
        formula=[f'LEFT(${C["w"]}3,1)="欠"'], font=Font(color='C00000', b=True)))
    ws.conditional_formatting.add(f"{C['w']}3:{C['w']}{NR}", FormulaRule(
        formula=[f'LEFT(${C["w"]}3,1)="✔"'], font=Font(color='006100')))
    ws.conditional_formatting.add(f"{C['yq']}3:{C['yq']}{NR}", FormulaRule(
        formula=[f'N(${C["yq"]}3)>0'], fill=red, font=Font(color='9C0006', b=True)))

    ws.freeze_panes = 'D3'
    if ws.auto_filter.ref:
        a, b = ws.auto_filter.ref.split(':')
        if CI(''.join(ch for ch in b if ch.isalpha())) >= qi:
            ws.auto_filter.ref = f"{a}:{C['lj']}{NR}"
    LOG.append(f'{sn}: 收款区插 6 列，末尾追加 {C["zt"]}~{C["kn"]}，付款周期按原备注预填 {len(pay0)} 行')
    print(f'  ✓ {sn}  收现={C["q"]} 收款不开票={C["x0"]} 开票应收={C["s"]} 应收余额={C["v"]} 追加 {C["zt"]}:{C["kn"]}')

# ══════════════════════════════════════════════════════════════
# 二、电梯档案
# ══════════════════════════════════════════════════════════════
ws = wb['电梯档案']
hsrc = ws['W1']
for col, txt, fill in [('X', '免保/质保\n到期日', FILL_NEW_IN), ('Y', '维保状态提醒', FILL_NEW_AUTO),
                       ('Z', '年检/校验提醒', FILL_NEW_AUTO), ('AA', '年检(认出的日期)', FILL_HELP),
                       ('AB', '限速器(认出的日期)', FILL_HELP), ('AC', '载重试验(认出的日期)', FILL_HELP)]:
    c = ws[f'{col}1']
    c.value = txt
    style_like(c, hsrc)
    c.fill = fill
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
wdim = None
for key, cd in list(ws.column_dimensions.items()):
    if (cd.min or 0) <= 23 <= (cd.max or 0):
        wdim = cd
dh = DimensionHolder(worksheet=ws)
for key, cd in ws.column_dimensions.items():
    lo, hi = cd.min or CI(key), cd.max or CI(key)
    if cd is wdim:
        dh['W'] = clone_dim(ws, cd, 23)
    else:
        dh[key] = cd
for col, wd, hid in [('X', 13, False), ('Y', 20, False), ('Z', 26, False), ('AA', 12, True),
                     ('AB', 12, True), ('AC', 12, True), ('AD', 8, True)]:
    dh[col] = clone_dim(ws, wdim, CI(col), width=wd, hidden=hid)
if wdim is not None:
    tail = clone_dim(ws, wdim, 31)
    tail.max = 16384
    dh['AE'] = tail
ws.column_dimensions = dh
c = ws['AD1']
c.value = '台量(计数)'
style_like(c, hsrc)
c.fill = FILL_HELP


def PD(x):
    """把 2026.09 / 2027.8 / "2026.10.1" / 真日期 都认成日期（只写到年月的，按当月最后一天算）"""
    return (f'IF(TRIM({x}&"")="","",IF(ISNUMBER({x}),IF({x}>10000,{x},'
            f'IFERROR(DATE(INT({x}),IF(ROUND(MOD({x},1)*100,0)>12,ROUND(MOD({x},1)*10,0),ROUND(MOD({x},1)*100,0))+1,0),"")),'
            f'IFERROR(IF(LEN({x})-LEN(SUBSTITUTE({x},".",""))>=2,DATEVALUE(SUBSTITUTE(TRIM({x}),".","-")),'
            f'DATE(VALUE(LEFT(TRIM({x}),4)),VALUE(MID(TRIM({x}),6,2))+1,0)),"")))')


def PART(lab, d):
    return (f'IF({d}="","",IF({d}<TODAY(),"{lab}已过期"&(TODAY()-{d})&"天 ",'
            f'IF({d}-TODAY()<=提醒天数,"{lab}剩"&({d}-TODAY())&"天 ","")))')


wsrc = ws['V3']
for r in range(2, LIFT_NR + 1):
    BLK = f'TRIM($B{r}&$C{r})=""'
    for col, f in [('AA', PD(f'$Q{r}')), ('AB', PD(f'$S{r}')), ('AC', PD(f'$T{r}'))]:
        c = ws[f'{col}{r}']
        c.value = f'={f}'
        style_like(c, wsrc, fmt='yyyy-mm-dd')
    c = ws[f'X{r}']
    style_like(c, wsrc, fmt='yyyy-mm-dd')
    c = ws[f'AD{r}']
    c.value = f'=IF({BLK},0,IF(N($N{r})=0,1,N($N{r})))'     # 台量空着的梯子也算 1 台
    style_like(c, wsrc, fmt='0')
    c = ws[f'Y{r}']
    c.value = (f'=IF({BLK},"",IF(OR($W{r}="安装免保",$W{r}="技术免保",$W{r}="质保期内"),'
               f'IF(NOT(ISNUMBER($X{r})),"▲免保到期日没填",IF($X{r}<TODAY(),"▲免保已到期，转收费",'
               f'IF($X{r}-TODAY()<=提醒天数,"免保剩"&($X{r}-TODAY())&"天","免保中"))),'
               f'IF($W{r}="暂停维保","暂停中",IF($W{r}="已解约","已解约",""))))')
    style_like(c, wsrc, fmt='General', auto=True)
    c = ws[f'Z{r}']
    c.value = f'=IF({BLK},"",TRIM({PART("年检", f"$AA{r}")}&{PART("限速器", f"$AB{r}")}&{PART("载重试验", f"$AC{r}")}))'
    style_like(c, wsrc, fmt='General', auto=True)
dv = DataValidation(type='list', formula1='"' + ','.join(STATES) + '"', allow_blank=True,
                    showErrorMessage=True, showInputMessage=True, promptTitle='维保状态',
                    prompt='安装免保 / 技术免保 / 质保期内 记得在 X 列填免保到期日。',
                    errorTitle='只能从下拉里选', error='汇总表按这几个字统计台量，请从下拉里选。')
ws.add_data_validation(dv); dv.add(f'W2:W{LIFT_NR}')
red = PatternFill('solid', fgColor='FFC7CE'); org = PatternFill('solid', fgColor='FFEB9C')
ws.conditional_formatting.add(f'Y2:Z{LIFT_NR}', FormulaRule(
    formula=['OR(LEFT(Y2,1)="▲",ISNUMBER(SEARCH("已过期",Y2)))'], fill=red, font=Font(color='9C0006')))
ws.conditional_formatting.add(f'Y2:Z{LIFT_NR}', FormulaRule(
    formula=['ISNUMBER(SEARCH("剩",Y2))'], fill=org, font=Font(color='9C5700')))
if ws.auto_filter.ref:
    ws.auto_filter.ref = f'A1:Z{max(ws.max_row, 304)}'
print('  ✓ 电梯档案  W 维保状态下拉 / X 免保到期 / Y 状态提醒 / Z 年检校验提醒')

exec(open(os.path.join(HERE, 'fix_原表公式_汇总.py'), encoding='utf-8').read())

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('\n'.join('  · ' + x for x in LOG))
print(f'\n✅ {OUT}')
