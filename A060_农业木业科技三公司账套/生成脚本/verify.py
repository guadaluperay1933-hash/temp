# -*- coding: utf-8 -*-
"""A060 · 独立复算校验：拿 model.py 的纯 Python 结果，逐项比对表里算出来的值"""
import sys
from openpyxl import load_workbook
import a060_data as D
import model

XL = sys.argv[1] if len(sys.argv) > 1 else '../三公司财务账套_A060.xlsx'
wb = load_workbook(XL, data_only=True)
M = model.build()
R2 = model.R2
GD, ACC, BANK = model.GD, model.ACC, model.BANK
ok = bad = 0
def chk(name, a, b, tol=0.02):
    global ok, bad
    try: d = abs(float(a or 0) - float(b or 0))
    except (TypeError, ValueError): d = None
    if d is not None and d <= tol:
        ok += 1; print(f'  ✓ {name}')
    else:
        bad += 1; print(f'  ✗ {name}   表内={a!r}  复算={b!r}')
def chk_txt(name, a, want):
    global ok, bad
    if str(a).startswith(want): ok += 1; print(f'  ✓ {name}  {a}')
    else: bad += 1; print(f'  ✗ {name}   表内={a!r}  应以「{want}」开头')

print('\n一、购销流水（不含税 / 税额 / 价税合计 / 结转成本）')
ws = wb[D.SH_BUY]
e = 0
for b in M['buy']:
    r = D.BY_0 + b['r']
    for col, key in (('S', 'net'), ('T', 'tax'), ('U', 'gross'), ('AA', 'cost')):
        v = ws[f'{col}{r}'].value
        if abs(float(v or 0) - b[key]) > 0.02:
            e += 1; print(f'  ✗ 第{r}行 {col} 表内={v} 复算={b[key]}')
chk('20 条购销全部列一致（差异条数）', e, 0)

print('\n二、进销存台账（加权平均单价 / 结存数量 / 结存金额）')
ws = wb[D.SH_INV]
eq = ea = ev = 0
for i, g in enumerate(D.GOODS):
    r = D.GD_0 + i
    k = (g[4], g[0])
    v = M['inv'].get(k)
    if not v: continue
    if abs(float(ws[f'N{r}'].value or 0) - v['avg']) > 0.01: ev += 1; print(f'  ✗ {k} 单价 {ws[f"N{r}"].value} vs {v["avg"]}')
    if abs(float(ws[f'S{r}'].value or 0) - v['eq']) > 0.001: eq += 1; print(f'  ✗ {k} 结存数量 {ws[f"S{r}"].value} vs {v["eq"]}')
    if abs(float(ws[f'U{r}'].value or 0) - v['ea']) > 0.05: ea += 1; print(f'  ✗ {k} 结存金额 {ws[f"U{r}"].value} vs {v["ea"]}')
chk('加权平均单价差异数', ev, 0); chk('结存数量差异数', eq, 0); chk('结存金额差异数', ea, 0)

print('\n三、生产加工（投入金额 / 产出成本分摊）')
ws = wb[D.SH_PROD]
e = 0
for p in M['prod']:
    r = D.PD_0 + p['r']
    v = float(ws[f'U{r}'].value or 0)
    if abs(v - p['amt']) > 0.02:
        e += 1; print(f'  ✗ 第{r}行 {p["bt"]} {p["kind"]} 表内={v} 复算={p["amt"]}')
chk('36 条生产加工记账金额一致（差异条数）', e, 0)

print('\n四、资金流水（金额拆分 + 逐行账户余额）')
ws = wb[D.SH_CASH]
e = 0
run = {k: v['open'] for k, v in BANK.items()}
for c in M['cash']:
    r = D.CS_0 + c['r']
    run[c['acct']] = R2(run[c['acct']] + c['cin'] - c['cout'])
    for col, want in (('Q', c['gross']), ('S', c['net']), ('T', c['tax']), ('V', run[c['acct']])):
        v = float(ws[f'{col}{r}'].value or 0)
        if abs(v - want) > 0.02:
            e += 1; print(f'  ✗ 第{r}行 {col} 表内={v} 复算={want}')
chk('29 条资金流水（含实时余额）一致（差异条数）', e, 0)

print('\n五、资金账户期末余额 + 账实对账')
ws = wb[D.SH_ACCT]
for i, a in enumerate(D.ACCOUNTS):
    chk(f'账户 {a[1]} 期末余额', ws[f'K{D.AC_0+i}'].value, M['acctbal'][a[1]])
tot = sum(M['acctbal'].values())
chk('全部账户余额合计', ws[f'K{D.AC_1+2}'].value, tot)
co_cash = sum(v['end'] for k, v in M['tb'].items() if k[1] in ('库存现金', '银行存款'))
chk('三家货币资金合计（应等于账户合计）', co_cash, tot)
chk_txt('账实对账结论', ws[f'G{D.AC_1+5}'].value, '√')

print('\n六、记账分录（每家公司借贷合计 + 逐科目余额）')
ws = wb[D.SH_TB]
tb = M['tb']
import collections
legsum = collections.defaultdict(lambda: [0.0, 0.0])
for x in M['legs']:
    legsum[x['co']][0] += x['dr']; legsum[x['co']][1] += x['cr']
wsv = wb[D.SH_VOU]
vd = collections.defaultdict(lambda: [0.0, 0.0])
for r in range(D.V_0, D.V_1 + 1):
    co = wsv[f'F{r}'].value
    if not co: continue
    vd[co][0] += float(wsv[f'I{r}'].value or 0); vd[co][1] += float(wsv[f'J{r}'].value or 0)
for co in D.CO_NAMES:
    chk(f'{co} 分录借方合计', vd[co][0], legsum[co][0])
    chk(f'{co} 分录贷方合计', vd[co][1], legsum[co][1])
    chk(f'{co} 借＝贷', vd[co][0], vd[co][1])
e = 0
pos = {}
for ci, co in enumerate(D.CO_NAMES):
    for ai, a in enumerate(D.ACCS):
        pos[(co, a[1])] = D.TB_0 + ci * (D.ACC_1 - D.ACC_0 + 1) + ai
for (co, acc), r in pos.items():
    v = tb.get((co, acc))
    want_open = v['open'] if v else 0.0
    want_end = v['end'] if v else 0.0
    if abs(float(ws[f'G{r}'].value or 0) - want_open) > 0.02:
        e += 1; print(f'  ✗ {co}/{acc} 年初 {ws[f"G{r}"].value} vs {want_open}')
    if abs(float(ws[f'L{r}'].value or 0) - want_end) > 0.02:
        e += 1; print(f'  ✗ {co}/{acc} 期末 {ws[f"L{r}"].value} vs {want_end}')
chk(f'科目余额表 {len(pos)} 个科目×公司 年初/期末一致（差异数）', e, 0)

print('\n七、三家利润表')
PL_ROW = {lab.strip('　'): 6 + i for i, (lab, _a, _b) in enumerate([
    ('一、营业收入',0,0),('减：营业成本',0,0),('税金及附加',0,0),('销售费用',0,0),('管理费用',0,0),
    ('研发费用',0,0),('财务费用',0,0),('加：投资收益',0,0),('二、营业利润',0,0),('加：营业外收入',0,0),
    ('减：营业外支出',0,0),('三、利润总额',0,0),('减：所得税费用',0,0),('四、净利润',0,0)])}
for co in D.CO_NAMES:
    ws = wb[f'{co}-利润表']; p = M['pl'][co]
    for lab, key in (('一、营业收入','营业收入'),('减：营业成本','营业成本'),('管理费用','管理费用'),
                     ('销售费用','销售费用'),('财务费用','财务费用'),('税金及附加','税金及附加'),
                     ('三、利润总额','利润总额'),('减：所得税费用','所得税费用'),('四、净利润','净利润')):
        chk(f'{co} {lab}', ws[f'C{PL_ROW[lab]}'].value, p.get(key, 0.0))

print('\n八、三家资产负债表')
BSL = {'货币资金':7,'应收账款':8,'预付款项':9,'其他应收款':10,'存货':11,'流动资产合计':12,
       '固定资产净值':16,'生产性生物资产净值':19,'无形资产净值':22,'非流动资产合计':24,'资产总计':25}
BSR = {'短期借款':7,'应付账款':8,'应付职工薪酬':10,'应交税费':11,'流动负债合计':13,'负债合计':16,
       '实收资本':18,'未分配利润':21,'所有者权益合计':22,'负债和所有者权益总计':23,'平衡':25}
for co in D.CO_NAMES:
    ws = wb[f'{co}-资产负债表']; b = M['bs'][co]
    chk(f'{co} 货币资金', ws[f'C{BSL["货币资金"]}'].value, b['item']['货币资金'])
    chk(f'{co} 存货', ws[f'C{BSL["存货"]}'].value, b['item']['存货'])
    chk(f'{co} 流动资产合计', ws[f'C{BSL["流动资产合计"]}'].value, b['流动资产'])
    chk(f'{co} 资产总计', ws[f'C{BSL["资产总计"]}'].value, b['资产总计'])
    chk(f'{co} 负债合计', ws[f'G{BSR["负债合计"]}'].value, b['负债合计'])
    chk(f'{co} 未分配利润', ws[f'G{BSR["未分配利润"]}'].value, b['未分配利润'])
    chk(f'{co} 所有者权益合计', ws[f'G{BSR["所有者权益合计"]}'].value, b['权益合计'])
    chk_txt(f'{co} 资产负债表平衡', ws[f'G{BSR["平衡"]}'].value, '√')

print('\n九、税费台账（全年合计） / 内部交易 / 校验中心 / 首页')
ws = wb[D.SH_TAX]
for ci, co in enumerate(D.CO_NAMES):
    r = 6 + ci * 14 + 12
    chk(f'{co} 全年销项', ws[f'C{r}'].value, M['vat'][co]['out'])
    chk(f'{co} 全年进项', ws[f'D{r}'].value, M['vat'][co]['inp'])
ws = wb[D.SH_IC]
bad_ic = sum(1 for r in range(6, 12) if str(ws[f'J{r}'].value or '').startswith('✗'))
chk('内部交易核对 ✗ 的对数', bad_ic, 0)
ws = wb[D.SH_CHK]
rows = [r for r in range(6, 60) if ws[f'B{r}'].value]
nbad = sum(1 for r in rows if str(ws[f'F{r}'].value or '').startswith('✗'))
print(f'  校验中心共 {len(rows)} 项')
for r in rows:
    if str(ws[f'F{r}'].value or '').startswith('✗'):
        print(f'    ✗ {ws[f"B{r}"].value} / {ws[f"C{r}"].value}: {ws[f"F{r}"].value}')
chk('校验中心未通过项数', nbad, 0)
ws = wb[D.SH_HOME]
chk('首页 三家营业收入合计', ws['E8'].value, sum(M['pl'][c].get('营业收入', 0) for c in D.CO_NAMES))
chk('首页 三家净利润合计', ws['E13'].value, sum(M['pl'][c]['净利润'] for c in D.CO_NAMES))

print(f'\n=========  {ok} 项一致，{bad} 项不一致  =========')
sys.exit(1 if bad else 0)
