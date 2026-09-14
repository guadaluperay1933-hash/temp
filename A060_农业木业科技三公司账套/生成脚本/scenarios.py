# -*- coding: utf-8 -*-
"""A060 · 筛选参数回归：换年度、换起止日期、换公司，重算后再验一遍结构恒等式和取数"""
import shutil, subprocess, sys, json, datetime as dt
from openpyxl import load_workbook
import a060_data as D
import model

SRC = sys.argv[1] if len(sys.argv) > 1 else '/tmp/out.xlsx'
TMP = '/tmp/scen.xlsx'
M = model.build()
R2 = model.R2
ACC = model.ACC

# 带筛选带的表：(表名, 是否带公司下拉)
BANDS = [(D.SH_TB, False), (D.SH_INVQ, False), (D.SH_MERGE, False),
         (D.SH_AR, True), (D.SH_EXP, False), (D.SH_IC, False)]
for co in D.CO_NAMES:
    BANDS += [(f'{co}-资产负债表', False), (f'{co}-利润表', False)]

SCEN = [
    ('全部期间',      None, None, None, '1900-01-01', '2199-12-31'),
    ('整年 2026',     2026, None, None, '2026-01-01', '2026-12-31'),
    ('单月 2026-08',  None, '2026-08-01', '2026-08-31', '2026-08-01', '2026-08-31'),
    ('上半年',        None, '2026-01-01', '2026-06-30', '2026-01-01', '2026-06-30'),
    ('三季度',        None, '2026-07-01', '2026-09-30', '2026-07-01', '2026-09-30'),
]
CO_PICK = '木业'
ok = bad = 0
def chk(name, a, b, tol=0.02):
    global ok, bad
    try: d = abs(float(a or 0) - float(b or 0))
    except (TypeError, ValueError): d = None
    if d is not None and d <= tol: ok += 1
    else: bad += 1; print(f'    ✗ {name}  表内={a!r} 复算={b!r}')
def chk_txt(name, a, want='√'):
    global ok, bad
    if str(a).startswith(want): ok += 1
    else: bad += 1; print(f'    ✗ {name}  {a!r}')

PL_ROW = {'一、营业收入':6,'减：营业成本':7,'税金及附加':8,'销售费用':9,'管理费用':10,
          '研发费用':11,'财务费用':12,'三、利润总额':17,'减：所得税费用':18,'四、净利润':19}
# 内账含税口径：营业收入/成本就是含税数，税金及附加＝当期实缴税费
BSL_TOT, BSR_TOT, BSR_CHK = 25, 23, 25

for nm, y, s_, e_, ds, de in SCEN:
    shutil.copy(SRC, TMP)
    wb = load_workbook(TMP)
    for sh, has_co in BANDS:
        ws = wb[sh]
        cy, cs, ce = ('D3', 'F3', 'H3') if has_co else ('B3', 'D3', 'F3')
        if has_co: ws['B3'] = CO_PICK
        ws[cy] = y
        ws[cs] = dt.date.fromisoformat(s_) if s_ else None
        ws[ce] = dt.date.fromisoformat(e_) if e_ else None
    wb.save(TMP)
    out = subprocess.run([sys.executable, '/mnt/skills/public/xlsx/scripts/recalc.py', TMP, '300'],
                         capture_output=True, text=True)
    j = json.loads(out.stdout)
    print(f'\n【{nm}】重算 {j.get("status")}  错误 {j.get("total_errors")}')
    ok0, bad0 = ok, bad
    if j.get('total_errors'): bad += 1
    wb = load_workbook(TMP, data_only=True)
    net = model.period_net(M, ds, de)

    for co in D.CO_NAMES:
        ws = wb[f'{co}-利润表']
        for lab, item in (('一、营业收入','营业收入'),('减：营业成本','营业成本'),
                          ('管理费用','管理费用'),('销售费用','销售费用'),
                          ('财务费用','财务费用'),('税金及附加','税金及附加'),
                          ('减：所得税费用','所得税费用')):
            want = R2(sum(v for k, v in net.items() if k[0] == co and ACC[k[1]]['item'] == item))
            chk(f'{co} 利润表本期 {lab}', ws[f'B{PL_ROW[lab]}'].value, want)
        chk(f'{co} 净利＝利润总额−所得税',
            ws[f'B{PL_ROW["四、净利润"]}'].value,
            R2(float(ws[f'B{PL_ROW["三、利润总额"]}'].value or 0)
               - float(ws[f'B{PL_ROW["减：所得税费用"]}'].value or 0)))
        bs = wb[f'{co}-资产负债表']
        chk_txt(f'{co} 资产负债表期末平衡', bs[f'G{BSR_CHK}'].value)
        chk_txt(f'{co} 资产负债表年初平衡', bs[f'F{BSR_CHK}'].value)

    tb = wb[D.SH_TB]
    e1 = e2 = 0
    for r in range(D.TB_0, D.TB_1 + 1):
        if not tb[f'C{r}'].value: continue
        h, k, l = (float(tb[f'{c}{r}'].value or 0) for c in 'HKL')
        if abs(h + k - l) > 0.02: e1 += 1
        g = float(tb[f'G{r}'].value or 0)
        co, acc = tb[f'A{r}'].value, tb[f'C{r}'].value
        want = net.get((co, acc), 0.0)
        if abs(k - want) > 0.02: e2 += 1
    chk('科目余额表 期初＋本期＝期末（差异行数）', e1, 0)
    chk('科目余额表 本期净额 vs 复算（差异行数）', e2, 0)

    iq = wb[D.SH_INVQ]
    e3 = 0
    for r in range(7, 7 + len(D.GOODS)):
        if not iq[f'B{r}'].value: continue
        E, G, I, K = (float(iq[f'{c}{r}'].value or 0) for c in 'EGIK')
        if abs(E + G - I - K) > 0.001: e3 += 1
    chk('收发存查询 期初＋入−出＝期末（差异行数）', e3, 0)

    ar = wb[D.SH_AR]
    e4 = 0
    for r in range(7, 7 + len(D.PARTNERS)):
        if not ar[f'A{r}'].value: continue
        C, Dv, E, F = (float(ar[f'{c}{r}'].value or 0) for c in 'CDEF')
        G, H, I, J = (float(ar[f'{c}{r}'].value or 0) for c in 'GHIJ')
        if abs(C + Dv - E - F) > 0.02 or abs(G + H - I - J) > 0.02: e4 += 1
    chk('往来台账 年初＋发生−收付＝期末（差异行数）', e4, 0)

    mg = wb[D.SH_MERGE]
    for c in 'BCDEFG':
        chk_txt(f'合并报表 {c} 列平衡', mg[f'{c}{51}'].value)
    ic = wb[D.SH_IC]
    chk('内部交易核对 ✗ 的对数',
        sum(1 for r in range(6, 12) if str(ic[f'I{r}'].value or '').startswith('✗')), 0)
    ex = wb[D.SH_EXP]
    e5 = 0
    for r in range(7, 7 + len(D.EXPENSES)):
        if not ex[f'A{r}'].value: continue
        if abs(sum(float(ex[f'{c}{r}'].value or 0) for c in 'BCD') - float(ex[f'E{r}'].value or 0)) > 0.02:
            e5 += 1
    chk('费用统计 三家合计＝各家之和（差异行数）', e5, 0)
    print(f'    校验 {ok-ok0}/{ok-ok0+bad-bad0}')

print(f'\n=========  合计 {ok} 项通过，{bad} 项不通过  =========')
sys.exit(1 if bad else 0)
