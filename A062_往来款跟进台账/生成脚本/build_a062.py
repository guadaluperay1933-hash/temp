# -*- coding: utf-8 -*-
"""A062【往来款跟进】改造：
   用户只录「销售金额 / 销售费用(赠送券) / 购货款」，
   已收款、已付款从【数据录入】自动取，应收款、尚欠款、未收总计全部自动算。

   跑法：python3 build_a062.py [输出文件]
"""
import os, sys, shutil
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from style import *

HERE = os.path.dirname(os.path.abspath(__file__))
SRC  = os.path.join(HERE, '..', '参考', '原A062_多账户收支台账.xlsx')
OUT  = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '..', 'A062_多账户收支台账.xlsx')

WLK  = '往来款跟进'
SJ   = '数据录入'
JC   = '基础资料'
REC  = '客户供应商往来对账'          # 新增的表

F0, F1 = 11, 510                     # 明细数据区
SJ0, SJ1 = 5, 4999                   # 数据录入的数据区（沿用原表其它表的写法）

# 数据录入的列：F收入 G支出 I产品 J客户 K供应商 B日期
SJ_IN   = f"{SJ}!$F${SJ0}:$F${SJ1}"
SJ_OUT  = f"{SJ}!$G${SJ0}:$G${SJ1}"
SJ_CUST = f"{SJ}!$J${SJ0}:$J${SJ1}"
SJ_SUPP = f"{SJ}!$K${SJ0}:$K${SJ1}"

print('读取原表', os.path.abspath(SRC))
wb = openpyxl.load_workbook(SRC)
ws = wb[WLK]

# ---------------------------------------------------------------- ①
# 先把用户已经手工填的「已收款」原样搬到备份列 T，一个数字都不能丢。
# 搬完 G 列才敢改成公式——否则那 38 行会全部归零。
bak_recv, bak_paid = {}, {}
for r in range(F0, F1 + 1):
    v = ws.cell(r, 7).value                 # G 已收款
    if isinstance(v, (int, float)) and v: bak_recv[r] = float(v)
    v = ws.cell(r, 16).value                # P 已付款
    if isinstance(v, (int, float)) and v: bak_paid[r] = float(v)
print(f'  备份 已收款 {len(bak_recv)} 行 合计 {sum(bak_recv.values()):,.2f}；'
      f'已付款 {len(bak_paid)} 行 合计 {sum(bak_paid.values()):,.2f}')

# 顺手记下改造前的几个数，改完要逐个对上
before = dict(
    F=sum(float(ws.cell(r,6).value or 0) for r in range(F0,F1+1) if isinstance(ws.cell(r,6).value,(int,float))),
    G=sum(bak_recv.values()),
    J=sum(float(ws.cell(r,10).value or 0) for r in range(F0,F1+1) if isinstance(ws.cell(r,10).value,(int,float))),
    O=sum(float(ws.cell(r,15).value or 0) for r in range(F0,F1+1) if isinstance(ws.cell(r,15).value,(int,float))),
)
print('  改造前：销售金额 {F:,.2f}  已收款 {G:,.2f}  销售费用 {J:,.2f}  购货款 {O:,.2f}'.format(**before))

# ---------------------------------------------------------------- ②
# 新增列的表头（T~AA）。A~S 的列位一个都不动，
# 这样原有的合并单元格、别的表的引用都不会错位。
NEW = [
    ('T', 16, '已收款\n(原手工记录)',      False, '你原来手工填在「已收款」里的数。\n数据录入里还没记这个客户的回款时，左边就用这个数顶着。'),
    ('U',  0, '本行应收',                  True,  '销售金额 − 销售费用'),
    ('V',  0, '同客户累计应收',            True,  '这个客户到本行为止一共该收多少'),
    ('W',  0, '该客户回款·数据录入',        True,  '从【数据录入】按客户名汇总的收入'),
    ('X', 14, '已付款\n(原手工记录)',      False, '同上，应付侧的手工备份'),
    ('Y',  0, '同供货商累计购货',          True,  ''),
    ('Z',  0, '该供货商付款·数据录入',      True,  '从【数据录入】按供货商汇总的支出'),
    ('AA',22, '核　对',                    False, '这一行有没有填错、名字对不对得上'),
]
for col, width, title, hidden, note in NEW:
    put(ws, f'{col}9',  title, font=F_HD0, fill=FILL_HD2, align=CLW)
    put(ws, f'{col}10', note or title, font=F_NOTE, fill=FILL_HD, align=CLW)
    ws.column_dimensions[col].hidden = hidden
    if width: ws.column_dimensions[col].width = width
ws.row_dimensions[10].height = 34

# ---------------------------------------------------------------- ③
# 明细区逐行写公式
MANUAL = {                     # 手工录入的格子：淡黄底 + 深蓝字
    'B': (F_IN, DATE), 'C': (F_IN, 'General'), 'D': (F_IN, 'General'), 'E': (F_IN, 'General'),
    'F': (F_IN, NUM),  'J': (F_IN, NUM), 'K': (F_IN, 'General'), 'L': (F_IN, 'General'),
    'M': (F_IN0, DATE),'N': (F_IN0, 'General'), 'O': (F_IN0, NUM), 'S': (F_IN0, 'General'),
    'T': (F_IN, NUM),  'X': (F_IN0, NUM),
}
for r in range(F0, F1 + 1):
    D, N_ = f'$D{r}', f'$N{r}'

    # —— 序号：跳过空行也不会乱；MAX 碰到表头文字算 0，所以第 11 行自然是 1
    put(ws, f'A{r}', f'=IF(AND({D}="",{N_}=""),"",MAX($A$10:A{r-1})+1)',
        font=F_AUTO, fill=FILL_AUTO, fmt='General', align=CL)

    # —— 手工录入格：只改样式，值一个都不碰
    for col, (fnt, fmt) in MANUAL.items():
        c = ws[f'{col}{r}']
        c.font, c.fill, c.number_format, c.border = fnt, FILL_IN, fmt, BD
        c.alignment = CL

    # —— 备份列：把刚才存下来的手工已收款/已付款写回 T / X
    if r in bak_recv: ws[f'T{r}'].value = bak_recv[r]
    if r in bak_paid: ws[f'X{r}'].value = bak_paid[r]

    # —— 隐藏的算料列
    put(ws, f'U{r}', f'=IF({D}="","",ROUND(N($F{r})-N($J{r}),2))', fmt=NUM)
    put(ws, f'V{r}', f'=IF({D}="","",ROUND(SUMIFS($F${F0}:F{r},$D${F0}:D{r},{D})'
                     f'-SUMIFS($J${F0}:J{r},$D${F0}:D{r},{D}),2))', fmt=NUM)
    put(ws, f'W{r}', f'=IF({D}="","",ROUND(SUMIFS({SJ_IN},{SJ_CUST},{D}),2))', fmt=NUM)
    put(ws, f'Y{r}', f'=IF({N_}="","",ROUND(SUMIFS($O${F0}:O{r},$N${F0}:N{r},{N_}),2))', fmt=NUM)
    put(ws, f'Z{r}', f'=IF({N_}="","",ROUND(SUMIFS({SJ_OUT},{SJ_SUPP},{N_}),2))', fmt=NUM)

    # —— 已收款：数据录入里记了这个客户的回款就按回款先来后到分摊；
    #    还没记的，就沿用 T 列那个手工数，免得你原来的账一夜之间全变 0。
    put(ws, f'G{r}',
        f'=IF({D}="","",IF(COUNTIF({SJ_CUST},{D})=0,N($T{r}),'
        f'ROUND(MIN(MAX(0,N($W{r})-(N($V{r})-N($U{r}))),N($U{r})),2)))',
        font=F_AUTO, fill=FILL_AUTO, fmt=NUM, align=CL)
    put(ws, f'H{r}', f'=IF({D}="","",ROUND(N($F{r})-N($G{r})-N($J{r}),2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY, align=CL)
    # 未收总计：直接对 H 列求和，不再一环扣一环，中间空一行也断不了
    put(ws, f'I{r}', f'=IF({D}="","",ROUND(SUM($H${F0}:H{r}),2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=NUM, align=CL)

    put(ws, f'P{r}',
        f'=IF({N_}="","",IF(COUNTIF({SJ_SUPP},{N_})=0,N($X{r}),'
        f'ROUND(MIN(MAX(0,N($Z{r})-(N($Y{r})-N($O{r}))),N($O{r})),2)))',
        font=F_AUTO0, fill=FILL_AUTO, fmt=NUM, align=CL)
    put(ws, f'Q{r}', f'=IF({N_}="","",ROUND(N($O{r})-N($P{r}),2))',
        font=F_AUTO0, fill=FILL_AUTO, fmt=NUM, align=CL)
    put(ws, f'R{r}', f'=IF({N_}="","",ROUND(SUM($Q${F0}:Q{r}),2))',
        font=F_AUTO0, fill=FILL_AUTO, fmt=MONEY, align=CL)

    put(ws, f'AA{r}',
        f'=IF(AND({D}="",{N_}=""),"",'
        f'IF(AND({D}<>"",COUNTIF({JC}!$E:$E,{D})=0),"客户名不在基础资料里",'
        f'IF(AND($E{r}<>"",COUNTIF({JC}!$B:$B,$E{r})=0),"产品名不在基础资料里",'
        f'IF(AND({N_}<>"",COUNTIF({JC}!$G:$G,{N_})=0),"供货商不在基础资料里",'
        f'IF(AND($C{r}<>"",ISNUMBER($C{r}),$C{r}<10000),"客户生日填得不对（不是日期）",'
        f'IF(N($J{r})>N($F{r}),"销售费用大于销售金额",'
        f'IF(AND({D}<>"",COUNTIF({SJ_CUST},{D})=0,N($T{r})<>0),"用的还是原手工已收款",'
        f'IF(AND({D}<>"",COUNTIF({SJ_CUST},{D})>0,ROUND(N($G{r})-N($T{r}),2)<>0,N($T{r})<>0),'
        f'"自动算的已收款和原手工数对不上","√"))))))))',
        font=F_T0, fill=FILL_CHK, fmt='General', align=CLW)
    ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 18) or 18

print(f'  明细区 {F0}~{F1} 行公式写好')

# ---------------------------------------------------------------- ④
# 左上角那块汇总：原来几个公式的区间各写各的（F11:F2001 / F11:F2000 / J11:J1001
# / O11:O1000 / P11:P1000），而且 S3/S4 的 SUMIFS 条件指到了 S2 这个表头文字
# 「供货商2」上，永远算出 0。这里全部重写成和数据区一致的区间。
SUM_BLOCK = [
    ('L2', '总销售金额', 'M2', f'=ROUND(SUM($F${F0}:$F${F1}),2)'),
    ('L3', '购货款',     'M3', f'=ROUND(SUM($O${F0}:$O${F1}),2)'),
    ('L4', '销售费用',   'M4', f'=ROUND(SUM($J${F0}:$J${F1}),2)'),
    ('L5', '毛　利',     'M5', '=ROUND($M$2-$M$3-$M$4,2)'),
    ('O2', '应收款项',   'P2', f'=ROUND(SUM($F${F0}:$F${F1})-SUM($J${F0}:$J${F1}),2)'),
    ('O3', '已收款',     'P3', f'=ROUND(SUM($G${F0}:$G${F1}),2)'),
    ('O4', '尚欠款',     'P4', '=ROUND($P$2-$P$3,2)'),
    ('R2', '购货款',     'S2', f'=ROUND(SUM($O${F0}:$O${F1}),2)'),
    ('R3', '已付款',     'S3', f'=ROUND(SUM($P${F0}:$P${F1}),2)'),
    ('R4', '尚欠款',     'S4', '=ROUND($S$2-$S$3,2)'),
]
for lab_c, lab, val_c, f in SUM_BLOCK:
    put(ws, lab_c, lab, font=F_HD, fill=FILL_SUM, align=CL, fmt='General')
    put(ws, val_c, f, font=F_SUM, fill=FILL_SUM, align=CL, fmt=MONEY)
# R5/S5 原来是「尚欠款」，上面挪到 R4/S4 了，第 5 行改成两边都看得到的核对
put(ws, 'R5', '应收-应付', font=F_HD, fill=FILL_SUM, align=CL, fmt='General')
put(ws, 'S5', '=ROUND($P$4-$S$4,2)', font=F_SUM, fill=FILL_SUM, align=CL, fmt=MONEY)

# A2:J5 原来是一大块空的合并区，拿来写使用说明
if 'A2:J5' not in [str(m) for m in ws.merged_cells.ranges]:
    try: ws.merge_cells('A2:J5')
    except Exception: pass
put(ws, 'A2',
    '填写说明：淡黄色格子是你要填的 —— 日期、客户名称、销售产品、销售金额、销售费用(赠送券)、'
    '支付方式、备注，以及右边应付侧的供货商、购货款。\n'
    '淡灰色格子全部自动算，不要手动改：已收款、尚欠款、未收总计、已付款、尚欠款总计。\n'
    '「已收款」「已付款」自动从【数据录入】取：在数据录入里按日记流水记账时，'
    '把「客户」或「供应商」那一列选上，这里就自动认领。\n'
    '数据录入里还没记到某个客户的回款之前，这里沿用你原来手工填的数（存在隐藏的 T 列），'
    '最右边「核对」列会提示是哪一种。',
    font=F_NOTE, fill=FILL_SUM, align=Alignment(horizontal='left', vertical='center', wrap_text=True))

# ---------------------------------------------------------------- ⑤
# 下拉：照抄【数据录入】的写法，直接指到基础资料整列，
# 以后基础资料里加客户加产品，这边不用动。
for sqref, src in ((f'D{F0}:D{F1}', f'{JC}!$E:$E'),
                   (f'E{F0}:E{F1}', f'{JC}!$B:$B'),
                   (f'K{F0}:K{F1}', f'{JC}!$A:$A'),
                   (f'N{F0}:N{F1}', f'{JC}!$G:$G')):
    dv = DataValidation(type='list', formula1=src, allow_blank=True, showDropDown=False)
    dv.error = '请从下拉里选；要加新的先去【基础资料】那张表加一行。'
    dv.errorTitle = '不在基础资料清单里'
    ws.add_data_validation(dv); dv.add(sqref)

# 核对列出问题就整行标红
ws.conditional_formatting.add(f'A{F0}:AA{F1}',
    FormulaRule(formula=[f'AND($AA{F0}<>"",$AA{F0}<>"√")'], fill=FILL_WARN))

ws.freeze_panes = f'D{F0}'
ws.column_dimensions['G'].width = 11.5
ws.column_dimensions['P'].width = 11.5
ws.row_dimensions[7].hidden = True

# ---------------------------------------------------------------- ⑥
# 新表【客户供应商往来对账】：按客户一行、按供货商一行，
# 左边是【往来款跟进】记的「该收多少」，右边是【数据录入】记的「实际收到多少」，
# 中间差多少一目了然 —— 这才是「自动得出应收款、尚欠款」真正好用的地方。
if REC in wb.sheetnames: del wb[REC]
rc = wb.create_sheet(REC, wb.sheetnames.index(WLK) + 1)

WF = lambda c: f"'{WLK}'!${c}${F0}:${c}${F1}"      # 往来款跟进某一列

CUST_H, CUST_0, CUST_N = 5, 6, 65                  # 客户块：表头行 / 首行 / 末行
SUPP_H, SUPP_0, SUPP_N = 69, 70, 89                # 供货商块

put(rc, 'A1', '客户 · 供货商 往来对账表', font=Font(name=YH, sz=16, bold=True, color='FF1F4E79'),
    align=LF, border=False)
rc.merge_cells('A1:J1')
put(rc, 'A2',
    '★ 全自动，不用填。左半边「该收/该付」来自【往来款跟进】，右半边「实收/实付」来自【数据录入】。\n'
    '　 差额 ≠ 0，说明两张表没对上：要么【数据录入】那笔回款没选客户名，要么【往来款跟进】的已收款还是原来的手工数。\n'
    '　 客户和供货商名单直接跟着【基础资料】走，你在基础资料里加一个客户，这里自动多一行。',
    font=F_NOTE, align=Alignment(horizontal='left', vertical='center', wrap_text=True), border=False)
rc.merge_cells('A2:J2'); rc.row_dimensions[2].height = 46

def block(h_row, r0, r1, kind):
    """kind='客户' 或 '供货商'"""
    is_c = kind == '客户'
    title = '一、按客户：应收款对账' if is_c else '二、按供货商：应付款对账'
    put(rc, f'A{h_row-1}', title, font=Font(name=YH, sz=12, bold=True, color='FFFFFFFF'),
        fill=PatternFill('solid', fgColor='FF4472C4'), align=LF)
    rc.merge_cells(f'A{h_row-1}:J{h_row-1}')
    rc.row_dimensions[h_row-1].height = 22

    heads = (['序号', '客户', '销售金额', '销售费用\n(赠送券)', '应收款', '已收款\n(往来款跟进)',
              '已收款\n(数据录入)', '尚欠款', '差额', '提　示']
             if is_c else
             ['序号', '供货商', '购货款', '', '应付款', '已付款\n(往来款跟进)',
              '已付款\n(数据录入)', '尚欠款', '差额', '提　示'])
    for i, t in enumerate(heads):
        put(rc, f'{get_column_letter(i+1)}{h_row}', t or None,
            font=F_HD, fill=FILL_HD, align=CLW)
    rc.row_dimensions[h_row].height = 34

    # 基础资料里客户在 E 列、供货商在 G 列，都是从第 2 行开始
    src_col, off = ('E', r0 - 2) if is_c else ('G', r0 - 2)
    for r in range(r0, r1 + 1):
        br = r - off                                  # 对应基础资料的行号
        nm = f'$B{r}'
        put(rc, f'A{r}', f'=IF($B{r}="","",MAX($A${h_row}:A{r-1})+1)',
            font=F_AUTO, fill=FILL_AUTO, fmt='General', align=CL)
        put(rc, f'B{r}', f'=IF({JC}!${src_col}{br}="","",{JC}!${src_col}{br})',
            font=F_T1 if is_c else F_T0, fill=FILL_AUTO, fmt='General', align=CL)
        if is_c:
            put(rc, f'C{r}', f'=IF({nm}="","",ROUND(SUMIFS({WF("F")},{WF("D")},{nm}),2))',
                font=F_AUTO, fill=FILL_AUTO, fmt=NUM, align=CL)
            put(rc, f'D{r}', f'=IF({nm}="","",ROUND(SUMIFS({WF("J")},{WF("D")},{nm}),2))',
                font=F_AUTO, fill=FILL_AUTO, fmt=NUM, align=CL)
            put(rc, f'E{r}', f'=IF({nm}="","",ROUND(N($C{r})-N($D{r}),2))',
                font=F_AUTO, fill=FILL_AUTO, fmt=NUM, align=CL)
            put(rc, f'F{r}', f'=IF({nm}="","",ROUND(SUMIFS({WF("G")},{WF("D")},{nm}),2))',
                font=F_AUTO, fill=FILL_AUTO, fmt=NUM, align=CL)
            put(rc, f'G{r}', f'=IF({nm}="","",ROUND(SUMIFS({SJ_IN},{SJ_CUST},{nm}),2))',
                font=F_AUTO, fill=FILL_AUTO, fmt=NUM, align=CL)
        else:
            put(rc, f'C{r}', f'=IF({nm}="","",ROUND(SUMIFS({WF("O")},{WF("N")},{nm}),2))',
                font=F_AUTO0, fill=FILL_AUTO, fmt=NUM, align=CL)
            put(rc, f'D{r}', None, fill=FILL_AUTO)
            put(rc, f'E{r}', f'=IF({nm}="","",N($C{r}))',
                font=F_AUTO0, fill=FILL_AUTO, fmt=NUM, align=CL)
            put(rc, f'F{r}', f'=IF({nm}="","",ROUND(SUMIFS({WF("P")},{WF("N")},{nm}),2))',
                font=F_AUTO0, fill=FILL_AUTO, fmt=NUM, align=CL)
            put(rc, f'G{r}', f'=IF({nm}="","",ROUND(SUMIFS({SJ_OUT},{SJ_SUPP},{nm}),2))',
                font=F_AUTO0, fill=FILL_AUTO, fmt=NUM, align=CL)
        put(rc, f'H{r}', f'=IF({nm}="","",ROUND(N($E{r})-N($F{r}),2))',
            font=F_SUM, fill=FILL_CHK, fmt=MONEY, align=CL)
        put(rc, f'I{r}', f'=IF({nm}="","",ROUND(N($F{r})-N($G{r}),2))',
            font=F_AUTO, fill=FILL_AUTO, fmt=NUM, align=CL)
        who = '客户' if is_c else '供货商'
        act = '回款' if is_c else '付款'
        put(rc, f'J{r}',
            f'=IF({nm}="","",'
            f'IF(AND(N($E{r})=0,N($G{r})=0),"没有往来",'
            f'IF(ROUND(N($I{r}),2)<>0,"差 "&TEXT(N($I{r}),"0.00")&" 元：'
            f'【数据录入】里这个{who}的{act}还没记全，或者没选{who}名",'
            f'IF(N($H{r})>0,"还欠 "&TEXT(N($H{r}),"0.00")&" 元","已结清"))))',
            font=F_T0, fill=FILL_CHK, fmt='General', align=CLW)
        rc.row_dimensions[r].height = 18

    # 合计行
    tr = r1 + 1
    put(rc, f'A{tr}', '合　计', font=F_HD, fill=FILL_SUM, align=CL)
    put(rc, f'B{tr}', None, font=F_HD, fill=FILL_SUM)
    for c in 'CDEFGHI':
        put(rc, f'{c}{tr}', f'=ROUND(SUM({c}{r0}:{c}{r1}),2)',
            font=F_SUM, fill=FILL_SUM, fmt=MONEY, align=CL)
    put(rc, f'J{tr}', None, fill=FILL_SUM)
    rc.conditional_formatting.add(f'A{r0}:J{r1}',
        FormulaRule(formula=[f'AND($B{r0}<>"",ROUND($I{r0},2)<>0)'], fill=FILL_WARN))
    return tr

block(CUST_H, CUST_0, CUST_N, '客户')
block(SUPP_H, SUPP_0, SUPP_N, '供货商')

for col, w in zip('ABCDEFGHIJ', (5.5, 14, 12, 12, 12, 14, 14, 13, 11, 40)):
    rc.column_dimensions[col].width = w
rc.freeze_panes = 'C6'
rc.sheet_view.showGridLines = False

# ---------------------------------------------------------------- ⑦
# 顺手把【数据录入】两个会闷声出错的容量缺口补上：
#   · H「账户余额」的公式只铺到第 100 行，录到第 101 笔余额就空白了；
#   · A「序号」只铺到第 249 行。
#   M「月份」铺到 4999 行，客户供应商汇总也按 $5:$4999 取数，所以统一补到 4999。
sj = wb[SJ]
H19 = sj['H19'].value
A5  = sj['A5'].value
fixed_h = fixed_a = 0
for r in range(19, SJ1 + 1):
    if not isinstance(sj.cell(r, 8).value, str):        # H 列，期初块 5~18 不动
        sj.cell(r, 8).value = (
            f'=IFERROR(VLOOKUP($C{r},$C$5:$H$18,6,0)'
            f'+SUMIF($C$19:$C{r},$C{r},$F$19:$F{r})'
            f'-SUMIF($C$19:$C{r},$C{r},$G$19:$G{r}),"")')
        sj.cell(r, 8).number_format = NUM
        fixed_h += 1
for r in range(5, SJ1 + 1):
    if not isinstance(sj.cell(r, 1).value, str):        # A 列
        sj.cell(r, 1).value = f'=IF(B{r}="","",COUNTA($B$5:B{r}))'
        fixed_a += 1
print(f'  【数据录入】补齐 账户余额公式 {fixed_h} 行、序号公式 {fixed_a} 行（原来只到第 100 / 249 行）')

# 日期和生日的显示格式：原表 B 列是美式 mm-dd-yy，C 列生日是 General（生日显示成 46113）
for r in range(F0, F1 + 1):
    ws.cell(r, 2).number_format = 'yyyy/mm/dd'          # B 日期
    ws.cell(r, 13).number_format = 'yyyy/mm/dd'         # M 应付日期
    # 生日只给「像日期」的格子套日期格式。C45 原来填的是数字 10，
    # 套上日期格式会显示成 1900-01-10，看着像个正经日期反而更糟，
    # 这种留着原样，让最右边的核对列去提醒。
    bd = ws.cell(r, 3).value
    if isinstance(bd, (int, float)) and bd >= 10000:
        ws.cell(r, 3).number_format = 'm"月"d"日"'

# ---------------------------------------------------------------- ⑧
# 「无主」提示：数据录入里有钱进出、但没挂客户名/供应商名的，
# 往来款跟进这边永远认领不到。把金额直接摆出来，让他知道还差多少没挂人头。
nr = SUPP_N + 4
put(rc, f'A{nr-1}', '三、还没挂到人头上的钱（这部分往来款跟进认领不到）',
    font=Font(name=YH, sz=12, bold=True, color='FFFFFFFF'),
    fill=PatternFill('solid', fgColor='FFC00000'), align=LF)
rc.merge_cells(f'A{nr-1}:J{nr-1}')
rc.row_dimensions[nr-1].height = 22
UNMATCHED = [
    ('收进来但没选客户名的钱', f'=ROUND(SUMIFS({SJ_IN},{SJ_CUST},""),2)',
     '在【数据录入】里把这些行的「客户」列选上，【往来款跟进】的已收款就会自动认领。'),
    ('付出去但没选供应商的钱（支出类别已选“供应商”）',
     f'=ROUND(SUMIFS({SJ_OUT},{SJ_SUPP},"",{SJ}!$D${SJ0}:$D${SJ1},"供应商"),2)',
     '原表第 32 行就是这样：支出类别选了「供应商」，但「供应商」那一列没填，1000 元没人认领。'),
    ('【数据录入】收入合计', f'=ROUND(SUM({SJ_IN}),2)', '供核对：上面第 1 行 ÷ 这一行 ＝ 还没挂人头的比例'),
    ('【数据录入】支出合计', f'=ROUND(SUM({SJ_OUT}),2)', ''),
]
for i, (lab, f, note) in enumerate(UNMATCHED):
    r = nr + i
    put(rc, f'A{r}', lab, font=F_T0, fill=FILL_SUM, align=LF)
    rc.merge_cells(f'A{r}:E{r}')
    put(rc, f'F{r}', f, font=F_SUM, fill=FILL_SUM, fmt=MONEY, align=CL)
    put(rc, f'G{r}', None, fill=FILL_SUM)
    put(rc, f'H{r}', None, fill=FILL_SUM)
    put(rc, f'I{r}', None, fill=FILL_SUM)
    put(rc, f'J{r}', note, font=F_NOTE, fill=FILL_SUM,
        align=Alignment(horizontal='left', vertical='center', wrap_text=True))
    rc.row_dimensions[r].height = 20

wb.save(OUT)
print('已保存', os.path.abspath(OUT))
nf = sum(1 for s in wb.worksheets for row in s.iter_rows()
         for c in row if isinstance(c.value, str) and c.value.startswith('='))
print(f'工作表 {len(wb.sheetnames)} 张，公式 {nf} 个')
