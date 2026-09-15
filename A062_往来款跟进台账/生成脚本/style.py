# -*- coding: utf-8 -*-
"""A062 的配色和字体——全部照抄原表，改出来的格子要跟原来的长得一样。"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

YH   = '微软雅黑'
F_T1 = Font(name=YH, sz=11, bold=True)                    # 明细区正文（应收侧）
F_T0 = Font(name=YH, sz=10)                               # 明细区正文（应付侧，原表就是 10 号）
F_HD = Font(name=YH, sz=11, bold=True, color='FF111827')  # 表头
F_HD0= Font(name=YH, sz=10, bold=True, color='FF111827')
F_SUM= Font(name=YH, sz=11, bold=True, color='FFC00000')  # 汇总块的数字
F_IN  = Font(name=YH, sz=11, bold=True, color='FF1F4E79')  # 手工录入：深蓝
F_IN0 = Font(name=YH, sz=10, color='FF1F4E79')
F_AUTO= Font(name=YH, sz=11, bold=True, color='FF7F7F7F')  # 自动计算：灰
F_AUTO0=Font(name=YH, sz=10, color='FF7F7F7F')
F_NOTE= Font(name=YH, sz=9,  color='FF808080')
F_WARN= Font(name=YH, sz=10, bold=True, color='FFC00000')

FILL_IN   = PatternFill('solid', fgColor='FFFFF7E0')   # 手工录入区：淡黄
FILL_AUTO = PatternFill('solid', fgColor='FFF2F2F2')   # 自动计算区：淡灰
FILL_HD   = PatternFill('solid', fgColor='FFD9E1F2')   # 表头：淡蓝
FILL_HD2  = PatternFill('solid', fgColor='FFBDD7EE')   # 二级表头
FILL_SUM  = PatternFill('solid', fgColor='FFFCE4D6')   # 汇总块：淡橙
FILL_CHK  = PatternFill('solid', fgColor='FFE2EFDA')   # 校验列：淡绿
FILL_WARN = PatternFill('solid', fgColor='FFFFC7CE')   # 警告：淡红

thin = Side(style='thin', color='FFBFBFBF')
BD   = Border(left=thin, right=thin, top=thin, bottom=thin)

CL = Alignment(horizontal='center', vertical='center')
CLW= Alignment(horizontal='center', vertical='center', wrap_text=True)
LF = Alignment(horizontal='left',   vertical='center')
RT = Alignment(horizontal='right',  vertical='center')

MONEY = '"￥"#,##0.00_);[Red]\\("￥"#,##0.00\\)'
NUM   = '#,##0.00_);[Red]\\(#,##0.00\\)'
PLAIN = '0.00_);[Red]\\(0.00\\)'
DATE  = 'yyyy/mm/dd'
PCT   = '0.0%'


def put(ws, coord, value, font=None, fill=None, fmt=None, align=None, border=True):
    c = ws[coord]
    c.value = value
    if font:  c.font = font
    if fill:  c.fill = fill
    if fmt:   c.number_format = fmt
    if align: c.alignment = align
    if border: c.border = BD
    return c
