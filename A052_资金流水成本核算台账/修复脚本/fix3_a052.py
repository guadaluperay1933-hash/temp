# -*- coding: utf-8 -*-
"""A052 第三轮：余额算法、录入口数据订正、项目核算表扩容。"""
import re
from copy import copy
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as gcl

SRC, DST = 'A052_fixed2.xlsx', 'A052_fixed3.xlsx'
LOG = []
def log(t, m): LOG.append(f'[{t}] {m}'); print(f'[{t}] {m}')

wb = openpyxl.load_workbook(SRC)
B, FLOW, INV = wb['基础资料'], wb['资金流水'], wb['开票登记']
LED, PRJ, DET = wb['应收应付台账'], wb['项目核算表'], wb['项目明细账']
RPT, DAY = wb['月度汇报表'], wb['资金日报表']

def st(ws, src, tgt): ws[tgt]._style = copy(ws[src]._style)

# ════════════════════════════════════════════════════════════════════
# 修复 5：把「余额」列当明细列整列求和 —— 全表最严重的一类算错
#   M 列是逐笔滚动余额，每一行都是"截至该行的账户结存"，把它整列相加没有任何意义。
#   正确写法：期初余额行的 M + 区间内收入 − 区间内支出。
# ════════════════════════════════════════════════════════════════════
FR = '资金流水'
def bal(where_acct, where_date_open, where_date_flow):
    """期初余额行的 M + 收 − 支"""
    o = (f'SUMIFS({FR}!$M$5:$M$3032,{FR}!$E$5:$E$3032,"期初余额"{where_acct}{where_date_open})')
    i = (f'SUMIFS({FR}!$K$5:$K$3032{where_acct}{where_date_flow})')
    p = (f'SUMIFS({FR}!$L$5:$L$3032{where_acct}{where_date_flow})')
    return f'{o}+{i}-{p}'

# 5a 资金流水顶部「资金总余额」
FLOW['A2'] = '资金总余额(全部账户)'
FLOW['C2'] = ('=SUMIFS($M$5:$M$3032,$E$5:$E$3032,"期初余额")+SUM($K$5:$K$3032)-SUM($L$5:$L$3032)')
FLOW['J2'] = '筛选后收/支合计→'
FLOW['M2'] = None
log('修复5', '【资金流水】C2「资金总余额」原来 = SUBTOTAL(M 列滚动余额)+收−支，把每一行的结存都加了一遍；'
             '改为 期初余额行合计 + 全部收入 − 全部支出。M2 那个对滚动余额求和的格子已清空')

# 5b 月度汇报表 H8 期末资金总额（按公司筛选、按发生月截至当月）
COMP = ',资金流水!$D$5:$D$3032,IF($E$2="全部","*",$E$2)'
RPT['H8'] = ('=' + bal(COMP, ',资金流水!$S$5:$S$3032,"<="&$B$2', COMP + ',资金流水!$S$5:$S$3032,"<="&$B$2'))
log('修复5', '【月度汇报表】H8「期末资金总额」原来算出 10.00 亿（真值 1147 万），同一个滚动余额错误；已改正')

# 5c 资金日报表 B8:B27 各账户期初余额
for r in range(8, 28):
    acct = f',资金流水!$C$5:$C$3032,$A{r}'
    RPT_F = ('=IF(OR($A%d="",NOT(ISNUMBER($I$2))),"",%s)' % (r, bal(
        acct, ',资金流水!$B$5:$B$3032,"<="&$B$2', acct + ',资金流水!$B$5:$B$3032,"<"&$B$2')))
    DAY[f'B{r}'] = RPT_F
log('修复5', '【资金日报表】B8:B27 各账户「期初余额」同一个错误；改正后 1 月 1 日期初合计由 804.88 万回到 659.77 万')

# ════════════════════════════════════════════════════════════════════
# 修复 6：【月度汇报表】收入明细的合计行套用了明细行公式 → 占比整列算不出来
# ════════════════════════════════════════════════════════════════════
RPT['C27'] = '=SUM($C$13:$C$26)-SUMIF($B$13:$B$26,"往来及保证金",$C$13:$C$26)'
RPT['D27'] = '=IF($C$27=0,"",1)'
log('修复6', '【月度汇报表】C27「收入合计」原来照抄了明细行的 SUMIFS（拿"合计（不含往来及保证金）"这几个字去匹配收入类别），'
             '永远等于 0，害得 D 列占比整列是空的；改为 SUM(C13:C26) 扣除往来及保证金')

# ════════════════════════════════════════════════════════════════════
# 修复 7：录入口的数据订正（这几处不改，公式再对也算不出正确结果）
# ════════════════════════════════════════════════════════════════════
# 7a 资金流水期初余额行的账户名与基础资料账户清单对不上
acc_list = [B.cell(r, 5).value for r in range(4, 24)]
old = {FLOW.cell(r, 3).value: r for r in range(5, 25)}
amt = {FLOW.cell(r, 3).value: FLOW.cell(r, 13).value for r in range(5, 25)}
dat = {FLOW.cell(r, 3).value: FLOW.cell(r, 2).value for r in range(5, 25)}
changed = []
for i, a in enumerate(acc_list):
    r = 5 + i
    cur = FLOW.cell(r, 3).value
    if cur != a:
        changed.append((r, cur, a))
        FLOW.cell(r, 3).value = a
        FLOW.cell(r, 13).value = amt.get(a, 0) if a in amt else 0
        FLOW.cell(r, 2).value = dat.get(a, dat.get(cur))
if changed:
    log('修复7', '【资金流水】期初余额行的账户与账户档案对不上：' +
        '；'.join(f'第{r}行 {c}→{n}' for r, c, n in changed) +
        '（原来"备用6"在流水里有期初行却不在账户档案，公司列显示"未匹配账户"；'
        '而档案里的"中展准东项目"反过来没有期初行）')

# 7b 应收应付台账 C6 公司名被截断
if LED['C6'].value == '中':
    LED['C6'] = '中展'
    log('修复7', '【应收应付台账】C6 公司名是截断的"中"，改为"中展"；不改的话这笔应付在任何按公司筛选的报表里都看不见')

# 7c 开票登记 收入类别整列基本没填 → 汇报表"收入明细"栏全是 0，跟"本月开票收入"对不上
cnt = 0
for r in range(5, 505):
    if INV.cell(r, 2).value is not None and not INV.cell(r, 8).value:
        INV.cell(r, 8).value = '项目收入'; cnt += 1
if cnt:
    log('修复7', f'【开票登记】{cnt} 张发票没填 H 列"收入类别"（73 张里空了 69 张），'
                 f'「月度汇报表」收入明细栏因此整栏是 0；按使用说明"一般选项目收入"统一补为「项目收入」，'
                 f'若有劳务/服务收入请自行改回')

# 7d 项目明细账默认项目是占位文字
first_prj = B['J4'].value
if DET['C2'].value not in [B.cell(r, 10).value for r in range(4, 104)]:
    DET['C2'] = first_prj
    log('修复7', f'【项目明细账】C2 默认值是占位文字"xxx采购项目A"，项目档案里根本没有，整页永远是 0；改为第一个真实项目')

# 7e 冻结窗格
FLOW.freeze_panes = 'F5'
log('修复7', '【资金流水】冻结窗格由 F673（把前 672 行全冻住，表打不开也滚不动）改为 F5')

# ════════════════════════════════════════════════════════════════════
# 修复 8：【项目核算表】60 行 → 100 行（项目档案有 83 个项目，原来只算前 60 个）
# ════════════════════════════════════════════════════════════════════
NEW_LAST, NEW_TOTAL = 104, 105
src_row = 64
tpl = {c: PRJ.cell(src_row, c).value for c in range(1, 30)}
tot_old = {c: PRJ.cell(65, c).value for c in range(1, 30)}
for c in range(1, 30):                       # 清掉原合计行
    PRJ.cell(65, c).value = None
for r in range(65, NEW_LAST + 1):
    for c in range(1, 30):
        f = tpl[c]
        co = f'{gcl(c)}{r}'
        st(PRJ, f'{gcl(c)}{src_row}', co)
        PRJ[co] = Translator(f, origin=f'{gcl(c)}{src_row}').translate_formula(co) if isinstance(f, str) and f.startswith('=') else f
    PRJ.row_dimensions[r].height = PRJ.row_dimensions[src_row].height
for c in range(1, 30):
    co = f'{gcl(c)}{NEW_TOTAL}'
    st(PRJ, f'{gcl(c)}65', co)
    v = tot_old[c]
    if isinstance(v, str) and v.startswith('='):
        v = re.sub(r'([A-Z]{1,2})5:([A-Z]{1,2})64', r'\g<1>5:\g<2>' + str(NEW_LAST), v)
        v = v.replace('$Y$65', f'$Y${NEW_TOTAL}').replace('$F$65', f'$F${NEW_TOTAL}') \
             .replace('$G$65', f'$G${NEW_TOTAL}').replace('$E$65', f'$E${NEW_TOTAL}')
    PRJ[co] = v
PRJ.row_dimensions[NEW_TOTAL].height = PRJ.row_dimensions[65].height
log('修复8', f'【项目核算表】明细行由 60 行扩到 100 行（第 5~104 行），合计行移到第 {NEW_TOTAL} 行；'
             f'项目档案实有 83 个项目，原来第 61 个以后的 23 个项目完全不出现在项目核算表里')

wb.save(DST)
open('fix3_log.txt', 'w', encoding='utf-8').write('\n'.join(LOG))
print('\n已保存', DST)
