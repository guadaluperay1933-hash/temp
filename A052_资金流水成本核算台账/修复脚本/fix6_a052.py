# -*- coding: utf-8 -*-
"""开账日落在查询区间内的账户，期初拿不到、逐日也看不到 —— 资金日报表会少一整笔。"""
import openpyxl
from openpyxl.utils import get_column_letter as gcl
SRC, DST = 'A052_fixed5.xlsx', 'A052_fixed6.xlsx'
LOG=[]
def log(t,m): LOG.append(f'[{t}] {m}'); print(f'[{t}] {m}')
wb = openpyxl.load_workbook(SRC); D = wb['资金日报表']; FR='资金流水'
n=0
for k in range(2, 32):                      # 第 2~31 天；第 1 天就是起始日，期初已经含它
    ci = 3 + 2*(k-1); col = gcl(ci)
    for r in range(8, 28):
        D.cell(r, ci).value = (
            f'=IF(OR($A{r}="",{col}$6=""),"",'
            f'SUMIFS({FR}!$K$5:$K$3032,{FR}!$C$5:$C$3032,$A{r},{FR}!$B$5:$B$3032,{col}$6)'
            f'+SUMIFS({FR}!$M$5:$M$3032,{FR}!$E$5:$E$3032,"期初余额",{FR}!$C$5:$C$3032,$A{r},{FR}!$B$5:$B$3032,{col}$6))')
        n+=1
log('修复10', f'【资金日报表】{n} 个逐日"收款"格补上开账余额：账户的开账日如果落在查询区间里（本表"中展农商行一般户"'
              f'就是 1 月 28 日开账、10 万），原来期初拿不到它、逐日也看不到它，期末凭空少 10 万；'
              f'现在开账那天会作为当日入账出现，期初+收−付与期末始终闭合')
wb.save(DST); open('fix6_log.txt','w',encoding='utf-8').write('\n'.join(LOG))
print('已保存', DST)
