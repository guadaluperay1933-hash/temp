# -*- coding: utf-8 -*-
"""A052 第二轮修复：容量、口径、余额算法、版式。"""
import re, collections
from copy import copy
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as gcl, column_index_from_string as cif

SRC, DST = 'A052_fixed.xlsx', 'A052_fixed2.xlsx'
LOG = []
def log(tag, msg):
    LOG.append(f'[{tag}] {msg}'); print(f'[{tag}] {msg}')

wb = openpyxl.load_workbook(SRC)
B   = wb['基础资料']; FLOW = wb['资金流水']; INV = wb['开票登记']
LED = wb['应收应付台账']; PRJ = wb['项目核算表']; DET = wb['项目明细账']
RPT = wb['月度汇报表']; DAY = wb['资金日报表']; QRY = wb['任意时段查询']; TRD = wb['月度趋势数据']

def style_from(ws, src_coord, tgt_coord):
    ws[tgt_coord]._style = copy(ws[src_coord]._style)

# ════════════════════════════════════════════════════════════════════
# 修复 3：【基础资料】项目清单去重 + 补编号；公司清单扩容位；客户/供应商清单换成真实名称
# ════════════════════════════════════════════════════════════════════
PRJ_COLS = list(range(9, 16))            # I..O 项目编号/名称/所属公司/客户/合同金额/签约日期/状态
rows_raw = []
for r in range(4, B.max_row + 1):
    vals = [B.cell(r, c).value for c in PRJ_COLS]
    if any(v not in (None, '') for v in vals):
        rows_raw.append(vals)
seen, cleaned, dups = set(), [], []
for vals in rows_raw:
    name = vals[1]
    if name in (None, ''):
        continue
    if name in seen:
        dups.append(name); continue
    seen.add(name); cleaned.append(vals)
for i, vals in enumerate(cleaned, 1):
    if vals[0] in (None, ''):
        vals[0] = f'P{i:03d}'
PRJ_CAP_ROW = 103                         # 项目清单容量 4..103 = 100 个
assert len(cleaned) <= PRJ_CAP_ROW - 3, len(cleaned)
for idx in range(PRJ_CAP_ROW - 3):
    r = 4 + idx
    vals = cleaned[idx] if idx < len(cleaned) else [None] * len(PRJ_COLS)
    for j, c in enumerate(PRJ_COLS):
        if r > 4:
            style_from(B, f'{gcl(c)}4', f'{gcl(c)}{r}')
        B.cell(r, c).value = vals[j]
log('修复3', f'【基础资料】项目清单去重 {len(dups)} 条重名（{"、".join(dups[:3])}…），规范为 {len(cleaned)} 个项目，'
             f'并为缺编号的项目补 P001~P{len(cleaned):03d}；清单容量由 60 扩到 100（第 4~103 行）')

# 公司清单容量 8 → 20（实际已有 11 家，原公式只取到第 11 行，卓卓通/财兴通/鸿禧被全表漏掉）
for r in range(5, 24):
    for c in (1, 2, 3):
        style_from(B, f'{gcl(c)}4', f'{gcl(c)}{r}')
log('修复3', '【基础资料】公司清单区域扩到第 4~23 行（20 家）')

# 客户 / 供应商清单：原来是"客户1…客户13"的占位名，实际发票上的购方一个都不在里面
cust = []
for r in range(5, 505):
    v = INV.cell(r, 6).value
    if v and v not in cust: cust.append(v)
for r in range(5, 3033):
    v = FLOW.cell(r, 14).value
    if v and v not in cust: cust.append(v)
supp = []
for r in range(5, 3033):
    v = FLOW.cell(r, 15).value
    if v and v not in supp: supp.append(v)
for idx in range(60):
    r = 4 + idx
    B.cell(r, 25).value = cust[idx] if idx < len(cust) else None
    B.cell(r, 27).value = supp[idx] if idx < len(supp) else None
    if r > 4:
        style_from(B, 'Y4', f'Y{r}'); style_from(B, 'AA4', f'AA{r}')
log('修复3', f'【基础资料】客户清单用发票实际购方重建（{len(cust)} 家），供应商清单 {len(supp)} 家；'
             f'原来全是"客户1…客户13"占位名，导致「任意时段查询」按客户排行永远是 0')

B.freeze_panes = 'A4'
log('修复3', '【基础资料】冻结窗格由 A91（冻住了前 90 行，表根本没法看）改为 A4')

# ════════════════════════════════════════════════════════════════════
# 修复 4：全表引用区间扩容（公司 11→20 位、项目 60→100 位），并补掉漏网的 3015
# ════════════════════════════════════════════════════════════════════
SUBS = [
    (re.compile(r'基础资料!\$A\$4:\$A\$11'), '基础资料!$A$4:$A$23'),
    (re.compile(r'基础资料!\$C\$4:\$C\$11'), '基础资料!$C$4:$C$23'),
    (re.compile(r'基础资料!(\$[I-O])\$4:(\$[I-O])\$63'), r'基础资料!\1$4:\2$103'),
    (re.compile(r'资金流水!([A-Z]{1,2})\$5:([A-Z]{1,2})\$3015'), r'资金流水!\1$5:\2$3032'),
]
n = 0
for sh in wb.sheetnames:
    ws = wb[sh]
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith('='):
                nv = v
                for pat, rep in SUBS:
                    nv = pat.sub(rep, nv)
                if nv != v:
                    c.value = nv; n += 1
log('修复4', f'{n} 个单元格的引用区间扩容：公司清单 $4:$11→$4:$23、项目清单 $4:$63→$4:$103，'
             f'并补上「项目明细账」1200 处漏改的 资金流水 X$5:X$3015')

# 数据验证同步
DV_FIX = {'基础资料!$A$4:$A$11': '基础资料!$A$4:$A$23',
          '基础资料!$J$4:$J$63': '基础资料!$J$4:$J$103',
          '基础资料!$J:$J':      '基础资料!$J$4:$J$103'}
m = 0
for sh in wb.sheetnames:
    for dv in wb[sh].data_validations.dataValidation:
        f1 = (dv.formula1 or '').strip()
        if f1 in DV_FIX:
            dv.formula1 = DV_FIX[f1]; m += 1
log('修复4', f'{m} 个下拉菜单的取值区间同步扩容（原来下拉里选不到卓卓通/财兴通/鸿禧，也选不到第 60 个以后的项目）')

# 数据验证作用范围：原来被导入数据打断成一段段，大半数据行根本没有下拉
from openpyxl.worksheet.datavalidation import DataValidation
DV_SQREF = {
    ('开票登记', '基础资料!$Y$4:$Y$63'):   'F5:F504',
    ('开票登记', '基础资料!$J$4:$J$103'):  'G5:G504',
    ('资金流水', '基础资料!$J$4:$J$103'):  'F5:F3032',
    ('资金流水', '基础资料!$Y$4:$Y$63'):   'N5:N3032',
    ('资金流水', '基础资料!$AA$4:$AA$63'): 'O5:O3032',
}
k = 0
for sh in ('开票登记', '资金流水'):
    for dv in wb[sh].data_validations.dataValidation:
        key = (sh, (dv.formula1 or '').strip())
        if key in DV_SQREF:
            dv.sqref = DV_SQREF[key]; k += 1
for dv in FLOW.data_validations.dataValidation:
    if dv.type == 'whole':
        dv.sqref = 'P5:P3032'; k += 1
log('修复4', f'{k} 个下拉菜单的作用行范围补全（原来「开票登记」购方/项目两列，第 15~58、60~66、68~72 行没有下拉；'
             f'「资金流水」项目列的下拉甚至盖到了标题行）')

wb.save(DST)
open('fix2_log.txt', 'w', encoding='utf-8').write('\n'.join(LOG))
print('\n已保存', DST)
