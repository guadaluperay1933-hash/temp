# -*- coding: utf-8 -*-
"""A052 第四轮：月度汇报表版式重排（分公司 / 开票汇总 8 行 → 20 行）+ 新增勾稽校验栏。"""
import re
from copy import copy
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as gcl
from openpyxl.formatting.rule import CellIsRule

SRC, DST = 'A052_fixed3.xlsx', 'A052_fixed4.xlsx'
LOG = []
def log(t, m): LOG.append(f'[{t}] {m}'); print(f'[{t}] {m}')

wb = openpyxl.load_workbook(SRC)
R = wb['月度汇报表']; P = wb['项目核算表']

# —— 先补上一轮遗漏的两格 ——
P['Z105'] = '=IFERROR($Y105/$F105,"")'
P['AB105'] = '=IFERROR($G105/$E105,"")'
log('修复8', '【项目核算表】合计行 Z105 毛利率、AB105 收款进度补正（原来还指着搬走前的第 65 行）')

# ════════════════════════════════════════════════════════════════════
# 修复 9：【月度汇报表】五、分公司汇总 和 七、开票汇总 只排了 8 行，
#        而公司档案实有 11 家 —— 卓卓通 / 财兴通 / 鸿禧 三家的收支被整栏漏掉，
#        合计与"一、核心指标"对不上（现金流出差 63,277.60）。两栏各扩到 20 行。
# ════════════════════════════════════════════════════════════════════
snap = {}
for r in range(89, 147):
    for c in range(1, 9):
        cell = R.cell(r, c)
        snap[(r, c)] = (cell.value, copy(cell._style))
heights = {r: (R.row_dimensions[r].height if r in R.row_dimensions else None) for r in range(89, 147)}

for rng in [m for m in list(R.merged_cells.ranges) if m.min_row >= 89]:
    R.unmerge_cells(str(rng))
for rng in list(R.conditional_formatting._cf_rules):
    if str(rng.sqref).startswith('F1'):
        del R.conditional_formatting[rng.sqref]
for r in range(89, 200):
    for c in range(1, 9):
        R.cell(r, c).value = None
    if r in R.row_dimensions:
        R.row_dimensions[r].height = None

def put(r, c, val, style_src, height=None):
    cell = R.cell(r, c)
    cell.value = val
    cell._style = copy(snap[style_src][1]) if isinstance(style_src, tuple) else style_src
    if height is not None:
        R.row_dimensions[r].height = height

def detail(tpl_row, i, new_row, col):
    """从模板行推出第 i 个明细行的公式：跨表引用顺着档案往下走，表内自引用改成新行号"""
    f = snap[(tpl_row, col)][0]
    if not (isinstance(f, str) and f.startswith('=')):
        return f
    g = Translator(f, origin=f'{gcl(col)}{tpl_row}').translate_formula(f'{gcl(col)}{tpl_row + i}')
    if new_row != tpl_row + i:
        g = re.sub(r'(?<![0-9])' + str(tpl_row + i) + r'(?![0-9])', str(new_row), g)
    return g

N = 20                                     # 公司位 / 账户位一律 20
merges = []

# —— 五、分公司汇总：89 标题 / 90 表头 / 91~110 明细 / 111 合计 ——
put(89, 1, snap[(89, 1)][0], (89, 1), heights[89]); merges.append('A89:H89')
for c in range(2, 9): put(89, c, None, (89, c))
for c in range(1, 9): put(90, c, snap[(90, c)][0], (90, c), heights[90])
for i in range(N):
    r = 91 + i
    for c in range(1, 9):
        put(r, c, detail(91, i, r, c), (91, c), heights[91])
put(111, 1, '合计', (99, 1), heights[99])
for c in range(2, 9):
    v = snap[(99, c)][0]
    put(111, c, f'=SUM({gcl(c)}91:{gcl(c)}110)' if isinstance(v, str) and v.startswith('=') else v, (99, c))

# —— 六、分账户资金余额：113 标题 / 114 表头 / 115~134 明细 / 135 合计 ——
put(112, 1, None, (100, 1))
put(113, 1, snap[(101, 1)][0], (101, 1), heights[101]); merges.append('A113:H113')
for c in range(2, 9): put(113, c, None, (101, c))
for c in range(1, 9): put(114, c, snap[(102, c)][0], (102, c), heights[102])
for i in range(N):
    r = 115 + i
    for c in range(1, 9):
        put(r, c, detail(103, i, r, c), (103, c), heights[103])
put(135, 1, '合计', (123, 1), heights[123])
for c in range(2, 9):
    v = snap[(123, c)][0]
    put(135, c, f'=SUM({gcl(c)}115:{gcl(c)}134)' if isinstance(v, str) and v.startswith('=') else v, (123, c))

# —— 七、本月开票汇总：137 标题 / 138 表头 / 139~158 明细 / 159~161 小计 ——
put(136, 1, None, (124, 1))
put(137, 1, snap[(125, 1)][0], (125, 1), heights[125]); merges.append('A137:H137')
for c in range(2, 9): put(137, c, None, (125, c))
for c in range(1, 9): put(138, c, snap[(126, c)][0], (126, c), heights[126])
for i in range(N):
    r = 139 + i
    for c in range(1, 9):
        put(r, c, detail(127, i, r, c), (127, c), heights[127])
for k, old in enumerate((135, 136, 137)):
    r = 159 + k
    for c in range(1, 9):
        v = snap[(old, c)][0]
        if isinstance(v, str) and v.startswith('='):
            v = re.sub(r'(?<![0-9])' + str(old) + r'(?![0-9])', str(r), v)
            if old == 137:
                v = f'=SUM({gcl(c)}159:{gcl(c)}160)'
        put(r, c, v, (old, c), heights.get(old))

# —— 八、跨期与往来提示：163 标题 / 164 表头 / 165~170 ——
put(162, 1, None, (138, 1))
put(163, 1, snap[(139, 1)][0], (139, 1), heights[139]); merges.append('A163:H163')
for c in range(2, 9): put(163, c, None, (139, c))
for c in range(1, 9): put(164, c, snap[(140, c)][0], (140, c), heights[140])
for k in range(6):
    old, r = 141 + k, 165 + k
    for c in range(1, 9):
        put(r, c, snap[(old, c)][0], (old, c), heights.get(old))
    merges.append(f'D{r}:H{r}')

# —— 九、勾稽校验（新增）——
HDR = (90, 1)     # 借用表头样式
TTL = (89, 1)
put(171, 1, None, (100, 1))
put(172, 1, '九、勾稽校验 —— 下面每行都应显示"√ 一致"。出现"✗"说明基础资料的公司/账户/类别清单没覆盖全部录入数据，请按差额去找。', TTL, 21.75)
for c in range(2, 9): put(172, c, None, (89, c))
merges.append('A172:H172')
for c, t in zip(range(1, 6), ('校验项', '本表算出', '应当等于', '差额', '结果')):
    put(173, c, t, (90, c), 31.5)
for c in range(6, 9): put(173, c, None, (90, c))
CHECKS = [
    ('分公司汇总·现金流入合计', '$C$111', '$D$6', '核心指标的现金流入'),
    ('分账户·期末余额合计',     '$F$135', '$H$8', '核心指标的期末资金总额'),
    ('开票汇总·价税合计',       '$F$161', '$B$6', '核心指标的本月开票收入'),
    ('收入明细·合计',           '$C$27',  '$B$7', '损益口径收入'),
    ('支出结构·合计',           '$B$37',  '$D$7', '损益口径成本费用'),
    ('支出明细·日常+项目成本',  '$C$87+$G$55', '$D$7', '损益口径成本费用'),
]
for k, (name, a, b, note) in enumerate(CHECKS):
    r = 174 + k
    put(r, 1, name, (141, 1), 19.5)
    put(r, 2, f'={a}', (141, 2))
    put(r, 3, f'={b}', (141, 2))
    put(r, 4, f'=N($B{r})-N($C{r})', (141, 2))
    put(r, 5, f'=IF(ABS($D{r})<0.01,"√ 一致","✗ 差 "&TEXT($D{r},"#,##0.00"))', (141, 1))
    put(r, 6, note, (141, 4))
    for c in (7, 8): put(r, c, None, (141, c))
    merges.append(f'F{r}:H{r}')

for m in merges:
    R.merge_cells(m)
R.conditional_formatting.add('F115:F134', CellIsRule(operator='lessThan', formula=['0'],
                                                     font=openpyxl.styles.Font(color='FFC00000')))
log('修复9', '【月度汇报表】"五、分公司汇总"和"七、本月开票汇总"由 8 行扩到 20 行；'
             '公司档案实有 11 家，原来第 9 家以后（卓卓通/财兴通/鸿禧）在这两栏里根本不出现，'
             '导致分公司合计的现金流出比核心指标少 63,277.60。六、八两栏顺延，条件格式同步搬到 F115:F134')
log('修复9', '【月度汇报表】新增"九、勾稽校验"六行：把分公司、分账户、开票汇总、收入明细、支出结构、支出明细'
             '这六处合计与核心指标逐一对差，全部显示"√ 一致"才算数对得上')

wb.save(DST)
open('fix4_log.txt', 'w', encoding='utf-8').write('\n'.join(LOG))
print('\n已保存', DST)
