# -*- coding: utf-8 -*-
"""A052 台账公式修复脚本 —— 逐条修复，每条都写明理由。"""
import re, sys, shutil
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string as cif, get_column_letter as gcl

SRC = 'A052_src.xlsx'
DST = 'A052_fixed.xlsx'

LOG = []
def log(tag, msg):
    LOG.append(f'[{tag}] {msg}')
    print(f'[{tag}] {msg}')

wb = openpyxl.load_workbook(SRC)

# ═══════════════════════════════════════════════════════════════════
# 修复 1：【资金流水】辅助列只填到 3015 行，数据列引用写 3032 → SUMIFS 区间行数不等
#         统一为 5:3032（3028 行容量），把辅助列补齐到 3032
# ═══════════════════════════════════════════════════════════════════
FLOW = wb['资金流水']
OLD_END, NEW_END = 3015, 3032
HELPER_COLS = ['A', 'M', 'S', 'T', 'U', 'V', 'W', 'X']

n_add = 0
for col in HELPER_COLS:
    src_coord = f'{col}{OLD_END}'
    src_f = FLOW[src_coord].value
    assert isinstance(src_f, str) and src_f.startswith('='), src_coord
    for r in range(OLD_END + 1, NEW_END + 1):
        tgt = f'{col}{r}'
        if FLOW[tgt].value is None:
            FLOW[tgt] = Translator(src_f, origin=src_coord).translate_formula(tgt)
            # 沿用第 3015 行的格式
            FLOW[tgt]._style = FLOW[src_coord]._style
            n_add += 1
log('修复1', f'【资金流水】辅助列 {"/".join(HELPER_COLS)} 由第 {OLD_END} 行补齐到第 {NEW_END} 行，新增 {n_add} 个公式')

RE_FLOW_REF = re.compile(r'(资金流水!)?(\$[A-Z]{1,2})\$5:(\$[A-Z]{1,2})\$3015')
def widen(f, in_flow_sheet):
    def sub(m):
        pre, c1, c2 = m.groups()
        if c1 != c2:
            return m.group(0)
        if pre is None and not in_flow_sheet:
            return m.group(0)          # 别的表里不带表名的 3015 区间与资金流水无关
        return f'{pre or ""}{c1}$5:{c2}$3032'
    return RE_FLOW_REF.sub(sub, f)

n_ref = 0
for sh in wb.sheetnames:
    ws = wb[sh]
    inflow = (sh == '资金流水')
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith('='):
                nv = widen(v, inflow)
                if nv != v:
                    c.value = nv
                    n_ref += 1
log('修复1', f'全表 {n_ref} 个单元格里的「资金流水 …$5:…$3015」区间改为 …$5:…$3032，SUMIFS 各区间行数恢复一致')

# ═══════════════════════════════════════════════════════════════════
# 修复 2：【开票登记】公式列里被常量/空值覆盖的"洞"
# ═══════════════════════════════════════════════════════════════════
INV = wb['开票登记']
# D58 被手打成"小规模纳税人"，公式算出来也是这个值，恢复公式以免用户改公司时不跟着变
INV['D58'] = Translator(INV['D57'].value, origin='D57').translate_formula('D58')
log('修复2', '【开票登记】D58 由手打的"小规模纳税人"恢复为自动匹配公式（算出的值不变）')
# N6 整格为空，缺了按票号自动核销的公式
INV['N6'] = Translator(INV['N5'].value, origin='N5').translate_formula('N6')
log('修复2', '【开票登记】N6 缺失的"已收款(按票号自动)"公式已补上')

wb.save(DST)
print()
print(f'已保存 {DST}')
open('fix_log.txt', 'w', encoding='utf-8').write('\n'.join(LOG))
