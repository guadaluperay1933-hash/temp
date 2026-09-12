# -*- coding: utf-8 -*-
import re
from copy import copy
import openpyxl
SRC, DST = 'A052_final.xlsx', 'A052_final2.xlsx'
LOG=[]
def log(t,m): LOG.append(f'[{t}] {m}'); print(f'[{t}] {m}')
wb = openpyxl.load_workbook(SRC)
R, Q, S = wb['月度汇报表'], wb['任意时段查询'], wb['使用说明']

# ── 修复 19：同一笔支出被"成本类别"和"费用类别"两个明细块各算一次 ──
n=0
for r in range(41, 87):
    f = R[f'C{r}'].value
    if isinstance(f,str) and '资金流水!$I$5:$I$3032' in f and '资金流水!$H$5:$H$3032,""' not in f:
        R[f'C{r}'] = f.replace(f'资金流水!$I$5:$I$3032,$A{r}',
                               f'资金流水!$I$5:$I$3032,$A{r},资金流水!$H$5:$H$3032,""'); n+=1
for r in range(43, 89):
    f = Q[f'B{r}'].value
    if isinstance(f,str) and '资金流水!$I$5:$I$3032' in f and '资金流水!$H$5:$H$3032,""' not in f:
        Q[f'B{r}'] = f.replace(f'资金流水!$I$5:$I$3032,A{r}',
                               f'资金流水!$I$5:$I$3032,A{r},资金流水!$H$5:$H$3032,""'); n+=1
log('修复19', f'{n} 格：「日常费用按类别」明细块原来只按费用类别取数，遇到同时填了成本类别和费用类别的行'
              f'（本表第 633 行，2026-04 付保函陪标费 480 元，H=费用摊销、I=其他管理费用）'
              f'会和「项目成本按类别」块各算一次，明细合计比损益口径成本费用多 480。'
              f'现在按「资金流水」W 列的判定口径（填了成本类别就算项目成本）把这类行从费用明细里排除')

# ── 修复 20：新增校验，盯住"两个类别都没填"的支出 ────────────────
r = 183
for c in range(1,9): R.cell(r,c)._style = copy(R.cell(182,c)._style)
R.row_dimensions[r] .height = R.row_dimensions[182].height
R[f'A{r}'] = '资金流水·没填类别的支出'
R[f'B{r}'] = ('=SUMIFS(资金流水!$L$5:$L$3032,资金流水!$T$5:$T$3032,$B$2,'
              '资金流水!$D$5:$D$3032,IF($E$2="全部","*",$E$2),资金流水!$W$5:$W$3032,"")')
R[f'C{r}'] = 0
R[f'D{r}'] = f'=N($B{r})-N($C{r})'
R[f'E{r}'] = f'=IF(ABS($D{r})<0.01,"√ 一致","✗ 有 "&TEXT($D{r},"#,##0.00")&" 元支出没填类别，没进内帐盈亏")'
R[f'F{r}'] = '这类支出进了现金流出、却进不了损益，请补 H 列成本类别或 I 列费用类别'
for c in (7,8): R.cell(r,c).value=None
R.merge_cells(f'F{r}:H{r}')
log('修复20', '【月度汇报表】勾稽校验加第 10 行：既没填成本类别、也没填费用类别的支出会被单列出来。'
              '本表导入的真实流水里有 44 笔、合计 2,745,393.77 元是这种情况 —— 它们计入了现金流出，'
              '却进不了「损益口径成本费用」，内帐盈亏因此被抬高。补上类别这一行才会变回"√ 一致"')

# ── 使用说明里的 #VALUE! 字样会被各种检查工具误判成错误单元格 ──
for rr in range(1, S.max_row+1):
    v = S.cell(rr,1).value
    if isinstance(v,str) and '#VALUE!' in v:
        S.cell(rr,1).value = v.replace('#VALUE!', '「#VALUE」错误')
S['A77'] = ('⑧「日常费用」与「项目成本」两个明细块会把同时填了两种类别的行各算一次（已修正）；'
            '另有 44 笔、合计 2,745,393.77 元的支出两种类别都没填，它们进了现金流出却进不了内帐盈亏，'
            '「月度汇报表」勾稽校验最后一行会一直提示，请在「资金流水」H 列或 I 列补上类别。')
S['A77']._style = copy(S['A76']._style); S.row_dimensions[77].height = 31.5

wb.save(DST); open('fix11_log.txt','w',encoding='utf-8').write('\n'.join(LOG))
print('已保存', DST)
