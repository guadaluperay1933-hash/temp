# -*- coding: utf-8 -*-
import datetime, re
from copy import copy
import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
SRC, DST = 'A052_final6.xlsx', 'A052_final7.xlsx'
LOG=[]
def log(t,m): LOG.append(f'[{t}] {m}'); print(f'[{t}] {m}')
wb = openpyxl.load_workbook(SRC)
I, L, F, P, D, R, S = (wb['开票登记'], wb['应收应付台账'], wb['资金流水'],
                       wb['项目核算表'], wb['项目明细账'], wb['月度汇报表'], wb['使用说明'])

# ── 修复 27：开票日期被录成文本 ──────────────────────────────────
assert I['B15'].value == '2026-2-2-'
I['B15'] = datetime.datetime(2026, 2, 2)
I['B15'].number_format = I['B14'].number_format
log('修复27', '【开票登记】B15 的开票日期是文本 "2026-2-2-"（末尾多一个横杠），不是日期。'
              'LibreOffice 能猜出来所以看不出问题，Excel/WPS 里 --TEXT(文本,"yyyymm") 直接 #VALUE，'
              'R15 归属月报错 → 这张 24 万的票在「月度汇报表」按归属月归集时整行取不到；'
              '「任意时段查询」按日期区间比较时文本永远不落在任何区间内。已改为真正的日期 2026-02-02')

# ── 修复 28：收款多于开票时，状态仍显示"已收齐" ────────────────────
for r in range(5, 505):
    I[f'Q{r}'] = (f'=IF($J{r}="","",IF($P{r}<-0.005,"多收款",IF($P{r}<=0.005,"已收齐",'
                  f'IF(IF($N{r}="",0,$N{r})+IF($O{r}="",0,$O{r})<=0,"未收款","部分收款"))))')
for r in range(5, 305):
    L[f'M{r}'] = (f'=IF($H{r}="","",IF($L{r}<-0.005,"多结算",IF($L{r}<=0.005,"已结清",'
                  f'IF($L{r}<$H{r},"部分结算","未结算"))))')
log('修复28', '【开票登记】Q 列收款状态只有"已收齐/未收款/部分收款"三档，收多了（未收款为负）也归进"已收齐"。'
              '现有数据第 42 行 M=45,000.01 却登了已收 162,144.81（正是上一行 O41 顺手复制下来的）、'
              '第 49 行多收 30,000，两行都显示"已收齐"，11.7 万的错录被状态列完全盖住。'
              '现在加了"多收款"一档；【应收应付台账】M 列同样加"多结算"')

# ── 修复 29：负的未收款在合计里冲减真实应收 ─────────────────────────
def only_pos(f):
    return f.replace('开票登记!$P$5:$P$504,开票登记!',
                     '开票登记!$P$5:$P$504,开票登记!$P$5:$P$504,">0",开票登记!')
n=0
for cell in ('D8','B170'):
    old=R[cell].value
    if isinstance(old,str) and '开票登记!$P$5:$P$504' in old:
        R[cell]=only_pos(old); n+=1
for r in range(5,155):
    P[f'I{r}'] = (f'=IF($B{r}="","",SUMIFS(开票登记!$P$5:$P$504,开票登记!$P$5:$P$504,">0",'
                  f'开票登记!$G$5:$G$504,$B{r}))'); n+=1
D['I4'] = '=SUMIFS(开票登记!$P$5:$P$504,开票登记!$P$5:$P$504,">0",开票登记!$G$5:$G$504,$C$2)'; n+=1
log('修复29', f'{n} 处「已开票未收款」改为只汇总真正欠款的行（P>0）。原来对 P 列整列求和，'
              f'上面那两笔多收款的负数（−117,144.80 和 −30,000）直接把别人的应收冲掉了：'
              f'真实应收 5,444,113.30，表上显示 5,296,968.50，凭空少 14.7 万')

# ── 修复 30：账龄超期的条件格式把 490 多个空行也标红 ────────────────
def fix_cf(ws, sqref, col, days):
    old=None
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            if rule.type=='cellIs' and rule.operator=='greaterThan' and rule.formula==[str(days)]:
                old=rule
    keep=[(cf.sqref, list(cf.rules)) for cf in ws.conditional_formatting]
    ws.conditional_formatting = ConditionalFormattingList()
    for sq, rules in keep:
        for rule in rules:
            if rule is old: continue
            ws.conditional_formatting.add(str(sq), rule)
    old.type='expression'; old.operator=None
    old.formula=[f'AND(${col}5<>"",${col}5>{days})']
    ws.conditional_formatting.add(sqref, old)
fix_cf(I, '$S$5:$S$504', 'S', 90)
fix_cf(L, '$N$5:$N$304', 'N', 60)
log('修复30', '【开票登记】S 列"账龄超 90 天标红"和【应收应付台账】N 列"超 60 天标红"用的是'
              '"单元格值 > 90"这种比较。可这两列对已收齐的行和全部空行返回的是空文本 ""，'
              '而 Excel/WPS/LibreOffice 一律认为"文本大于任何数字"，于是 500 行里 490 多行全被标红，'
              '真正超期的那几张淹没在满屏红色里，预警等于没有。改成公式规则 AND(S5<>"",S5>90)')

# ── 修复 31：三行核销示例的单号挂在了收支方向相反的行上 ────────────
for r, want in ((176,'收'),(177,'付'),(178,'收')):
    F.cell(r,17).value = None
    F.cell(r,18).value = (str(F.cell(r,18).value or '') +
                          '（原来这里挂了一个关联单号，但本行是' +
                          ('支出' if want=='收' else '收入') + '行，方向反了，已清除）')
log('修复31', '【资金流水】第 176/177/178 行的"关联单号"挂反了方向：FP2025110001、FP20260601 是应收发票号，'
              '却贴在两笔支出行上（缴纳医疗保险 1,196.30、租赁费 77,000）；YF202512001 是应付单号，'
              '却贴在一笔 192 万的收入行上。核销的取数方向本身是对的（应收取收入列、应付取支出列），'
              '所以这三处永远核销出 0，用户照着使用说明操作会以为"按票号自动核销这个功能是坏的"。'
              '已清除这三个错挂的单号（金额一分未动），并在备注里写明原因')

# ── 修复 32：勾稽校验加第 12 行：多收款的发票 ─────────────────────
r = 185
for c in range(1,9): R.cell(r,c)._style = copy(R.cell(184,c)._style)
R.row_dimensions[r].height = R.row_dimensions[184].height
R[f'A{r}'] = '开票登记·收款多于开票额'
R[f'B{r}'] = '=-SUMIFS(开票登记!$P$5:$P$504,开票登记!$P$5:$P$504,"<0")'
R[f'C{r}'] = 0
R[f'D{r}'] = f'=N($B{r})-N($C{r})'
R[f'E{r}'] = f'=IF(ABS($D{r})<0.01,"√ 一致","✗ 有 "&TEXT($D{r},"#,##0.00")&" 元收款超过了开票额，多半是已收款登错行")'
R[f'F{r}'] = '在「开票登记」筛选 Q 列="多收款"即可定位'
for c in (7,8): R.cell(r,c).value=None
R.merge_cells(f'F{r}:H{r}')
log('修复32', '【月度汇报表】勾稽校验加第 12 行：收款多于开票额的发票会被单列出来')

S['A80'] = ('⑪ 开票登记 B15 的日期是文本（末尾多一个横杠）已改正；收款状态加了"多收款"一档；'
            '"已开票未收款"只汇总真正欠款的行，不再被多收的负数冲减；账龄超期的条件格式原来把'
            '490 多个空行也标红，已改成只标真正超期的行；资金流水第 176~178 行挂反方向的关联单号已清除。')
S['A80']._style = copy(S['A79']._style); S.row_dimensions[80].height = 31.5
wb.save(DST); open('fix14_log.txt','w',encoding='utf-8').write('\n'.join(LOG))
print('已保存', DST)
