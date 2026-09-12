# -*- coding: utf-8 -*-
import re, collections
from copy import copy
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as gcl

SRC, DST = 'A052_fixed8.xlsx', 'A052_fixed9.xlsx'
VAL = 'A052_r8.xlsx'
LOG=[]
def log(t,m): LOG.append(f'[{t}] {m}'); print(f'[{t}] {m}')
wb  = openpyxl.load_workbook(SRC)
wv  = openpyxl.load_workbook(VAL, data_only=True)
B, I, F, P, R = wb['基础资料'], wb['开票登记'], wb['资金流水'], wb['项目核算表'], wb['月度汇报表']
Iv, Fv = wv['开票登记'], wv['资金流水']
def st(ws,s,t): ws[t]._style = copy(ws[s]._style)

# ── 修复 16：把只在发票/流水里出现、却不在项目档案里的项目补进档案 ──
arch = [B.cell(r,10).value for r in range(4,104) if B.cell(r,10).value]
used, meta, amt = [], {}, collections.defaultdict(float)
for r in range(5,505):
    g = Iv.cell(r,7).value
    if g:
        if g not in used and g not in arch: used.append(g)
        meta.setdefault(g, (Iv.cell(r,3).value, Iv.cell(r,6).value))
        m = Iv.cell(r,13).value
        amt[g] += m if isinstance(m,(int,float)) else 0
for r in range(5,3033):
    f = Fv.cell(r,6).value
    if f:
        if f not in used and f not in arch: used.append(f)
        meta.setdefault(f, (Fv.cell(r,4).value, None))
miss = [g for g in used if g not in arch]
CAP = 153                                         # 项目档案容量 4..153 = 150 个
assert len(arch)+len(miss) <= CAP-3
for i,name in enumerate(miss):
    r = 4 + len(arch) + i
    comp, cust = meta.get(name, (None,None))
    for c,val in ((9,f'P{len(arch)+i+1:03d}'),(10,name),(11,comp),(12,cust),(15,'待整理')):
        B.cell(r,c).value = val
for r in range(4, CAP+1):
    for c in range(9,16):
        if r>4: st(B, f'{gcl(c)}4', f'{gcl(c)}{r}')
log('修复16', f'【基础资料】项目档案补进 {len(miss)} 个"只在发票/流水里出现、档案里却没有"的项目'
              f'（合计开票 {sum(amt[g] for g in miss):,.2f} 元，占全部开票的 56%）。'
              f'不补的话这半壁江山在「项目核算表」里根本不存在，累计开票、成本、毛利全都看不到。'
              f'这批行的"项目状态"统一标成「待整理」，其中有几条明显是把摘要打进了项目名'
              f'（如"2套黑板""维护费""以前公司的项目…"），请在项目档案里改名或合并，'
              f'并同步改「开票登记」G 列 /「资金流水」F 列')

# 区间由 103 扩到 153
n=0
pat = re.compile(r'基础资料!(\$[I-O])\$4:(\$[I-O])\$103')
for sh in wb.sheetnames:
    for row in wb[sh].iter_rows():
        for c in row:
            v=c.value
            if isinstance(v,str) and v.startswith('='):
                nv = pat.sub(r'基础资料!\1$4:\2$153', v)
                if nv!=v: c.value=nv; n+=1
for sh in wb.sheetnames:
    for dv in wb[sh].data_validations.dataValidation:
        if (dv.formula1 or '').strip()=='基础资料!$J$4:$J$103':
            dv.formula1='基础资料!$J$4:$J$153'
log('修复16', f'{n} 个单元格 + 4 个下拉的项目区间由 $4:$103 扩到 $4:$153（容量 150 个项目，已用 {len(arch)+len(miss)}）')

# ── 项目核算表 100 行 → 150 行 ──────────────────────────────────
SRC_ROW, LAST, TOTAL = 104, 154, 155
tpl = {c: P.cell(SRC_ROW,c).value for c in range(1,30)}
tot = {c: P.cell(105,c).value for c in range(1,30)}
for c in range(1,30): P.cell(105,c).value=None
for r in range(105, LAST+1):
    for c in range(1,30):
        co=f'{gcl(c)}{r}'; st(P, f'{gcl(c)}{SRC_ROW}', co)
        f_=tpl[c]
        P[co] = Translator(f_, origin=f'{gcl(c)}{SRC_ROW}').translate_formula(co) if isinstance(f_,str) and f_.startswith('=') else f_
    P.row_dimensions[r].height = P.row_dimensions[SRC_ROW].height
for c in range(1,30):
    co=f'{gcl(c)}{TOTAL}'; st(P,f'{gcl(c)}105',co); v=tot[c]
    if isinstance(v,str) and v.startswith('='):
        v = re.sub(r'([A-Z]{1,2})5:([A-Z]{1,2})104', r'\g<1>5:\g<2>'+str(LAST), v)
        v = re.sub(r'(\$[A-Z]{1,2})105', r'\g<1>'+str(TOTAL), v)
    P[co]=v
P.row_dimensions[TOTAL].height = P.row_dimensions[105].height
log('修复16', f'【项目核算表】明细行扩到 150 行（第 5~154 行），合计行移到第 {TOTAL} 行')

# ── 修复 17：勾稽校验的口径自适应 + 补两条项目级校验 ───────────────
R['B176'] = '=IF(基础资料!$AH$4="不含税金额",$D$161,$F$161)'
R['B180'] = ('=$B$6-IF(基础资料!$AH$4="不含税金额",'
             'SUMIFS(开票登记!$J$5:$J$504,开票登记!$R$5:$R$504,$B$2,开票登记!$C$5:$C$504,IF($E$2="全部","*",$E$2),开票登记!$H$5:$H$504,"项目收入")'
             '+SUMIFS(开票登记!$J$5:$J$504,开票登记!$R$5:$R$504,$B$2,开票登记!$C$5:$C$504,IF($E$2="全部","*",$E$2),开票登记!$H$5:$H$504,"劳务/服务收入"),'
             'SUMIFS(开票登记!$M$5:$M$504,开票登记!$R$5:$R$504,$B$2,开票登记!$C$5:$C$504,IF($E$2="全部","*",$E$2),开票登记!$H$5:$H$504,"项目收入")'
             '+SUMIFS(开票登记!$M$5:$M$504,开票登记!$R$5:$R$504,$B$2,开票登记!$C$5:$C$504,IF($E$2="全部","*",$E$2),开票登记!$H$5:$H$504,"劳务/服务收入"))')
log('修复17', '【月度汇报表】勾稽校验第 3、7 行原来拿"跟着含税开关走的数"去比"固定含税的数"，'
              '一旦把 基础资料!AH4 切到不含税就会误报✗；改为两边同口径')

NEW = [
    ('项目核算表·累计开票合计',
     '=项目核算表!$F$155',
     '=IF(基础资料!$AH$4="不含税金额",SUMIFS(开票登记!$J$5:$J$504,开票登记!$G$5:$G$504,"<>"),SUMIFS(开票登记!$M$5:$M$504,开票登记!$G$5:$G$504,"<>"))',
     '开票登记里填了项目名的发票合计（累计，不看月份）'),
    ('项目核算表·成本合计',
     '=项目核算表!$X$155',
     '=SUMIFS(资金流水!$L$5:$L$3032,资金流水!$W$5:$W$3032,"项目成本",资金流水!$F$5:$F$3032,"<>")',
     '资金流水里填了项目名的项目成本合计（累计，不看月份）'),
]
for k,(name,a,b,note) in enumerate(NEW):
    r = 181+k
    for c in range(1,9): R.cell(r,c)._style = copy(R.cell(180,c)._style)
    R.row_dimensions[r].height = R.row_dimensions[180].height
    R[f'A{r}']=name; R[f'B{r}']=a; R[f'C{r}']=b
    R[f'D{r}']=f'=N($B{r})-N($C{r})'
    R[f'E{r}']=f'=IF(ABS($D{r})<0.01,"√ 一致","✗ 差 "&TEXT($D{r},"#,##0.00")&"（有项目名不在项目档案里）")'
    R[f'F{r}']=note
    for c in (7,8): R.cell(r,c).value=None
    R.merge_cells(f'F{r}:H{r}')
log('修复17', '【月度汇报表】勾稽校验再加两行：项目核算表的累计开票、成本合计要等于'
              '发票/流水里所有挂了项目名的金额 —— 只要有人往项目列打了档案里没有的名字，这两行立刻报✗')

wb.save(DST); open('fix9_log.txt','w',encoding='utf-8').write('\n'.join(LOG))
print('已保存', DST)
