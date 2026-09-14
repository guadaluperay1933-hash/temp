# -*- coding: utf-8 -*-
import openpyxl
from copy import copy
SRC, DST = 'A052_fixed7.xlsx', 'A052_fixed8.xlsx'
LOG=[]
def log(t,m): LOG.append(f'[{t}] {m}'); print(f'[{t}] {m}')
wb = openpyxl.load_workbook(SRC)
B, I, R = wb['基础资料'], wb['开票登记'], wb['月度汇报表']

# ── 修复 14：收入口径开关填错时全表静默按含税算 ──────────────────
for dv in B.data_validations.dataValidation:
    if dv.type == 'list' and '价税合计' in (dv.formula1 or ''):
        dv.allow_blank = False
        dv.showErrorMessage = True
        dv.errorTitle = '收入口径只能二选一'
        dv.error = '只能选「价税合计(含税)」或「不含税金额」，请用右边的下拉箭头选择。'
        dv.showInputMessage = True
        dv.promptTitle = '收入口径'
        dv.prompt = '这一格决定全表"开票收入"取含税还是不含税。改它，月度汇报表 / 项目核算表 / 趋势表 / 图表一起变。'
I['C2'] = ('=IF(基础资料!$AH$4="不含税金额","本表口径收入(不含税)",'
           'IF(基础资料!$AH$4="价税合计(含税)","本表口径收入(含税)",'
           '"⚠ 基础资料!AH4 收入口径填错，全表正按含税计算"))')
log('修复14', '【基础资料】AH4「收入口径」只给了下拉、没开出错拦截，手打成"不含税"或带个空格都不会报错；'
              '而全表 174 条公式是 IF(AH4="不含税金额",不含税,含税) 的二值判断，填错就一律按含税默默算下去。'
              '现在下拉改为强制拦截 + 输入提示，【开票登记】C2 的口径标签也会在填错时显示告警')

# ── 修复 15：新增一行勾稽校验，盯住"不该进开票登记的发票" ────────
r = 180
for c in range(1, 9):
    R.cell(r, c)._style = copy(R.cell(179, c)._style)
R.row_dimensions[r].height = R.row_dimensions[179].height
R[f'A{r}'] = '开票登记·非经营类发票'
R[f'B{r}'] = ('=$B$6-SUMIFS(开票登记!$M$5:$M$504,开票登记!$R$5:$R$504,$B$2,开票登记!$C$5:$C$504,'
              'IF($E$2="全部","*",$E$2),开票登记!$H$5:$H$504,"项目收入")'
              '-SUMIFS(开票登记!$M$5:$M$504,开票登记!$R$5:$R$504,$B$2,开票登记!$C$5:$C$504,'
              'IF($E$2="全部","*",$E$2),开票登记!$H$5:$H$504,"劳务/服务收入")')
R[f'C{r}'] = 0
R[f'D{r}'] = f'=N($B{r})-N($C{r})'
R[f'E{r}'] = f'=IF(ABS($D{r})<0.01,"√ 一致","✗ 有 "&TEXT($D{r},"#,##0.00")&" 元发票的收入类别不是经营收入")'
R[f'F{r}'] = '开票登记只登经营收入的票；理财/利息/补助/保证金/股东往来不开票，走资金流水'
for c in (7, 8): R.cell(r, c).value = None
R.merge_cells(f'F{r}:H{r}')
R['A172'] = ('九、勾稽校验 —— 下面每行都应显示"√ 一致"。出现"✗"说明基础资料的公司/账户/类别清单没覆盖全部录入数据，'
             '或者有不该登进开票登记的票，请按差额去找。')
log('修复15', '【月度汇报表】勾稽校验加第 7 行：开票登记里出现"非经营收入"类别的发票会被立刻标出来。'
              '本来的隐患是 —— 开票登记 H 列的下拉把 理财/利息/补助/保证金退回/股东投入 也列出来了，'
              '一旦选中，这笔钱会被"本月开票收入"按开票计一次、又被"损益口径收入"按银行到账计一次，重复计入利润，'
              '而第二栏收入明细那边是剔除往来及保证金的，两栏对不上还查不出原因')

wb.save(DST); open('fix8_log.txt','w',encoding='utf-8').write('\n'.join(LOG))
print('已保存', DST)
