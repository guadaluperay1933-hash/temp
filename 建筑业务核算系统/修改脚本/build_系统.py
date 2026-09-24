# -*- coding: utf-8 -*-
"""在客户 9.21《项目总台账》上设置建筑业务核算系统（保持原框架，只做小调整）

原表结构（不动）：
  目录       —— 项目名称 → 项目简称（几个项目合一张分表就填同一个简称）
  发票统计   —— 总台账，所有业务逐笔录在这里（第 2 行表头、第 3 行合计、第 4 行起数据）
  各项目分表 —— 表名 = 项目简称，从发票统计自动取本项目的行

这次改动：
  ① 发票统计：公式统一铺到第 1003 行；L/P/T 三个余额改成「本项目累计到这一行」（原来是全表滚动，
     中间插行后引用错位了好几处）；合计行 P3 / T3 口径改正；Q 列误填的 =H 公式清掉；
     新增 V「校验」列、隐藏 W「取数键」列；C 列下拉选目录里的项目名称；冻结表头
  ② 各项目分表：原来用的是 WPS 专用的 FILTER + SHEETSNAME（Excel 打不开），
     改成按表名逐行 INDEX/MATCH 取数，WPS / Excel 都能用；L/P/T 本表累计；合计行改正；
     合川瑞翔的标题更正；删掉引用铜城表的错公式
  ③ 目录：补上漏掉的序号 10-12；加一列隐藏的「汇总序号」给项目汇总用
  ④ 新增《项目汇总》《使用说明》两张表

跑法：python3 build_系统.py [源.xlsx] [输出.xlsx]
"""
import sys, os, copy
import openpyxl
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.formatting.rule import FormulaRule

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, '参考', '原表_9.21.xlsx')
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, '建筑业务核算系统.xlsx')

LED = '发票统计'
R0, R1 = 4, 1003          # 发票统计数据行
OLD_LAST = 98             # 原表铺了格式的最后一行
TPL = 90                  # 原表第 90 行：铺了格式、没有数据，新行照它的格式
P0, P1 = 4, 303           # 分表数据行（每张分表最多显示 300 行）
D0, D1 = 3, 102           # 目录可用行
S0, S1 = 4, 53            # 项目汇总的项目行（50 个）
ACC = '_ * #,##0.00_ ;_ * \\-#,##0.00_ ;_ * "-"??_ ;_ @_ '
PROJ = ['蒲吕片区A段', '淮阳站阳高线', '合川表计', '云龙社区', '铜城2回、玉璟尚线、铁井一回、淮拦线4', '合川瑞翔']
RED = Font(color='C00000', bold=True)
ORANGE = Font(color='C65911')

# 表名 → 取数用的简称。Excel / WPS 的 CELL 返回 “路径[文件]表名”，LibreOffice 老版本返回 “'文件'#$表名”
KEYF = ('=IFERROR(MID(CELL("filename",$A$1),FIND("]",CELL("filename",$A$1))+1,99),'
        'IFERROR(SUBSTITUTE(MID(CELL("filename",$A$1),FIND("#$",CELL("filename",$A$1))+2,99),"\'",""),""))')


def set_col(ws, idx, width=None, hidden=None):
    """单独设某一列。原表 V 列的列定义一直管到最右边，要先把它拆开"""
    dims = ws.column_dimensions
    hit = None
    for key, cd in list(dims.items()):
        lo = cd.min or CI(key)
        hi = cd.max or lo
        if lo <= idx <= hi:
            hit = (key, cd, lo, hi)
            break

    def clone(cd, lo, hi):
        nd = copy.copy(cd)
        nd._style = copy.copy(cd._style)
        nd.index, nd.min, nd.max = CL(lo), lo, hi
        return nd

    if hit is None:
        nd = ColumnDimension(ws, index=CL(idx))
        nd.min = nd.max = idx
    else:
        key, cd, lo, hi = hit
        del dims[key]
        if lo < idx:
            dims[CL(lo)] = clone(cd, lo, idx - 1)
        if hi > idx:
            dims[CL(idx + 1)] = clone(cd, idx + 1, hi)
        nd = clone(cd, idx, idx)
    if width is not None:
        nd.width = width
    if hidden is not None:
        nd.hidden = hidden
    dims[CL(idx)] = nd


def style_like(dst, src, fmt=None, h=None, wrap=None, sz=None):
    if src.has_style:
        dst._style = copy.copy(src._style)
    if fmt:
        dst.number_format = fmt
    if h is not None or wrap is not None:
        a = dst.alignment
        dst.alignment = Alignment(horizontal=h if h is not None else a.horizontal, vertical='center',
                                  wrap_text=wrap if wrap is not None else a.wrap_text)
    if sz:
        f = copy.copy(dst.font)
        f.sz = sz
        dst.font = f


def warn_cf(ws, rng, first):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},1)="⚠"'], font=RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},1)="△"'], font=ORANGE))


wb = openpyxl.load_workbook(SRC, rich_text=True)
led = wb[LED]
ml = wb['目录']
NAME = {ml[f'C{r}'].value: ml[f'B{r}'].value for r in range(D0, 23) if ml[f'C{r}'].value}

# ───────────────────────── ③ 目录 ─────────────────────────
for r, n in ((12, 10), (13, 11), (14, 12)):
    if ml[f'A{r}'].value is None:
        ml[f'A{r}'] = n
style_like(ml['F2'], ml['D2'])
ml['F2'] = '汇总序号(辅助)'
for r in range(D0, D1 + 1):
    # 简称第一次出现时编号，项目汇总按这个顺序一项目一行
    ml[f'F{r}'] = f'=IF(C{r}="","",IF(COUNTIF(C${D0}:C{r},C{r})=1,MAX(F${D0 - 1}:F{r - 1})+1,""))'
set_col(ml, 6, hidden=True)

# ───────────────────────── ① 发票统计 ─────────────────────────
TXT = '+'.join(f'ISTEXT(${c}{{r}})' for c in 'GHKNOQRS')


def vcheck(r):
    return ('=IF(AND($C{r}="",COUNTA($D{r}:$H{r},$K{r},$N{r}:$O{r},$Q{r}:$S{r},$U{r})=0),"",'
            'IF($C{r}="","⚠ 没填项目名称",'
            'IF($A{r}="⚠目录里没有","⚠ 项目名称在目录里找不到（先加到目录）",'
            'IF($A{r}="","⚠ 目录里这个项目没填简称",'
            'IF(' + TXT + '>0,"⚠ 金额列里有文字，算不进合计",'
            'IF(N($G{r})>=1,"⚠ 管理费点要填百分比，如 10%",'
            'IF(AND($D{r}<>"",$B{r}=""),"△ 摘要开头没写日期（如 2026.9.21），月份取不到",'
            'IF(AND(N($H{r})<>0,COUNTIFS($A${R0}:$A${R1},$A{r},$H${R0}:$H${R1},$H{r},'
            '$D${R0}:$D${R1},LEFT($D{r},9)&"*")>1),"△ 同一天同项目有一样的销售金额，看看是不是录重了",'
            '"✓"))))))))').format(r=r, R0=R0, R1=R1)


for r in range(R0, R1 + 1):
    if r > OLD_LAST:
        for c in range(1, 22):
            style_like(led.cell(r, c), led.cell(TPL, c))
    led[f'A{r}'] = f'=IF($C{r}="","",IFERROR(VLOOKUP($C{r},目录!$B:$C,2,0)&"","⚠目录里没有"))'
    led[f'B{r}'] = (f'=IF($D{r}="","",IFERROR(MID($D{r},FIND(".",$D{r})+1,'
                    f'FIND(".",$D{r},FIND(".",$D{r})+1)-FIND(".",$D{r})-1)&"月",""))')
    led[f'I{r}'] = f'=IF($H{r}="","",N($H{r})*N($G{r}))'
    led[f'J{r}'] = f'=IF($H{r}="","",N($H{r})-N($I{r}))'
    led[f'M{r}'] = f'=IF($H{r}="","",N($H{r}))'
    k = f'$A${R0}:$A{r},$A{r}'
    s = lambda c: f'SUMIF({k},${c}${R0}:${c}{r})'
    led[f'L{r}'] = f'=IF($A{r}="","",{s("J")}-{s("K")})'
    led[f'P{r}'] = f'=IF($A{r}="","",{s("M")}-{s("N")}-{s("O")})'
    led[f'T{r}'] = f'=IF($A{r}="","",{s("Q")}-{s("R")}-{s("S")})'
    q = led[f'Q{r}'].value
    if isinstance(q, str) and q.startswith('=H'):
        led[f'Q{r}'] = None                     # 原表 Q56:Q72 误拉了 =H，结果都是 0
    led[f'V{r}'] = vcheck(r)
    style_like(led[f'V{r}'], led[f'U{r}'], fmt='General', h='left', wrap=False)
    led[f'W{r}'] = f'=IF($A{r}="","",$A{r}&"#"&COUNTIF($A${R0}:$A{r},$A{r}))'

for c in 'HIJKMNOQRS':
    led[f'{c}3'] = f'=SUM({c}{R0}:{c}{R1})'
led['L3'] = '=J3-K3'
led['P3'] = '=M3-N3-O3'            # 原来是 M3-O3，漏减质保金
led['T3'] = '=Q3-R3-S3'            # 原来是 SUM(T4:T318)，把逐行余额加在一起了
led['U3'] = None                   # 备注列不用合计
style_like(led['V2'], led['U2'])
led['V2'] = '校验'
style_like(led['V3'], led['U3'], fmt='General', wrap=True, sz=10)
led['V3'] = (f'=IF(COUNTA(C{R0}:C{R1})>COUNTIF(W{R0}:W{R1},"?*"),'
             f'"⚠ 有 "&(COUNTA(C{R0}:C{R1})-COUNTIF(W{R0}:W{R1},"?*"))&" 行缺公式或缺简称",'
             f'IF(COUNTIF(V{R0}:V{R1},"⚠*")>0,"⚠ "&COUNTIF(V{R0}:V{R1},"⚠*")&" 行有错，往下看",'
             f'IF(COUNTIF(V{R0}:V{R1},"△*")>0,"△ "&COUNTIF(V{R0}:V{R1},"△*")&" 行请留意","✓ 全部通过")))')
led['W2'] = '取数键(辅助)'
set_col(led, 22, width=34)
set_col(led, 23, hidden=True)
warn_cf(led, f'V{R0}:V{R1}', f'$V{R0}')     # V3 在深色表头上，保持白字
dv = DataValidation(type='list', formula1=f'目录!$B${D0}:$B${D1}', allow_blank=True,
                    showErrorMessage=True, errorStyle='warning', errorTitle='项目名称',
                    error='要跟「目录」里的项目名称一字不差；新项目先加到目录里。')
dv.add(f'C{R0}:C{R1}')
led.add_data_validation(dv)
led.freeze_panes = 'A4'

# ───────────────────────── ② 各项目分表 ─────────────────────────
for name in PROJ:
    ps = wb[name]
    styled = ps.max_row                      # 原表铺了格式的最后一行（48 或 83）
    if name == '合川瑞翔':
        ps['A1'] = NAME['合川瑞翔']          # 原来误抄了云龙社区的标题
    for r in range(P0, max(styled, P1) + 1):
        for c in range(1, 22):
            ps.cell(r, c).value = None       # 清掉 FILTER 溢出区和旧余额公式
    for r in range(P0, P1 + 1):
        if r > styled:
            for c in range(1, 22):
                style_like(ps.cell(r, c), ps.cell(P0, c))
        ps[f'W{r}'] = (f'=IF($V$3="","",IFERROR(MATCH($V$3&"#"&ROWS(W${P0}:W{r}),'
                       f'{LED}!$W${R0}:$W${R1},0),""))')
        for c in 'ABCDEFU':
            ps[f'{c}{r}'] = f'=IF($W{r}="","",INDEX({LED}!{c}${R0}:{c}${R1},$W{r})&"")'
        g = f'INDEX({LED}!G${R0}:G${R1},$W{r})'
        ps[f'G{r}'] = f'=IF($W{r}="","",IF({g}="","",{g}))'
        ps[f'G{r}'].number_format = '0%'
        for c in 'HIJKMNOQRS':
            ps[f'{c}{r}'] = f'=IF($W{r}="","",N(INDEX({LED}!{c}${R0}:{c}${R1},$W{r})))'
        ps[f'L{r}'] = f'=IF($W{r}="","",SUM(J${P0}:J{r})-SUM(K${P0}:K{r}))'
        ps[f'P{r}'] = f'=IF($W{r}="","",SUM(M${P0}:M{r})-SUM(N${P0}:N{r})-SUM(O${P0}:O{r}))'
        ps[f'T{r}'] = f'=IF($W{r}="","",SUM(Q${P0}:Q{r})-SUM(R${P0}:R{r})-SUM(S${P0}:S{r}))'
    for c in 'HIJKMNOQRS':
        ps[f'{c}3'] = f'=SUM({c}{P0}:{c}{P1})'
    ps['L3'] = '=J3-K3'
    ps['P3'] = '=M3-N3-O3'
    ps['T3'] = '=Q3-R3-S3'
    ps['U3'] = None
    style_like(ps['V2'], ps['U2'], sz=10)
    ps['V2'] = '本表项目简称\n（自动取表名）'
    style_like(ps['V3'], ps['U3'], fmt='General', wrap=True, sz=10)
    ps['V3'] = KEYF
    style_like(ps['V4'], ps['U4'], fmt='General', h='left', wrap=True, sz=10)
    cnt = f'COUNTIF({LED}!$A${R0}:$A${R1},$V$3)'
    ps['V4'] = (f'=IF($V$3="","⚠ 取不到表名，请在 V3 手填项目简称",'
                f'IF({cnt}=0,"⚠ 发票统计里没有「"&$V$3&"」的行，表名要和目录里的简称一样",'
                f'IF({cnt}>{P1 - P0 + 1},"⚠ 本项目超过 {P1 - P0 + 1} 行，后面的没显示","共 "&{cnt}&" 笔")))')
    warn_cf(ps, 'V4', '$V4')
    set_col(ps, 22, width=24)
    set_col(ps, 23, hidden=True)
    ps.freeze_panes = 'A4'

# ───────────────────────── ④ 项目汇总 ─────────────────────────
sm = wb.create_sheet('项目汇总', 2)
PS4 = wb['合川表计']
sm.sheet_view.zoomScale = 85
HEAD = ['序号', '项目简称\n（点开分表）', '笔数', '最后一笔摘要', '销售金额', '管理费', '应开\n成本票', '已开\n成本票',
        '未开\n成本票', '应收款', '质保金', '已收款', '未收款', '应付\n工程款', '扣除\n管理费', '已付\n工程款',
        '未付\n工程款', '分表核对']
# 项目汇总 E..Q 与发票统计 H..T 一一对应（中间三列余额自己算）
MAPC = dict(zip('EFGHIJKLMNOPQ', 'HIJKLMNOPQRST'))
sm.merge_cells('A1:R1')
style_like(sm['A1'], led['A1'])
sm['A1'] = '项目汇总'
sm.row_dimensions[1].height = 48
sm.row_dimensions[2].height = 33
for i, h in enumerate(HEAD, 1):
    style_like(sm.cell(2, i), led['H2'], fmt='General')
    sm.cell(2, i).value = h
sm.merge_cells('A3:D3')
for i in range(1, 19):
    style_like(sm.cell(3, i), led['H3'], fmt=ACC if 5 <= i <= 17 else 'General')
sm['A3'] = '合计'
for c in 'EFGHJKLNOP':
    sm[f'{c}3'] = f'=SUM({c}{S0}:{c}{S1})'
sm['I3'] = '=G3-H3'
sm['M3'] = '=J3-K3-L3'
sm['Q3'] = '=N3-O3-P3'
style_like(sm['R3'], led['H3'], fmt='General', wrap=True, sz=10)
sm['R3'] = (f'=IF(SUMPRODUCT(ABS(E3:Q3-{LED}!H3:T3))<0.01,"✓ 与发票统计合计一致",'
            f'"⚠ 与发票统计合计对不上：有行没归到项目，看发票统计 V 列")')
LA, LD = f'{LED}!$A${R0}:$A${R1}', f'{LED}!$W${R0}:$W${R1}'
for r in range(S0, S1 + 1):
    for i in range(1, 19):
        # 数据行照分表的样式（白底细框），不用发票统计的彩色底，空行不显眼
        src = PS4[f'D{P0}'] if i <= 4 or i == 18 else PS4[f'H{P0}']
        style_like(sm.cell(r, i), src, fmt=ACC if 5 <= i <= 17 else 'General',
                   h='center' if i in (1, 3) else ('left' if i in (2, 4, 18) else None), wrap=(i == 2))
    nm = f'INDEX(目录!$C${D0}:$C${D1},MATCH(ROWS(B${S0}:B{r}),目录!$F${D0}:$F${D1},0))'
    sm[f'B{r}'] = f'=IFERROR(HYPERLINK("#\'"&{nm}&"\'!A1",{nm}),"")'
    sm[f'B{r}'].font = Font(name='微软雅黑', sz=11, color='0563C1', underline='single')
    sm[f'A{r}'] = f'=IF($B{r}="","",ROWS(A${S0}:A{r}))'
    sm[f'C{r}'] = f'=IF($B{r}="","",COUNTIF({LA},$B{r}))'
    sm[f'D{r}'] = f'=IF(N($C{r})=0,"",INDEX({LED}!$D${R0}:$D${R1},MATCH($B{r}&"#"&$C{r},{LD},0))&"")'
    for c in 'EFGHJKLNOP':
        lc = MAPC[c]
        sm[f'{c}{r}'] = f'=IF($B{r}="","",SUMIF({LA},$B{r},{LED}!${lc}${R0}:${lc}${R1}))'
    sm[f'I{r}'] = f'=IF($B{r}="","",G{r}-H{r})'
    sm[f'M{r}'] = f'=IF($B{r}="","",J{r}-K{r}-L{r})'
    sm[f'Q{r}'] = f'=IF($B{r}="","",N{r}-O{r}-P{r})'
    sm[f'R{r}'] = (f'=IF($B{r}="","",IFERROR(IF(SUMPRODUCT(ABS(INDIRECT("\'"&$B{r}&"\'!H3:T3")-$E{r}:$Q{r}))<0.01,'
                   f'"✓","⚠ 分表合计对不上"),"⚠ 还没建分表"))')
warn_cf(sm, f'R{S0}:R{S1}', f'$R{S0}')
for c, w in zip('ABCDEFGHIJKLMNOPQR', [6, 24, 6, 50] + [15] * 13 + [22]):
    sm.column_dimensions[c].width = w
sm.freeze_panes = 'C4'

# ───────────────────────── 使用说明 ─────────────────────────
TEXT = [
    ('h', '一、表的结构'),
    ('', '1. 目录：项目名称 → 项目简称。几个项目合在一张分表里算，就给它们填同一个简称（如铜城2回、玉璟尚线、铁井一回、淮拦线4 这 4 个项目）。'),
    ('', '2. 发票统计：总台账，所有业务都只录在这里。'),
    ('', '3. 项目汇总：每个项目一行，自动汇总；点项目简称可以跳到那张分表。'),
    ('', '4. 各项目分表（表名 = 项目简称）：自动从发票统计取本项目的行，按录入先后排列。不用手录，也不要在分表上改数——要改就改发票统计。'),
    ('h', '二、发票统计怎么录'),
    ('', '1. 手填的列：C 项目名称（下拉选）、D 摘要（开头写日期，如「2026.9.21……」）、E 收票单位、F 发票性质、G 管理费点、H 销售金额、'
         'K 已开成本票、N 质保金、O 已收款、Q 应付工程款、R 扣除管理费、S 已付工程款、U 备注。'),
    ('', '2. 自动算的列（不要手改）：A 项目简称、B 月份、I 管理费＝H×G、J 应开成本票＝H－I、M 应收款＝H、L 未开成本票、P 未收款、T 未付工程款、V 校验。'),
    ('', '3. L / P / T 是「这个项目累计到这一行」的余额，和分表上同一笔的数一样；第 3 行是全部项目的合计。'),
    ('', f'4. 新业务直接往下面的空行录，不用插行（公式已经铺到第 {R1} 行），分表会按录入先后自动排进去。'
         '想插在项目中间也可以：先复制上面一整行，右键「插入复制的单元格」，再改内容——公式会跟着过来。'),
    ('', '5. 录完看 V 列：✓ 正常；⚠ 是错了要改（项目名称不在目录里、金额里有文字、管理费点填成了整数等）；'
         '△ 是提醒（摘要开头没写日期、同一天同项目同金额可能录重了）。V3 是整张表的检查结果。'),
    ('h', '三、常见业务的录法（照原来的习惯）'),
    ('', '1. 开销售票：H 填开票金额，G 填管理费点（没有管理费就空着）→ 管理费、应开成本票、应收款自动算。'),
    ('', '2. 收到成本票（包括工资表挂账）：K 填金额，可以跟开票写在同一行，也可以单独录一行。'),
    ('', '3. 收到款：O 填金额；回款里直接扣掉的税费、管理费、残保金、个税等，也录在 O，摘要写清楚扣的是什么。'),
    ('', '4. 质保金：建议录在 N 列（质保金），这样项目汇总里能单独看到扣了多少质保金；录在 O 也会减未收款，只是看不出是质保金。'),
    ('', '5. 付工程款 / 代付工资：S 填金额；应付工程款录在 Q。Q 不录的话，未付工程款就是负数（＝已付合计）。'),
    ('', '6. 红冲或更正：不删原来那一行，另录一行负数（H、K 或 O 填负数），摘要写「红冲……」，合计和分表自动冲掉。'),
    ('h', '四、新增项目'),
    ('', '1. 目录里加一行：项目名称 + 项目简称。'),
    ('', '2. 要单独一张分表：右键任意一张项目分表 →「移动或复制」→ 勾「建立副本」，把新表名改成这个简称（要和目录里的简称一字不差），'
         '再改一下第 1 行标题，数据自动出来。'),
    ('', '3. 分表右上角 V3 是本表取数用的简称（自动取表名，也可以直接手填）；V4 显示共几笔，或者提示哪里不对。'),
    ('', '4. 项目汇总会自动多出这个项目；「分表核对」显示「⚠ 还没建分表」，就是第 2 步还没做。'),
    ('h', '五、核对'),
    ('', '1. 发票统计 V3 显示「✓ 全部通过」（△ 只是提醒，看一眼没问题就行）。'),
    ('', '2. 项目汇总第 3 行显示「✓ 与发票统计合计一致」，每个项目的「分表核对」都是 ✓。'),
    ('', f'3. 每张分表最多显示 {P1 - P0 + 1} 笔，超过会在 V4 提示；发票统计公式铺到第 {R1} 行，快用完时把最后一行往下拉。'),
    ('h', '六、这次整理改了什么'),
    ('', '1. 分表原来用 WPS 专用的 FILTER + SHEETSNAME 取数，Excel 打开是错的；现在换成通用公式，WPS、Excel 都能用，表名 = 简称的规则不变。'),
    ('', '2. 发票统计的余额原来是整张表往下滚，中间插过行，好几处引用错位（如第 10、81、82 行）；现在改成按项目累计，插行、排序都不会错。'),
    ('', '3. 合计行：未收款原来漏减质保金，未付工程款原来把每一行的余额加在一起，都已改正。'),
    ('', '4. 原来有几行没拉公式（第 50、51 行），应开成本票、应收款少算了；现在所有行统一用同一套公式。'),
    ('', '5. 新增 V 列校验、C 列下拉、项目汇总、本说明；合川瑞翔分表标题原来抄成了云龙社区，已改正；目录补上序号 10-12。'),
]
us = wb.create_sheet('使用说明')
us.sheet_view.showGridLines = False
us.column_dimensions['A'].width = 3
us.column_dimensions['B'].width = 120
us['B1'] = '使用说明'
us['B1'].font = Font(name='微软雅黑', sz=18, b=True)
us.row_dimensions[1].height = 36
for i, (kind, t) in enumerate(TEXT, 3):
    c = us.cell(i, 2, t)
    c.font = Font(name='微软雅黑', sz=12 if kind == 'h' else 11, b=(kind == 'h'),
                  color='1F4E79' if kind == 'h' else None)
    c.alignment = Alignment(wrap_text=True, vertical='center')
    if kind == 'h' and i > 3:
        us.row_dimensions[i].height = 26

wb.active = wb.sheetnames.index('项目汇总')
for w in wb.worksheets:
    w.sheet_view.tabSelected = (w.title == '项目汇总')
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('已生成', OUT)
