# -*- coding: utf-8 -*-
"""独立复算 + 场景实测

一、复算：只读原表的手填列，用 Python 按原表口径（I=H×G、J=H−I、M=H，余额按项目累计）自己算一遍，
   跟新表的发票统计、6 张分表、项目汇总逐格比；再查原表手填内容和底色一格没动。
二、场景（在副本上改完用 LibreOffice 重算）：往下录新行、加新项目并复制分表、目录有项目但没建分表、
   项目名称不在目录、没填项目名称、管理费点填成整数、金额里有文字、插行没带公式、红冲负数。

跑法：python3 verify.py [新表.xlsx] [原表.xlsx] [临时目录]
"""
import sys, os, re, shutil, subprocess
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, '建筑业务核算系统.xlsx')
SRC = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, '参考', '原表_9.21.xlsx')
TMP = sys.argv[3] if len(sys.argv) > 3 else '/tmp/建筑业务核算_verify'
RECALC = os.path.join(ROOT, '..', '工具', 'recalc_safe.py')
R0, R1, P0, P1 = 4, 1003, 4, 303
PROJ = ['蒲吕片区A段', '淮阳站阳高线', '合川表计', '云龙社区', '铜城2回、玉璟尚线、铁井一回、淮拦线4', '合川瑞翔']
BAD = []
N = lambda v: float(v) if isinstance(v, (int, float)) else 0.0


def chk(name, got, exp, tol=0.005, quiet=True):
    if isinstance(exp, str) or exp is None:
        ok = (got if got is not None else '') == (exp if exp is not None else '')
    else:
        ok = isinstance(got, (int, float)) and abs(got - exp) <= tol
    if not ok or not quiet:
        print(('  ✓ ' if ok else '  ✗ ') + f'{name}: {got!r}  应为 {exp!r}')
    if not ok:
        BAD.append(name)
    return ok


def month(d):
    m = re.match(r'[^.]*\.([^.]*)\.', d or '')
    return (m.group(1) + '月') if m else ''


# ───────── 原表手填数据 ─────────
sf = openpyxl.load_workbook(SRC)
sv = openpyxl.load_workbook(SRC, data_only=True)
dirmap = {sf['目录'][f'B{r}'].value: sf['目录'][f'C{r}'].value for r in range(3, 23) if sf['目录'][f'B{r}'].value}
INP = 'CDEFGHKNOQRSU'
rows = []
for r in range(R0, 99):
    rec = {'r': r}
    for c in INP:
        f, v = sf['发票统计'][f'{c}{r}'].value, sv['发票统计'][f'{c}{r}'].value
        if c == 'Q' and isinstance(f, str) and f.startswith('=H'):
            v = None                                   # 原表误拉的 =H，本来就是 0
        rec[c] = v
    if all(rec[c] in (None, '') for c in INP):
        continue
    h, g = rec['H'], rec['G']
    rec['A'] = dirmap.get(rec['C'], '⚠目录里没有') if rec['C'] else ''
    rec['B'] = month(rec['D'])
    rec['I'] = N(h) * N(g) if h is not None else None
    rec['J'] = N(h) - rec['I'] if h is not None else None
    rec['M'] = N(h) if h is not None else None
    rows.append(rec)
cum = {}
for rec in rows:
    a = cum.setdefault(rec['A'], dict(L=0.0, P=0.0, T=0.0))
    a['L'] += N(rec['J']) - N(rec['K'])
    a['P'] += N(rec['M']) - N(rec['N']) - N(rec['O'])
    a['T'] += N(rec['Q']) - N(rec['R']) - N(rec['S'])
    rec.update(L=a['L'], P=a['P'], T=a['T'])
print(f'原表有数据的行：{len(rows)}')

wf = openpyxl.load_workbook(OUT)
wv = openpyxl.load_workbook(OUT, data_only=True)
lf, lv = wf['发票统计'], wv['发票统计']

print('一、原表手填内容、底色没动')
for r in range(R0, 99):
    for c in INP:
        a, b = sf['发票统计'][f'{c}{r}'].value, lf[f'{c}{r}'].value
        if c == 'Q' and isinstance(a, str) and a.startswith('=H'):
            a = None
        chk(f'发票统计 {c}{r} 原值', b, a) if isinstance(a, str) or a is None else chk(f'发票统计 {c}{r} 原值', b, a, tol=0)
    for c in 'ABCDEFGHIJKLMNOPQRSTU':
        f0, f1 = sf['发票统计'][f'{c}{r}'].fill.fgColor, lf[f'{c}{r}'].fill.fgColor
        if (f0.type, f0.rgb if f0.type == 'rgb' else f0.theme, round(f0.tint or 0, 3)) != \
           (f1.type, f1.rgb if f1.type == 'rgb' else f1.theme, round(f1.tint or 0, 3)):
            chk(f'发票统计 {c}{r} 底色', 'changed', 'same')
for r in range(3, 23):
    for c in 'BCD':
        chk(f'目录 {c}{r}', wf['目录'][f'{c}{r}'].value, sf['目录'][f'{c}{r}'].value)
for name in PROJ:
    for c in 'ABCDEFGHIJKLMNOPQRSTU':
        if c not in wf[name].column_dimensions and c not in sf[name].column_dimensions:
            continue
    for c in ('B', 'C', 'Q', 'R', 'S', 'T'):
        def hid(ws, col):
            idx = openpyxl.utils.column_index_from_string(col)
            for k, d in ws.column_dimensions.items():
                if (d.min or 0) <= idx <= (d.max or 0):
                    return bool(d.hidden)
            return False
        chk(f'{name} {c} 列隐藏状态', str(hid(wf[name], c)), str(hid(sf[name], c)))

print('二、发票统计逐行复算')
byrow = {rec['r']: rec for rec in rows}
for r in range(R0, R1 + 1):
    rec = byrow.get(r)
    for c in 'ABIJLMPT':
        got = lv[f'{c}{r}'].value
        if rec is None:
            chk(f'发票统计 {c}{r} 空行', got if got != '' else None, None)
        elif c in 'AB':
            chk(f'发票统计 {c}{r}', got, rec[c])
        elif rec[c] is None:
            chk(f'发票统计 {c}{r}', got if got != '' else None, None)
        else:
            chk(f'发票统计 {c}{r}', got, rec[c])
    v = lv[f'V{r}'].value
    if rec is None:
        chk(f'发票统计 V{r} 空行', v if v != '' else None, None)
    else:
        chk(f'发票统计 V{r}', v, '△ 同一天同项目有一样的销售金额，看看是不是录重了' if r in (50, 51) else '✓')
tot = {c: sum(N(rec[c]) for rec in rows) for c in 'HIJKMNOQRS'}
tot['L'], tot['P'], tot['T'] = tot['J'] - tot['K'], tot['M'] - tot['N'] - tot['O'], tot['Q'] - tot['R'] - tot['S']
for c, v in tot.items():
    chk(f'发票统计 {c}3 合计', lv[f'{c}3'].value, v)
chk('发票统计 V3', lv['V3'].value, '△ 2 行请留意')
# 没改口径的列，合计应与原表缓存值一致
for c in 'HKOS':
    chk(f'发票统计 {c}3 与原表一致', lv[f'{c}3'].value, sv['发票统计'][f'{c}3'].value)
print(f'  原表 J3 {sv["发票统计"]["J3"].value:,.2f} → 现在 {lv["J3"].value:,.2f}；'
      f'M3 {sv["发票统计"]["M3"].value:,.2f} → {lv["M3"].value:,.2f}（第 50、51 行补了公式）')

print('三、分表逐行复算')
TEXTC = 'ABCDEFU'
for name in PROJ:
    ps = wv[name]
    mine = [rec for rec in rows if rec['A'] == name]
    chk(f'{name} V3 简称', ps['V3'].value, name)
    chk(f'{name} V4', ps['V4'].value, f'共 {len(mine)} 笔')
    for i in range(P1 - P0 + 1):
        r = P0 + i
        if i < len(mine):
            rec = mine[i]
            for c in 'ABCDEFGHIJKLMNOPQRSTU':
                got = ps[f'{c}{r}'].value
                exp = rec.get(c)
                if c in TEXTC:
                    chk(f'{name} {c}{r}', got, '' if exp is None else str(exp))
                elif c == 'G':
                    chk(f'{name} G{r}', got if got != '' else None, exp) if exp is None else chk(f'{name} G{r}', got, exp)
                else:
                    chk(f'{name} {c}{r}', got, N(exp))
        else:
            for c in 'ADHLPT':
                chk(f'{name} {c}{r} 空行', ps[f'{c}{r}'].value if ps[f'{c}{r}'].value != '' else None, None)
    for c in 'HIJKMNOQRS':
        chk(f'{name} {c}3', ps[f'{c}3'].value, sum(N(rec[c]) for rec in mine))
    chk(f'{name} L3', ps['L3'].value, sum(N(x['J']) - N(x['K']) for x in mine))
    chk(f'{name} P3', ps['P3'].value, sum(N(x['M']) - N(x['N']) - N(x['O']) for x in mine))
    chk(f'{name} T3', ps['T3'].value, sum(N(x['Q']) - N(x['R']) - N(x['S']) for x in mine))
    # 分表每一行的余额 = 发票统计同一笔的余额
    for i, rec in enumerate(mine):
        for c in 'LPT':
            chk(f'{name} {c}{P0 + i} 与发票统计第 {rec["r"]} 行', ps[f'{c}{P0 + i}'].value, lv[f'{c}{rec["r"]}'].value)
chk('合川瑞翔 标题', wv['合川瑞翔']['A1'].value, '高新区厂房及配套设施提质改造项目开票到瑞翔厂区电力土建(管廊)建设工程')

print('四、项目汇总')
sm = wv['项目汇总']
order = []
for r in range(3, 23):
    c = sf['目录'][f'C{r}'].value
    if c and c not in order:
        order.append(c)
MAPC = dict(zip('EFGHIJKLMNOPQ', 'HIJKLMNOPQRST'))
for i, name in enumerate(order):
    r = 4 + i
    mine = [rec for rec in rows if rec['A'] == name]
    chk(f'汇总 B{r}', sm[f'B{r}'].value, name)
    chk(f'汇总 C{r} 笔数', sm[f'C{r}'].value, len(mine))
    chk(f'汇总 D{r} 最后一笔', sm[f'D{r}'].value, mine[-1]['D'] if mine else '')
    for sc, lc in MAPC.items():
        if lc in 'LPT':
            exp = mine[-1][lc] if mine else 0
        else:
            exp = sum(N(rec[lc]) for rec in mine)
        chk(f'汇总 {sc}{r}（{name}）', sm[f'{sc}{r}'].value, exp)
    chk(f'汇总 R{r}', sm[f'R{r}'].value, '✓')
for r in range(4 + len(order), 54):
    chk(f'汇总 B{r} 空', sm[f'B{r}'].value if sm[f'B{r}'].value != '' else None, None)
for sc, lc in MAPC.items():
    chk(f'汇总 {sc}3', sm[f'{sc}3'].value, tot[lc])
chk('汇总 R3', sm['R3'].value, '✓ 与发票统计合计一致')

print('五、全表无错误值')
for ws in wv.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('#') and c.value.rstrip('!?0/').upper() in (
                    '#REF', '#VALUE', '#N/A', '#NAME', '#DIV', '#NUM', '#NULL'):
                chk(f'{ws.title}!{c.coordinate} 错误值', c.value, '')
print(f'复算完毕：不一致 {len(BAD)} 项')
n_base = len(BAD)

# ───────── 场景 ─────────
print('六、场景实测')
os.makedirs(TMP, exist_ok=True)
X = os.path.join(TMP, 'scenario.xlsx')
shutil.copy(OUT, X)
wb = openpyxl.load_workbook(X)
L = wb['发票统计']
ml = wb['目录']
YUN = dirmap and [k for k, v in dirmap.items() if v == '云龙社区'][0]
HCB = [k for k, v in dirmap.items() if v == '合川表计'][0]
ml['B12'], ml['C12'] = '测试新项目（全称）', '测试新项目'
ml['B13'], ml['C13'] = '没建表的项目（全称）', '没建表'
cp = wb.copy_worksheet(wb['云龙社区'])
cp.title = '测试新项目'


def put(r, **kw):
    for c, v in kw.items():
        L[f'{c}{r}'] = v


put(90, C=YUN, D='2026.9.24测试开票', E='新德润', F='3%专票', G=0.1, H=1000, K=500)
put(91, C='测试新项目（全称）', D='2026.9.24新项目开票', H=2000, K=2000)
put(92, C='没建表的项目（全称）', D='2026.9.24没建表的项目开票', H=300)
put(93, C='目录里没有的项目', D='2026.9.24名称不在目录', H=100)
put(94, D='2026.9.24没填项目名称', H=50)
put(95, C=YUN, D='2026.9.24管理费点填成整数', G=10, H=10)
put(96, C=YUN, D='测试金额里有文字', H='一千')
put(97, C=HCB, D='2026.9.24红冲第48行', H=-150000, K=-150000)
put(98, C=YUN, D='2026.9.24插行没带公式', O=1)
for c in 'ABIJLMPTVW':
    L[f'{c}98'] = None                                  # 模拟插了一行、没复制公式
wb.save(X)
subprocess.run([sys.executable, RECALC, X, '900'], check=True, capture_output=True)
v = openpyxl.load_workbook(X, data_only=True)
L2, S2 = v['发票统计'], v['项目汇总']
yl = v['云龙社区']
chk('① 新行进了云龙社区分表（第 5 笔）', yl['D8'].value, '2026.9.24测试开票', quiet=False)
chk('① 分表 L8 = 本项目累计未开成本票', yl['L8'].value, 900 - 500, quiet=False)
chk('① 发票统计 L90 = 同一个数', L2['L90'].value, 900 - 500, quiet=False)
srow = {S2[f'B{r}'].value: r for r in range(4, 54) if S2[f'B{r}'].value}
chk('① 项目汇总 云龙 销售金额 +1000（"一千"是文字不算，管理费点 10 那行 +10）', S2[f'E{srow["云龙社区"]}'].value, 227260.8 + 1000 + 10, quiet=False)
nv = v['测试新项目']
chk('② 复制出来的分表自动认表名', nv['V3'].value, '测试新项目', quiet=False)
chk('② 新分表 V4', nv['V4'].value, '共 1 笔', quiet=False)
chk('② 新分表第 1 笔', nv['D4'].value, '2026.9.24新项目开票', quiet=False)
chk('② 新分表合计 H3', nv['H3'].value, 2000, quiet=False)
chk('② 项目汇总多出新项目且分表核对 ✓', S2[f'R{srow.get("测试新项目", 4)}'].value if '测试新项目' in srow else None, '✓', quiet=False)
chk('③ 目录有、没建分表', S2[f'R{srow.get("没建表", 4)}'].value if '没建表' in srow else None, '⚠ 还没建分表', quiet=False)
chk('④ 名称不在目录 → 校验', L2['V93'].value, '⚠ 项目名称在目录里找不到（先加到目录）', quiet=False)
chk('④ 项目汇总合计对不上 → 报出来', (S2['R3'].value or '')[:1], '⚠', quiet=False)
chk('⑤ 没填项目名称', L2['V94'].value, '⚠ 没填项目名称', quiet=False)
chk('⑥ 管理费点填成整数', L2['V95'].value, '⚠ 管理费点要填百分比，如 10%', quiet=False)
chk('⑦ 金额里有文字', L2['V96'].value, '⚠ 金额列里有文字，算不进合计', quiet=False)
chk('⑧ 插行没带公式 → V3 报出来', L2['V3'].value, '⚠ 有 1 行缺公式或缺简称', quiet=False)
hb = v['合川表计']
chk('⑨ 红冲：合川表计 销售金额 −150000', hb['H3'].value, 895460 - 150000, quiet=False)
chk('⑨ 红冲：合川表计 已开成本票 −150000', hb['K3'].value, 795460 - 150000, quiet=False)
chk('⑨ 红冲行排在合川表计最后', hb[f'D{P0 + 42}'].value, '2026.9.24红冲第48行', quiet=False)
print(f'\n不一致：复算 {n_base} 项，场景 {len(BAD) - n_base} 项')
sys.exit(1 if BAD else 0)
