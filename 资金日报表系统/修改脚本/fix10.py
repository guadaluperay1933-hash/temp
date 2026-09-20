# -*- coding: utf-8 -*-
"""资金日报表 · 第十轮

① 《汇报表》② 货币资金日报表：加「内部转入 / 内部转出」两列，跟 ③ 多帐户资金汇总表一个口径。
   加完之后 ② 块占 A:H，③ 块整体右移两列到 J:Q（原来在 H:O）。
   本日收入/本日支出也跟着改成「不含内部转账」，账户之间互转单独走那两列，
   本日余额 = 昨日余额 + 本日收入 − 本日支出 + 内部转入 − 内部转出，数还是原来那个数。

② 《收支报表》：在「资金支出（按收支类别）」下面再加一块「资金支出（按供应商）」，
   12 个月 × 40 行，跟上面收入按客户那块一模一样的算法。

③ 《任意时间段收支报表》：右边再加一栏「资金支出（按供应商）」，跟旁边两栏并排。

跑法：python3 fix10.py <入.xlsx> <出.xlsx>
"""
import sys, re
import openpyxl
from copy import copy

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)

NOTIN = '"<>内部转账"'
ISIN = '"内部转账"'


def st(ws, addr):
    return copy(ws[addr]._style)


# ══════════════════════════════════════════════════════════════
# ① 汇报表：日报表加内部转入/转出，月度块右移两列
# ══════════════════════════════════════════════════════════════
ws = wb['汇报表']
R0, R1, RT = 43, 72, 73          # 数据行 / 合计行

# —— 先把要复用的样式抄下来（等会儿原格就被覆盖了）——
S2 = {c: st(ws, f'{c}42') for c in 'ABCDEF'}            # ② 表头
S2D = {c: st(ws, f'{c}43') for c in 'ABCDEF'}           # ② 数据行
S2T = {c: st(ws, f'{c}{RT}') for c in 'ABCDEF'}         # ② 合计行
S3 = {c: st(ws, f'{c}42') for c in 'HIJKLMNO'}          # ③ 表头
S3D = {c: st(ws, f'{c}43') for c in 'HIJKLMNO'}         # ③ 数据行
S3T = {c: st(ws, f'{c}{RT}') for c in 'HIJKLMNO'}       # ③ 合计行
S_T40 = st(ws, 'A40'); S_T40b = st(ws, 'H40')
S_L41 = st(ws, 'A41'); S_V41 = st(ws, 'B41'); S_N41 = st(ws, 'C41')
S_L41b = st(ws, 'H41'); S_V41b = st(ws, 'I41'); S_N41b = st(ws, 'J41')
OLD_MONTH = ws['I41'].value
OLD_NOTE2 = ws['C41'].value
OLD_NOTE3 = ws['J41'].value

for rng in ('A40:F40', 'C41:F41', 'H40:O40', 'J41:O41'):
    if rng in [str(x) for x in ws.merged_cells.ranges]:
        ws.unmerge_cells(rng)

# —— 清空 40~73 行 H:O 的老内容（马上要重排）——
for r in range(40, RT + 1):
    for c in 'HIJKLMNO':
        ws[f'{c}{r}'].value = None

# —— ② 块：A:H ——
ws['A40'] = '② 货币资金日报表（内部转入/转出单列，跟右边 ③ 一个口径）'
ws['A40']._style = S_T40
for c in 'BCDEFGH':
    ws[f'{c}40']._style = S_T40
ws.merge_cells('A40:H40')

ws['A41'] = '日期：'; ws['A41']._style = S_L41
ws['B41']._style = S_V41                      # B41 公式原样保留
ws['C41'] = OLD_NOTE2
for c in 'CDEFGH':
    ws[f'{c}41']._style = S_N41
ws.merge_cells('C41:H41')

H2 = ['序号', '账户', '昨日余额', '本日收入', '本日支出', '内部转入', '内部转出', '本日余额']
for c, v in zip('ABCDEFGH', H2):
    ws[f'{c}42'] = v
    ws[f'{c}42']._style = S2[c] if c in S2 else S2['F']
ws['F42']._style = copy(S2['E']); ws['G42']._style = copy(S2['E']); ws['H42']._style = copy(S2['F'])

for r in range(R0, R1 + 1):
    k = r - 42
    br = r - 38                                # 基础资料 账户行
    blank = f'OR($B{r}="",$B$41="")'
    day = f'流水账户,$B{r},流水日期,$B$41'
    ws[f'A{r}'] = f'=IF($B{r}="","",{k})'
    ws[f'B{r}'] = f'=IF(基础资料!$B${br}="","",基础资料!$B${br})'
    ws[f'C{r}'] = (f'=IF({blank},"",ROUND(SUMIF(账户表,$B{r},账户期初)'
                   f'+SUMIFS(流水收入,流水账户,$B{r},流水日期,"<"&$B$41)'
                   f'-SUMIFS(流水支出,流水账户,$B{r},流水日期,"<"&$B$41),2))')
    ws[f'D{r}'] = f'=IF({blank},"",ROUND(SUMIFS(流水收入,{day},流水属性,{NOTIN}),2))'
    ws[f'E{r}'] = f'=IF({blank},"",ROUND(SUMIFS(流水支出,{day},流水属性,{NOTIN}),2))'
    ws[f'F{r}'] = f'=IF({blank},"",ROUND(SUMIFS(流水收入,{day},流水属性,{ISIN}),2))'
    ws[f'G{r}'] = f'=IF({blank},"",ROUND(SUMIFS(流水支出,{day},流水属性,{ISIN}),2))'
    ws[f'H{r}'] = f'=IF({blank},"",ROUND($C{r}+$D{r}-$E{r}+$F{r}-$G{r},2))'
    for c in 'ABCDE':
        ws[f'{c}{r}']._style = copy(S2D[c])
    ws[f'F{r}']._style = copy(S2D['E'])
    ws[f'G{r}']._style = copy(S2D['E'])
    ws[f'H{r}']._style = copy(S2D['F'])

ws[f'A{RT}'] = '合计'; ws[f'A{RT}']._style = copy(S2T['A'])
ws[f'B{RT}']._style = copy(S2T['B'])
for c in 'CDEFGH':
    ws[f'{c}{RT}'] = f'=ROUND(SUM({c}{R0}:{c}{R1}),2)'
ws[f'C{RT}']._style = copy(S2T['C']); ws[f'D{RT}']._style = copy(S2T['D'])
ws[f'E{RT}']._style = copy(S2T['E']); ws[f'F{RT}']._style = copy(S2T['E'])
ws[f'G{RT}']._style = copy(S2T['E']); ws[f'H{RT}']._style = copy(S2T['F'])

# —— ③ 块：整体搬到 J:Q ——
ws['J40'] = '③ 多帐户资金汇总表·月度（经营口径，内部转入/转出单列）'
for c in 'JKLMNOPQ':
    ws[f'{c}40']._style = copy(S_T40b)
ws.merge_cells('J40:Q40')

ws['J41'] = '月份：'; ws['J41']._style = copy(S_L41b)
ws['K41'] = OLD_MONTH; ws['K41']._style = copy(S_V41b)
ws['L41'] = OLD_NOTE3
for c in 'LMNOPQ':
    ws[f'{c}41']._style = copy(S_N41b)
ws.merge_cells('L41:Q41')

H3 = ['序号', '账号名称', '期初余额', '本月收入', '本月支出', '内部转入', '内部转出', '结余']
for c, src, v in zip('JKLMNOPQ', 'HIJKLMNO', H3):
    ws[f'{c}42'] = v
    ws[f'{c}42']._style = copy(S3[src])

for r in range(R0, R1 + 1):
    k = r - 42
    br = r - 38
    blank = f'OR($K{r}="",$K$41="")'
    mon = f'流水账户,$K{r},流水日期,">="&$K$41,流水日期,"<="&EOMONTH($K$41,0)'
    ws[f'J{r}'] = f'=IF($K{r}="","",{k})'
    ws[f'K{r}'] = f'=IF(基础资料!$B${br}="","",基础资料!$B${br})'
    ws[f'L{r}'] = (f'=IF({blank},"",ROUND(SUMIF(账户表,$K{r},账户期初)'
                   f'+SUMIFS(流水收入,流水账户,$K{r},流水日期,"<"&$K$41)'
                   f'-SUMIFS(流水支出,流水账户,$K{r},流水日期,"<"&$K$41),2))')
    ws[f'M{r}'] = f'=IF({blank},"",ROUND(SUMIFS(流水收入,{mon},流水属性,{NOTIN}),2))'
    ws[f'N{r}'] = f'=IF({blank},"",ROUND(SUMIFS(流水支出,{mon},流水属性,{NOTIN}),2))'
    ws[f'O{r}'] = f'=IF({blank},"",ROUND(SUMIFS(流水收入,{mon},流水属性,{ISIN}),2))'
    ws[f'P{r}'] = f'=IF({blank},"",ROUND(SUMIFS(流水支出,{mon},流水属性,{ISIN}),2))'
    ws[f'Q{r}'] = f'=IF({blank},"",ROUND($L{r}+$M{r}-$N{r}+$O{r}-$P{r},2))'
    for c, src in zip('JKLMNOPQ', 'HIJKLMNO'):
        ws[f'{c}{r}']._style = copy(S3D[src])

ws[f'J{RT}'] = '合计'
for c, src in zip('JKLMNOPQ', 'HIJKLMNO'):
    ws[f'{c}{RT}']._style = copy(S3T[src])
for c in 'LMNOPQ':
    ws[f'{c}{RT}'] = f'=ROUND(SUM({c}{R0}:{c}{R1}),2)'

# 中间留一列做间隔
for r in range(40, RT + 1):
    ws[f'I{r}'].value = None
    ws[f'I{r}'].border = openpyxl.styles.Border()
    ws[f'I{r}'].fill = openpyxl.styles.PatternFill()

# 顶上那行总说明跟着改一句
top = ws['A2'].value
if isinstance(top, str) and '含内部划转' in top:
    ws['A2'] = top.replace('这一块是「账户口径」　——　收款/付款按钱进出账户算，内部划转也算在里面（因为账户余额确实动了）',
                           '这一块是「账户口径」　——　收款/付款按钱进出账户算，内部划转也算在里面（因为账户余额确实动了）')
ws['A2'] = re.sub(r'②\s*货币资金日报表：指定某一天。',
                  '② 货币资金日报表：指定某一天，内部转入/转出跟 ③ 一样单列两列。', str(ws['A2'].value))

# 核对表跟着改引用
ck = wb['核对表']
v = ck['D9'].value
if isinstance(v, str):
    ck['D9'] = v.replace('汇报表!$O$73', '汇报表!$Q$73').replace('汇报表!$I$41', '汇报表!$K$41')

print('  ✓ 汇报表：② 日报表加内部转入/转出（A:H），③ 月度块右移到 J:Q，核对表引用同步')


# ══════════════════════════════════════════════════════════════
# ② 收支报表：加「资金支出（按供应商）」块
# ══════════════════════════════════════════════════════════════
ws = wb['收支报表']
GAP, TOT, HDR, D0, D1, NIL = 91, 92, 93, 94, 133, 134
MON = 'CDEFGHIJKLMN'

for r in range(GAP, NIL + 1):
    ws.row_dimensions[r].height = ws.row_dimensions[{GAP: 47, TOT: 48, HDR: 49, NIL: 90}.get(r, 50)].height

for c in list('AB') + list(MON) + ['O', 'P']:
    ws[f'{c}{GAP}']._style = st(ws, f'{c}47')

# 支出总额（再放一次，好让下面的「未列示」补差）
ws[f'A{TOT}'] = ''
ws[f'B{TOT}'] = '支出总额（不含内部转账）'
for i, c in enumerate(MON, 1):
    ws[f'{c}{TOT}'] = f'=ROUND(SUMIFS(流水支出,流水月份,(管理年度*100+{i}),流水属性,{NOTIN}),2)'
ws[f'O{TOT}'] = f'=ROUND(SUM(C{TOT}:N{TOT}),2)'
ws[f'P{TOT}'] = ''
for c in list('AB') + list(MON) + ['O', 'P']:
    ws[f'{c}{TOT}']._style = st(ws, f'{c}48')

# 表头
ws[f'A{HDR}'] = '序'
ws[f'B{HDR}'] = '资金支出（按供应商）'
for i, c in enumerate(MON, 1):
    ws[f'{c}{HDR}'] = f'{i}月'
ws[f'O{HDR}'] = '全年合计'
ws[f'P{HDR}'] = '占比'
for c in list('AB') + list(MON) + ['O', 'P']:
    ws[f'{c}{HDR}']._style = st(ws, f'{c}49')

# 数据行
SUP = [wb['基础资料'][f'M{r}'].value for r in range(5, 5 + (D1 - D0 + 1))]
for i, r in enumerate(range(D0, D1 + 1)):
    ws[f'A{r}'] = f'=IF($B{r}="","",COUNTA($B${D0}:$B{r}))'
    ws[f'B{r}'] = SUP[i] if i < len(SUP) and SUP[i] else None
    for m, c in enumerate(MON, 1):
        ws[f'{c}{r}'] = (f'=IF($B{r}="","",ROUND(SUMIFS(流水支出,流水供应商,$B{r},'
                         f'流水月份,(管理年度*100+{m}),流水属性,{NOTIN}),2))')
    ws[f'O{r}'] = f'=IF($B{r}="","",ROUND(SUM(C{r}:N{r}),2))'
    ws[f'P{r}'] = f'=IF(OR($B{r}="",N($O${TOT})=0),"",$O{r}/$O${TOT})'
    for c in list('AB') + list(MON) + ['O', 'P']:
        ws[f'{c}{r}']._style = st(ws, f'{c}50')

# 未列示
ws[f'A{NIL}'] = ''
ws[f'B{NIL}'] = '未列示 / 没挂供应商（自动补差，别删）'
for c in MON:
    ws[f'{c}{NIL}'] = f'=ROUND({c}{TOT}-SUM({c}{D0}:{c}{D1}),2)'
ws[f'O{NIL}'] = f'=ROUND(SUM(C{NIL}:N{NIL}),2)'
ws[f'P{NIL}'] = f'=IF(N($O${TOT})=0,"",$O{NIL}/$O${TOT})'
for c in list('AB') + list(MON) + ['O', 'P']:
    ws[f'{c}{NIL}']._style = st(ws, f'{c}90')

ws['A2'] = ('收入按「客户」归集、支出分两块：按「收支类别」和按「供应商」各来一遍，同一笔钱两种口径各看一次。'
            '每一块都留了 40 行，空着的行直接往下打字就行，合计范围已经把整块都算进去了，不用改公式。'
            '最后一行「未列示」自动补差 —— 按供应商那一块里，工资、税费这些没挂供应商的支出就落在那一行，属正常。')
print('  ✓ 收支报表：第 91~134 行加「资金支出（按供应商）」块（12 个月 × 40 行）')


# ══════════════════════════════════════════════════════════════
# ③ 任意时间段收支报表：右边并排再加一栏「按供应商」
# ══════════════════════════════════════════════════════════════
ws = wb['任意时间段收支报表']
r0, r1, rn, rt = 7, 46, 47, 48

for rng in ('A1:H1', 'A2:H2'):
    if rng in [str(x) for x in ws.merged_cells.ranges]:
        ws.unmerge_cells(rng)
ws.merge_cells('A1:L1')
ws.merge_cells('A2:L2')
for c in 'IJKL':
    ws[f'{c}1']._style = st(ws, 'H1')
    ws[f'{c}2']._style = st(ws, 'H2')

for c, w in (('I', 5), ('J', 20), ('K', 15), ('L', 10)):
    ws.column_dimensions[c].width = w

for c, src, v in zip('IJKL', 'EFGH', ['序', '资金支出（按供应商）', '金额', '支出占比']):
    ws[f'{c}6'] = v
    ws[f'{c}6']._style = st(ws, f'{src}6')

SUP = [wb['基础资料'][f'M{r}'].value for r in range(5, 5 + (r1 - r0 + 1))]
for i, r in enumerate(range(r0, r1 + 1)):
    ws[f'I{r}'] = f'=IF($J{r}="","",COUNTA($J${r0}:$J{r}))'
    ws[f'J{r}'] = SUP[i] if i < len(SUP) and SUP[i] else None
    ws[f'K{r}'] = (f'=IF($J{r}="","",ROUND(SUMIFS(流水支出,流水供应商,$J{r},'
                   f'流水日期,">="&$B$3,流水日期,"<="&$B$4,流水属性,{NOTIN}),2))')
    ws[f'L{r}'] = f'=IF(OR($J{r}="",N($F$3)=0),"",K{r}/$F$3)'
    for c, src in zip('IJKL', 'EFGH'):
        ws[f'{c}{r}']._style = st(ws, f'{src}{r}')

ws[f'I{rn}'] = ''
ws[f'J{rn}'] = '未列示 / 没挂供应商（自动补差，别删）'
ws[f'K{rn}'] = f'=ROUND($F$3-SUM(K{r0}:K{r1}),2)'
ws[f'L{rn}'] = f'=IF(N($F$3)=0,"",K{rn}/$F$3)'
ws[f'I{rt}'] = ''
ws[f'J{rt}'] = '合计'
ws[f'K{rt}'] = f'=ROUND(SUM(K{r0}:K{rn}),2)'
ws[f'L{rt}'] = ''
for r in (rn, rt):
    for c, src in zip('IJKL', 'EFGH'):
        ws[f'{c}{r}']._style = st(ws, f'{src}{r}')

ws['A2'] = ('左边收入按客户，中间支出按收支类别，右边支出按供应商 —— 支出的两栏是同一笔钱的两种看法，合计相等。'
            '各留了 40 行，空行直接往下打字，合计已经把整块算进去了。最后一行「未列示」自动补差：'
            '按供应商那一栏里，工资、税费这些没挂供应商的支出就落在那一行。'
            '收入总额和支出总额都不含内部转账，内部转账净额单独显示。')
print('  ✓ 任意时间段收支报表：右边并排加「资金支出（按供应商）」栏（I:L）')

# ══════════════════════════════════════════════════════════════
# ④ 使用说明跟着改几句
# ══════════════════════════════════════════════════════════════
ws = wb['使用说明']
ws['C10'] = ('【月度汇报表】、【汇报表】② 日报表和 ③ 月度汇总表，各有「内部转入 / 内部转出」两列；'
             '【收支分类报表】单列一行「内部转账净额」。收入合计、支出合计里都不含它。')
ws['C16'] = ('指定某一天，各账户的 昨日余额 / 本日收入 / 本日支出 / 内部转入 / 内部转出 / 本日余额。'
             '跟 ③ 一个口径：本日收入、本日支出不含内部划转，账户之间互转单走那两列，'
             '本日余额 = 昨日余额 + 收入 − 支出 + 转入 − 转出。')
ws['C18'] = '收入按客户一块；支出两块：按收支类别一块、按供应商一块。同一笔支出两种口径各看一次，合计相等。'
ws['C19'] = '三栏并排：收入按客户 / 支出按收支类别 / 支出按供应商。'
ws['C25'] = ('【基础资料】对应清单接着写。【收支报表】和【任意时间段收支报表】的清单是手填的，'
             '记得把新客户、新供应商也补到那几块的空行里（不补也不会少数，会落到「未列示」那一行）。')
print('  ✓ 使用说明：口径说明同步')

# 交付件不走 LibreOffice 重算（它不认 FILTER，会把账户子表的函数名改小写），
# 所以让 Excel/WPS 一打开就全表重算一次，新加的格子立刻有数。
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(f'已写 {OUT}')
