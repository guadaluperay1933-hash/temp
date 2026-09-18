# -*- coding: utf-8 -*-
"""专门验三件事（都是审计在原模板上实测出来、这次重做时特意堵住的）：

  ① 日期录成文本（"2026.6.28"）—— 原表：K 列 #VALUE!，三张查询报表静默少算，没有任何提示
  ② 流水不在管理年度内 —— 原表：收支报表月份表头写死 202601~202612，跨年整表归零
  ③ 汇报表查询区间超过 31 天 / 起止日期填反 —— 原表：期末结余整列 #VALUE!、期初显示 0

用法: python3 test_guards.py <交付件xlsx> <工作目录>
"""
import os, sys, shutil, datetime
import openpyxl

SRC, WORK = sys.argv[1], sys.argv[2]
R0, N_ROW, B0 = 5, 5000, 5
os.makedirs(WORK, exist_ok=True)
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__)))), '工具'))
import recalc_safe                                                   # noqa: E402

ok = bad = 0


def chk(n, g, w):
    global ok, bad
    if g == w:
        ok += 1
    else:
        bad += 1
        print(f'  ✗ {n}: 实际 {g!r}  应为 {w!r}')


def N(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


def load(tag):
    t = os.path.join(WORK, f'{tag}.xlsx')
    shutil.copy(SRC, t)
    return t, openpyxl.load_workbook(t)


def last_row(ws):
    return max(r for r in range(R0, R0 + N_ROW) if ws.cell(r, 3).value) + 1


# ══ ① 文本日期 + ② 跨年度 ══════════════════════════════════════════
t, wb = load('g1')
de = wb['数据录入']
r = last_row(de)
for c, v in ((2, '2026.6.28'), (3, '现金'), (4, '销售回款'), (6, 1234), (9, '客户1')):
    de.cell(r, c).value = v                                    # ① 文本日期
for c, v in ((2, datetime.datetime(2027, 1, 15)), (3, '现金'), (4, '销售回款'),
             (6, 5678), (9, '客户1')):
    de.cell(r + 1, c).value = v                                # ② 2027 年，不在管理年度
wb.save(t)
print(f'塞进去：第 {r} 行日期录成文本 "2026.6.28" 收 1234，第 {r+1} 行 2027-01-15 收 5678')
recalc_safe.recalc(t, 900)
g = openpyxl.load_workbook(t, data_only=True)
de, ck = g['数据录入'], g['核对表']
print('  第 %d 行核对列：%s' % (r, de.cell(r, 16).value))
print('  第 %d 行核对列：%s' % (r + 1, de.cell(r + 1, 16).value))
chk('文本日期 → 核对列点名', str(de.cell(r, 16).value).startswith('※日期是文本'), True)
chk('文本日期 → 月份列留空不报 #VALUE!', de.cell(r, 11).value, None)
chk('跨年度 → 核对列点名', str(de.cell(r + 1, 16).value).startswith('△不是管理年度'), True)
rows = {str(ck.cell(x, 3).value): x for x in range(5, 30) if ck.cell(x, 3).value}
k1 = next(x for k, x in rows.items() if k.startswith('12 个月加起来'))
k2 = next(x for k, x in rows.items() if k.startswith('日期没填'))
print('  勾稽「12 个月加起来 = 全部流水」= %s（%s）' % (ck.cell(k1, 4).value, ck.cell(k1, 5).value))
print('  待办「日期没填/文本/不在管理年度」= %s' % ck.cell(k2, 4).value)
chk('勾稽把这两笔抓出来了（差额 = 1234+5678）', round(N(ck.cell(k1, 4).value), 2), -6912.0)
chk('待办笔数 = 2', N(ck.cell(k2, 4).value), 2)

# 把管理年度改成 2027，看那笔 2027 的能不能归位
t2, wb2 = load('g2')
de2 = wb2['数据录入']
r2 = last_row(de2)
for c, v in ((2, datetime.datetime(2027, 1, 15)), (3, '现金'), (4, '销售回款'),
             (6, 5678), (9, '客户1')):
    de2.cell(r2, c).value = v
wb2['基础资料'].cell(B0, 18).value = 2027                      # 管理年度 2026 → 2027
wb2.save(t2)
print('\n把【基础资料】管理年度改成 2027，那笔 2027-01-15 应该归位')
recalc_safe.recalc(t2, 900)
g2 = openpyxl.load_workbook(t2, data_only=True)
mr2, kr2, sr2 = g2['月度汇报表'], g2['收支分类报表'], g2['收支报表']
chk('月度汇报表 1 月收入 = 5678', N(mr2.cell(5, 2).value), 5678)
chk('月度汇报表 6 月收入 = 0（2026 年的已不在管理年度）', N(mr2.cell(10, 2).value), 0)
chk('月度汇报表 A5 显示 2027 年 1 月', mr2['A5'].value, datetime.datetime(2027, 1, 1))
chk('收支分类报表 1 月「销售回款」= 5678', N(kr2.cell(12, 3).value), 5678)
chk('收支报表 1 月收入总额 = 5678', N(sr2.cell(4, 3).value), 5678)
chk('收支报表 1 月「客户1」= 5678', N(sr2.cell(6, 3).value), 5678)
ck2 = g2['核对表']
k1b = next(x for k, x in {str(ck2.cell(x, 3).value): x
                          for x in range(5, 30) if ck2.cell(x, 3).value}.items()
           if k.startswith('12 个月加起来'))
print('  换年后勾稽差额 = %s（2026 那 25 笔现在落在年度外，属正常）' % ck2.cell(k1b, 4).value)

# ══ ③ 汇报表区间超 31 天 / 日期填反 ═════════════════════════════════
for tag, b, e, want in (('g3', datetime.datetime(2026, 6, 1), datetime.datetime(2026, 7, 31),
                         '※超过 31 天，请分段查'),
                        ('g4', datetime.datetime(2026, 6, 30), datetime.datetime(2026, 6, 1),
                         '※截止早于起始')):
    t3, wb3 = load(tag)
    wb3['汇报表']['B3'], wb3['汇报表']['E3'] = b, e
    wb3.save(t3)
    recalc_safe.recalc(t3, 900)
    hb = openpyxl.load_workbook(t3, data_only=True)['汇报表']
    print('\n汇报表查 %s ~ %s：H3 = %s' % (b.date(), e.date(), hb['H3'].value))
    chk(f'{tag} 查询天数给出提示', hb['H3'].value, want)
    for cell, nm in (('B5', '公司期初余额'), ('D5', '期间收款合计'),
                     ('F5', '期间付款合计'), ('H5', '期末结余金额')):
        v = hb[cell].value
        chk(f'{tag} {nm} 不是错误值', isinstance(v, str) and v.startswith('#'), False)
    bad_cells = [f'{c}{x}' for x in range(8, 38) for c in ('B', 'BM', 'BN', 'BO')
                 if isinstance(hb[f'{c}{x}'].value, str) and hb[f'{c}{x}'].value.startswith('#')]
    chk(f'{tag} 明细区一个错误值都没有', bad_cells, [])

print(f'\n══ 三道防线：{ok} 项通过，{bad} 项不通过 ══')
sys.exit(1 if bad else 0)
