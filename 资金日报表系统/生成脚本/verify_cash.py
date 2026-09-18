# -*- coding: utf-8 -*-
"""核对《资金日报表_混合录入版.xlsx》。用法: python3 verify_cash.py <重算件> [交付件] [原模板]

重算件里账户子表是 #NAME?（LibreOffice 没有 FILTER 函数），那是重算工具的毛病，
交付件里公式是大写的、没动过。这里核对所有能核对的。
"""
import sys, datetime, collections
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

got = openpyxl.load_workbook(sys.argv[1], data_only=True)
out = openpyxl.load_workbook(sys.argv[2]) if len(sys.argv) > 2 else None
ba, de = got['基础资料'], got['数据录入']
hb, mr, kr, sr, qr, ck = (got['汇报表'], got['月度汇报表'], got['收支分类报表'],
                          got['收支报表'], got['任意时间段收支报表'], got['核对表'])
B0, R0, N_ACC, N_ROW, N_CAT, N_ITEM = 5, 5, 30, 5000, 40, 40
YEAR = 2026
ok = bad = 0


def chk(n, g, w, tol=0.02):
    global ok, bad
    good = abs(g - w) <= tol if isinstance(g, (int, float)) and isinstance(w, (int, float)) else g == w
    if good:
        ok += 1
    else:
        bad += 1
        print(f'  ✗ {n}: 实际 {g!r}  应为 {w!r}')


def N(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


# ── 读成 Python 结构 ──────────────────────────────────────────────────
accs = [dict(name=ba.cell(r, 2).value, ty=ba.cell(r, 3).value, open=N(ba.cell(r, 4).value))
        for r in range(B0, B0 + N_ACC) if ba.cell(r, 2).value]
cats = {ba.cell(r, 8).value: ba.cell(r, 9).value
        for r in range(B0, B0 + N_CAT) if ba.cell(r, 8).value}
flows = []
for r in range(R0, R0 + N_ROW):
    if not de.cell(r, 3).value:
        continue
    flows.append(dict(row=r, d=de.cell(r, 2).value, acc=de.cell(r, 3).value,
                      cat=de.cell(r, 4).value, inc=N(de.cell(r, 6).value),
                      exp=N(de.cell(r, 7).value), cus=de.cell(r, 9).value,
                      sup=de.cell(r, 10).value, attr=cats.get(de.cell(r, 4).value),
                      prj=de.cell(r, 14).value))
print(f'账户 {len(accs)} 个，收支类别 {len(cats)} 个，流水 {len(flows)} 笔')
OPEN = {a['name']: a['open'] for a in accs}
NOT_T = [f for f in flows if f['attr'] != '内部转账']
IS_T = [f for f in flows if f['attr'] == '内部转账']

print('══ 1. 数据录入：余额逐行结出来 ══')
run = dict(OPEN)
e = 0
for f in flows:
    run[f['acc']] = round(run.get(f['acc'], 0) + f['inc'] - f['exp'], 2)
    if abs(N(de.cell(f['row'], 8).value) - run[f['acc']]) > 0.02:
        e += 1
chk('账户余额逐行', e, 0)
chk('类别属性逐行',
    [de.cell(f['row'], 12).value for f in flows], [f['attr'] for f in flows])
chk('月份列逐行', [de.cell(f['row'], 11).value for f in flows],
    [int(f['d'].strftime('%Y%m')) for f in flows])
chk('顶上「收入总额」(不含内部转账)', N(de['B3'].value), round(sum(f['inc'] for f in NOT_T), 2))
chk('顶上「支出总额」(不含内部转账)', N(de['D3'].value), round(sum(f['exp'] for f in NOT_T), 2))
chk('顶上「经营净额」', N(de['F3'].value),
    round(sum(f['inc'] for f in NOT_T) - sum(f['exp'] for f in NOT_T), 2))
chk('顶上「内部转账净额」', N(de['H3'].value),
    round(sum(f['inc'] - f['exp'] for f in IS_T), 2))
chk('顶上「资金总额」', N(de['J3'].value),
    round(sum(OPEN.values()) + sum(f['inc'] - f['exp'] for f in flows), 2))

print('══ 2. 汇报表 ① 资金日报统计表（账户口径，含内部划转）══')
d0, d1 = hb['B3'].value, hb['E3'].value
days = (d1 - d0).days + 1
chk('查询天数', hb['H3'].value, days)
e1 = e2 = e3 = 0
for i, a in enumerate(accs):
    r = 8 + i
    op = round(OPEN[a['name']] + sum(f['inc'] - f['exp'] for f in flows
                                     if f['acc'] == a['name'] and f['d'] < d0), 2)
    if abs(N(hb.cell(r, 2).value) - op) > 0.02:
        e1 += 1
    rcv = round(sum(f['inc'] for f in flows if f['acc'] == a['name'] and d0 <= f['d'] <= d1), 2)
    pay = round(sum(f['exp'] for f in flows if f['acc'] == a['name'] and d0 <= f['d'] <= d1), 2)
    if abs(N(hb.cell(r, 65).value) - rcv) > 0.02:
        e2 += 1
    if abs(N(hb.cell(r, 66).value) - pay) > 0.02:
        e3 += 1
    if abs(N(hb.cell(r, 67).value) - round(op + rcv - pay, 2)) > 0.02:
        e3 += 1
chk('各账户期初余额', e1, 0)
chk('期间收款', e2, 0)
chk('期间付款 / 期末余额', e3, 0)
# 逐天对：第 1 天和最后一天
for k in (0, min(days, 31) - 1):
    day = d0 + datetime.timedelta(days=k)
    c = 3 + k * 2
    w_in = round(sum(f['inc'] for f in flows if f['acc'] == accs[0]['name'] and f['d'] == day), 2)
    chk(f'{accs[0]["name"]} {day:%m-%d} 收款', N(hb.cell(8, c).value), w_in)
chk('公司期初余额合计', N(hb['B5'].value),
    round(sum(OPEN[a['name']] + sum(f['inc'] - f['exp'] for f in flows
                                    if f['acc'] == a['name'] and f['d'] < d0) for a in accs), 2))
chk('期末结余金额', N(hb['H5'].value),
    round(sum(OPEN.values()) + sum(f['inc'] - f['exp'] for f in flows if f['d'] <= d1), 2))

print('══ 3. 汇报表 ② 货币资金日报表 ══')
D2R = 8 + N_ACC + 2        # ①数据 8..37、合计 38、空 39、②③标题 40
day = hb.cell(D2R + 1, 2).value
for i, a in enumerate(accs[:5]):
    r = D2R + 3 + i
    prev = round(OPEN[a['name']] + sum(f['inc'] - f['exp'] for f in flows
                                       if f['acc'] == a['name'] and f['d'] < day), 2)
    inc = round(sum(f['inc'] for f in flows if f['acc'] == a['name'] and f['d'] == day), 2)
    exp = round(sum(f['exp'] for f in flows if f['acc'] == a['name'] and f['d'] == day), 2)
    chk(f'{a["name"]} 昨日余额', N(hb.cell(r, 3).value), prev)
    chk(f'{a["name"]} 本日余额', N(hb.cell(r, 6).value), round(prev + inc - exp, 2))

print('══ 4. 汇报表 ③ 多帐户资金汇总表·月度（经营口径）══')
ms = hb.cell(D2R + 1, 9).value
me = (ms.replace(day=28) + datetime.timedelta(days=4)).replace(day=1) - datetime.timedelta(days=1)
inm = [f for f in flows if ms <= f['d'] <= me]
for i, a in enumerate(accs):
    r = D2R + 3 + i
    nm = a['name']
    w = dict(
        j=round(OPEN[nm] + sum(f['inc'] - f['exp'] for f in flows if f['acc'] == nm and f['d'] < ms), 2),
        k=round(sum(f['inc'] for f in inm if f['acc'] == nm and f['attr'] != '内部转账'), 2),
        l=round(sum(f['exp'] for f in inm if f['acc'] == nm and f['attr'] != '内部转账'), 2),
        m=round(sum(f['inc'] for f in inm if f['acc'] == nm and f['attr'] == '内部转账'), 2),
        n=round(sum(f['exp'] for f in inm if f['acc'] == nm and f['attr'] == '内部转账'), 2))
    for j, key in enumerate('jklmn'):
        chk(f'{nm} ③{key}', N(hb.cell(r, 10 + j).value), w[key])
    chk(f'{nm} ③结余', N(hb.cell(r, 15).value),
        round(w['j'] + w['k'] - w['l'] + w['m'] - w['n'], 2))

print('══ 5. 月度汇报表 ══')
for m in range(1, 13):
    r = 4 + m
    ym = YEAR * 100 + m
    fm = [f for f in flows if int(f['d'].strftime('%Y%m')) == ym]
    chk(f'{m}月 收入', N(mr.cell(r, 2).value),
        round(sum(f['inc'] for f in fm if f['attr'] != '内部转账'), 2))
    chk(f'{m}月 支出', N(mr.cell(r, 3).value),
        round(sum(f['exp'] for f in fm if f['attr'] != '内部转账'), 2))
    chk(f'{m}月 内部转入', N(mr.cell(r, 5).value),
        round(sum(f['inc'] for f in fm if f['attr'] == '内部转账'), 2))
    chk(f'{m}月 内部转出', N(mr.cell(r, 6).value),
        round(sum(f['exp'] for f in fm if f['attr'] == '内部转账'), 2))
    eom = (datetime.datetime(YEAR, m, 1).replace(day=28) + datetime.timedelta(days=4)).replace(day=1) \
        - datetime.timedelta(days=1)
    chk(f'{m}月 月末资金总额', N(mr.cell(r, 7).value),
        round(sum(OPEN.values()) + sum(f['inc'] - f['exp'] for f in flows if f['d'] <= eom), 2))

print('══ 6. 收支分类报表 ══')
K0 = 12
names = list(cats)
e = 0
for i, c in enumerate(names):
    r = K0 + i
    for m in range(1, 13):
        ym = YEAR * 100 + m
        sign = -1 if cats[c] == '支出' else 1
        w = round(sign * sum(f['inc'] - f['exp'] for f in flows
                             if f['cat'] == c and int(f['d'].strftime('%Y%m')) == ym), 2)
        if abs(N(kr.cell(r, 2 + m).value) - w) > 0.02:
            e += 1
chk('类别 × 12 个月 逐格', e, 0)
w_inc = round(sum(f['inc'] - f['exp'] for f in flows if f['attr'] == '收入'), 2)
w_exp = round(sum(f['exp'] - f['inc'] for f in flows if f['attr'] == '支出'), 2)
w_trf = round(sum(f['inc'] - f['exp'] for f in flows if f['attr'] == '内部转账'), 2)
chk('收入合计', N(kr['O5'].value), w_inc)
chk('支出合计', N(kr['O6'].value), w_exp)
chk('内部转账净额', N(kr['O7'].value), w_trf)
chk('经营净额', N(kr['O8'].value), round(w_inc - w_exp, 2))

print('══ 7. 收支报表（收支分列 + 未列示补差）══')
IN_0, IN_E, IN_X = 6, 6 + N_ITEM - 1, 6 + N_ITEM
EX_T = IN_X + 2
EX_0, EX_E, EX_X = EX_T + 2, EX_T + 1 + N_ITEM, EX_T + 2 + N_ITEM
for m in (6,):
    cc = 2 + m
    ym = YEAR * 100 + m
    fm = [f for f in flows if int(f['d'].strftime('%Y%m')) == ym and f['attr'] != '内部转账']
    chk(f'{m}月 收入总额', N(sr.cell(4, cc).value), round(sum(f['inc'] for f in fm), 2))
    chk(f'{m}月 支出总额', N(sr.cell(EX_T, cc).value), round(sum(f['exp'] for f in fm), 2))
    chk(f'{m}月 收支盈亏', N(sr.cell(3, cc).value),
        round(sum(f['inc'] for f in fm) - sum(f['exp'] for f in fm), 2))
    # 未列示补差：列示合计 + 未列示 = 总额
    chk(f'{m}月 收入 列示+未列示 = 总额',
        round(sum(N(sr.cell(r, cc).value) for r in range(IN_0, IN_X + 1)), 2),
        N(sr.cell(4, cc).value))
    chk(f'{m}月 支出 列示+未列示 = 总额',
        round(sum(N(sr.cell(r, cc).value) for r in range(EX_0, EX_X + 1)), 2),
        N(sr.cell(EX_T, cc).value))
e = 0
for r in range(IN_0, IN_E + 1):
    nm = sr.cell(r, 2).value
    if not nm:
        continue
    for m in range(1, 13):
        ym = YEAR * 100 + m
        w = round(sum(f['inc'] for f in flows if f['cus'] == nm
                      and int(f['d'].strftime('%Y%m')) == ym and f['attr'] != '内部转账'), 2)
        if abs(N(sr.cell(r, 2 + m).value) - w) > 0.02:
            e += 1
chk('收入区（按客户）逐格', e, 0)
e = 0
for r in range(EX_0, EX_E + 1):
    nm = sr.cell(r, 2).value
    if not nm:
        continue
    for m in range(1, 13):
        ym = YEAR * 100 + m
        w = round(sum(f['exp'] for f in flows if f['cat'] == nm
                      and int(f['d'].strftime('%Y%m')) == ym and f['attr'] != '内部转账'), 2)
        if abs(N(sr.cell(r, 2 + m).value) - w) > 0.02:
            e += 1
chk('支出区（按类别）逐格', e, 0)
chk('收入区预留行数', sum(1 for r in range(IN_0, IN_E + 1)), N_ITEM)
chk('支出区预留行数', sum(1 for r in range(EX_0, EX_E + 1)), N_ITEM)

print('══ 8. 任意时间段收支报表 ══')
q0, q1 = qr['B3'].value, qr['B4'].value
sel = [f for f in flows if q0 <= f['d'] <= q1]
chk('收入总额', N(qr['D3'].value),
    round(sum(f['inc'] for f in sel if f['attr'] != '内部转账'), 2))
chk('支出总额', N(qr['F3'].value),
    round(sum(f['exp'] for f in sel if f['attr'] != '内部转账'), 2))
chk('经营盈亏', N(qr['H3'].value),
    round(sum(f['inc'] for f in sel if f['attr'] != '内部转账')
          - sum(f['exp'] for f in sel if f['attr'] != '内部转账'), 2))
chk('内部转账净额', N(qr['D4'].value),
    round(sum(f['inc'] - f['exp'] for f in sel if f['attr'] == '内部转账'), 2))
Q0, QE, QX = 7, 7 + N_ITEM - 1, 7 + N_ITEM
chk('收入侧 列示+未列示 = 总额',
    round(sum(N(qr.cell(r, 3).value) for r in range(Q0, QX + 1)), 2), N(qr['D3'].value))
chk('支出侧 列示+未列示 = 总额',
    round(sum(N(qr.cell(r, 7).value) for r in range(Q0, QX + 1)), 2), N(qr['F3'].value))

print('══ 9. 核对表 ══')
for r in range(5, 5 + 30):
    if ck.cell(r, 2).value != '勾稽':
        continue
    chk('勾稽 · ' + str(ck.cell(r, 3).value)[:30], round(N(ck.cell(r, 4).value), 2), 0)
print('   待办项：', {str(ck.cell(r, 3).value)[:22]: ck.cell(r, 4).value
                   for r in range(5, 5 + 30)
                   if ck.cell(r, 2).value == '待办' and N(ck.cell(r, 4).value) != 0})

if out is not None:
    print('══ 10. 交付件结构 ══')
    subs = [a['name'] for a in accs if a['name'] in out.sheetnames]
    chk('账户子表张数 = 账户数', len(subs), len(accs))
    for s in subs:
        v = out[s]['A5'].value
        chk(f'{s} 是 FILTER 且套了 IF',
            isinstance(v, ArrayFormula) and 'IF(数据录入!$A$5:$P$5004="","",' in v.text, True)
        chk(f'{s} B2 = 本表账户名', out[s]['B2'].value, s)
    chk('FILTER 没被改小写',
        sum(1 for s in subs if '_xlfn._xlws.FILTER' in out[s]['A5'].value.text), len(subs))
    od = out['数据录入']
    chk('数据录入 A~K 列名跟原表一字不差',
        [od.cell(4, c).value for c in range(1, 12)],
        ['序号', '日期 ★', '公司账户 ★', '收/支类别 ★', '摘要内容', '收入', '支出',
         '账户余额', '客户', '供应商', '月份'])
    for c, nm in ((1, '账户'), (2, '期初余额'), (65, '期间收款'), (67, '期末余额')):
        chk(f'汇报表①第 {c} 列表头', out['汇报表'].cell(7, c).value, nm)
    chk('汇报表③ 有内部转入/内部转出两列',
        [out['汇报表'].cell(8 + N_ACC + 4, c).value for c in (13, 14)], ['内部转入', '内部转出'])

if len(sys.argv) > 3:
    print('══ 11. 跟原模板对照：资料一条不少 ══')
    import openpyxl as _o
    from openpyxl.workbook.external_link.external import ExternalSheetNames as _E
    _oi = _E.__init__
    _E.__init__ = lambda self, sheetName=(), count=None: _oi(self, sheetName)
    src = _o.load_workbook(sys.argv[3], data_only=True)
    sb, sd = src['基础资料'], src['数据录入']
    chk('账户 14 个全在', [a['name'] for a in accs][:14],
        [sb.cell(r, 1).value for r in range(2, 16)])
    chk('账户期初余额跟原表一致', [OPEN[a['name']] for a in accs][:14],
        [N(sd.cell(r, 8).value) for r in range(5, 19)])
    old_exp = [sb.cell(r, 3).value for r in range(2, 15) if sb.cell(r, 3).value]
    chk('原来 12 个支出类别一个没丢', [c for c in old_exp if c in cats], old_exp)
    old_cus = [sb.cell(r, 5).value for r in range(2, 16) if sb.cell(r, 5).value]
    cus = [ba.cell(r, 11).value for r in range(B0, B0 + 200) if ba.cell(r, 11).value]
    chk('原来 13 个客户一个没丢', [c for c in old_cus if c in cus], old_cus)
    old_n = sum(1 for r in range(19, 5000) if sd.cell(r, 3).value)
    chk('原来 23 笔流水都搬过来了（另加 2 笔内部转账示例）', len(flows), old_n + 2)

print(f'\n══ 结果：{ok} 项通过，{bad} 项不通过 ══')
sys.exit(1 if bad else 0)
