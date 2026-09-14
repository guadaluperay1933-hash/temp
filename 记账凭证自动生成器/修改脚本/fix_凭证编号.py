# -*- coding: utf-8 -*-
"""记账凭证自动生成器 · 凭证号改成手工编，自动号挪到右边辅助列

做三件事：
1. 三张凭证页右上角「字第 N 号」中间那个数字留空（打印出来自己手写），
   由【参数】新增的「凭证号是否打印」开关控制，默认「否」。
2. 自动排的号挪到每张凭证右边的辅助列 M/N/O：
   M＝打印第几页　N＝凭证号　O＝完整字号（含共几页第几页）。
3. 三张凭证页固定打印区域 B:I，辅助列在打印区域外，不影响出纸。
顺带把漏掉没隐藏的【助手报销】隐藏了，跟另外两张助手页一致。
"""
import re, sys
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

SRC = sys.argv[1] if len(sys.argv) > 1 else '../参考/原表_2026.9.9.xlsx'
DST = sys.argv[2] if len(sys.argv) > 2 else '../记账凭证自动生成器_2026.9.9.xlsx'

BLOCK = 15                       # 每张凭证 15 行
BLANK = '　' * 4             # 留给手写的四个全角空格
SW = '参数!$B$9'                 # 凭证号是否打印
SHEETS = [('凭证收入', '助手收入', 66), ('凭证回款', '助手回款', 66), ('凭证报销', '助手报销', 121)]

F_LBL = Font(name='微软雅黑', size=9, bold=True, color='1F4E79')
F_TXT = Font(name='微软雅黑', size=9, color='808080')
F_NUM = Font(name='微软雅黑', size=10, bold=True, color='C00000')
CEN = Alignment(horizontal='center', vertical='center')
LEF = Alignment(horizontal='left', vertical='center')

wb = load_workbook(SRC)
log = []

# ---------------- 1) 参数页加开关 ----------------
ws = wb['参数']
for col, src_row, val in (('A', 6, '凭证号是否打印'), ('B', 6, '否'),
                          ('C', 6, '选「否」＝凭证右上角留空，打印出来自己手写号；选「是」＝把自动排的号直接印出来')):
    c = ws[f'{col}9']; s = ws[f'{col}{src_row}']
    c._style = copy(s._style); c.value = val
dv = DataValidation(type='list', formula1='"否,是"', allow_blank=False,
                    showErrorMessage=True, errorStyle='stop', errorTitle='只能选 否 / 是',
                    error='否＝右上角留空自己手写；是＝把自动排的号印出来')
ws.add_data_validation(dv); dv.add('B9')
log.append('参数!A9:C9 新增「凭证号是否打印」开关，默认 否')

# ---------------- 2) 三张凭证页 ----------------
NUMPAT = re.compile(r'&"  字第 "&INDEX\((助手[^!]+)!\$P\$2:\$P\$(\d+),\1!\$R\$(\d+)\)&" 号"')
for vn, hn, last in SHEETS:
    ws = wb[vn]
    nblk = (ws.max_row - 1) // BLOCK + 1
    hit = 0
    # 辅助列表头
    for col, txt, w in (('M', '打印第几页', 11), ('N', '凭证号', 9), ('O', '凭证字号（核对用）', 34)):
        c = ws[f'{col}1']; c.value = txt; c.font = F_LBL; c.alignment = CEN
        ws.column_dimensions[col].width = w
    ws['P1'] = '← 这三列是核对用的辅助列，打印区域已设成 B:I，印不出来'
    ws['P1'].font = F_TXT; ws['P1'].alignment = LEF
    ws.column_dimensions['P'].width = 46

    for b in range(1, ws.max_row + 1, BLOCK):
        k = (b - 1) // BLOCK + 2                       # 对应助手表的行号
        hdr = ws.cell(row=b + 2, column=6)             # F 列，合并 F:I
        if not isinstance(hdr.value, str) or '字第' not in hdr.value:
            continue
        new = NUMPAT.sub(f'&"  字第 "&IF({SW}="是",INDEX(\\1!$P$2:$P$\\2,\\1!$R$\\3),"{BLANK}")&" 号"',
                         hdr.value)
        assert new != hdr.value, f'{vn} F{b+2} 没匹配上：{hdr.value[:80]}'
        hdr.value = new
        hit += 1
        idx = f'INDEX({hn}!$P$2:$P${last},{hn}!$R${k})'
        pg = f'INDEX({hn}!$L$2:$L${last},{hn}!$R${k})'
        m = ws.cell(row=b + 2, column=13)
        m.value = f'=IF({hn}!$R${k}="","",{hn}!$Q${k})'; m.font = F_TXT; m.alignment = CEN
        n = ws.cell(row=b + 2, column=14)
        n.value = f'=IF({hn}!$R${k}="","",{idx})'; n.font = F_NUM; n.alignment = CEN
        o = ws.cell(row=b + 2, column=15)
        o.value = (f'=IF({hn}!$R${k}="","",INDEX({hn}!$C$2:$C${last},{hn}!$R${k})&"  字第 "&{idx}&" 号"'
                   f'&IF({pg}>1,"  (共"&{pg}&"页 第"&{hn}!$S${k}&"页)",""))')
        o.font = F_TXT; o.alignment = LEF
    ws.print_area = f'$B$1:$I${ws.max_row}'
    log.append(f'{vn}: {hit}/{nblk} 张凭证的号改成留空＋辅助列 M/N/O，打印区域固定为 B:I')

# ---------------- 3) 助手报销 补隐藏 ----------------
if wb['助手报销'].sheet_state != 'hidden':
    wb['助手报销'].sheet_state = 'hidden'
    log.append('助手报销 页签补隐藏（原来漏了，另外两张助手页都是隐藏的）')

# ---------------- 4) 说明页补一段 ----------------
ws = wb['说明']
src = ws['B10']
NOTE = [
    '',
    '【凭证号自己手工编】',
    '   凭证右上角的“字第 ＿＿ 号”已经留空，打印出来自己按顺序手写；',
    '   自动排的号挪到了每张凭证右边的辅助列（M / N / O 三列）：',
    '      M＝这是打印的第几页　　N＝凭证号（原来印在“字第”后面的那个数）　　O＝完整字号（含共几页第几页）',
    '   核对办法：打出来的第 7 页，就到 M 列找 7，右边 N 列写着 5，这页就手写“第 5 号”。',
    '   三张凭证页的打印区域已固定为 B:I 列，辅助列在打印区域外，印不出来、也不占版面。',
    '   想改回“把号直接印出来”，到 参数 页把“凭证号是否打印”选成 是 就行。',
]
r0 = 63
for i, t in enumerate(NOTE):
    c = ws.cell(row=r0 + i, column=2)
    c._style = copy(src._style); c.value = t or None
    ws.row_dimensions[r0 + i].height = ws.row_dimensions[10].height
log.append(f'说明 B{r0}:B{r0+len(NOTE)-1} 补了「凭证号自己手工编」一段')

wb.save(DST)
print('\n'.join('· ' + x for x in log))
print(f'\n已生成 {DST}')
