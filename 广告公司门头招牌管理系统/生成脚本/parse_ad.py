# -*- coding: utf-8 -*-
"""把散在 22 张表里的原始数据抽出来，清洗成一份 ad_data.json。

原表有两种项目表版式：
  std      : 摘要在 M 列，金额在 T 列（13 张）
  shifted  : 少了「序号」列，整体左移一格，摘要在 L、金额在 S（4 张）
日记帐前半段（3~4 月）人工分过类，后半段是微信账单原样粘贴，科目/项目全空，
这部分按摘要关键字自动归类，认不出来的一律标「待分类」，不猜。
"""
import json, re, os, sys, datetime, collections
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else '/root/.claude/uploads/9101e7cb-4f82-53a7-bbe8-41b0f93b12b0/b5653653-1111.xlsx'
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ad_data.json')

# ── 项目名归并：同一个项目在原表里有好几种叫法 ────────────────────────────
ALIAS = {
    '兰溪牛肉面': '兰溪手擀面', '津华钢铁': '德清津华钢铁', '巴比贝克': '芭比贝克',
    '芭比贝克1': '芭比贝克', '汕味杏花': '汕味杏花室内项目', '双林面馆': '双林面馆-滨河路',
    '螺有名': '螺有名灯箱门头', '手工鲜饺店': '饺子店', '玖熠装饰': '九熠装饰',
    '塞港文化': '赛港文化', '坚蛋健身房': '坚蛋健身', '坚蛋健身炳秀': '坚蛋健身',
    '万达每甜': '每甜', '每甜女装': '每甜', '铁柱烧烤': '铁柱烧烤店', '铁柱烧烤3m膜': '铁柱烧烤店',
    '织里店项目': '织里店', '织里项目': '织里店', '校服基地软膜': '校服基地',
    '定金-梓灿': '梓灿软装', '君澜校服': '君澜服饰', '睿芯行-个人': '睿芯行',
    '南浔古镇美食街': '南浔美食城', '万达kkv': '万达kvv', 'kkv': '万达kvv',
    '物料-万达kvv': '万达kvv', 'vip宠物店': 'vip宠物乐园', '大东北烧烤': '东北烧烤大转盘店',
    '钟管熟食店': '钟管熟食店', '德清钟管菜市场': '德清钟管菜市场',
}
# 这些不是项目，是费用科目 / 往来对象，落到项目列上属于串列
NOT_PROJECT = {
    '工资', '往来款', '公司', '办公用品', '团建费', '工作机话费', '分类？', '项目？',
    '报销费用', '付-付久华', '章凡', '凌宇借款', '徐可欣借款', '应收押金40元', '合计',
    '物料', '项目其他费用类', '项目？-工资', '赵总项目', '力高',
}
SUPPLIER_SUFFIX = ('炳秀', '钮')          # 「织里店-炳秀」= 项目 织里店 + 供应商 炳秀

def norm_proj(raw):
    """返回 (项目名, 从名字里剥出来的供应商)。不是项目就返回 ('', '')。"""
    s = str(raw or '').strip()
    if not s or s in NOT_PROJECT:
        return '', ''
    supp = ''
    for sfx in SUPPLIER_SUFFIX:
        if s.endswith('-' + sfx) or s.endswith(sfx) and '-' in s:
            if s.endswith('-' + sfx):
                s, supp = s[:-len(sfx) - 1], sfx
                break
    s = s.strip()
    if s.startswith('项目采购-') or s.startswith('定金-') and s not in ALIAS:
        if s.startswith('项目采购-'):
            return '', supp
    s = ALIAS.get(s, s)
    if s in NOT_PROJECT:
        return '', supp
    return s, supp

# ── 微信账单原样粘贴段的关键字归类 ──────────────────────────────────────
KEYWORD_RULES = [
    (('铝单板', '铝方通', '铝板', '灯条', '电线', '型材', '方管', '发光字', '软膜', '亚克力',
      '雪弗板', '灯箱', '蚀刻', '标识标牌', '印刷', '写真', '喷绘', '背胶', '不锈钢', '钢材',
      '镀锌', '五金', '结构胶', '玻璃胶', '刻字', '激光切割', 'UV', 'uv'), '主材采购'),
    (('工资', '老姚'), '安装工资'),
    (('报销',), '其他管理费用'),
    (('设计',), '设计费'),
    (('视频', '代剪', '剪辑'), '广告推广'),
    (('炳秀', '货拉拉', '顺丰', '京邦达', '德邦', '中通', '圆通', '申通', '韵达', '极兔', '快递', '闪送'), '运费/快递费'),
    (('滴滴', '加油', '停车', '高速', '过路', '中石化', '中石油', '洗车', '车险', '出行'), '车费/油费/过路费'),
    (('巨量', '抖音', '推广', '本地推', '广告投放'), '广告推广'),
    (('中国移动', '中国电信', '中国联通', '话费', '流量'), '通讯费'),
    (('美团', '古茗', '蜜雪', '瑞幸', '星巴克', '饿了么', '餐饮', 'food'), '业务招待/餐费'),
    (('拼多多', '淘宝', '天猫', '1688', '京东', '阿里巴巴'), '网购材料耗材'),
    (('手续费', '服务费收益', '账户服务'), '银行手续费'),
    (('提现', '零钱'), '账户互转'),
]
def guess_cat(memo):
    m = str(memo or '')
    for keys, cat in KEYWORD_RULES:
        for k in keys:
            if k in m:
                return cat
    return '待分类'

def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0

def as_date(v):
    """原表日期有 datetime、Excel 序列号、'3.11'、'26.2.27' 好几种写法。"""
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, datetime.date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, (int, float)) and 40000 < v < 60000:
        return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=int(v))).strftime('%Y-%m-%d')
    s = str(v or '').strip()
    m = re.match(r'^(?:(\d{2,4})\.)?(\d{1,2})\.(\d{1,2})$', s)
    if m:
        y = m.group(1)
        y = 2026 if not y else (2000 + int(y) if int(y) < 100 else int(y))
        try:
            return datetime.date(y, int(m.group(2)), int(m.group(3))).strftime('%Y-%m-%d')
        except ValueError:
            return ''
    return ''

# ══════════════════════════════════════════════════════════════════════
wb = openpyxl.load_workbook(SRC, data_only=True)
wbf = openpyxl.load_workbook(SRC, data_only=False)

projects = collections.OrderedDict()   # 项目名 -> dict
purchases, labour, orders, journal = [], [], [], []
suppliers, persons = collections.Counter(), collections.Counter()
notes = []

def touch(name, **kw):
    if not name:
        return None
    p = projects.setdefault(name, dict(
        name=name, customer=name, ptype='门头招牌', owner='', order_date='',
        done_date='', contract=0.0, opening_recv=0.0, memo='', src=''))
    for k, v in kw.items():
        if v not in (None, '', 0, 0.0) and not p.get(k):
            p[k] = v
    return p

# ── 1. 17 张项目表 ────────────────────────────────────────────────────
CAT2KIND = {
    '主材': ('采购', '主材'), '辅材': ('采购', '辅材'),
    '人工': ('人工', '安装工时'),
    '运费/快递费': ('人工', '运费/搬运'), '运费': ('人工', '运费/搬运'),
    '交通费/过路费': ('人工', '吊车/租车'), '交通费': ('人工', '吊车/租车'),
    '吊车费': ('人工', '吊车/租车'), '小车费': ('人工', '吊车/租车'),
    '其它补助': ('人工', '其他人工'),
}
LABOUR_WORD = {'设计费': '设计费', '安装工资': '安装工时', '加班': '加班费',
               '其它小工工资': '外叫小工', '制作': '制作工时'}

sheet_re = re.compile(r'^(\d+)-(.+)$')
for sn in wb.sheetnames:
    m = sheet_re.match(sn)
    if not m:
        continue
    ws, wsf = wb[sn], wbf[sn]
    pname = ALIAS.get(m.group(2).strip(), m.group(2).strip())
    # 版式判定：摘要落在 L 还是 M
    memo_col = 13
    for c in (12, 13, 14):
        if ws.cell(3, c).value == '摘要' or ws.cell(4, c).value == '摘要':
            memo_col = c
            break
    off = memo_col - 13                      # shifted 版式 off = -1
    C = dict(memo=memo_col, supp=14 + off, pay=15 + off, payer=16 + off,
             unit=17 + off, qty=18 + off, price=19 + off, amt=20 + off)

    first_date, contract, recv = '', 0.0, 0.0
    for r in range(5, 46):
        f_e, f_f = wsf.cell(r, 5).value, wsf.cell(r, 6).value
        if isinstance(f_e, str) and f_e.startswith('='):
            continue
        e, fv = num(ws.cell(r, 5).value), num(ws.cell(r, 6).value)
        if isinstance(f_f, str) and f_f.startswith('='):
            fv = 0.0
        contract += e
        recv += fv
        d = as_date(ws.cell(r, 1).value)
        if d and not first_date:
            first_date = d
    touch(pname, order_date=first_date, contract=round(contract, 2), src=sn)
    projects[pname]['contract'] = round(contract, 2)
    projects[pname]['opening_recv'] = round(recv, 2)

    cur = ''
    for r in range(5, 53):
        k = str(ws.cell(r, 11).value or '').strip()
        if k and not k.startswith('合计') and not k.startswith('总合计'):
            cur = k
        memo = str(ws.cell(r, C['memo']).value or '').strip()
        if memo.startswith('合计') or memo.startswith('总合计'):
            continue
        fa = wsf.cell(r, C['amt']).value
        amt = num(ws.cell(r, C['amt']).value)
        if isinstance(fa, str) and fa.upper().startswith('=SUM('):
            continue
        qty, price = num(ws.cell(r, C['qty']).value), num(ws.cell(r, C['price']).value)
        if round(amt, 2) == 0:
            continue      # 原表预留的空行，只有品名没有金额
        supp = str(ws.cell(r, C['supp']).value or '').strip()
        if supp in ('项目', '公司', '合计', '-', '／', '/'):
            supp = ''
        payer = str(ws.cell(r, C['payer']).value or '').strip()
        unit = str(ws.cell(r, C['unit']).value or '').strip()
        kind, sub = CAT2KIND.get(cur, ('采购', '其他材料'))
        if kind == '人工':
            for w, t in LABOUR_WORD.items():
                if w in memo:
                    sub = t
                    break
        if supp:
            suppliers[supp] += 1
        if payer:
            persons[payer] += 1
        note = '原表·' + (cur or '')
        if abs(round(qty * price, 2) - round(amt, 2)) > 0.004:
            # 原表的「金额」不等于 数量×单价（当时是手打进去的），以金额为准
            if qty or price:
                note += '　原表数量%s×单价%s，与金额对不上，已按金额入账' % (qty, price)
            qty, price = 1.0, round(amt, 2)
        row = dict(date=first_date, project=pname, memo=memo or sub, spec='', unit=unit,
                   qty=qty, price=price, amount=round(amt, 2), supplier=supp, payer=payer,
                   note=note)
        if kind == '采购':
            row['mcat'] = sub
            purchases.append(row)
        else:
            row['wcat'] = sub
            row['person'] = supp or payer
            labour.append(row)

# ── 2. 东尼电子 / Sheet1 的制作清单 → 订单明细 ─────────────────────────
for sn, pname in (('东尼电子', '东尼电子26.3-4'), ('Sheet1', '东尼电子26.5')):
    if sn not in wb.sheetnames:
        continue
    ws, wsf = wb[sn], wbf[sn]
    dates = []
    for r in range(3, ws.max_row + 1):
        if isinstance(wsf.cell(r, 7).value, str) and wsf.cell(r, 7).value.startswith('='):
            continue
        nm = str(ws.cell(r, 2).value or '').strip()
        if not nm or nm in ('合计', '税票'):
            continue
        d = as_date(ws.cell(r, 1).value)
        if d:
            dates.append(d)
        orders.append(dict(date=d, project=pname, item=nm,
                           spec=str(ws.cell(r, 3).value or '').strip(),
                           qty=num(ws.cell(r, 4).value),
                           unit=str(ws.cell(r, 5).value or '').strip(),
                           price=num(ws.cell(r, 6).value),
                           amount=round(num(ws.cell(r, 7).value), 2),
                           note=str(ws.cell(r, 9).value or '').strip()))
    p = touch(pname, order_date=min(dates) if dates else '', ptype='广告物料')
    p['customer'] = '东尼新能源'
    p['done_date'] = max(dates) if dates else ''
    p['memo'] = '原「%s」制作清单' % sn
    # 原表 H 列的成本是整批估的，按批次挂到「其他材料」一行，不摊到每个品名
    cost = sum(num(ws.cell(r, 8).value) for r in range(3, ws.max_row + 1)
               if not (isinstance(wsf.cell(r, 8).value, str) and str(wsf.cell(r, 8).value).startswith('=')))
    if cost:
        purchases.append(dict(date=max(dates) if dates else '', project=pname, mcat='其他材料',
                              memo='原表整批成本(未分品名)', spec='', unit='批', qty=1,
                              price=round(cost, 2), amount=round(cost, 2), supplier='',
                              payer='', note='原表 H 列'))

# ── 3. 日记帐 940 行 ─────────────────────────────────────────────────
SUBJ2CAT = {
    ('主营业务收入', '收入'): '项目收款',
    ('其它业务收入', '收入'): '零星物料销售',
    ('其他业务支出', '支出'): '其他直接费用',
    ('主营业务成本', '支出'): '主材采购',
    ('管理费用', '支出'): '其他管理费用',
    ('管理费用', '收入'): '其他管理费用',
    ('财务费用', '收入'): '利息收入',
    ('财务费用', '支出'): '银行手续费',
    ('其他应收款', '支出'): '借出款',
    ('其他应收款', '收入'): '收回借出款',
    ('其他应付款', '支出'): '归还借入款',
    ('其他应付款', '收入'): '借入款',
}
FINE = [  # 在「主营业务成本 / 管理费用」内部再按摘要细分
    (('工资', '预支', '章凡', '老姚'), '安装工资'),
    (('设计',), '设计费'),
    (('炳秀', '货拉拉', '顺丰', '京邦达', '快递', '闪送'), '运费/快递费'),
    (('加油', '停车', '高速', '过路', '滴滴', '312'), '车费/油费/过路费'),
    (('话费', '移动', '电信', '联通'), '通讯费'),
    (('团建',), '团建费'),
    (('办公', '昵享网'), '办公用品'),
    (('餐', '饭', '午餐', '美团', '古茗'), '业务招待/餐费'),
    (('巨量', '推广', '抖音'), '广告推广'),
]
jws = wb['日记帐']
last_row = 0
for r in range(2, jws.max_row + 1):
    if any(jws.cell(r, c).value not in (None, '') for c in (1, 3, 4, 5, 8)):
        last_row = r
n_auto = n_wait = 0
RUN = [0.0, 0.0]      # [收支累计, 已补平的差额]
gaps = []
for r in range(2, last_row + 1):
    d = as_date(jws.cell(r, 1).value)
    direction = str(jws.cell(r, 3).value or '').strip()
    subj = str(jws.cell(r, 4).value or '').strip()
    praw = str(jws.cell(r, 5).value or '').strip()
    memo = str(jws.cell(r, 8).value or '').strip()
    inc, out = num(jws.cell(r, 9).value), num(jws.cell(r, 10).value)
    note = str(jws.cell(r, 12).value or '').strip()
    if praw == '合计' or praw.endswith('合计') or memo.endswith('合计') or '本年累计' in memo:
        continue
    if inc == 0 and out == 0:
        continue          # 没金额的行是原表的分隔行 / 备注行
    if not direction:
        direction = '收入' if inc > out else '支出'
    proj, supp = norm_proj(praw)
    cat = SUBJ2CAT.get((subj, direction), '')
    if cat and subj in ('主营业务成本', '管理费用'):
        for keys, c2 in FINE:
            if any(k in memo or k in praw for k in keys):
                cat = c2
                break
    if not cat:
        cat = guess_cat(memo)
        if cat == '待分类':
            n_wait += 1
        else:
            n_auto += 1
    # 项目列串了费用科目的，按科目补类别
    if not proj and praw in ('工资', '项目？-工资'):
        cat = '安装工资'
    elif not proj and praw in ('办公用品',):
        cat = '办公用品'
    elif not proj and praw in ('团建费',):
        cat = '团建费'
    elif not proj and praw in ('工作机话费',):
        cat = '通讯费'
    elif praw in ('往来款', '付-付久华', '章凡', '凌宇借款', '徐可欣借款'):
        cat = '借出款' if direction == '支出' else '收回借出款'
        supp = supp or praw.replace('付-', '')
    if proj:
        touch(proj, order_date=d)
        if cat == '项目收款':
            projects[proj].setdefault('_jrecv', 0.0)
            projects[proj]['_jrecv'] = round(projects[proj].get('_jrecv', 0.0) + inc - out, 2)
    party = supp or ''
    if not party and memo:
        mm = re.match(r'^(?:付|收)?([一-龥A-Za-z0-9]{2,12}?)(?:配送|公司|平台商户)?(?:\s|\(|（|$)', memo)
        if mm and cat in ('运费/快递费', '主材采购', '辅材采购', '网购材料耗材', '外发加工'):
            party = mm.group(1)
    if party:
        suppliers[party] += 1
    # 原表 K 列自己写的「结余」和收支累计对不上时，说明那里漏记了一笔。
    # 不猜是什么钱，原样补一行标出来，这样余额能跟他们手机里的对上。
    stated = jws.cell(r, 11).value
    if isinstance(stated, (int, float)):
        gap = round(stated - (RUN[0] + inc - out), 2)
        if abs(gap - RUN[1]) > 0.004:
            delta = round(gap - RUN[1], 2)
            journal.append(dict(
                date=d, account='微信1', direction='收入' if delta > 0 else '支出',
                cat='待分类', project='', party='',
                memo='※原表这里余额跳了 %.2f，但没有对应的收支行' % delta,
                income=abs(delta) if delta > 0 else 0.0,
                outgo=abs(delta) if delta < 0 else 0.0,
                note='导入时按原表「结余」列补平，请核实这笔钱是什么再改类别',
                raw_subject='', raw_project=''))
            gaps.append((r, d, delta))
            RUN[1] = gap
        RUN[0] = round(RUN[0] + inc - out, 2)
    else:
        RUN[0] = round(RUN[0] + inc - out, 2)
    journal.append(dict(date=d, account='微信1', direction=direction, cat=cat,
                        project=proj, party=party, memo=memo or praw or subj,
                        income=round(inc, 2), outgo=round(out, 2), note=note,
                        raw_subject=subj, raw_project=praw))
notes.append('日记帐共取 %d 行；关键字自动归类 %d 行，认不出留「待分类」%d 行。'
             % (len(journal), n_auto, n_wait))
for r, d, delta in gaps:
    notes.append('原表第 %d 行（%s）余额跳了 %.2f 却没有收支行，已补一行标出来。' % (r, d or '无日期', delta))

# ── 3.5 每个项目最后一次有动静的日期 ────────────────────────────────
for src in (journal,):
    for x in src:
        if x['project'] and x['date']:
            p = projects.get(x['project'])
            if p and x['date'] > p.get('last_date', ''):
                p['last_date'] = x['date']
for p in projects.values():
    p.setdefault('last_date', p.get('order_date', ''))
    if p.get('done_date'):
        continue
    # 只有「收齐了钱」才敢认定已完工——没收齐的一律留空，由用户自己填
    if p.get('contract', 0) > 0 and p.get('opening_recv', 0) >= p['contract'] - 0.01:
        p['done_date'] = max(p['last_date'] or '', p.get('order_date', '') or '')

# ── 4. 期初已收：项目表的已收款 减掉 日记账里已经有的收款，避免重复 ──────
for p in projects.values():
    jr = p.pop('_jrecv', 0.0)
    if p.get('opening_recv'):
        p['opening_recv'] = round(max(0.0, p['opening_recv'] - jr), 2)

# ── 5. 客户 / 供应商 / 人员名单 ───────────────────────────────────────
DROP_PARTY = {'', '合计', '公司', '往来款', '项目', '收入', '支出'}
supplier_list = [s for s, n in suppliers.most_common()
                 if s not in DROP_PARTY and len(s) >= 2 and '借款' not in s][:80]
person_list = [s for s, n in persons.most_common() if s not in DROP_PARTY][:30]
for extra in ('凌宇', '杨德', '章凡', '徐可欣', '老姚', '小章', '小姜'):
    if extra not in person_list:
        person_list.append(extra)
customer_list = []
for p in projects.values():
    if p['customer'] not in customer_list:
        customer_list.append(p['customer'])

data = dict(
    projects=list(projects.values()), purchases=purchases, labour=labour,
    orders=orders, journal=journal,
    suppliers=supplier_list, persons=person_list, customers=customer_list, notes=notes)

with open(OUT, 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=1)

print('项目 %d  采购行 %d  人工行 %d  订单明细 %d  日记账 %d' %
      (len(projects), len(purchases), len(labour), len(orders), len(journal)))
print('供应商 %d  客户 %d  人员 %d' % (len(supplier_list), len(customer_list), len(person_list)))
print('采购金额 %.2f  人工金额 %.2f  订单金额 %.2f' %
      (sum(x['amount'] for x in purchases), sum(x['amount'] for x in labour),
       sum(x['amount'] for x in orders)))
print('日记账 收 %.2f  支 %.2f' % (sum(x['income'] for x in journal), sum(x['outgo'] for x in journal)))
for n in notes:
    print('·', n)
