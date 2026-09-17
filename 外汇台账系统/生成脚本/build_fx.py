# -*- coding: utf-8 -*-
"""生成《外汇资金台账.xlsx》——多币种（人民币本位）资金台账模板。

一张【流水录入】混合录入，其余全部自动：
  · 银行子表按账户名自动拆（FILTER 动态数组，跟资金台账一个路子）
  · 汇率按「交易日期所在月」自动匹配；当月没填就往前找最近一个填过的月份
  · 本位币口径分两套，各自算各自的，不混：
      账面人民币成本 = 期初成本 + Σ(每笔原币 × 该笔应用汇率)
      期末重估人民币 = 期末原币余额 × 期末汇率
      汇兑损益      = 重估 − 账面
  · 内部转账单独一类，不计入收入/支出合计（原模板把它算进去了，两头都虚增）

跑法：python3 build_fx.py
"""
import os, sys, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, '外汇资金台账.xlsx')

YEAR = 2026
N_ACC, N_ROW, N_CUR, N_CAT, N_PRJ, N_PTY = 30, 2000, 8, 30, 40, 200
A0, R0, SUB = 11, 6, 400          # 参数各清单起始行 / 流水起始行 / 子表预留行
AE = A0 + N_ACC - 1               # 账户档案末行
RE_ = R0 + N_ROW - 1              # 流水末行

# ── 样式 ──────────────────────────────────────────────────────────────
C_MAIN, C_SUB, C_RPT, C_BASE, C_CHK = '1F4E79', '806000', '375623', '7030A0', 'C00000'
F_IN = PatternFill('solid', fgColor='FFFBEA')     # 要人填
F_AUTO = PatternFill('solid', fgColor='F2F2F2')   # 公式算
F_SUM = PatternFill('solid', fgColor='FFF2CC')
FT_IN = Font(name='微软雅黑', size=10)
FT_AUTO = Font(name='微软雅黑', size=10, color='7F7F7F')
FT_HDR = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
FT_TITLE = Font(name='微软雅黑', size=15, bold=True, color='FFFFFF')
FT_SUM = Font(name='微软雅黑', size=10, bold=True, color='C00000')
FT_TIP = Font(name='微软雅黑', size=9, color='8B5E00')
THIN = Side('thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CEN = Alignment('center', 'center', wrap_text=True)
LEFT = Alignment('left', 'center')
MONEY, NUM0, RATE, DATEF, PCT = '#,##0.00', '#,##0', '0.0000', 'yyyy-mm-dd', '0.0%'
DATEQ = 'yyyy-mm-dd;;;@'          # 自动区专用：FILTER 取到空格子会给 0，这个格式不把 0 画成 1900-01-00

MONTHS = ['%d月' % m for m in range(1, 13)]

# ══════════════════════════════════════════════════════════════════════
# 示例数据（第一次打开能直接看到效果；正式用之前整段删掉即可）
# ══════════════════════════════════════════════════════════════════════
CURS = [('CNY', '人民币'), ('USD', '美元'), ('HKD', '港币'), ('EUR', '欧元'),
        ('JPY', '日元'), ('GBP', '英镑'), ('SGD', '新加坡元'), ('AUD', '澳元')]

ACCOUNTS = [
    ('工商银行', '工商银行基本户-8888', '1202...8888', 'CNY', '基本户', 500000, None),
    ('中国银行', '中国银行美元结算户-6688', '4567...6688', 'USD', '外币结算户', 800000, 7.15),
    ('汇丰银行', '汇丰银行港币结算户-5566', '0048...5566', 'HKD', '外币结算户', 300000, 0.915),
    ('招商银行', '招商银行一般户-2388', '7712...2388', 'CNY', '一般户', 200000, None),
    ('汇丰银行', '汇丰银行美元结算户-7788', '0048...7788', 'USD', '外币结算户', 500000, 7.15),
    ('中国银行', '中国银行欧元结算户-9900', '4567...9900', 'EUR', '外币结算户', 120000, None),
]

# 类别 → 属性（收入 / 支出 / 内部转账）。内部转账不进收支合计。
CATS = [('销售回款', '收入'), ('其他业务收入', '收入'), ('投资款', '收入'), ('借款', '收入'),
        ('利息收入', '收入'), ('退款收回', '收入'),
        ('采购付款', '支出'), ('工资', '支出'), ('运营费用', '支出'), ('税费', '支出'),
        ('手续费', '支出'), ('利息支出', '支出'), ('还款', '支出'), ('资产购买', '支出'),
        ('股利分配', '支出'),
        ('内部转账', '内部转账'), ('购汇/结汇', '内部转账')]

PROJECTS = ['项目A', '出口业务', '进口业务', '香港业务', '欧洲业务', '融资', '行政', '税务',
            '海外服务', '内部调拨', '其他业务']

PARTIES = ['华东客户A', 'Overseas Client A', 'Overseas Client C', '香港客户B', '欧洲客户D',
           '香港供应商A', 'Overseas Supplier B', '欧洲供应商E', '员工薪酬', '税务机关',
           '银行借款', '股东A', 'Service Provider', '香港办公室', '其他客户']

# 记账汇率：只有 8 月是原模板里给的（USD 7.18 / HKD 0.92）。其余月份、以及欧元，
# 都空着——空着不是坏了，是等你填；核对表会点名提醒。
RATES = {(YEAR, 8): {'USD': 7.18, 'HKD': 0.92}}

# 示例流水：日期, 账户, 类别, 收入, 支出, 实际汇率, 摘要, 项目, 往来单位
FLOWS = [
    ('2026-08-03', '工商银行基本户-8888', '销售回款', 120000, None, None, '销售回款', '项目A', '华东客户A'),
    ('2026-08-04', '中国银行美元结算户-6688', '销售回款', 25000, None, None, '美元货款', '出口业务', 'Overseas Client A'),
    ('2026-08-05', '汇丰银行港币结算户-5566', '采购付款', None, 4500, None, '港币采购付款', '进口业务', '香港供应商A'),
    ('2026-08-06', '招商银行一般户-2388', '工资', None, 30000, None, '工资发放', '行政', '员工薪酬'),
    ('2026-08-07', '汇丰银行美元结算户-7788', '投资款', 12000, None, None, '美元投资款', '融资', '股东A'),
    ('2026-08-08', '工商银行基本户-8888', '税费', None, 18000, None, '税费缴纳', '税务', '税务机关'),
    ('2026-08-10', '中国银行美元结算户-6688', '采购付款', None, 8000, 7.16, '美元采购款（银行实际成交价）', '进口业务', 'Overseas Supplier B'),
    ('2026-08-11', '汇丰银行港币结算户-5566', '销售回款', 80000, None, None, '港币销售回款', '香港业务', '香港客户B'),
    ('2026-08-12', '招商银行一般户-2388', '借款', 60000, None, None, '短期借款到账', '融资', '银行借款'),
    ('2026-08-13', '汇丰银行美元结算户-7788', '运营费用', None, 5000, None, '美元服务费', '海外服务', 'Service Provider'),
    ('2026-08-14', '工商银行基本户-8888', '内部转账', None, 50000, None, '转至招商银行', '内部调拨', '招商银行一般户-2388'),
    ('2026-08-14', '招商银行一般户-2388', '内部转账', 50000, None, None, '收到工商银行转入', '内部调拨', '工商银行基本户-8888'),
    ('2026-08-18', '中国银行美元结算户-6688', '销售回款', 15000, None, None, '第二笔美元回款', '出口业务', 'Overseas Client C'),
    ('2026-08-19', '汇丰银行港币结算户-5566', '运营费用', None, 20000, None, '香港运营费', '行政', '香港办公室'),
    ('2026-08-20', '中国银行欧元结算户-9900', '销售回款', 30000, None, None, '欧元货款（欧元汇率还没填，看核对列）', '欧洲业务', '欧洲客户D'),
    ('2026-08-21', '工商银行基本户-8888', '其他业务收入', 35000, None, None, '其他经营收入', '其他业务', '其他客户'),
    ('2026-08-25', '中国银行美元结算户-6688', '购汇/结汇', None, 20000, 7.17, '结汇：美元转出', '内部调拨', '工商银行基本户-8888'),
    ('2026-08-25', '工商银行基本户-8888', '购汇/结汇', 143400, None, None, '结汇：人民币到账', '内部调拨', '中国银行美元结算户-6688'),
]

# ══════════════════════════════════════════════════════════════════════
wb = openpyxl.Workbook()
wb.remove(wb.active)
NAMES = {}


def newsheet(name, title, tip, color, width, merge2=True):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = color
    ws['A1'] = title
    ws['A1'].font, ws['A1'].fill, ws['A1'].alignment = FT_TITLE, PatternFill('solid', fgColor='1F3864'), LEFT
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.row_dimensions[2].height = 24
    if merge2:
        ws['A2'] = tip
        ws['A2'].font, ws['A2'].fill = FT_TIP, PatternFill('solid', fgColor='FFF7E6')
        ws['A2'].alignment = LEFT
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws.sheet_view.showGridLines = False
    return ws


def put(ws, addr, v, fmt=None, font=None, fill=None, align=None, border=True):
    c = ws[addr]
    c.value = v
    if fmt: c.number_format = fmt
    c.font = font or FT_IN
    if fill: c.fill = fill
    c.alignment = align or CEN
    if border: c.border = BOX
    return c


def hdr(ws, row, col0, cols, color):
    fill = PatternFill('solid', fgColor=color)
    for i, (t, w, fmt) in enumerate(cols, col0):
        c = ws.cell(row, i, t)
        c.font, c.fill, c.alignment, c.border = FT_HDR, fill, CEN, BOX
        if w: ws.column_dimensions[CL(i)].width = w
    ws.row_dimensions[row].height = 30


def band(ws, r0, r1, col0, cols, auto=False):
    for r in range(r0, r1 + 1):
        for i, (t, w, fmt) in enumerate(cols, col0):
            c = ws.cell(r, i)
            c.border = BOX
            c.font = FT_AUTO if auto else FT_IN
            c.fill = F_AUTO if auto else F_IN
            c.alignment = CEN
            if fmt: c.number_format = fmt


def title_bar(ws, row, col0, ncol, text, color=None):
    put(ws, f'{CL(col0)}{row}', text,
        font=Font(name='微软雅黑', size=11, bold=True, color='FFFFFF'),
        fill=PatternFill('solid', fgColor=color or C_RPT), align=LEFT)
    ws.merge_cells(start_row=row, start_column=col0, end_row=row, end_column=col0 + ncol - 1)
    ws.row_dimensions[row].height = 22


def dv(ws, f1, rng, stop=False):
    d = DataValidation(type='list', formula1=f1, allow_blank=True, showErrorMessage=True,
                       errorStyle='stop' if stop else 'warning', errorTitle='不在清单里',
                       error='先去【参数设置】把它加上，或者选一个已有的。')
    ws.add_data_validation(d)
    d.add(rng)


def nb(rng):
    """FILTER 取数区先过一道：源表空格子会被当成 0 带回来，日期列就成了 1900-01-00。"""
    return f'IF({rng}="","",{rng})'


# 清单里第 i 行那一格的**直接地址**。
# 为什么不用 INDEX(名字,i)：INDEX 落到空格子，Excel 返回的是数字 0 而不是空文本，
# 于是外面的 IF(...="","",...) 判不出来，几十行没用上的空账户会跟着算出一堆 0。
# 直接写单元格地址就没这个毛病 —— 空格子和 "" 比较为真。
PICK = {'账户名': ('C', A0), '账户银行': ('B', A0), '账户币种': ('E', A0),
        '账户期初': ('G', A0), '账户期初本位': ('I', A0),
        '币种表': ('L', 11), '类别表': ('O', 11), '类别属性': ('P', 11),
        '项目表': ('R', 11), '单位表': ('T', 11)}


K_0 = '账户名'
K_1 = '账户银行'
K_2 = '账户币种'
K_3 = '账户期初'
K_4 = '账户期初本位'
K_5 = '币种表'
K_6 = '类别表'
K_7 = '类别属性'
K_8 = '项目表'
K_9 = '单位表'


def pick(key, i):
    col, r0 = PICK[key]
    return f'参数设置!${col}${r0 + i - 1}'

def rate_at(d, cur, eom=False):
    """取汇率：按「交易日期所在月」匹配；当月没填就往前找最近一个填过的月份。

    LOOKUP(2,1/(条件),取值列) 是「最后一个满足条件的值」的老写法 ——
    1/FALSE 得 #DIV/0!，LOOKUP 会跳过错误，于是取到最后一个 TRUE 的位置。
    原模板用的是 INDEX+MATCH 精确匹配当月，当月没填就直接空白，
    整张表跟着空 —— 12 个月里 11 个月没填，等于这模板只有 8 月能用。"""
    cb = f'INDEX(记账汇率区,0,MATCH({cur},汇率币种头,0))'
    book = f'LOOKUP(2,1/((汇率月份<={d})*({cb}<>"")),{cb})'
    if not eom:
        return f'IF({cur}=本位币,1,IFERROR({book},""))'
    ce = f'INDEX(月末汇率区,0,MATCH({cur},汇率币种头,0))'
    end = f'LOOKUP(2,1/((汇率月份<={d})*({ce}<>"")),{ce})'
    return f'IF({cur}=本位币,1,IFERROR({end},IFERROR({book},"")))'


# ══════════════════════════════════════════════════════════════════════
# 1  参数设置
# ══════════════════════════════════════════════════════════════════════
ws = newsheet('参数设置', '参数设置 · 账户、币种、类别、项目都在这里维护',
              '淡黄色格子是要你填的，灰色是公式算的。账户名称必须唯一，'
              '建议写成「银行 + 用途 + 账号尾号」。这一页改了，后面所有表跟着变。', C_BASE, 21)

BASIC = [('本位币', 'CNY', '所有汇报表都折算成它。改成别的币种，整套跟着走。'),
         ('管理年度', YEAR, '【月度汇率】和各张年报按这个年度展开。'),
         ('报表月份', datetime.datetime(YEAR, 8, 1), '汇报表默认看哪个月。填该月 1 号。'),
         ('期末重估基准日', None, '自动 = 报表月份的月末。外币余额按这一天的汇率重估。')]
for i, (lab, v, memo) in enumerate(BASIC, 4):
    put(ws, f'A{i}', lab, font=Font(name='微软雅黑', size=10, bold=True), fill=F_SUM, align=LEFT)
    put(ws, f'B{i}', v, fmt=(DATEF if isinstance(v, datetime.datetime) else None),
        font=Font(name='微软雅黑', size=11, bold=True, color='C00000'), fill=F_IN)
    put(ws, f'C{i}', memo, font=FT_TIP, align=LEFT, border=False)
put(ws, 'B7', '=IF($B$6="","",EOMONTH($B$6,0))', fmt=DATEF, font=FT_AUTO, fill=F_AUTO)
for col, w in (('A', 17), ('B', 15), ('C', 46)):
    ws.column_dimensions[col].width = w

# ── 账户档案 ──────────────────────────────────────────────────────────
ACOL = [('序号', 6, None), ('银行 ★', 13, None), ('账户名称 ★（唯一）', 24, None),
        ('账号（可不填）', 15, None), ('币种 ★', 9, None), ('账户类型', 12, None),
        ('期初余额(原币)', 15, MONEY), ('期初汇率', 10, RATE),
        ('期初人民币成本', 16, MONEY), ('核对', 26, None)]
title_bar(ws, 9, 1, len(ACOL), '一、账户档案　——　加一个账户：在这里接着写一行，再复制一张银行子表改 B2', C_MAIN)
hdr(ws, 10, 1, ACOL, C_MAIN)
band(ws, A0, AE, 1, ACOL)
for r in range(A0, AE + 1):
    for c in (1, 9, 10):
        cc = ws.cell(r, c); cc.font, cc.fill = FT_AUTO, F_AUTO
    g = f'IF($C{r}="","",'
    put(ws, f'A{r}', f'={g}COUNTA($C${A0}:$C{r}))', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'I{r}', f'={g}IFERROR(ROUND(N($G{r})*IF($H{r}<>"",$H{r},'
                     f'{rate_at("重估基准日", f"$E{r}", eom=True)}),2),""))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'J{r}', f'={g}'
        f'IF(COUNTIF($C${A0}:$C${AE},$C{r})>1,"※账户名称重复了，后面按名字取数会串",'
        f'IF(AND($E{r}<>"",COUNTIF(币种表,$E{r})=0),"※币种不在右边币种清单里",'
        f'IF($G{r}="","※期初余额没填（没有就填 0）",'
        f'IF(AND($E{r}<>本位币,$H{r}="",$I{r}=""),"△这个币种一个汇率都没填，期初人民币成本算不出来",'
        f'IF(AND($E{r}<>本位币,$H{r}=""),"△没填期初汇率，先拿期末汇率折的",""))))))',
        font=FT_AUTO, fill=F_AUTO, align=LEFT)
for i, a in enumerate(ACCOUNTS):
    r = A0 + i
    for c, v in zip((2, 3, 4, 5, 6, 7, 8), a):
        ws.cell(r, c).value = v

# ── 币种 / 类别 / 项目 / 往来单位 ─────────────────────────────────────
def list_block(col0, row_title, cols, n, title, seed):
    title_bar(ws, row_title, col0, len(cols), title, C_BASE)
    hdr(ws, row_title + 1, col0, cols, C_BASE)
    band(ws, row_title + 2, row_title + 1 + n, col0, cols)
    for i, row in enumerate(seed):
        for j, v in enumerate(row if isinstance(row, (list, tuple)) else [row]):
            ws.cell(row_title + 2 + i, col0 + j).value = v

list_block(12, 9, [('币种代码', 10, None), ('币种名称', 12, None)], N_CUR, '二、币种', CURS)
list_block(15, 9, [('收/支类别', 14, None), ('属性 ★', 11, None)], N_CAT,
           '三、收支类别（属性=内部转账的，不计入收入/支出合计）', CATS)
list_block(18, 9, [('项目 / 业务', 15, None)], N_PRJ, '四、项目', PROJECTS)
list_block(20, 9, [('往来单位', 20, None)], N_PTY, '五、往来单位', PARTIES)
dv(ws, '"收入,支出,内部转账"', f'P11:P{10 + N_CAT}', stop=True)
dv(ws, '=币种表', f'E{A0}:E{AE}')
ws.freeze_panes = 'A11'
ws.conditional_formatting.add(f'J{A0}:J{AE}', FormulaRule(
    formula=[f'LEFT(J{A0},1)="※"'], font=Font(color='C00000', bold=True),
    fill=PatternFill('solid', fgColor='FFD7D7')))

NAMES.update({
    '本位币': '参数设置!$B$4', '管理年度': '参数设置!$B$5',
    '报表月份': '参数设置!$B$6', '重估基准日': '参数设置!$B$7',
    '账户银行': f'参数设置!$B${A0}:$B${AE}', '账户名': f'参数设置!$C${A0}:$C${AE}',
    '账户币种': f'参数设置!$E${A0}:$E${AE}', '账户期初': f'参数设置!$G${A0}:$G${AE}',
    '账户期初本位': f'参数设置!$I${A0}:$I${AE}',
    '币种表': f'参数设置!$L$11:$L${10 + N_CUR}',
    '类别表': f'参数设置!$O$11:$O${10 + N_CAT}',
    '类别属性': f'参数设置!$P$11:$P${10 + N_CAT}',
    '项目表': f'参数设置!$R$11:$R${10 + N_PRJ}',
    '单位表': f'参数设置!$T$11:$T${10 + N_PTY}',
})

# ══════════════════════════════════════════════════════════════════════
# 2  月度汇率
# ══════════════════════════════════════════════════════════════════════
ws = newsheet('月度汇率', '月度汇率 · 每月填一次，填了就不用再动',
              '左边【记账汇率】是流水折人民币用的；右边【月末汇率】是月末重估用的，'
              '空着就自动用左边那个。某月没填？自动往前找最近一个填过的月份，并在核对表点名。', C_BASE, 18)
title_bar(ws, 4, 1, 1 + N_CUR, '记账汇率　——　流水按「交易日期所在月」自动匹配', C_MAIN)
title_bar(ws, 4, 11, N_CUR, '月末汇率　——　月末重估用；空着=用左边记账汇率', C_SUB)
RCOL = [('月份', 12, DATEF)] + [(c, 10, RATE) for c, _ in CURS]
hdr(ws, 5, 1, RCOL, C_MAIN)
hdr(ws, 5, 11, [(c, 10, RATE) for c, _ in CURS], C_SUB)
band(ws, 6, 17, 1, RCOL)
band(ws, 6, 17, 11, [(c, 10, RATE) for c, _ in CURS])
for m in range(12):
    r = 6 + m
    put(ws, f'A{r}', '=DATE(管理年度,1,1)' if m == 0 else f'=EDATE(A{r-1},1)',
        fmt=DATEF, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'B{r}', 1, fmt=RATE, font=FT_AUTO, fill=F_AUTO)      # 本位币恒为 1
    put(ws, f'K{r}', 1, fmt=RATE, font=FT_AUTO, fill=F_AUTO)
    got = RATES.get((YEAR, m + 1), {})
    for j, (code, _) in enumerate(CURS):
        if code in got:
            ws.cell(r, 2 + j).value = got[code]
put(ws, 'T5', '这一列别删：给核对表数「哪几个月还没填」用的',
    font=FT_TIP, align=LEFT, border=False)
for m in range(12):
    r = 6 + m
    put(ws, f'T{r}', f'=IF($A{r}="","",COUNTA($C{r}:$I{r}))', fmt='0', font=FT_AUTO, fill=F_AUTO)
ws.column_dimensions['J'].width = 2.2
ws.column_dimensions['T'].width = 12
ws.freeze_panes = 'B6'
NAMES.update({'汇率月份': '月度汇率!$A$6:$A$17', '记账汇率区': '月度汇率!$B$6:$I$17',
              '汇率币种头': '月度汇率!$B$5:$I$5', '月末汇率区': '月度汇率!$K$6:$R$17'})


# ══════════════════════════════════════════════════════════════════════
# 3  流水录入（全套只有这一张要填）
# ══════════════════════════════════════════════════════════════════════
FCOL = [('日期 ★', 12, DATEF), ('账户名称 ★', 23, None), ('银行', 12, None), ('币种', 8, None),
        ('收/支类别 ★', 13, None), ('类别属性', 10, None),
        ('收入金额(原币)', 15, MONEY), ('支出金额(原币)', 15, MONEY),
        ('实际成交汇率\n(选填)', 12, RATE), ('应用汇率', 10, RATE),
        ('原币净额', 14, MONEY), ('折人民币', 15, MONEY),
        ('账户余额(原币)', 15, MONEY), ('账面人民币余额', 16, MONEY),
        ('摘要', 26, None), ('项目/业务', 13, None), ('往来单位', 16, None),
        ('备注', 14, None), ('核对', 32, None)]
AUTO_F = {3, 4, 6, 10, 11, 12, 13, 14, 19}        # 这些列是公式，别手填

ws = newsheet('流水录入', '流水录入 · 全套表只有这一张要填',
              '选账户名称 → 银行、币种、汇率自动带出来；收入和支出分两列，只填一个。'
              '银行有实际成交价（购汇、结汇）就填「实际成交汇率」，它优先于当月记账汇率。'
              '最后一列「核对」会自己挑毛病，红字是错、琥珀字是待办。', C_MAIN, len(FCOL))
hdr(ws, 5, 1, FCOL, C_MAIN)
band(ws, R0, RE_, 1, FCOL)
for r in range(R0, RE_ + 1):
    for c in AUTO_F:
        cc = ws.cell(r, c); cc.font, cc.fill = FT_AUTO, F_AUTO
    for c in (2, 15, 17, 18):          # 账户名称/摘要/往来单位/备注 靠左，居中挤成一团不好读
        ws.cell(r, c).alignment = LEFT
    g = f'IF($B{r}="","",'
    put(ws, f'C{r}', f'={g}IFERROR(INDEX(账户银行,MATCH($B{r},账户名,0)),""))', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'D{r}', f'={g}IFERROR(INDEX(账户币种,MATCH($B{r},账户名,0)),""))', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'F{r}', f'=IF($E{r}="","",IFERROR(INDEX(类别属性,MATCH($E{r},类别表,0)),""))',
        font=FT_AUTO, fill=F_AUTO)
    put(ws, f'J{r}', f'=IF(OR($A{r}="",$B{r}=""),"",IF($I{r}<>"",$I{r},{rate_at(f"$A{r}", f"$D{r}")}))',
        fmt=RATE, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'K{r}', f'=IF(AND($G{r}="",$H{r}=""),"",ROUND(N($G{r})-N($H{r}),2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'L{r}', f'=IF(OR($K{r}="",$J{r}=""),"",ROUND($K{r}*$J{r},2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    # 即时余额：期初 + 本账户到本行为止的累计。原币一套、人民币账面成本一套，各算各的。
    put(ws, f'M{r}', f'={g}ROUND(SUMIF(账户名,$B{r},账户期初)+SUMIFS($K${R0}:K{r},$B${R0}:B{r},$B{r}),2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    # 期初人民币成本取不到（这个币种一个汇率都没填）时，INDEX 返回 ""，""+数字 = #VALUE!，
    # 外面 IFERROR 兜住 → 整列留空。比画成 0.00 诚实：0 会被读成「这个账户人民币是 0」。
    put(ws, f'N{r}', f'={g}IFERROR(ROUND(INDEX(账户期初本位,MATCH($B{r},账户名,0))'
                     f'+SUMIFS($L${R0}:L{r},$B${R0}:B{r},$B{r}),2),""))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'S{r}', f'={g}'
        f'IF(ISNA(MATCH($B{r},账户名,0)),"※这个账户不在【参数设置】的账户档案里",'
        f'IF($A{r}="","※没填日期，这一笔不会进任何月份的报表",'
        f'IF(AND($G{r}<>"",$H{r}<>""),"※收入和支出不能同时填",'
        f'IF(AND(N($G{r})=0,N($H{r})=0),"※收入和支出都没填",'
        f'IF($E{r}="","※没填收支类别，分类报表会漏掉这一笔",'
        f'IF(ISNA(MATCH($E{r},类别表,0)),"※收支类别不在【参数设置】的清单里",'
        f'IF($J{r}="","△"&$D{r}&" 还没填汇率，去【月度汇率】补上就自动算了",'
        f'IF(YEAR($A{r})<>管理年度,"△不是管理年度的单据，年报里看不到",'
        f'IF(AND($F{r}="内部转账",$Q{r}=""),"△内部转账建议在「往来单位」写上对方账户",'
        f'IF($M{r}<0,"△这一笔之后该账户余额成负数了，核一下","")))))))))))',
        font=FT_AUTO, fill=F_AUTO, align=LEFT)
for i, f in enumerate(FLOWS):
    r = R0 + i
    d = datetime.datetime.strptime(f[0], '%Y-%m-%d')
    for c, v in ((1, d), (2, f[1]), (5, f[2]), (7, f[3]), (8, f[4]), (9, f[5]),
                 (15, f[6]), (16, f[7]), (17, f[8])):
        ws.cell(r, c).value = v
dv(ws, '=账户名', f'B{R0}:B{RE_}', stop=True)
dv(ws, '=类别表', f'E{R0}:E{RE_}', stop=True)
dv(ws, '=项目表', f'P{R0}:P{RE_}')
dv(ws, '=单位表', f'Q{R0}:Q{RE_}')
ws.conditional_formatting.add(f'S{R0}:S{RE_}', FormulaRule(
    formula=[f'LEFT(S{R0},1)="※"'], font=Font(color='C00000', bold=True),
    fill=PatternFill('solid', fgColor='FFD7D7')))
ws.conditional_formatting.add(f'S{R0}:S{RE_}', FormulaRule(
    formula=[f'LEFT(S{R0},1)="△"'], font=Font(color='8B5E00')))
ws.conditional_formatting.add(f'M{R0}:M{RE_}', FormulaRule(
    formula=[f'N(M{R0})<0'], font=Font(color='C00000', bold=True)))
ws.freeze_panes = f'C{R0}'
ws.auto_filter.ref = f'A5:S{RE_}'
NAMES.update({
    '流水日期': f'流水录入!$A${R0}:$A${RE_}', '流水账户': f'流水录入!$B${R0}:$B${RE_}',
    '流水币种': f'流水录入!$D${R0}:$D${RE_}', '流水类别': f'流水录入!$E${R0}:$E${RE_}',
    '流水属性': f'流水录入!$F${R0}:$F${RE_}', '流水收入': f'流水录入!$G${R0}:$G${RE_}',
    '流水支出': f'流水录入!$H${R0}:$H${RE_}', '流水净额': f'流水录入!$K${R0}:$K${RE_}',
    '流水本位': f'流水录入!$L${R0}:$L${RE_}', '流水项目': f'流水录入!$P${R0}:$P${RE_}',
    '流水单位': f'流水录入!$Q${R0}:$Q${RE_}',
})


# ══════════════════════════════════════════════════════════════════════
# 4  账户余额表（含汇兑损益）
# ══════════════════════════════════════════════════════════════════════
BCOL = [('序号', 6, '0'), ('银行', 12, None), ('账户名称', 23, None), ('币种', 8, None),
        ('期初余额(原币)', 15, MONEY), ('累计收入(原币)', 15, MONEY), ('累计支出(原币)', 15, MONEY),
        ('期末余额(原币)', 15, MONEY), ('期末汇率', 10, RATE),
        ('期末人民币(重估)', 17, MONEY), ('账面人民币成本', 16, MONEY),
        ('汇兑损益', 15, MONEY), ('核对', 30, None)]
B0 = 6
BE = B0 + N_ACC - 1
ws = newsheet('账户余额表', '账户余额表 · 每个账户一行，原币和人民币各算各的',
              '「账面人民币成本」= 期初成本 + 每笔按当时汇率折的人民币；'
              '「期末人民币(重估)」= 期末原币 × 期末汇率。两者之差就是汇兑损益 —— '
              '这是外汇台账跟普通资金台账唯一的区别，别把两个口径混着加。', C_RPT, len(BCOL))
hdr(ws, 5, 1, BCOL, C_RPT)
band(ws, B0, BE, 1, BCOL, auto=True)
for i in range(1, N_ACC + 1):
    r = B0 + i - 1
    g = f'IF($C{r}="","",'
    put(ws, f'C{r}', f'=IF({pick(K_0, i)}="","",{pick(K_0, i)})', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'A{r}', f'={g}{i})', fmt='0', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'B{r}', f'={g}{pick(K_1, i)})', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'D{r}', f'={g}{pick(K_2, i)})', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'E{r}', f'={g}ROUND(N({pick(K_3, i)}),2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'F{r}', f'={g}ROUND(SUMIF(流水账户,$C{r},流水收入),2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'G{r}', f'={g}ROUND(SUMIF(流水账户,$C{r},流水支出),2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'H{r}', f'={g}ROUND($E{r}+$F{r}-$G{r},2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'I{r}', f'={g}IFERROR({rate_at("重估基准日", f"$D{r}", eom=True)},""))',
        fmt=RATE, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'J{r}', f'=IF(OR($C{r}="",$I{r}=""),"",ROUND($H{r}*$I{r},2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'K{r}', f'={g}IFERROR(ROUND({pick(K_4, i)}+SUMIF(流水账户,$C{r},流水本位),2),""))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'L{r}', f'=IF(OR($C{r}="",$J{r}=""),"",ROUND($J{r}-$K{r},2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'M{r}', f'={g}'
        f'IF($H{r}<0,"※期末余额是负数，查一下是不是漏记了收入",'
        f'IF($I{r}="","△"&$D{r}&" 还没填汇率，人民币这几列算不出来",'
        f'IF(AND($D{r}<>本位币,ABS($L{r})>ABS($K{r})*0.2),"△汇兑损益超过账面成本 20%，核一下汇率",""))))',
        font=FT_AUTO, fill=F_AUTO, align=LEFT)
TR = BE + 1
put(ws, f'A{TR}', '合计', font=FT_HDR, fill=PatternFill('solid', fgColor=C_RPT))
for c in 'BCDEFGHI':
    put(ws, f'{c}{TR}', '' if c not in 'D' else '—', font=FT_HDR,
        fill=PatternFill('solid', fgColor=C_RPT))
put(ws, f'D{TR}', '原币不同，不能相加', font=Font(name='微软雅黑', size=9, color='FFFFFF'),
    fill=PatternFill('solid', fgColor=C_RPT))
ws.merge_cells(f'D{TR}:I{TR}')
for c in 'JKL':
    put(ws, f'{c}{TR}', f'=ROUND(SUM({c}{B0}:{c}{BE}),2)', fmt=MONEY, font=FT_SUM, fill=F_SUM)
put(ws, f'M{TR}', '「期末人民币合计」就是公司的资金总额（人民币口径）',
    font=FT_TIP, align=LEFT)

# ── 按币种小计 ────────────────────────────────────────────────────────
CC = [('币种', 10, None), ('账户数', 9, '0'), ('期末余额(原币)', 16, MONEY), ('期末汇率', 10, RATE),
      ('折人民币', 16, MONEY), ('账面人民币成本', 16, MONEY), ('汇兑损益', 14, MONEY), ('占资金比重', 11, PCT)]
CR = TR + 2
title_bar(ws, CR, 1, len(CC), '按币种小计　——　外币敞口一眼看完', C_SUB)
hdr(ws, CR + 1, 1, CC, C_SUB)
band(ws, CR + 2, CR + 1 + N_CUR, 1, CC, auto=True)
for i in range(1, N_CUR + 1):
    r = CR + 1 + i
    g = f'IF($A{r}="","",'
    put(ws, f'A{r}', f'=IF({pick(K_5, i)}="","",{pick(K_5, i)})', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'B{r}', f'={g}COUNTIF($D${B0}:$D${BE},$A{r}))', fmt='0', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'C{r}', f'={g}ROUND(SUMIF($D${B0}:$D${BE},$A{r},$H${B0}:$H${BE}),2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'D{r}', f'={g}IFERROR({rate_at("重估基准日", f"$A{r}", eom=True)},""))',
        fmt=RATE, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'E{r}', f'={g}ROUND(SUMIF($D${B0}:$D${BE},$A{r},$J${B0}:$J${BE}),2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'F{r}', f'={g}ROUND(SUMIF($D${B0}:$D${BE},$A{r},$K${B0}:$K${BE}),2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'G{r}', f'={g}ROUND($E{r}-$F{r},2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'H{r}', f'=IF(OR($A{r}="",N($J${TR})=0),"",$E{r}/$J${TR})',
        fmt=PCT, font=FT_AUTO, fill=F_AUTO)
CT = CR + 2 + N_CUR
put(ws, f'A{CT}', '合计', font=FT_HDR, fill=PatternFill('solid', fgColor=C_SUB))
for c, f in (('B', f'=SUM(B{CR+2}:B{CT-1})'), ('C', '—'),
             ('D', '—'), ('E', f'=ROUND(SUM(E{CR+2}:E{CT-1}),2)'),
             ('F', f'=ROUND(SUM(F{CR+2}:F{CT-1}),2)'), ('G', f'=ROUND(SUM(G{CR+2}:G{CT-1}),2)'),
             ('H', f'=IF(N($J${TR})=0,"",SUM(H{CR+2}:H{CT-1}))')):
    put(ws, f'{c}{CT}', f, fmt=(PCT if c == 'H' else '0' if c == 'B' else MONEY),
        font=FT_SUM, fill=F_SUM)
ws.conditional_formatting.add(f'M{B0}:M{BE}', FormulaRule(
    formula=[f'LEFT(M{B0},1)="※"'], font=Font(color='C00000', bold=True),
    fill=PatternFill('solid', fgColor='FFD7D7')))
ws.conditional_formatting.add(f'L{B0}:L{TR}', FormulaRule(
    formula=[f'N(L{B0})<0'], font=Font(color='C00000')))
ws.freeze_panes = f'A{B0}'
NAMES['资金总额本位'] = f'账户余额表!$J${TR}'

# ══════════════════════════════════════════════════════════════════════
# 5  银行子表（每个账户一张，自动拆）
# ══════════════════════════════════════════════════════════════════════
SUBSUM = [('币种', f'IFERROR(INDEX(账户币种,MATCH($B$2,账户名,0)),"")', None),
          ('期初(原币)', f'ROUND(SUMIF(账户名,$B$2,账户期初),2)', MONEY),
          ('累计收入', f'ROUND(SUMIF(流水账户,$B$2,流水收入),2)', MONEY),
          ('累计支出', f'ROUND(SUMIF(流水账户,$B$2,流水支出),2)', MONEY),
          ('期末(原币)', f'ROUND(SUMIF(账户名,$B$2,账户期初)+SUMIF(流水账户,$B$2,流水净额),2)', MONEY),
          ('期末人民币', f'IFERROR(ROUND(VLOOKUP($B$2,账户余额表!$C${B0}:$L${BE},8,0),2),"")', MONEY),
          ('汇兑损益', f'IFERROR(ROUND(VLOOKUP($B$2,账户余额表!$C${B0}:$L${BE},10,0),2),"")', MONEY)]

for acc in ACCOUNTS:
    name = acc[1]
    ws = newsheet(name, f'【{name}】流水明细 · 自动从【流水录入】拆出来，不用填', '',
                  C_SUB, len(FCOL), merge2=False)
    put(ws, 'A2', '账户名称', font=Font(name='微软雅黑', size=10, bold=True, color='FFFFFF'),
        fill=PatternFill('solid', fgColor=C_SUB))
    put(ws, 'B2', name, font=Font(name='微软雅黑', size=11, bold=True, color='C00000'), fill=F_AUTO)
    put(ws, 'C2', '← 这格是拆分依据，别改。下面整块是一条公式的结果，'
                  '在里面打字会把公式顶掉；要改数据去【流水录入】改。',
        font=FT_TIP, fill=PatternFill('solid', fgColor='FFF7E6'), align=LEFT)
    ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=len(FCOL))
    for j, (lab, f, fmt) in enumerate(SUBSUM):
        c1, c2 = CL(1 + j * 2), CL(2 + j * 2)
        put(ws, f'{c1}3', lab, font=Font(name='微软雅黑', size=9, bold=True, color='FFFFFF'),
            fill=PatternFill('solid', fgColor=C_SUB))
        put(ws, f'{c2}3', f'=IF($B$2="","",{f})', fmt=fmt, font=FT_SUM, fill=F_SUM)
    ws.row_dimensions[3].height = 20
    hdr(ws, 5, 1, FCOL, C_SUB)
    band(ws, R0, R0 + SUB - 1, 1, [(t, w, DATEQ if fmt == DATEF else fmt) for t, w, fmt in FCOL],
         auto=True)
    ws[f'A{R0}'] = ArrayFormula(
        ref=f'A{R0}:S{R0 + SUB - 1}',
        text=f'=_xlfn._xlws.FILTER({nb(f"流水录入!$A${R0}:$S${RE_}")},'
             f'流水录入!$B${R0}:$B${RE_}=$B$2,"这个账户还没有流水")')
    ws.freeze_panes = f'A{R0}'


# ══════════════════════════════════════════════════════════════════════
# 6  月度汇报表（本位币）
# ══════════════════════════════════════════════════════════════════════
def MS(m): return f'DATE(管理年度,{m},1)'
def ME(m): return f'EOMONTH(DATE(管理年度,{m},1),0)'
def RNG(m): return f'流水日期,">="&{MS(m)},流水日期,"<="&{ME(m)}'


MCOL = [('月份', 10, None), ('收入(人民币)', 16, MONEY), ('支出(人民币)', 16, MONEY),
        ('净流入(人民币)', 16, MONEY), ('内部转入', 14, MONEY), ('内部转出', 14, MONEY),
        ('月末资金总额(人民币账面)', 22, MONEY), ('说明', 30, None)]
ws = newsheet('月度汇报表', '月度汇报表 · 全部折成人民币',
              '收入/支出按「钱的方向」认（折人民币为正=收，为负=支），'
              '内部转账（账户之间调头寸、购汇结汇）单独两列列出来，不算进收入和支出 —— '
              '原模板把内部转账算进去了，收入和支出两头都虚增。', C_RPT, len(MCOL))
title_bar(ws, 4, 1, len(MCOL), '一、按月汇总（人民币）', C_RPT)
hdr(ws, 5, 1, MCOL, C_RPT)
band(ws, 6, 17, 1, MCOL, auto=True)
NOTIN = '流水属性,"<>内部转账"'
for m in range(1, 13):
    r = 5 + m
    put(ws, f'A{r}', f'={MS(m)}', fmt='yyyy年m月', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'B{r}', f'=ROUND(SUMIFS(流水本位,{RNG(m)},{NOTIN},流水本位,">0"),2)',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'C{r}', f'=-ROUND(SUMIFS(流水本位,{RNG(m)},{NOTIN},流水本位,"<0"),2)',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'D{r}', f'=ROUND($B{r}-$C{r},2)', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'E{r}', f'=ROUND(SUMIFS(流水本位,{RNG(m)},流水属性,"内部转账",流水本位,">0"),2)',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'F{r}', f'=-ROUND(SUMIFS(流水本位,{RNG(m)},流水属性,"内部转账",流水本位,"<0"),2)',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'G{r}', f'=ROUND(SUM(账户期初本位)+SUMIFS(流水本位,流水日期,"<="&{ME(m)}),2)',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'H{r}', f'=IF(ROUND($E{r}-$F{r},2)=0,"",'
                     f'"内部转账净额 "&TEXT($E{r}-$F{r},"#,##0.00")&"，是换汇价差")',
        font=FT_TIP, align=LEFT)
MT = 18
put(ws, f'A{MT}', '全年合计', font=FT_HDR, fill=PatternFill('solid', fgColor=C_RPT))
for c in 'BCDEF':
    put(ws, f'{c}{MT}', f'=ROUND(SUM({c}6:{c}17),2)', fmt=MONEY, font=FT_SUM, fill=F_SUM)
put(ws, f'G{MT}', '=G17', fmt=MONEY, font=FT_SUM, fill=F_SUM)
put(ws, f'H{MT}', '年末资金总额 = 12 月那一行', font=FT_TIP, align=LEFT)

# ── 二、报表月份 · 各账户 ─────────────────────────────────────────────
ACOL2 = [('账户名称', 23, None), ('币种', 8, None), ('月初余额(原币)', 15, MONEY),
         ('本月收入(原币)', 15, MONEY), ('本月支出(原币)', 15, MONEY), ('月末余额(原币)', 15, MONEY),
         ('本月收入(人民币)', 16, MONEY), ('本月支出(人民币)', 16, MONEY),
         ('月末余额(人民币账面)', 19, MONEY)]
M2 = MT + 2
title_bar(ws, M2, 1, len(ACOL2), '二、报表月份 · 各账户（报表月份在【参数设置】B6 改）', C_MAIN)
hdr(ws, M2 + 1, 1, ACOL2, C_MAIN)
band(ws, M2 + 2, M2 + 1 + N_ACC, 1, ACOL2, auto=True)
M_S, M_E = '报表月份', 'EOMONTH(报表月份,0)'
for i in range(1, N_ACC + 1):
    r = M2 + 1 + i
    g = f'IF($A{r}="","",'
    put(ws, f'A{r}', f'=IF({pick(K_0, i)}="","",{pick(K_0, i)})', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'B{r}', f'={g}{pick(K_2, i)})', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'C{r}', f'={g}ROUND(N({pick(K_3, i)})'
                     f'+SUMIFS(流水净额,流水账户,$A{r},流水日期,"<"&{M_S}),2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'D{r}', f'={g}ROUND(SUMIFS(流水收入,流水账户,$A{r},流水日期,">="&{M_S},'
                     f'流水日期,"<="&{M_E}),2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'E{r}', f'={g}ROUND(SUMIFS(流水支出,流水账户,$A{r},流水日期,">="&{M_S},'
                     f'流水日期,"<="&{M_E}),2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'F{r}', f'={g}ROUND($C{r}+$D{r}-$E{r},2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'G{r}', f'={g}ROUND(SUMIFS(流水本位,流水账户,$A{r},流水日期,">="&{M_S},'
                     f'流水日期,"<="&{M_E},流水本位,">0"),2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'H{r}', f'={g}-ROUND(SUMIFS(流水本位,流水账户,$A{r},流水日期,">="&{M_S},'
                     f'流水日期,"<="&{M_E},流水本位,"<0"),2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'I{r}', f'={g}IFERROR(ROUND({pick(K_4, i)}'
                     f'+SUMIFS(流水本位,流水账户,$A{r},流水日期,"<="&{M_E}),2),""))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
M2T = M2 + 2 + N_ACC
put(ws, f'A{M2T}', '合计', font=FT_HDR, fill=PatternFill('solid', fgColor=C_MAIN))
put(ws, f'B{M2T}', '原币不能加', font=Font(name='微软雅黑', size=9, color='FFFFFF'),
    fill=PatternFill('solid', fgColor=C_MAIN))
ws.merge_cells(f'B{M2T}:F{M2T}')
for c in 'GHI':
    put(ws, f'{c}{M2T}', f'=ROUND(SUM({c}{M2+2}:{c}{M2T-1}),2)', fmt=MONEY, font=FT_SUM, fill=F_SUM)
ws.freeze_panes = 'A6'
NAMES['月报收入合计'] = f'月度汇报表!$B${MT}'
NAMES['月报支出合计'] = f'月度汇报表!$C${MT}'

# ══════════════════════════════════════════════════════════════════════
# 7  收支分类报表（本位币）
# ══════════════════════════════════════════════════════════════════════
KCOL = ([('收/支类别', 15, None), ('属性', 10, None)]
        + [(f'{m}月', 12, MONEY) for m in range(1, 13)]
        + [('全年合计', 15, MONEY), ('占比', 9, PCT)])
ws = newsheet('收支分类报表', '收支分类报表 · 按收支类别，全部折成人民币',
              '支出类别显示为正数（好读）。内部转账单独一行，不进收入/支出合计；'
              '它的净额如果不是 0，那就是购汇结汇的价差。', C_RPT, len(KCOL))
hdr(ws, 4, 1, KCOL, C_RPT)
SUMR = [('收入合计', '"收入"'), ('支出合计', '"支出"'), ('内部转账净额', '"内部转账"')]
K0 = 12
KE = K0 + N_CAT - 1
for j, (lab, attr) in enumerate(SUMR):
    r = 5 + j
    put(ws, f'A{r}', lab, font=FT_HDR, fill=PatternFill('solid', fgColor=C_RPT), align=LEFT)
    put(ws, f'B{r}', attr.strip('"'), font=FT_HDR, fill=PatternFill('solid', fgColor=C_RPT))
    for m in range(1, 13):
        cc = CL(2 + m)
        put(ws, f'{cc}{r}', f'=ROUND(SUMIF($B${K0}:$B${KE},{attr},{cc}${K0}:{cc}${KE}),2)',
            fmt=MONEY, font=FT_SUM, fill=F_SUM)
    put(ws, f'O{r}', f'=ROUND(SUM(C{r}:N{r}),2)', fmt=MONEY, font=FT_SUM, fill=F_SUM)
    put(ws, f'P{r}', '', fill=F_SUM)
put(ws, 'A8', '经营净额', font=FT_HDR, fill=PatternFill('solid', fgColor=C_MAIN), align=LEFT)
put(ws, 'B8', '收入−支出', font=FT_HDR, fill=PatternFill('solid', fgColor=C_MAIN))
for m in range(1, 13):
    cc = CL(2 + m)
    put(ws, f'{cc}8', f'=ROUND({cc}5-{cc}6,2)', fmt=MONEY, font=FT_SUM, fill=F_SUM)
put(ws, 'O8', '=ROUND(SUM(C8:N8),2)', fmt=MONEY, font=FT_SUM, fill=F_SUM)
put(ws, 'P8', '', fill=F_SUM)
title_bar(ws, 10, 1, len(KCOL), '按类别明细　——　类别清单在【参数设置】维护，这里自动跟着走', C_SUB)
hdr(ws, 11, 1, KCOL, C_SUB)
band(ws, K0, KE, 1, KCOL, auto=True)
for i in range(1, N_CAT + 1):
    r = K0 + i - 1
    g = f'IF($A{r}="","",'
    put(ws, f'A{r}', f'=IF({pick(K_6, i)}="","",{pick(K_6, i)})', font=FT_AUTO, fill=F_AUTO, align=LEFT)
    put(ws, f'B{r}', f'={g}{pick(K_7, i)})', font=FT_AUTO, fill=F_AUTO)
    for m in range(1, 13):
        cc = CL(2 + m)
        put(ws, f'{cc}{r}', f'={g}ROUND(IF($B{r}="支出",-1,1)*'
                            f'SUMIFS(流水本位,流水类别,$A{r},{RNG(m)}),2))',
            fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'O{r}', f'={g}ROUND(SUM(C{r}:N{r}),2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'P{r}', f'=IF(OR($A{r}="",$B{r}="内部转账"),"",'
                     f'IFERROR($O{r}/IF($B{r}="收入",$O$5,$O$6),""))',
        fmt=PCT, font=FT_AUTO, fill=F_AUTO)
ws.freeze_panes = 'C12'
NAMES['分类收入合计'] = '收支分类报表!$O$5'
NAMES['分类支出合计'] = '收支分类报表!$O$6'


# ══════════════════════════════════════════════════════════════════════
# 8  资金日报表
# ══════════════════════════════════════════════════════════════════════
N_DA = 15                      # 按日走势里排几个账户的余额列
DCOL = ([('日期', 12, DATEF), ('当日收款(人民币)', 16, MONEY), ('当日付款(人民币)', 16, MONEY),
         ('当日净额', 15, MONEY), ('当日资金总额(人民币账面)', 22, MONEY)]
        + [('', 17, MONEY) for _ in range(N_DA)])
ws = newsheet('资金日报表', '资金日报表 · 改上面两个日期就行',
              '上半张是按日走势（最多 31 天），下半张是某一天的出纳日报。'
              '收付都不含内部转账。账户余额列按原币显示，列头带币种。', C_RPT, len(DCOL))
for lab, addr, v in (('起始日期', 'B3', datetime.datetime(YEAR, 8, 1)),
                     ('截止日期', 'D3', datetime.datetime(YEAR, 8, 31))):
    put(ws, f'{CL(ord(addr[0])-64-1)}3', lab, font=Font(name='微软雅黑', size=10, bold=True),
        fill=F_SUM)
    put(ws, addr, v, fmt=DATEF, font=Font(name='微软雅黑', size=11, bold=True, color='C00000'),
        fill=F_IN)
put(ws, 'E3', '天数', font=Font(name='微软雅黑', size=10, bold=True), fill=F_SUM)
put(ws, 'F3', '=IF(OR($B$3="",$D$3=""),"",IF($D$3<$B$3,"※截止早于起始",'
              'IF($D$3-$B$3+1>31,"※超过 31 天，请分段查",$D$3-$B$3+1)))',
    font=FT_SUM, fill=F_AUTO)
put(ws, 'G3', '超过 31 天就按月分段查；月度数据直接看【月度汇报表】。', font=FT_TIP, align=LEFT)
title_bar(ws, 5, 1, len(DCOL), '一、按日走势（人民币，不含内部转账）', C_RPT)
hdr(ws, 6, 1, DCOL, C_RPT)
for i in range(1, N_DA + 1):
    cc = CL(5 + i)
    c = ws.cell(6, 5 + i)
    c.value = (f'=IF(IF({pick(K_0, i)}="","",{pick(K_0, i)})="","",'
               f'{pick(K_0, i)}&"("&{pick(K_2, i)}&")")')
    c.font, c.fill, c.alignment, c.border = FT_HDR, PatternFill('solid', fgColor=C_RPT), CEN, BOX
band(ws, 7, 37, 1, DCOL, auto=True)
NOT_T = '流水属性,"<>内部转账"'
for d in range(31):
    r = 7 + d
    put(ws, f'A{r}', f'=IF(OR($B$3="",NOT(ISNUMBER($F$3))),"",IF({d + 1}>$F$3,"",$B$3+{d}))',
        fmt=DATEQ, font=FT_AUTO, fill=F_AUTO)
    g = f'IF($A{r}="","",'
    put(ws, f'B{r}', f'={g}ROUND(SUMIFS(流水本位,流水日期,$A{r},{NOT_T},流水本位,">0"),2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'C{r}', f'={g}-ROUND(SUMIFS(流水本位,流水日期,$A{r},{NOT_T},流水本位,"<0"),2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'D{r}', f'={g}ROUND($B{r}-$C{r},2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'E{r}', f'={g}ROUND(SUM(账户期初本位)+SUMIFS(流水本位,流水日期,"<="&$A{r}),2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    for i in range(1, N_DA + 1):
        put(ws, f'{CL(5 + i)}{r}',
            f'=IF(OR($A{r}="",IF({pick(K_0, i)}="","",{pick(K_0, i)})=""),"",'
            f'ROUND(N({pick(K_3, i)})+SUMIFS(流水净额,流水账户,{pick(K_0, i)},'
            f'流水日期,"<="&$A{r}),2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
put(ws, 'A38', '期间合计', font=FT_HDR, fill=PatternFill('solid', fgColor=C_RPT))
for c in 'BCD':
    put(ws, f'{c}38', f'=ROUND(SUM({c}7:{c}37),2)', fmt=MONEY, font=FT_SUM, fill=F_SUM)
put(ws, 'E38', '← 右边各列是当日余额，不能纵向相加', font=FT_TIP, align=LEFT)

# ── 二、某一天的出纳日报 ──────────────────────────────────────────────
DD = [('序号', 6, '0'), ('账户名称', 23, None), ('币种', 8, None), ('昨日余额(原币)', 15, MONEY),
      ('本日收入(原币)', 15, MONEY), ('本日支出(原币)', 15, MONEY), ('本日余额(原币)', 15, MONEY),
      ('汇率', 10, RATE), ('本日余额(人民币)', 17, MONEY)]
D2 = 40
title_bar(ws, D2, 1, len(DD), '二、某一天的出纳日报', C_MAIN)
put(ws, f'A{D2+1}', '日报日期', font=Font(name='微软雅黑', size=10, bold=True), fill=F_SUM)
put(ws, f'B{D2+1}', '=IF($D$3="","",$D$3)', fmt=DATEF,
    font=Font(name='微软雅黑', size=11, bold=True, color='C00000'), fill=F_AUTO)
put(ws, f'C{D2+1}', '默认取上面的截止日期；要看别的天，直接把这格改成那一天。',
    font=FT_TIP, align=LEFT)
ws.merge_cells(start_row=D2 + 1, start_column=3, end_row=D2 + 1, end_column=len(DD))
hdr(ws, D2 + 2, 1, DD, C_MAIN)
band(ws, D2 + 3, D2 + 2 + N_ACC, 1, DD, auto=True)
DAY = f'$B${D2+1}'
for i in range(1, N_ACC + 1):
    r = D2 + 2 + i
    g = f'IF(OR($B{r}="",{DAY}=""),"",'
    put(ws, f'B{r}', f'=IF({pick(K_0, i)}="","",{pick(K_0, i)})', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'A{r}', f'=IF($B{r}="","",{i})', fmt='0', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'C{r}', f'=IF($B{r}="","",{pick(K_2, i)})', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'D{r}', f'={g}ROUND(N({pick(K_3, i)})'
                     f'+SUMIFS(流水净额,流水账户,$B{r},流水日期,"<"&{DAY}),2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'E{r}', f'={g}ROUND(SUMIFS(流水收入,流水账户,$B{r},流水日期,{DAY}),2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'F{r}', f'={g}ROUND(SUMIFS(流水支出,流水账户,$B{r},流水日期,{DAY}),2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'G{r}', f'={g}ROUND($D{r}+$E{r}-$F{r},2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'H{r}', f'={g}IFERROR({rate_at(DAY, f"$C{r}", eom=True)},""))',
        fmt=RATE, font=FT_AUTO, fill=F_AUTO)
    put(ws, f'I{r}', f'=IF(OR($B{r}="",$H{r}=""),"",ROUND($G{r}*$H{r},2))',
        fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
DT = D2 + 3 + N_ACC
put(ws, f'A{DT}', '合计', font=FT_HDR, fill=PatternFill('solid', fgColor=C_MAIN))
put(ws, f'B{DT}', '原币不能加', font=Font(name='微软雅黑', size=9, color='FFFFFF'),
    fill=PatternFill('solid', fgColor=C_MAIN))
ws.merge_cells(f'B{DT}:H{DT}')
put(ws, f'I{DT}', f'=ROUND(SUM(I{D2+3}:I{DT-1}),2)', fmt=MONEY, font=FT_SUM, fill=F_SUM)
ws.freeze_panes = 'A7'


# ══════════════════════════════════════════════════════════════════════
# 9  区间查询
# ══════════════════════════════════════════════════════════════════════
QR = '流水日期,">="&$B$3,流水日期,"<="&$D$3'
ws = newsheet('区间查询', '区间查询 · 任意起止日期，全部折成人民币',
              '改上面两个日期就行。收入/支出都不含内部转账；内部转账单独列在右上角。', C_RPT, 6)
for lab, addr, v in (('起始日期', 'B3', datetime.datetime(YEAR, 1, 1)),
                     ('截止日期', 'D3', datetime.datetime(YEAR, 12, 31))):
    put(ws, f'{CL(ord(addr[0]) - 65)}3', lab, font=Font(name='微软雅黑', size=10, bold=True), fill=F_SUM)
    put(ws, addr, v, fmt=DATEF, font=Font(name='微软雅黑', size=11, bold=True, color='C00000'), fill=F_IN)
put(ws, 'E3', '=IF(OR($B$3="",$D$3=""),"",IF($D$3<$B$3,"※截止日期早于起始日期",'
              '"共 "&($D$3-$B$3+1)&" 天，"&COUNTIFS(' + QR + ')&" 笔"))',
    font=FT_TIP, align=LEFT)
ws.merge_cells('E3:F3')
QSUM = [('收入总额', f'ROUND(SUMIFS(流水本位,{QR},{NOT_T},流水本位,">0"),2)'),
        ('支出总额', f'-ROUND(SUMIFS(流水本位,{QR},{NOT_T},流水本位,"<0"),2)'),
        ('经营净额', '=ROUND($B$5-$D$5,2)'),
        ('内部转账净额', f'ROUND(SUMIFS(流水本位,{QR},流水属性,"内部转账"),2)')]
for j, (lab, f) in enumerate(QSUM):
    put(ws, f'{CL(1 + j * 2)}5', lab, font=Font(name='微软雅黑', size=10, bold=True, color='FFFFFF'),
        fill=PatternFill('solid', fgColor=C_RPT))
    put(ws, f'{CL(2 + j * 2)}5', f if f.startswith('=') else f'=IF($B$3="","",{f})',
        fmt=MONEY, font=FT_SUM, fill=F_SUM)
for c, w in zip('ABCDEFGH', (20, 16, 16, 16, 14, 14, 16, 16)):
    ws.column_dimensions[c].width = w

QB = [(7, '一、按收支类别', [('收/支类别', None), ('属性', None), ('金额(人民币)', MONEY), ('占比', PCT)],
       N_CAT, '类别表'),
      (0, '二、按项目 / 业务', [('项目/业务', None), ('收入(人民币)', MONEY), ('支出(人民币)', MONEY),
                             ('净额', MONEY), ('占收入比', PCT)], N_PRJ, '项目表'),
      (0, '三、按币种', [('币种', None), ('收入(人民币)', MONEY), ('支出(人民币)', MONEY),
                       ('净额', MONEY), ('笔数', '0')], N_CUR, '币种表'),
      (0, '四、按往来单位', [('往来单位', None), ('收入(人民币)', MONEY), ('支出(人民币)', MONEY),
                         ('净额', MONEY), ('笔数', '0')], N_PTY, '单位表')]
KEYCOL = {'类别表': '流水类别', '项目表': '流水项目', '币种表': '流水币种', '单位表': '流水单位'}
row = 7
for _, title, cols, n, src in QB:
    cc = [(t, None, fmt) for t, fmt in cols]
    title_bar(ws, row, 1, len(cols), title, C_SUB)
    hdr(ws, row + 1, 1, cc, C_SUB)
    band(ws, row + 2, row + 1 + n, 1, cc, auto=True)
    key = KEYCOL[src]
    for i in range(1, n + 1):
        r = row + 1 + i
        g = f'IF($A{r}="","",'
        put(ws, f'A{r}', f'=IF({pick(src, i)}="","",{pick(src, i)})',
            font=FT_AUTO, fill=F_AUTO, align=LEFT)
        if src == '类别表':
            put(ws, f'B{r}', f'={g}{pick(K_7, i)})', font=FT_AUTO, fill=F_AUTO)
            put(ws, f'C{r}', f'={g}ROUND(IF($B{r}="支出",-1,1)*SUMIFS(流水本位,{key},$A{r},{QR}),2))',
                fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
            put(ws, f'D{r}', f'=IF(OR($A{r}="",$B{r}="内部转账"),"",'
                             f'IFERROR($C{r}/IF($B{r}="收入",$B$5,$D$5),""))',
                fmt=PCT, font=FT_AUTO, fill=F_AUTO)
        else:
            put(ws, f'B{r}', f'={g}ROUND(SUMIFS(流水本位,{key},$A{r},{QR},流水本位,">0"),2))',
                fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
            put(ws, f'C{r}', f'={g}-ROUND(SUMIFS(流水本位,{key},$A{r},{QR},流水本位,"<0"),2))',
                fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
            put(ws, f'D{r}', f'={g}ROUND($B{r}-$C{r},2))', fmt=MONEY, font=FT_AUTO, fill=F_AUTO)
            if src == '项目表':
                put(ws, f'E{r}', f'=IF($A{r}="","",IFERROR($B{r}/$B$5,""))',
                    fmt=PCT, font=FT_AUTO, fill=F_AUTO)
            else:
                put(ws, f'E{r}', f'={g}COUNTIFS({key},$A{r},{QR}))',
                    fmt='0', font=FT_AUTO, fill=F_AUTO)
    tot = row + 2 + n
    put(ws, f'A{tot}', '小计', font=FT_HDR, fill=PatternFill('solid', fgColor=C_SUB))
    span = 'CD' if src == '类别表' else 'BCD'
    for c in span:
        if c == 'D' and src == '类别表':
            put(ws, f'D{tot}', '', fill=F_SUM); continue
        put(ws, f'{c}{tot}', f'=ROUND(SUM({c}{row+2}:{c}{tot-1}),2)', fmt=MONEY,
            font=FT_SUM, fill=F_SUM)
    if src == '类别表':
        put(ws, f'B{tot}', '收入+支出(取正)+内部转账', font=Font(name='微软雅黑', size=9, color='FFFFFF'),
            fill=PatternFill('solid', fgColor=C_SUB))
    elif src in ('币种表', '单位表'):
        put(ws, f'E{tot}', f'=SUM(E{row+2}:E{tot-1})', fmt='0', font=FT_SUM, fill=F_SUM)
    else:
        put(ws, f'E{tot}', '', fill=F_SUM)
    row = tot + 2
ws.freeze_panes = 'A7'

# ══════════════════════════════════════════════════════════════════════
# 10  月度汇率补两列辅助数（核对表要用）
# ══════════════════════════════════════════════════════════════════════
ws = wb['月度汇率']
put(ws, 'U5', '本月外币流水笔数', font=FT_HDR, fill=PatternFill('solid', fgColor=C_SUB))
ws.column_dimensions['U'].width = 14
for m in range(12):
    r = 6 + m
    put(ws, f'U{r}', f'=IF($A{r}="","",COUNTIFS(流水日期,">="&$A{r},流水日期,"<="&EOMONTH($A{r},0),'
                     f'流水币种,"<>"&本位币))', fmt='0', font=FT_AUTO, fill=F_AUTO)
put(ws, 'A18', '全年填了几个月', font=FT_HDR, fill=PatternFill('solid', fgColor=C_MAIN), align=LEFT)
for j in range(N_CUR):
    put(ws, f'{CL(2 + j)}18', f'=COUNT({CL(2 + j)}6:{CL(2 + j)}17)', fmt='0',
        font=FT_SUM, fill=F_SUM)
NAMES.update({'汇率填写数': '月度汇率!$T$6:$T$17', '外币流水数': '月度汇率!$U$6:$U$17',
              '汇率列计数': '月度汇率!$B$18:$I$18'})


# ══════════════════════════════════════════════════════════════════════
# 11  核对表
# ══════════════════════════════════════════════════════════════════════
NAMES['流水核对'] = f'流水录入!$S${R0}:$S${RE_}'
NAMES['档案核对'] = f'参数设置!$J${A0}:$J${AE}'
N_SUB = len(ACCOUNTS)

CHECKS = [
    ('勾稽', '各账户「账面人民币成本」合计 = 期初成本 + 全部流水折人民币',
     f'=ROUND(账户余额表!$K${TR}-SUM(账户期初本位)-SUM(流水本位),2)', MONEY,
     '两边算法不同，必须分毫不差。不为 0 说明有账户名对不上。'),
    ('勾稽', '各账户「期末原币」合计 = 期初原币 + 全部流水原币净额',
     f'=ROUND(SUM(账户余额表!$H${B0}:$H${BE})-SUM(账户期初)-SUM(流水净额),2)', MONEY,
     '同上。（跨币种相加没有意义，但这个恒等式成立）'),
    ('勾稽', '分类报表「收入合计」= 流水里属性为收入的折人民币合计',
     '=ROUND(分类收入合计-SUMIFS(流水本位,流水属性,"收入"),2)', MONEY,
     '不为 0 说明有收入类别没在【参数设置】登记。'),
    ('勾稽', '分类报表「支出合计」= 流水里属性为支出的折人民币合计',
     '=ROUND(分类支出合计+SUMIFS(流水本位,流水属性,"支出"),2)', MONEY,
     '支出在流水里是负数，所以这里是加。'),
    ('勾稽', '按币种小计的「折人民币」= 各账户「期末人民币」合计',
     f'=ROUND(账户余额表!$E${CT}-账户余额表!$J${TR},2)', MONEY,
     '不为 0 说明有账户的币种不在【参数设置】的币种清单里，按币种那块会漏掉它。'),
    ('勾稽', '银行子表张数 = 账户档案里的账户数',
     f'={N_SUB}-COUNTA(账户名)', '0',
     f'现在有 {N_SUB} 张子表。加了账户没加子表，这里就不是 0 —— 复制一张子表改 B2 即可。'),
    ('待办', '流水里的红字（※，是错，必须改）',
     '=COUNTIF(流水核对,"※*")', '0', '去【流水录入】最后一列看，红的都要改。'),
    ('待办', '流水里的琥珀字（△，是待补的数据）',
     '=COUNTIF(流水核对,"△*")', '0', '不影响公式，但补齐了数才准。'),
    ('待办', '账户档案里的红字（※）',
     '=COUNTIF(档案核对,"※*")', '0', '账户名重复、币种不对、期初没填。'),
    ('待办', '有外币流水、但当月记账汇率一个都没填的月份数',
     '=SUMPRODUCT((汇率填写数=0)*(外币流水数>0))', '0',
     '这些月份的外币流水会去借用**更早**月份的汇率。去【月度汇率】补上。'),
    ('待办', '在用的外币里，全年一个汇率都没填的个数',
     '=SUMPRODUCT((汇率币种头<>本位币)*(COUNTIF(账户币种,汇率币种头)>0)*(汇率列计数=0))', '0',
     '这些币种的流水折不出人民币，所有人民币报表都会少这一块。'),
    ('待办', '内部转账折人民币净额（购汇结汇的价差）',
     '=ROUND(SUMIFS(流水本位,流水属性,"内部转账"),2)', MONEY,
     '同币种之间调头寸应该是 0；不是 0 的部分就是购汇/结汇的买卖价差，属于正常。'),
    ('待办', '内部转账「收」「支」笔数差',
     '=COUNTIFS(流水属性,"内部转账",流水本位,">0")-COUNTIFS(流水属性,"内部转账",流水本位,"<0")', '0',
     '内部转账要两头都录（转出账户记支出、转入账户记收入），不为 0 说明漏了一头。'),
    ('待办', '没填收支类别的流水笔数',
     f'=COUNTIFS(流水账户,"<>",流水类别,"")', '0', '没类别的笔数不会进分类报表。'),
    ('待办', '账户数 −【资金日报表】按日走势预留的余额列数',
     f'=MAX(0,COUNTA(账户名)-{N_DA})', '0',
     f'按日走势那块只排了 {N_DA} 个账户的余额列，超出的看【账户余额表】，或者告诉我加列。'),
    ('待办', '流水已用行数 / 预留行数',
     f'=COUNTA(流水账户)&" / {N_ROW}"', None, '快满了就告诉我加行。'),
    ('待办', '类别清单里的重复项',
     '=SUMPRODUCT((类别表<>"")*(COUNTIF(类别表,类别表&"")>1))', '0',
     '重复会在分类报表里出现两行一样的，金额被算两遍 —— 原模板就栽在这里。'),
    ('待办', '项目清单里的重复项',
     '=SUMPRODUCT((项目表<>"")*(COUNTIF(项目表,项目表&"")>1))', '0', '同上。'),
    ('待办', '汇兑损益合计（提示，不是错）',
     f'=ROUND(账户余额表!$L${TR},2)', MONEY,
     '＝期末按期末汇率重估 − 账面历史成本。外币余额越大、汇率波动越大，这个数越大。'),
]
CKC = [('序号', 6, '0'), ('类别', 8, None), ('检查项', 52, None), ('结果', 16, None),
       ('判断', 14, None), ('说明', 62, None)]
ws = newsheet('核对表', '核对表 · 每次录完看一眼这张就够了',
              '【勾稽】是算法自检，必须全是 0，不是 0 就是表算错了要告诉我；'
              '【待办】是数据没填齐，你们自己补。', C_CHK, len(CKC))
hdr(ws, 4, 1, CKC, C_CHK)
band(ws, 5, 4 + len(CHECKS), 1, CKC, auto=True)
for i, (kind, name, f, fmt, memo) in enumerate(CHECKS, 1):
    r = 4 + i
    put(ws, f'A{r}', i, fmt='0', font=FT_AUTO, fill=F_AUTO)
    put(ws, f'B{r}', kind, font=Font(name='微软雅黑', size=10, bold=True,
                                     color=('C00000' if kind == '勾稽' else '8B5E00')), fill=F_AUTO)
    put(ws, f'C{r}', name, font=FT_AUTO, fill=F_AUTO, align=LEFT)
    put(ws, f'D{r}', f, fmt=fmt, font=FT_SUM, fill=F_SUM)
    if fmt is None:
        put(ws, f'E{r}', '—', font=FT_AUTO, fill=F_AUTO)
    elif kind == '勾稽':
        put(ws, f'E{r}', f'=IF(ROUND(N($D{r}),2)=0,"✓ 通过","✗ 算不平")', font=FT_SUM, fill=F_AUTO)
    else:
        put(ws, f'E{r}', f'=IF(ROUND(N($D{r}),2)=0,"✓ 没有","△ 有")', font=FT_SUM, fill=F_AUTO)
    put(ws, f'F{r}', memo, font=FT_TIP, align=LEFT)
CKN = 4 + len(CHECKS)
put(ws, f'A{CKN+2}', '勾稽通过',
    font=Font(name='微软雅黑', size=11, bold=True, color='FFFFFF'),
    fill=PatternFill('solid', fgColor=C_CHK), align=LEFT)
ws.merge_cells(f'A{CKN+2}:C{CKN+2}')
NG = sum(1 for c in CHECKS if c[0] == '勾稽')
put(ws, f'D{CKN+2}', f'=COUNTIF($E$5:$E${CKN},"✓ 通过")&" / {NG}"', font=FT_SUM, fill=F_SUM)
put(ws, f'E{CKN+2}', '待办', font=Font(name='微软雅黑', size=11, bold=True, color='FFFFFF'),
    fill=PatternFill('solid', fgColor=C_SUB))
put(ws, f'F{CKN+2}', f'=COUNTIF($E$5:$E${CKN},"△ 有")&" 项没填齐"', font=FT_SUM, fill=F_SUM, align=LEFT)
ws.conditional_formatting.add(f'E5:E{CKN}', FormulaRule(
    formula=['LEFT(E5,1)="✗"'], font=Font(color='C00000', bold=True),
    fill=PatternFill('solid', fgColor='FFD7D7')))
ws.conditional_formatting.add(f'E5:E{CKN}', FormulaRule(
    formula=['LEFT(E5,1)="△"'], font=Font(color='8B5E00', bold=True),
    fill=PatternFill('solid', fgColor='FFF2CC')))
ws.freeze_panes = 'A5'


# ══════════════════════════════════════════════════════════════════════
# 12  使用说明（放到最前面）
# ══════════════════════════════════════════════════════════════════════
ws = wb.create_sheet('使用说明', 0)
ws.sheet_properties.tabColor = '1F3864'
ws.sheet_view.showGridLines = False
for i, w in enumerate([20, 16, 74, 20], 1):
    ws.column_dimensions[CL(i)].width = w
put(ws, 'A1', '外汇资金台账 · 使用说明（本位币：人民币）', font=FT_TITLE,
    fill=PatternFill('solid', fgColor='1F3864'), align=LEFT)
ws.merge_cells('A1:D1'); ws.row_dimensions[1].height = 30


def sec(r, t, color=C_MAIN):
    put(ws, f'A{r}', t, font=Font(name='微软雅黑', size=12, bold=True, color='FFFFFF'),
        fill=PatternFill('solid', fgColor=color), align=LEFT)
    ws.merge_cells(f'A{r}:D{r}'); ws.row_dimensions[r].height = 24


def ln(r, a, b='', c=''):
    put(ws, f'A{r}', a, font=Font(name='微软雅黑', size=10, bold=True), align=LEFT)
    put(ws, f'B{r}', b, font=Font(name='微软雅黑', size=10), align=LEFT)
    put(ws, f'C{r}', c, font=Font(name='微软雅黑', size=10),
        align=Alignment('left', 'center', wrap_text=True))
    ws.merge_cells(f'C{r}:D{r}')


def para(r, t, color='595959'):
    put(ws, f'A{r}', t, font=Font(name='微软雅黑', size=9, color=color), align=LEFT)
    ws.merge_cells(f'A{r}:D{r}')


r = 3
sec(r, '一、三步就能用起来'); r += 1
for a, b, c in [('第 1 步', '【参数设置】', '把账户档案改成你们真实的账户：银行、账户名称（唯一）、币种、期初余额。'
                                      '外币账户把「期初汇率」也填上 —— 那是这笔期初余额当初入账的汇率。'),
                ('第 2 步', '【月度汇率】', '每月填一次当月记账汇率。只填你们在用的币种（美元、港币、欧元）就行。'),
                ('第 3 步', '【流水录入】', '每天在这一张录。选账户名称 → 银行、币种、汇率自动带出来；'
                                      '收入和支出分两列，只填一个。其余所有表全自动。')]:
    ln(r, a, b, c); r += 1
r += 1

sec(r, '二、表都干什么用'); r += 1
ln(r, '表名', '谁用', '干什么'); r += 1
for a, b, c in [
        ('参数设置', '管理员', '账户档案、币种、收支类别、项目、往来单位。这一页改了，后面全部跟着变。'),
        ('月度汇率', '财务', '每月一行。左边记账汇率（流水折人民币用），右边月末汇率（月末重估用，空着=用左边）。'),
        ('流水录入', '出纳', '★ 全套表唯一要填的一张。最后一列「核对」会自己挑毛病。'),
        ('银行子表（每账户一张）', '看', '自动从流水拆出来，不用填。表头 B2 是拆分依据。'),
        ('账户余额表', '财务', '每账户一行：期末原币、期末人民币、账面成本、汇兑损益；下面还有按币种小计。'),
        ('月度汇报表', '老板', '按月的人民币收入/支出/净流入/月末资金总额；下半张是报表月份的各账户明细。'),
        ('收支分类报表', '老板', '收支类别 × 12 个月，人民币。支出显示成正数，好读。'),
        ('资金日报表', '出纳/老板', '上半张按日走势（最多 31 天），下半张是某一天的出纳日报。'),
        ('区间查询', '财务', '任意起止日期 → 按类别、按项目、按币种、按往来单位四块。'),
        ('核对表', '所有人', '录完看这一张就够了。勾稽必须全 0，待办是数据没填齐。')]:
    ln(r, a, b, c); r += 1
r += 1

sec(r, '三、汇率是怎么取的（这是外汇台账的核心）', C_SUB); r += 1
ln(r, '优先级 1', '实际成交汇率', '流水里填了「实际成交汇率」就用它。购汇、结汇、银行有实际成交价的都填这个。'); r += 1
ln(r, '优先级 2', '当月记账汇率', '按交易日期所在的月份，去【月度汇率】取。'); r += 1
ln(r, '优先级 3', '往前最近一个月', '当月还没填？自动往前找最近一个填过的月份，并在【核对表】点名提醒你补。'); r += 1
para(r, '　 原模板是「当月没填就留空」，而 12 个月里只填了 8 月一个月 —— 于是除了 8 月，'
        '整套人民币报表全是空的。这一条是这次改动里最关键的一处。'); r += 1
r += 1

sec(r, '四、两套人民币口径，各算各的，千万别混着加', C_SUB); r += 1
ln(r, '账面人民币成本', '历史成本', '＝期初人民币成本 + 每一笔原币 × 该笔当时的汇率。钱进来时值多少就是多少，以后汇率怎么变都不动它。'); r += 1
ln(r, '期末人民币(重估)', '现在值多少', '＝期末原币余额 × 期末汇率。'); r += 1
ln(r, '汇兑损益', '两者之差', '＝重估 − 账面成本。外币放着不动，汇率涨了就是赚、跌了就是亏，这一栏就是那个数。'); r += 1
para(r, '　 原模板把「期初余额 × 当前月汇率」和「每笔按当时汇率折的人民币」直接相加，'
        '两个口径混在一格里，加出来的数在会计上没有含义，而且会随报表月份来回跳。'); r += 1
r += 1

sec(r, '五、内部转账不算收入也不算支出', C_SUB); r += 1
ln(r, '怎么录', '两头都录', '转出账户记一笔支出，转入账户记一笔收入，类别都选「内部转账」。'
                        '购汇结汇选「购汇/结汇」，也是内部转账属性。'); r += 1
ln(r, '报表里怎么处理', '单列，不进合计', '月度汇报表单列「内部转入 / 内部转出」两列；'
                                   '分类报表单列一行「内部转账净额」。收入合计、支出合计里都不含它。'); r += 1
ln(r, '净额不是 0 怎么办', '那是换汇价差', '同币种之间调头寸净额应该是 0。购汇结汇跨币种，净额就是买卖价差，正常。'); r += 1
para(r, '　 原模板把「内部调拨」同时放进了收入清单和支出清单，收入和支出两头一起虚增；'
        '而且收入清单里「融资」「出口业务」各出现了两次、支出清单里「行政」出现了两次，'
        '这几项的金额在合计里被算了两遍。'); r += 1
r += 1

sec(r, '六、加东西怎么加'); r += 1
ln(r, '加一个账户', '两步', '① 在【参数设置】账户档案接着写一行；② 右键复制任意一张银行子表，'
                       '改表名，把 B2 改成新账户名。核对表第 6 项会盯着你有没有漏。'); r += 1
ln(r, '加一个币种', '一步', '在【参数设置】币种清单加一行，再去【月度汇率】填它的汇率。'
                       '（汇率表预留了 8 个币种的列）'); r += 1
ln(r, '加类别 / 项目 / 往来单位', '一步', '在【参数设置】对应清单接着写，所有报表自动跟着长。'
                                  '注意别写重复 —— 核对表会数。'); r += 1
ln(r, '行不够了', '告诉我', f'流水预留 {N_ROW} 行、账户 {N_ACC} 个、子表每张 {SUB} 行。'
                        '核对表第 16 项显示用了多少。'); r += 1
r += 1

sec(r, '七、颜色和符号', C_BASE); r += 1
ln(r, '淡黄色格子', '要你填', '这是人工输入区。'); r += 1
ln(r, '灰色格子', '公式自动算', '别手动改，改了就把公式顶掉了。'); r += 1
ln(r, '※ 红字', '是错，必须改', '账户不在档案、收支同时填、类别不在清单……'); r += 1
ln(r, '△ 琥珀字', '待补的数据', '汇率没填、内部转账没写对方账户……不影响公式，补齐了数才准。'); r += 1
r += 1

sec(r, '八、示例数据', C_BASE); r += 1
para(r, f'· 【流水录入】第 {R0} ~ {R0 + len(FLOWS) - 1} 行是示例（原模板那 15 笔 + 3 笔补充），'
        '正式用之前整段选中删掉即可。删了之后所有报表会自动归零。'); r += 1
para(r, '· 【月度汇率】只填了 2026 年 8 月的美元 7.18、港币 0.92 —— 这是原模板里给的，'
        '其余月份和欧元都空着，等你填。'); r += 1
para(r, '· 示例里特意留了一笔欧元收款，而欧元汇率空着，你可以看到「△EUR 还没填汇率」是怎么提示的 —— '
        '那不是坏了，是这套表在提醒你补数。'); r += 1
para(r, '· 顺带说清楚：某个币种一个汇率都没填时，它的余额**进不了人民币合计数**（宁可不算，也不瞎算），'
        '那个账户的人民币几列会是空白而不是 0。【核对表】第 11 项专门盯这件事。', 'C00000'); r += 1
para(r, '· 示例账户里有 1 个欧元户、2 个美元户、1 个港币户、2 个人民币户，'
        '正好把你说的美金 / 港币 / 欧元三种都覆盖到。'); r += 1

# ══════════════════════════════════════════════════════════════════════
for nm, ref in NAMES.items():
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
for s in wb.worksheets:
    s.page_setup.orientation = 'landscape'
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 0
    s.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
wb.save(OUT)
print('已生成:', OUT)
print('工作表 %d 张:' % len(wb.worksheets), ' / '.join(s.title for s in wb.worksheets))
n = sum(1 for s in wb.worksheets for row in s.iter_rows()
        for c in row if isinstance(c.value, str) and c.value.startswith('='))
print('公式格子 %d 个' % n)

sys.path.insert(0, os.path.join(os.path.dirname(ROOT), '工具'))
import dyn_array, fix_sheet_selection, check_formula                 # noqa: E402
dyn_array.install(OUT)              # 补回 metadata.xml + cm="1"，FILTER 才是真·动态数组
fix_sheet_selection.fix(OUT)        # 一张表都没选中的话 WPS 判成「工作组」
if check_formula.scan(OUT):
    raise SystemExit('公式括号不配对，先修了再交付')
