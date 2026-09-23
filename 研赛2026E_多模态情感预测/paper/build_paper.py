"""论文初稿生成器：paper/论文.md（正文模板）+ 各问题运行结果 → 论文初稿.docx。

正文模板用 Markdown 书写（公式用 LaTeX，pandoc 转成 Word 原生公式），其中可以引用运行结果：

1. 数值占位符  {{q2.metrics_test.complete.after_projection.acc3|pct2}}
   键名规则：<问题>.<结果文件名(不含扩展名)>.<JSON 内部路径>，列表下标用数字，如 q2.degradation_fit.T.beta。
   过滤器：f2/f3/f4（小数位）、pct1/pct2（百分数）、int、raw。另有 derived.* 来自 paper/derived.py 的派生量。
   取不到的值渲染为醒目的【待填：键名】，便于在拿到真实结果前先检查版面。
2. 表格块（从 CSV / JSON 列表生成三线表）：
   ```table
   source: q2/ablation.csv            # 相对于 --q2 运行目录；也可 q1/… q3/…
   query: protocol == "mixed"         # 可选，pandas query
   columns: variant=模型, acc3=Acc, f1_weighted=F1   # 可选，选列并改名（顺序即列序）
   format: acc3=f4, f1_weighted=f4    # 可选
   caption: 表 6  消融实验结果
   max_rows: 120                      # 可选
   ```
3. 图片  ![图 5  缺失率–性能曲线](fig:q2/figures/missing_type_ratio.png){width=95%}
   找不到图片时渲染为【图待生成：路径】段落。

用法：
    python -m paper.build_paper --q1 outputs/q1/run1 --q2 outputs/q2/run1 --q3 outputs/q3/run1 \
        --out 论文初稿.docx
    （缺哪个就不传哪个；--draft-note "合成数据测试" 会在页眉打上醒目标注）
"""
from __future__ import annotations

import argparse
import importlib.util
import json
import math
import re
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
PLACEHOLDER = re.compile(r"\{\{\s*([^}|]+?)\s*(?:\|\s*([a-z0-9_]+)\s*)?\}\}")
TABLE_BLOCK = re.compile(r"```table\n(.*?)```", re.S)
FIG = re.compile(r"!\[([^\]]*)\]\(fig:([^)]+)\)(\{[^}]*\})?")


# ----------------------------------------------------------------------------- 结果上下文

def _flatten(obj, prefix: str, out: dict) -> None:
    if isinstance(obj, dict):
        for k, v in obj.items():
            _flatten(v, f"{prefix}.{k}" if prefix else str(k), out)
    elif isinstance(obj, list):
        out[prefix] = obj
        for i, v in enumerate(obj):
            _flatten(v, f"{prefix}.{i}", out)
    else:
        out[prefix] = obj


def load_context(run_dirs: dict[str, Path | None]) -> dict:
    """把每个运行目录下（含一级子目录）的所有 JSON 展平成 '问题.文件名.路径' → 值。"""
    ctx: dict = {}
    for q, d in run_dirs.items():
        if d is None or not Path(d).exists():
            continue
        for p in sorted(Path(d).glob("*.json")) + sorted(Path(d).glob("*/*.json")):
            if p.parent.name in ("alignment", "cards"):
                continue
            try:
                data = json.loads(p.read_text(encoding="utf-8"))
            except Exception:  # noqa: BLE001
                continue
            _flatten(data, f"{q}.{p.stem}", ctx)
    return ctx


def load_derived(ctx: dict, run_dirs: dict) -> None:
    """paper/derived.py 若存在，调用其 compute(ctx, run_dirs) -> dict，结果以 derived.* 注入。"""
    p = HERE / "derived.py"
    if not p.exists():
        return
    spec = importlib.util.spec_from_file_location("paper_derived", p)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    try:
        extra = mod.compute(ctx, run_dirs) or {}
    except Exception as e:  # noqa: BLE001 - 派生量失败不能拖垮整篇论文
        print(f"[warn] derived.compute 失败：{e}", file=sys.stderr)
        extra = {}
    for k, v in extra.items():
        _flatten(v, f"derived.{k}", ctx)


def fmt_value(v, f: str | None) -> str:
    if v is None or (isinstance(v, float) and (math.isnan(v) or math.isinf(v))):
        return "—"
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False)
    if f in (None, "raw"):
        if isinstance(v, float):
            return f"{v:.4f}"
        return str(v)
    try:
        x = float(v)
    except (TypeError, ValueError):
        return str(v)
    if f == "int":
        return f"{int(round(x))}"
    if f.startswith("pct"):
        nd = int(f[3:] or 1)
        return f"{x * 100:.{nd}f}%"
    if f.startswith("f"):
        return f"{x:.{int(f[1:])}f}"
    if f.startswith("e"):
        return f"{x:.{int(f[1:])}e}"
    return str(v)


def sub_placeholders(text: str, ctx: dict, missing: list) -> str:
    def rep(m):
        key, f = m.group(1).strip(), m.group(2)
        if key not in ctx:
            missing.append(key)
            return f"**【待填：{key}】**"
        return fmt_value(ctx[key], f)

    return PLACEHOLDER.sub(rep, text)


# ----------------------------------------------------------------------------- 表格与图片

def _parse_block(body: str) -> dict:
    spec = {}
    for line in body.strip().splitlines():
        line = line.split("  #", 1)[0].rstrip()
        if not line.strip() or ":" not in line:
            continue
        k, v = line.split(":", 1)
        spec[k.strip()] = v.strip()
    return spec


def _kv_list(s: str) -> list[tuple[str, str]]:
    out = []
    for part in s.split(","):
        part = part.strip()
        if not part:
            continue
        if "=" in part:
            a, b = part.split("=", 1)
            out.append((a.strip(), b.strip()))
        else:
            out.append((part, part))
    return out


def _md_escape(x: str) -> str:
    return str(x).replace("|", "\\|").replace("\n", " ")


def render_table(spec: dict, run_dirs: dict, missing: list) -> str:
    import pandas as pd

    caption = spec.get("caption", "")
    src = spec.get("source", "")
    q, _, rel = src.partition("/")
    base = run_dirs.get(q)
    cols = _kv_list(spec["columns"]) if "columns" in spec else None
    path = Path(base) / rel if base else None
    df = None
    if path is not None and path.exists():
        if path.suffix == ".csv":
            df = pd.read_csv(path, encoding="utf-8-sig")
        elif path.suffix in (".xlsx", ".xls"):
            df = pd.read_excel(path)
        elif path.suffix == ".json":
            data = json.loads(path.read_text(encoding="utf-8"))
            key = spec.get("json_key")
            if key:
                for k in key.split("."):
                    data = data[int(k)] if isinstance(data, list) else data[k]
            df = pd.DataFrame(data if isinstance(data, list) else [data])
    if df is None:
        missing.append(f"table:{src}")
        heads = [b for _, b in cols] if cols else ["（待生成）"]
        lines = [f"Table: {caption}", "", "| " + " | ".join(heads) + " |", "|" + "|".join(["---"] * len(heads)) + "|",
                 "| " + " | ".join([f"【待运行后生成：{src}】"] + [""] * (len(heads) - 1)) + " |", ""]
        return "\n".join(lines)
    if "query" in spec:
        df = df.query(spec["query"])
    if "sort" in spec:
        key, _, order = spec["sort"].partition(" ")
        df = df.sort_values(key, ascending=order.strip().lower() != "desc")
    if cols:
        keep = [(a, b) for a, b in cols if a in df.columns]
        df = df[[a for a, _ in keep]].rename(columns=dict(keep))
    fmts = dict(_kv_list(spec.get("format", "")))
    rename = dict(cols) if cols else {}
    for a, f in fmts.items():
        c = rename.get(a, a)
        if c in df.columns:
            df[c] = df[c].map(lambda v, f=f: fmt_value(v, f))
    max_rows = int(spec.get("max_rows", 400))
    note = ""
    if len(df) > max_rows:
        note = f"（共 {len(df)} 行，此处列出前 {max_rows} 行，全表见附件。）"
        df = df.head(max_rows)
    heads = [str(c) for c in df.columns]
    lines = [f"Table: {caption}{note}", "", "| " + " | ".join(_md_escape(h) for h in heads) + " |",
             "|" + "|".join([":---:"] * len(heads)) + "|"]
    for _, row in df.iterrows():
        cells = []
        for v in row.tolist():
            cells.append(_md_escape(fmt_value(v, None) if isinstance(v, float) else v))
        lines.append("| " + " | ".join(cells) + " |")
    lines.append("")
    return "\n".join(lines)


def render_figures(text: str, run_dirs: dict, missing: list) -> str:
    def rep(m):
        cap, rel, attr = m.group(1), m.group(2), m.group(3) or "{width=90%}"
        q, _, r = rel.partition("/")
        base = run_dirs.get(q)
        p = (Path(base) / r) if base else None
        if p is None or not p.exists():
            missing.append(f"fig:{rel}")
            return f"\n**【图待生成：{rel}】** {cap}\n"
        return f"![{cap}]({p.resolve().as_posix()}){attr}"

    return FIG.sub(rep, text)


# ----------------------------------------------------------------------------- Word 样式

def make_reference_docx(pandoc: str, path: Path) -> None:
    """生成 pandoc 参考样式文档，并改成中文论文常用版式：宋体/Times New Roman 小四、1.5 倍行距、首行缩进 2 字符、黑体标题。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    subprocess.run([pandoc, "-o", str(path), "--print-default-data-file", "reference.docx"], check=True)
    doc = Document(str(path))

    def set_font(style, east: str, west: str, size: float | None, bold: bool | None = None):
        style.font.name = west
        if size:
            style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(a), west)
        rfonts.set(qn("w:eastAsia"), east)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(a)) is not None:
                del rfonts.attrib[qn(a)]
        style.font.color.rgb = None

    styles = doc.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        if name in [s.name for s in styles]:
            st = styles[name]
            set_font(st, "宋体", "Times New Roman", 12)
            pf = st.paragraph_format
            pf.line_spacing = 1.5
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            if name in ("Body Text", "First Paragraph"):
                pf.first_line_indent = Pt(24)
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 14), ("Heading 3", 12), ("Heading 4", 12)):
        if name in [s.name for s in styles]:
            st = styles[name]
            set_font(st, "黑体", "Times New Roman", size, True)
            st.paragraph_format.space_before = Pt(12 if name != "Heading 3" else 6)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.line_spacing = 1.3
            if name in ("Title", "Heading 1"):
                st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ("Image Caption", "Table Caption", "Caption"):
        if name in [s.name for s in styles]:
            st = styles[name]
            set_font(st, "宋体", "Times New Roman", 10.5, False)
            st.font.italic = False
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            st.paragraph_format.space_before = Pt(3)
            st.paragraph_format.space_after = Pt(6)
    if "Abstract" in [s.name for s in styles]:
        set_font(styles["Abstract"], "宋体", "Times New Roman", 12)
    # A4 页面、常规页边距
    from docx.shared import Cm

    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.5)
        sec.top_margin = sec.bottom_margin = Cm(2.5)
        sec.header_distance = sec.footer_distance = Cm(1.5)
        sec.gutter = Cm(0)
    doc.save(str(path))


def _insert_ordered(parent, el, later_tags) -> None:
    """按 OOXML 模式规定的子元素顺序插入：放在第一个"应当排在它后面"的兄弟元素之前。"""
    from docx.oxml.ns import qn

    for tag in later_tags:
        sib = parent.find(qn(f"w:{tag}"))
        if sib is not None:
            sib.addprevious(el)
            return
    parent.append(el)


def postprocess_docx(path: Path, draft_note: str | None) -> None:
    """三线表、表格字号、图片居中、页眉标注、页码。"""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document(str(path))

    def border(tag, sz):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "single" if sz else "nil")
        if sz:
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        return el

    for tbl in doc.tables:
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = tbl._tbl.tblPr
        old = tblPr.find(qn("w:tblBorders"))
        if old is not None:
            tblPr.remove(old)
        b = OxmlElement("w:tblBorders")
        for tag, sz in (("top", 12), ("left", 0), ("bottom", 12), ("right", 0), ("insideH", 0), ("insideV", 0)):
            b.append(border(tag, sz))
        _insert_ordered(tblPr, b, ("shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"))
        if tbl.rows:
            for cell in tbl.rows[0].cells:
                tcPr = cell._tc.get_or_add_tcPr()
                old = tcPr.find(qn("w:tcBorders"))
                if old is not None:
                    tcPr.remove(old)
                tb = OxmlElement("w:tcBorders")
                tb.append(border("bottom", 6))
                _insert_ordered(tcPr, tb, ("shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"))
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.line_spacing = 1.15
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.size = Pt(9)
    for p in doc.paragraphs:
        if p._p.xpath(".//pic:pic") or p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
        if "【待填" in p.text or "【图待生成" in p.text or "【待运行后生成" in p.text:
            for r in p.runs:
                if "【" in r.text:
                    r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
    hp.text = draft_note or ""
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in hp.runs:
        r.font.size = Pt(9)
        if draft_note:
            r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
            r.font.bold = True
    # 页脚页码
    fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    doc.save(str(path))


# ----------------------------------------------------------------------------- 主流程

def build(md_path: Path, run_dirs: dict, out: Path, draft_note: str | None = None, keep_md: bool = True) -> list:
    import pypandoc

    pandoc = pypandoc.get_pandoc_path()
    ctx = load_context(run_dirs)
    load_derived(ctx, run_dirs)
    text = md_path.read_text(encoding="utf-8")
    missing: list = []
    text = TABLE_BLOCK.sub(lambda m: render_table(_parse_block(m.group(1)), run_dirs, missing), text)
    text = render_figures(text, run_dirs, missing)
    text = sub_placeholders(text, ctx, missing)
    out.parent.mkdir(parents=True, exist_ok=True)
    md_out = out.with_suffix(".rendered.md")
    md_out.write_text(text, encoding="utf-8")
    with tempfile.TemporaryDirectory() as td:
        ref = Path(td) / "reference.docx"
        make_reference_docx(pandoc, ref)
        cmd = [pandoc, str(md_out), "-f", "markdown+tex_math_dollars+pipe_tables+table_captions+implicit_figures",
               "-t", "docx", "--reference-doc", str(ref), "-o", str(out), "--resource-path", str(ROOT)]
        subprocess.run(cmd, check=True)
    postprocess_docx(out, draft_note)
    if not keep_md:
        md_out.unlink(missing_ok=True)
    return missing


def main():
    ap = argparse.ArgumentParser(description="由 paper/论文.md 与运行结果生成论文初稿 docx")
    ap.add_argument("--md", default=str(HERE / "论文.md"))
    ap.add_argument("--q1", default=None)
    ap.add_argument("--q2", default=None)
    ap.add_argument("--q3", default=None)
    ap.add_argument("--out", default=str(ROOT / "论文初稿.docx"))
    ap.add_argument("--draft-note", default=None, help="页眉醒目标注，如 '合成数据测试版，数值无效'")
    a = ap.parse_args()
    run_dirs = {"q1": Path(a.q1) if a.q1 else None, "q2": Path(a.q2) if a.q2 else None,
                "q3": Path(a.q3) if a.q3 else None}
    missing = build(Path(a.md), run_dirs, Path(a.out), a.draft_note)
    uniq = sorted(set(missing))
    print(f"已生成：{a.out}")
    print(f"未填占位：{len(uniq)} 处" + ("" if not uniq else "（前 30 个）\n  " + "\n  ".join(uniq[:30])))


if __name__ == "__main__":
    main()
