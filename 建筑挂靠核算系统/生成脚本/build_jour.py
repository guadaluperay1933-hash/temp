# -*- coding: utf-8 -*-
"""生成《出纳资金日记账.xlsx》—— 出纳单独用这一本登记，录完整块复制回主表。

列序跟主册【资金日记账】**一模一样**（A 序号 / B 日期 / C 资金账户 / D 项目编号 /
E 费用类型 / F 往来单位 / G 经办人 / H 摘要 / I 收入 / J 支出 / K 该账户余额 /
L 三账户合计 / M 工程回款 / N 校验 / O 项目简称），
所以回流就是「选中 B:J 和 M 两块 → 复制 → 到主册选择性粘贴数值」，不用对列。

项目编号、往来单位这些下拉，在这本册子里是**本地清单**（从主册抄成死值放在
【对照清单】那张表），这样出纳电脑上没有主册也能正常用下拉。

跑法：python3 build_jour.py [主册xlsx] [输出xlsx]
"""
import os, sys, datetime as dt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

HERE = os.path.dirname(os.path.abspath(__file__))
SRC  = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '..', '建筑挂靠业务核算系统.xlsx')
OUT  = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, '..', '出纳资金日记账.xlsx')

YH = '微软雅黑'
F_T  = Font(name=YH, sz=10)
F_H  = Font(name=YH, sz=10, bold=True, color='FFFFFFFF')
F_H2 = Font(name=YH, sz=10, bold=True, color='FF1F4E79')
F_IN = Font(name=YH, sz=10, bold=True, color='FF1F4E79')
F_AU = Font(name=YH, sz=10, color='FF7F7F7F')
F_NT = Font(name=YH, sz=9,  color='FF808080')
F_TT = Font(name=YH, sz=16, bold=True, color='FF1F4E79')
FILL_IN  = PatternFill('solid', fgColor='FFFFF7E0')
FILL_AU  = PatternFill('solid', fgColor='FFF2F2F2')
FILL_HD  = PatternFill('solid', fgColor='FF4472C4')
FILL_HD2 = PatternFill('solid', fgColor='FF8EA9DB')
FILL_CHK = PatternFill('solid', fgColor='FFE2EFDA')
FILL_WRN = PatternFill('solid', fgColor='FFFFC7CE')
FILL_NT  = PatternFill('solid', fgColor='FFFCE4D6')
_t = Side(style='thin', color='FFBFBFBF')
BD = Border(left=_t, right=_t, top=_t, bottom=_t)
C  = Alignment(horizontal='center', vertical='center')
CW = Alignment(horizontal='center', vertical='center', wrap_text=True)
LF = Alignment(horizontal='left', vertical='center')
LW = Alignment(horizontal='left', vertical='center', wrap_text=True)
MONEY = '"¥"#,##0.00_);[Red]\\("¥"#,##0.00\\)'
DATE  = 'yyyy/mm/dd'

def put(ws, coord, v, font=None, fill=None, fmt=None, align=None, border=True):
    c = ws[coord]; c.value = v
    if font: c.font = font
    if fill: c.fill = fill
    if fmt:  c.number_format = fmt
    if align: c.alignment = align
    if border: c.border = BD
    return c

# ---------------------------------------------------------------- 从主册抄清单
print('读取主册', os.path.abspath(SRC))
src = openpyxl.load_workbook(SRC, data_only=True)
sp, su, sj = src['项目档案'], src['单位档案'], src['资金日记账']

PROJ = []                       # (编号, 简称, 全称)
for r in range(6, 400):
    code = sp.cell(r, 1).value
    if not code: continue
    PROJ.append((str(code), str(sp.cell(r, 3).value or ''), str(sp.cell(r, 2).value or '')))
UNITS = [str(su.cell(r, 1).value) for r in range(6, 60) if su.cell(r, 1).value]
ACCTS = ['泓普', '仟茂', '现金']
ETYPES = ['工程回款','货款','借款','还借款','工资','社保','材料费','运费','餐费','耗材费','办公费用',
          '维修费','青苗费','饮用水费用','转备用金','手续费','燃油费','税费','管理费','租赁费',
          '考核款','账户年费','其他应收收回','其他应付支付','其他']
# 主册【资金日记账】第 2 行的三个期初余额
OPEN = {}
for i, a in enumerate(ACCTS):
    cell = sj.cell(2, 3 + i * 2).value
    OPEN[a] = float(cell) if isinstance(cell, (int, float)) else 0.0
print(f'  项目 {len(PROJ)} 个 / 往来单位 {len(UNITS)} 家 / 期初 ' +
      '、'.join(f'{a} {OPEN[a]:,.2f}' for a in ACCTS))

J0, JN = 5, 800
J1 = J0 + JN - 1
wb = openpyxl.Workbook()

# ---------------------------------------------------------------- ② 对照清单（先建，下拉要指过来）
ref = wb.active; ref.title = '对照清单'
put(ref, 'A1', '对　照　清　单（从主册抄过来的死值，出纳不用改）', font=F_TT, align=LF, border=False)
ref.merge_cells('A1:F1')
put(ref, 'A2', '★ 这张表只是给左边【资金日记账】的下拉用的。主册那边加了新项目或新单位，'
               '让财务重新发一本给你，或者直接在这里照着补一行。',
    font=F_NT, align=LW, border=False)
ref.merge_cells('A2:F2'); ref.row_dimensions[2].height = 30
for i, t in enumerate(['项目编号', '项目简称', '项目全称', '往来单位', '费用类型', '资金账户']):
    put(ref, f'{chr(65+i)}4', t, font=F_H, fill=FILL_HD, align=C)
for i, (code, short, full) in enumerate(PROJ):
    put(ref, f'A{5+i}', code, font=F_T, align=C)
    put(ref, f'B{5+i}', short, font=F_T, align=LF)
    put(ref, f'C{5+i}', full, font=F_T, align=LF)
for i, u in enumerate(UNITS):   put(ref, f'D{5+i}', u, font=F_T, align=C)
for i, e in enumerate(ETYPES):  put(ref, f'E{5+i}', e, font=F_T, align=C)
for i, a in enumerate(ACCTS):   put(ref, f'F{5+i}', a, font=F_T, align=C)
for col, w in zip('ABCDEF', (10, 16, 46, 14, 14, 10)): ref.column_dimensions[col].width = w
ref.freeze_panes = 'A5'

P_CODE = f"'对照清单'!$A$5:$A${4+len(PROJ)}"
P_SN   = f"'对照清单'!$B$5:$B${4+len(PROJ)}"
U_NAME = f"'对照清单'!$D$5:$D${4+len(UNITS)}"
E_LIST = f"'对照清单'!$E$5:$E${4+len(ETYPES)}"
A_LIST = f"'对照清单'!$F$5:$F${4+len(ACCTS)}"

# ---------------------------------------------------------------- ① 资金日记账
ws = wb.create_sheet('资金日记账', 0)
put(ws, 'A1', '出　纳　资　金　日　记　账（泓普 · 仟茂 · 现金 混合录入）',
    font=F_TT, align=LF, border=False)
ws.merge_cells('A1:Q1'); ws.row_dimensions[1].height = 28

put(ws, 'A2', '期初余额', font=F_H2, align=Alignment(horizontal='right', vertical='center'), border=False)
for i, a in enumerate(ACCTS):
    put(ws, f'{chr(66+i*2)}2', a, font=F_NT,
        align=Alignment(horizontal='right', vertical='center'), border=False)
    put(ws, f'{chr(67+i*2)}2', OPEN[a], font=F_IN, fill=FILL_IN, fmt=MONEY, align=C)
OPB = {a: f'${chr(67+i*2)}$2' for i, a in enumerate(ACCTS)}
OPB_ALL = '+'.join(OPB.values())

put(ws, 'H2',
    '三个账户混在一起、按日期顺着往下录。淡黄色格子是你要填的，灰色的自动算。\n'
    '录完怎么交回财务：选中 B 列到 J 列这一整块 → 复制 → 打开主册《建筑挂靠业务核算系统》的'
    '【资金日记账】→ 点到第一行空白行的 B 列 → 右键「选择性粘贴 → 数值」；'
    '有工程回款的，再把 M 列同样粘一次。两边列序完全一样，不用对列。',
    font=F_NT, align=LW, border=False)
ws.merge_cells('H2:Q2'); ws.row_dimensions[2].height = 46

put(ws, 'A3', '↓ 淡黄色 = 出纳要填的', font=F_NT, fill=FILL_IN, align=C)
ws.merge_cells('A3:J3')
put(ws, 'K3', '↓ 灰色 = 自动算的，别手工改', font=F_NT, fill=FILL_AU, align=C)
ws.merge_cells('K3:Q3')

HEAD = ['序号','日期','资金账户','项目编号','费用类型','往来单位','经办人','摘要','收入','支出',
        '该账户余额','三账户合计','工程\n回款','校验','项目简称','年月','年度']
for i, t in enumerate(HEAD):
    put(ws, f'{openpyxl.utils.get_column_letter(i+1)}4', t,
        font=F_H, fill=(FILL_HD if i < 10 else FILL_HD2), align=CW)
ws.row_dimensions[4].height = 30

for r in range(J0, J1 + 1):
    for col, fmt, txt in (('B', DATE, 0), ('C', None, 0), ('D', None, 0), ('E', None, 0),
                          ('F', None, 0), ('G', None, 0), ('H', None, 1),
                          ('I', MONEY, 0), ('J', MONEY, 0)):
        put(ws, f'{col}{r}', None, font=F_IN, fill=FILL_IN, fmt=fmt, align=LF if txt else C)
    put(ws, f'M{r}', None, font=F_IN, fill=FILL_IN, align=C)
    put(ws, f'A{r}', f'=IF($B{r}="","",ROW()-{J0-1})', font=F_AU, fill=FILL_AU, align=C)
    put(ws, f'K{r}',
        f'=IF($B{r}="","",IF($C{r}="泓普",{OPB["泓普"]},IF($C{r}="仟茂",{OPB["仟茂"]},'
        f'IF($C{r}="现金",{OPB["现金"]},0)))'
        f'+SUMIFS($I${J0}:$I{r},$C${J0}:$C{r},$C{r})-SUMIFS($J${J0}:$J{r},$C${J0}:$C{r},$C{r}))',
        font=F_AU, fill=FILL_AU, fmt=MONEY, align=C)
    put(ws, f'L{r}', f'=IF($B{r}="","",{OPB_ALL}+SUM($I${J0}:$I{r})-SUM($J${J0}:$J{r}))',
        font=F_AU, fill=FILL_AU, fmt=MONEY, align=C)
    put(ws, f'O{r}', f'=IF($D{r}="","",IFERROR(INDEX({P_SN},MATCH($D{r},{P_CODE},0)),"⚠编号不存在"))',
        font=F_AU, fill=FILL_AU, align=LF)
    put(ws, f'P{r}', f'=IF($B{r}="","",IF(ISNUMBER($B{r}),TEXT($B{r},"yyyy-mm"),"日期非法"))',
        font=F_AU, fill=FILL_AU, align=C)
    put(ws, f'Q{r}', f'=IF($B{r}="","",IF(ISNUMBER($B{r}),YEAR($B{r}),"日期非法"))',
        font=F_AU, fill=FILL_AU, align=C)
    put(ws, f'N{r}',
        f'=IF($B{r}="","",'
        f'IF(NOT(ISNUMBER($B{r})),"日期格式不对",'
        f'IF($C{r}="","未选资金账户",'
        f'IF(AND(N($I{r})=0,N($J{r})=0),"收支都为空",'
        f'IF(AND(N($I{r})>0,N($J{r})>0),"收支不能同时填",'
        f'IF(AND($D{r}<>"",LEFT($O{r},1)="⚠"),"项目编号不存在",'
        f'IF(AND($M{r}="是",$F{r}=""),"工程回款必须填往来单位","√")))))))',
        font=F_T, fill=FILL_CHK, align=C)
    ws.row_dimensions[r].height = 16

for sq, f1 in ((f'C{J0}:C{J1}', f'={A_LIST}'), (f'D{J0}:D{J1}', f'={P_CODE}'),
               (f'E{J0}:E{J1}', f'={E_LIST}'), (f'F{J0}:F{J1}', f'={U_NAME}'),
               (f'M{J0}:M{J1}', '"是,否"')):
    dv = DataValidation(type='list', formula1=f1, allow_blank=True, showDropDown=False)
    dv.error = '请从下拉里选'; dv.errorTitle = '不在清单里'
    ws.add_data_validation(dv); dv.add(sq)
for sq in (f'I{J0}:I{J1}', f'J{J0}:J{J1}'):
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
    dv.error = '金额要填 0 或正数；付出去的填「支出」那一列。'; dv.errorTitle = '金额不对'
    ws.add_data_validation(dv); dv.add(sq)

ws.conditional_formatting.add(f'A{J0}:Q{J1}',
    FormulaRule(formula=[f'AND($N{J0}<>"",$N{J0}<>"√")'], fill=FILL_WRN))
for col, w in zip('ABCDEFGHIJKLMNOPQ',
                  (7, 11, 10, 10, 13, 12, 10, 40, 13, 13, 14, 14, 9, 15, 26, 9, 7)):
    ws.column_dimensions[col].width = w
ws.auto_filter.ref = f'A4:Q{J1}'
ws.freeze_panes = 'C5'

put(ws, f'A{J1+2}',
    '交回财务的步骤：① 选中 B5:J' + str(J1) + ' 整块 → 复制；② 打开主册《建筑挂靠业务核算系统》'
    '的【资金日记账】，点到第一行空白行的 B 列 → 右键「选择性粘贴 → 数值」；'
    '③ 有工程回款的，再把 M5:M' + str(J1) + ' 同样粘到主册 M 列对应位置；'
    '④ 粘完在主册上看一眼「校验」列全是 √ 就完事。'
    '⑤ 本册不参与主册的任何汇总，主册只认它自己那张【资金日记账】，所以不会重复算。',
    font=F_NT, fill=FILL_NT, align=LW)
ws.merge_cells(f'A{J1+2}:Q{J1+2}')
ws.row_dimensions[J1+2].height = 46

wb.save(OUT)
nf = sum(1 for s_ in wb.worksheets for row in s_.iter_rows()
         for c in row if isinstance(c.value, str) and c.value.startswith('='))
print('已保存', os.path.abspath(OUT))
print(f'  工作表 {len(wb.sheetnames)} 张（资金日记账 {JN} 行 + 对照清单），公式 {nf} 个')
