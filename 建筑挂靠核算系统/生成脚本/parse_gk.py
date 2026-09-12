# -*- coding: utf-8 -*-
"""把 8 张手工对账表逐行还原成业务事件"""
import openpyxl, re, json, datetime as dt
SRC='/root/.claude/uploads/9101e7cb-4f82-53a7-bbe8-41b0f93b12b0/c2938ae7-____9.041.xlsx'

# 每张表的列映射：sheet -> {字段: 列号}
M = {
 '德誉嘉 ': dict(unit='德誉嘉', proj=2, memo=3, biz=4, inv=5, to=6, sale=7, mfee=8, cost_due=9,
                cost_done=10, recv_up=13, mfee_set=16, recv_us=17),
 '迅驰':   dict(unit='迅驰', proj=2, memo=3, biz=4, inv=5, to=6, sale=7, mfee=8, cost_due=9,
                cost_done=10, tax_v=12, tax_s=13, tax_y=14, tax_i=15, tax_due=16, tax_paid=17,
                recv_up=20, mfee_set=23, bond=24, recv_us=25),
 '华城':   dict(unit='华城', proj=2, memo=3, biz=4, inv=5, to=6, sale=7, mfee=8, cost_due=9,
                cost_done=10, tax_v=12, tax_s=13, tax_y=14, tax_i=15, tax_due=16, tax_paid=17,
                recv_up=20, mfee_set=23, recv_us=24),
 '康欣':   dict(unit='康欣', proj=2, memo=3, biz=4, inv=5, to=6, sale=7, mfee=8, cost_due=9,
                cost_done=10, tax_v=12, tax_s=13, tax_y=14, tax_i=15, tax_due=16, tax_paid=17,
                recv_up=20, bond=23, mfee_set=24, recv_us=25),
 '金沁':   dict(unit='金沁', proj=2, memo=3, biz=4, inv=5, to=6, sale=7, mfee=8, cost_due=9,
                cost_done=10, tax_v=12, tax_s=13, tax_y=14, tax_i=15, tax_due=16, tax_paid=17,
                bond_up=20, recv_up=21, bond=24, mfee_set=25, recv_us=26),
 '安锐':   dict(unit='安锐', proj=2, memo=3, biz=5, inv=4, to=6, sale=7, mfee=8, cost_due=9,
                cost_done=10, wage=11, tax_v=13, tax_s=14, tax_y=15, tax_i=16, tax_due=17,
                tax_paid=18, recv_up=21, mfee_set=24, bond=25, recv_us=26),
 '湖南锦泰': dict(unit='湖南锦泰', proj=2, memo=3, biz=4, inv=5, to=6, sale=7, mfee=8, cost_due=9,
                cost_done=10, wage=11, tax_v=13, tax_s=14, tax_y=15, tax_i=16, tax_due=17,
                tax_paid=18, recv_up=21, mfee_set=24, bond=25, recv_us=26),
 '大太线老旧线路': dict(unit='杰华', proj=2, memo=3, biz=4, inv=5, to=6, sale=7, mfee=9, cost_due=10,
                cost_done=11, wage=12, tax_v=14, tax_s=15, tax_y=16, tax_i=17, tax_due=18,
                tax_paid=19, recv_up=22, mfee_set=25, bond=26, recv_us=27),
}
OWNERS = {'民能','铜梁供电','重庆民能','国网','铜梁供电分公司'}
US = {'泓普','仟茂'}
NORM = {'德誉佳':'德誉嘉','德誉嘉':'德誉嘉','金沁':'金沁','康欣':'康欣','迅驰':'迅驰',
        '华城':'华城','安锐':'安锐','湖南锦泰':'湖南锦泰','杰华':'杰华',
        '民能':'民能','铜梁供电':'铜梁供电','泓普':'泓普','仟茂':'仟茂'}

def num(ws,r,c):
    if not c: return 0.0
    v=ws.cell(row=r,column=c).value
    return float(v) if isinstance(v,(int,float)) else 0.0
def txt(ws,r,c):
    if not c: return ''
    v=ws.cell(row=r,column=c).value
    return str(v).strip() if v is not None else ''
ALLU = ['泓普','仟茂','民能','铜梁供电','德誉佳','德誉嘉','迅驰','华城','康欣','金沁','安锐','湖南锦泰','杰华']
def parse_pair(memo):
    """从摘要里抓「X开…到Y」的开票方与收票方；原表收票单位列填错时以摘要为准"""
    if not memo: return None, None
    hits = []
    for u in ALLU:
        i = memo.find(u)
        while i >= 0:
            hits.append((i, NORM.get(u, u)))
            i = memo.find(u, i + 1)
    hits.sort()
    seen = []
    for _, u in hits:
        if not seen or seen[-1] != u: seen.append(u)
    if len(seen) >= 2: return seen[0], seen[1]
    if len(seen) == 1: return seen[0], None
    return None, None

def norm_unit(s):
    s=(s or '').strip()
    for k,v in NORM.items():
        if k in s: return v
    return s
DATE_RE=re.compile(r'(20\d{2})[.\-年/]\s*(\d{1,2})[.\-月/]\s*(\d{1,2})')
def grab_date(memo, fallback):
    m=DATE_RE.search(memo or '')
    if m:
        y,mo,d=int(m.group(1)),int(m.group(2)),int(m.group(3))
        try: return dt.date(y,mo,d)
        except ValueError: pass
    m2=re.search(r'(20\d{2})[.\-年/]\s*(\d{1,2})月', memo or '')
    if m2:
        try: return dt.date(int(m2.group(1)),int(m2.group(2)),1)
        except ValueError: pass
    return fallback

wb=openpyxl.load_workbook(SRC, data_only=True)
events=[]; projects={}; last_proj={}
for sh,mp in M.items():
    ws=wb[sh]; unit=mp['unit']; cur_proj=''; fb=dt.date(2026,1,1)
    for r in range(5, ws.max_row+1):
        proj=txt(ws,r,mp['proj'])
        memo=txt(ws,r,mp['memo'])
        if proj: cur_proj=proj; projects.setdefault(proj,set()).add(unit)
        if not memo and not proj: continue
        d=grab_date(memo, fb); fb=d
        to=norm_unit(txt(ws,r,mp.get('to')))
        inv=txt(ws,r,mp.get('inv'))
        biz=txt(ws,r,mp.get('biz'))
        base=dict(date=d.isoformat(), unit=unit, proj=cur_proj, memo=memo, inv=inv, biz=biz, src=f'{sh}!{r}')
        sale=num(ws,r,mp['sale']); cost=num(ws,r,mp['cost_done'])
        m_iss, m_rcv = parse_pair(memo)
        if abs(sale)>0.004:
            si = m_iss if (m_iss and m_rcv) else unit
            sr = m_rcv if (m_iss and m_rcv) else (to or '民能')
            events.append({**base, 'kind':'销项开票', 'payer':si, 'payee':sr,
                           'amt':round(sale,2), 'mfee':round(num(ws,r,mp['mfee']),2),
                           'cost_due':round(num(ws,r,mp['cost_due']),2)})
        # 同一行两列填的是同一张票时不要记两遍
        if abs(cost)>0.004 and abs(cost-sale)>0.004:
            iss = m_iss or '泓普'
            rcv = m_rcv or to or unit
            kind = '成本票' if iss in US else '销项开票'
            events.append({**base, 'kind':kind, 'payer':iss, 'payee':rcv,
                           'amt':round(cost,2), 'mfee':0, 'cost_due':0, 'from_cost_col':True})
        for key,label in [('tax_paid','已交税')]:
            v=num(ws,r,mp.get(key))
            if abs(v)>0.004:
                events.append({**base,'kind':label,'payer':unit,'payee':'税局','amt':round(v,2),'mfee':0,'cost_due':0})
        tv=sum(num(ws,r,mp.get(k)) for k in ('tax_v','tax_s','tax_y','tax_i'))
        if abs(tv)>0.004:
            events.append({**base,'kind':'预提税费','payer':unit,'payee':'税局','amt':round(tv,2),
                           'mfee':0,'cost_due':0,
                           'tax_v':round(num(ws,r,mp.get('tax_v')),2),'tax_s':round(num(ws,r,mp.get('tax_s')),2),
                           'tax_y':round(num(ws,r,mp.get('tax_y')),2),'tax_i':round(num(ws,r,mp.get('tax_i')),2)})
        ru=num(ws,r,mp.get('recv_up'))
        if abs(ru)>0.004:
            events.append({**base,'kind':'挂靠单位代收','payer':'民能','payee':unit,'amt':round(ru,2),'mfee':0,'cost_due':0})
        rs=num(ws,r,mp.get('recv_us'))
        if abs(rs)>0.004:
            events.append({**base,'kind':'我方收款','payer':unit,'payee':'泓普','amt':round(rs,2),'mfee':0,'cost_due':0})
        ms=num(ws,r,mp.get('mfee_set'))
        if abs(ms)>0.004:
            events.append({**base,'kind':'管理费结算','payer':'泓普','payee':unit,'amt':round(ms,2),'mfee':0,'cost_due':0})
        bd=num(ws,r,mp.get('bond'))
        if abs(bd)>0.004:
            events.append({**base,'kind':'扣质保金','payer':unit,'payee':'泓普','amt':round(bd,2),'mfee':0,'cost_due':0})
# 跨表重复：同一张票在两张表里各记一次，第二次起标记为不计入汇总
key_seen = {}
for e in events:
    if e['kind'] not in ('成本票','销项开票'):
        e['count_in']='是'; continue
    sheet = e['src'].split('!')[0]
    k = (e['kind'], e['payer'], e['payee'], round(e['amt'], 2))
    if k in key_seen:
        first_sheet, first_src = key_seen[k]
        if first_sheet != sheet:      # 只有来自不同表才算同一张票被记了两次
            e['count_in'] = '否'; e['dup_of'] = first_src
        else:
            e['count_in'] = '是'      # 同一张表内同额多笔＝确实是不同的票
    else:
        key_seen[k] = (sheet, e['src']); e['count_in'] = '是'
print('跨表重复标记', sum(1 for e in events if e.get('count_in')=='否'), '笔')
json.dump(events, open('/tmp/claude-0/-home-user-temp/9101e7cb-4f82-53a7-bbe8-41b0f93b12b0/scratchpad/events.json','w'), ensure_ascii=False)
import collections
print('事件总数', len(events))
print('按类型:', dict(collections.Counter(e['kind'] for e in events)))
print('按单位:', dict(collections.Counter(e['unit'] for e in events)))
print('项目数', len(projects))
print('\n各单位核对（按开票方/收票方口径，已排除跨表重复）：')
ORIG={'德誉嘉':(651402.95,611391.37),'迅驰':(219510.99,210730.55),'华城':(472843.99,463387.12),
      '康欣':(1020420.18,943853.98),'金沁':(600622.61,581152.54),'安锐':(141749.77,133690.89),
      '湖南锦泰':(163893.26,0.0),'杰华':(122305.72,0.0)}
V=[e for e in events if e.get('count_in')=='是']
for u,(os_,oc) in ORIG.items():
    sale=sum(e['amt'] for e in V if e['kind']=='销项开票' and e['payer']==u)
    recv=sum(e['amt'] for e in V if e['payee']==u and e['kind'] in ('成本票','销项开票'))
    f1='OK ' if abs(sale-os_)<1 else '差异'
    f2='OK ' if abs(recv-oc)<1 else '差异'
    print(f'  {u:<6} 开票 {sale:>12,.2f} vs {os_:>12,.2f} {f1} | 收到成本票 {recv:>12,.2f} vs {oc:>12,.2f} {f2}')
