# -*- coding: utf-8 -*-
import openpyxl, re, json, datetime as dt
SRC='/root/.claude/uploads/9101e7cb-4f82-53a7-bbe8-41b0f93b12b0/4ae06a40-08______1.xlsx'
wb=openpyxl.load_workbook(SRC, data_only=True)
rows=[]
DATE_RE=re.compile(r'(20\d{2})[.\-年/]\s*0?(\d{1,2})[.\-月/]\s*0?(\d{1,2})')
def gd(t,fb):
    m=DATE_RE.search(t or '')
    if m:
        try: return dt.date(int(m.group(1)),int(m.group(2)),int(m.group(3)))
        except ValueError: pass
    return fb
ws=wb['现金']; fb=dt.date(2026,8,1)
for r in range(5, ws.max_row+1):
    memo=ws.cell(row=r,column=7).value
    if not memo: continue
    d=gd(str(memo),fb); fb=d
    rows.append(dict(date=d.isoformat(), acct='现金', proj=str(ws.cell(row=r,column=4).value or ''),
                     etype=str(ws.cell(row=r,column=5).value or ''), who=str(ws.cell(row=r,column=6).value or ''),
                     memo=str(memo)[:120],
                     inc=float(ws.cell(row=r,column=8).value or 0), exp=float(ws.cell(row=r,column=9).value or 0)))
ws=wb['银行存款2个']
for acct,(r0,r1) in [('泓普',(5,10)),('仟茂',(23,ws.max_row))]:
    fb=dt.date(2026,8,1)
    for r in range(r0, r1+1):
        memo=ws.cell(row=r,column=4).value
        if not memo or '期初' in str(memo): continue
        d=gd(str(memo),fb); fb=d
        inc=ws.cell(row=r,column=5).value; exp=ws.cell(row=r,column=6).value
        if not isinstance(inc,(int,float)): inc=0
        if not isinstance(exp,(int,float)): exp=0
        if inc==0 and exp==0: continue
        rows.append(dict(date=d.isoformat(), acct=acct, proj='',
                         etype=str(ws.cell(row=r,column=2).value or ''), who=str(ws.cell(row=r,column=3).value or ''),
                         memo=str(memo)[:120], inc=float(inc), exp=float(exp)))
rows.sort(key=lambda x:(x['date'],x['acct']))
json.dump(rows, open('/tmp/claude-0/-home-user-temp/9101e7cb-4f82-53a7-bbe8-41b0f93b12b0/scratchpad/jour.json','w'), ensure_ascii=False)
import collections
print('日记账导入', len(rows), '笔 |', dict(collections.Counter(r['acct'] for r in rows)))
print('收入合计', round(sum(r['inc'] for r in rows),2), '支出合计', round(sum(r['exp'] for r in rows),2))
