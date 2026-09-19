# -*- coding: utf-8 -*-
"""电商一体化账务模板 —— 一键生成

结构：档案 → 业务登记 → 自动凭证＋手工凭证 → 科目余额表 → 三表 → 六张管理表
跑法：python3 build_ec.py <输出.xlsx>
"""
import sys, os, datetime as dt
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from ec_style import *
from ec_data import *
import ec_test as T

OUT = sys.argv[1] if len(sys.argv) > 1 else '电商一体化账务模板_2026.xlsx'
wb = openpyxl.Workbook()
wb.remove(wb.active)

SH = {}
def sheet(name, hidden=False):
    ws = wb.create_sheet(name)
    SH[name] = ws
    if hidden:
        ws.sheet_state = 'hidden'
    return ws

def dv(ws, sqref, formula, msg=None, block=False):
    d = DataValidation(type='list', formula1=formula, allow_blank=True,
                       showDropDown=False, showInputMessage=True,
                       showErrorMessage=block)
    if msg:
        d.promptTitle, d.prompt = '选一个', msg
    if block:
        d.errorTitle, d.error = '不在档案里', '请先到对应的档案表里加一行，再回来选'
    ws.add_data_validation(d)
    d.add(sqref)
    return d

def name(nm, ref):
    if nm in wb.defined_names:
        del wb.defined_names[nm]
    wb.defined_names.add(DefinedName(nm, attr_text=ref))

def dynlist(sheet_name, header, hdr_row, r0, r1, last_col):
    """按表头名字定位列、有多少条就多长 —— 以后插列/加内容都不用回来改公式"""
    m = f'MATCH("{header}",{sheet_name}!$A${hdr_row}:${last_col}${hdr_row},0)'
    c = f'COUNTA(INDEX({sheet_name}!$A${r0}:${last_col}${r1},0,{m}))'
    return f'OFFSET({sheet_name}!$A${r0},0,{m}-1,MAX(1,{c}),1)'

# 各登记表的行范围
N_CO_MAX = 8                         # 配对池按最多几家公司铺
BR0, BR1 = 4, 63                     # 基础资料的数据行
P0, P1 = D0, D0 + N_PUR - 1          # 采购
S0, S1 = D0, D0 + N_SAL - 1          # 销售
B0, B1 = D0, D0 + N_RBT - 1          # 返利
F0, F1 = D0, D0 + N_DRP - 1          # 一件代发
K0, K1 = D0, D0 + N_CASH - 1         # 资金流水
V0, V1 = D0, D0 + N_INV - 1          # 发票
E0, E1 = D0, D0 + N_EXP - 1          # 费用
M0, M1 = D0, D0 + N_MAN - 1          # 手工凭证

# ════════════════════════════════════════════════════════════
# ① 查询设置
# ════════════════════════════════════════════════════════════
ws = sheet('查询设置')
title(ws, '查 询 设 置',  'F',
      '★ 整套表只认这一页。改这里的【年度 / 公司 / 月份】，利润表、资产负债表、现金流、'
      '综合往来、店铺利润、老板看板会一起跟着换，不用一张张去改。')
widths(ws, {'A': 4, 'B': 18, 'C': 22, 'D': 40, 'E': 4, 'F': 60})

rows = [
    ('账套年度', YEAR, '0', '一年一个账套。2027 年把这个文件复制一份，把年度改成 2027，'
                            '再把《期初余额》按 2026 年末数重填即可。'),
    ('查询公司', '全部', 'General', '从下拉里选。「全部」＝几家公司合并看；选某一家＝只看那一家。'
                         '注意：往来是按「公司＋往来单位」分开算的，不会把甲公司的应付跟乙公司的应收抵掉。'),
    ('查询月份', '全年', 'General', '从下拉里选。「全年」＝1~12 月累计；'
                         '选某个月＝只看那一个月的发生额，期初自动滚到上月末。'),
]
r = 4
for lbl, val, fmt_, note in rows:
    put(ws, f'B{r}', lbl, font=F_SEC, fill=FILL_SEC, align=CR)
    put(ws, f'C{r}', val, font=F_IN, fill=FILL_IN, fmt=fmt_)
    put(ws, f'D{r}', note, font=F_NOTE, align=CL)
    ws.row_dimensions[r].height = 34
    r += 1
dv(ws, 'C5', '=公司下拉', '全部 或 某一家公司', block=True)
dv(ws, 'C6', '"全年,1,2,3,4,5,6,7,8,9,10,11,12"', '全年 或 1~12 月', block=True)

put(ws, 'B8', '—— 下面是自动算的，别动 ——', font=F_NOTE, align=CR, border=None)
AUTO = [
    ('起始年月', '=IFERROR(VALUE($C$4),0)*100+IF($C$6="全年",1,IFERROR(VALUE($C$6),1))',
     YM, '本次查询的第一个月'),
    ('截止年月', '=IFERROR(VALUE($C$4),0)*100+IF($C$6="全年",12,IFERROR(VALUE($C$6),12))',
     YM, '本次查询的最后一个月'),
    ('起始日期', '=DATE(IFERROR(VALUE($C$4),2026),IF($C$6="全年",1,IFERROR(VALUE($C$6),1)),1)', DATEF, ''),
    ('截止日期', '=EOMONTH(DATE(IFERROR(VALUE($C$4),2026),'
                 'IF($C$6="全年",12,IFERROR(VALUE($C$6),12)),1),0)', DATEF, ''),
    ('公司条件', '=IF(OR($C$5="",$C$5="全部"),"*",$C$5)', TXT, 'SUMIFS 用的通配条件'),
    ('本期标题', '=IFERROR(VALUE($C$4),0)&" 年 "&IF($C$6="全年","1-12 月",'
                 'IFERROR(VALUE($C$6),"?")&" 月")&"　｜　"&$C$5',
     TXT, '报表标题里显示的那一行字'),
    ('年初年月', '=IFERROR(VALUE($C$4),0)*100+1', YM, '本年度第一个月，算「本年累计」用'),
    ('参数自检', '=IF(AND(IFERROR(VALUE($C$4),0)>=2000,OR($C$6="全年",'
                 'AND(IFERROR(VALUE($C$6),0)>=1,IFERROR(VALUE($C$6),0)<=12))),'
                 '"✔ 参数正常","✘ 年度或月份填错了 —— 报表会全是 0，请从下拉里重选")',
     TXT, '★ 这一格不是 ✔ 的话，下面所有报表都不能信'),
]
r = 9
for lbl, f, fmt, note in AUTO:
    put(ws, f'B{r}', lbl, font=F_TOT, fill=FILL_AUTO, align=CR)
    put(ws, f'C{r}', f, font=F_AUTO, fill=FILL_AUTO, fmt=fmt)
    put(ws, f'D{r}', note, font=F_NOTE, align=CL)
    r += 1
# 「查询公司」下拉的选项来源：第一项固定是「全部」，后面跟着《基础资料》的公司列
put(ws, 'H19', '公司下拉选项（自动）', font=F_NOTE, align=CL, border=None)
put(ws, 'H20', '全部', font=F_AUTO, border=None)
put(ws, 'I20', 1, font=F_AUTO, border=None, fmt='0')
for i in range(1, N_CO_MAX + 1):
    rr = 20 + i
    put(ws, f'H{rr}', f'=IF(基础资料!$A{BR0 + i - 1}="","",基础资料!$A{BR0 + i - 1})',
        font=F_AUTO, border=None)
    put(ws, f'I{rr}', f'=N(I{rr-1})+IF($H{rr}="",0,1)', font=F_AUTO, border=None, fmt='0')
ws.column_dimensions['H'].hidden = True
ws.column_dimensions['I'].hidden = True
name('公司下拉', f'OFFSET(查询设置!$H$20,0,0,MAX(1,MAX(查询设置!$I$20:$I${20 + N_CO_MAX})),1)')

name('账套年度', '查询设置!$C$4')
name('查询公司', '查询设置!$C$5')
name('查询月份', '查询设置!$C$6')
name('起始年月', '查询设置!$C$9')
name('截止年月', '查询设置!$C$10')
name('起始日期', '查询设置!$C$11')
name('截止日期', '查询设置!$C$12')
name('公司条件', '查询设置!$C$13')
name('本期标题', '查询设置!$C$14')
name('年初年月', '查询设置!$C$15')
name('参数自检', '查询设置!$C$16')

put(ws, 'F4', '常见问题：\n'
              '· 报表全是 0？先看这一页的【公司】是不是选了一家没有业务的公司。\n'
              '· 只想看 3 月？【月份】填 3；想看全年就填「全年」。\n'
              '· 换年度只改【年度】没用 —— 期初余额要按上年末重新结转一次。',
    font=F_NOTE, align=CT, fill=FILL_KPI)
ws.merge_cells('F4:F12')
page(ws, landscape=False)

# ════════════════════════════════════════════════════════════
# ② 基础资料（各种下拉选项）
# ════════════════════════════════════════════════════════════
ws = sheet('基础资料')
title(ws, '基 础 资 料（下拉选项总表）', 'O',
      '★ 每一列就是一个下拉菜单的来源。要加新选项，直接在那一列最后一个非空格子下面接着写，'
      '登记表的下拉会自动多出来 —— 不用回来改公式（下拉区域是按表头名字动态找列的）。')
BASE_COLS = [
    ('公司', COMPANIES, 16),
    ('平台', PLATFORMS, 12),
    ('计量单位', ['个', '条', '台', '件', '箱', '套', '包', '双', '盒', 'kg'], 10),
    ('税率', [0.13, 0.09, 0.06, 0.03, 0.01, 0.00], 8),
    ('付款状态', PAY_STATUS, 12),
    ('回款状态', REC_STATUS, 12),
    ('票据状态', INV_STATUS, 12),
    ('票据类型', INV_KIND, 10),
    ('发票状态', INV_STATE, 12),
    ('是否', ['是', '否'], 8),
    ('账户类型', ['银行', '平台', '现金', '其他'], 10),
    ('手工凭证用途', MANUAL_TAG, 14),
    ('费用项目', [x[0] for x in EXPENSE_ITEMS], 20),
    ('费用对应科目', [x[1] for x in EXPENSE_ITEMS], 18),
    ('费用贷方科目', [x[2] for x in EXPENSE_ITEMS], 20),
]
headers(ws, HDR, [c[0] for c in BASE_COLS])
widths(ws, {L(i): c[2] for i, c in enumerate(BASE_COLS, 1)})
for i, (h, vals, _) in enumerate(BASE_COLS, 1):
    col = L(i)
    for k in range(BR0, BR1 + 1):
        v = vals[k - BR0] if k - BR0 < len(vals) else None
        fmt = PCT if h == '税率' else TXT
        put(ws, f'{col}{k}', v, font=F_IN, fill=FILL_IN, fmt=fmt)
put(ws, f'A{BR1+2}',
    '注：最后三列是一组 ——「费用项目」＋「费用对应科目」（借方，进费用还是进资产）＋「费用贷方科目」（钱先挂在哪）。'
    '《费用及其他》选了费用项目，凭证就按这一行自动做：借 费用对应科目 ／ 贷 费用贷方科目。'
    '想让某笔支出进固定资产就把对应科目改成「固定资产」；工资挂「应付职工薪酬」、税费挂「应交税费」，'
    '这样《资金流水》付钱时才冲得对。',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells(f'A{BR1+2}:O{BR1+2}')
page(ws)

for h in ['公司', '平台', '计量单位', '税率', '付款状态', '回款状态', '票据状态',
          '票据类型', '发票状态', '是否', '账户类型', '手工凭证用途', '费用项目']:
    name(h + '表', dynlist('基础资料', h, HDR, BR0, BR1, 'O'))

# ════════════════════════════════════════════════════════════
# ③ 往来单位主档（多角色）
# ════════════════════════════════════════════════════════════
ws = sheet('往来单位主档')
title(ws, '往 来 单 位 主 档（同一家只建一次，可以同时是供应商/客户/资方/店铺合作方）', 'N',
      '★ 这是整套表的往来单位唯一入口。一家公司只建一个编码，右边四个「是否」勾上它的身份。'
      '勾了供应商，采购登记的下拉里才有它；勾了客户，销售登记的下拉里才有它 —— 互不影响。\n'
      '★ 身份归身份，账还是分开记的：采购走应付账款、销售走应收账款、返利走其他应收款、借款走短期借款/其他应付款。'
      '想把应收应付抵掉，必须在《手工凭证》里做一张正式的抵销凭证（用途选「往来抵销」）。')
PC = ['编码', '往来单位名称', '简称', '是否供应商', '是否客户', '是否资方', '是否店铺合作方',
      '信用期(天)', '默认税率', '联系人', '联系电话', '开户行/账号', '状态', '备注']
headers(ws, HDR, PC)
widths(ws, {'A': 10, 'B': 26, 'C': 10, 'D': 11, 'E': 10, 'F': 10, 'G': 13, 'H': 11,
            'I': 10, 'J': 10, 'K': 14, 'L': 20, 'M': 10, 'N': 34})
PR0, PR1 = 4, 4 + N_PARTY - 1
for i, p in enumerate(PARTIES):
    r = PR0 + i
    vals = [p[0], p[1], p[2], '是' if p[3] else '否', '是' if p[4] else '否',
            '是' if p[5] else '否', '是' if p[6] else '否', p[7], p[8], p[9], p[10], '', '启用', p[11]]
    for c, v in enumerate(vals, 1):
        fmt = PCT if c == 9 else ('0' if c == 8 else TXT)
        put(ws, f'{L(c)}{r}', v, font=F_IN, fill=FILL_IN,
            align=CL if c in (2, 14) else C, fmt=fmt)
for r in range(PR0 + len(PARTIES), PR1 + 1):
    for c in range(1, 15):
        put(ws, f'{L(c)}{r}', None, font=F_IN, fill=FILL_IN,
            align=CL if c in (2, 14) else C,
            fmt=PCT if c == 9 else ('0' if c == 8 else TXT))
for col in ('D', 'E', 'F', 'G'):
    dv(ws, f'{col}{PR0}:{col}{PR1}', '=是否表', '这家单位是不是这个身份')
dv(ws, f'M{PR0}:M{PR1}', '"启用,停用"', '停用的不会出现在下拉里')
ws.auto_filter.ref = f'A{HDR}:N{PR1}'
ws.freeze_panes = 'C4'
page(ws, titles=f'{HDR}:{HDR}')
name('往来单位表', dynlist('往来单位主档', '往来单位名称', HDR, PR0, PR1, 'N'))

# ════════════════════════════════════════════════════════════
# ④ 商品档案
# ════════════════════════════════════════════════════════════
ws = sheet('商品档案')
title(ws, '商 品 档 案', 'H',
      '★ 采购、销售、一件代发都只填「商品编码」，名称/规格/单位会自动带出来。')
IC = ['商品编码', '商品名称', '规格', '计量单位', '默认税率', '商品类别', '状态', '备注']
headers(ws, HDR, IC)
widths(ws, {'A': 12, 'B': 20, 'C': 18, 'D': 10, 'E': 10, 'F': 12, 'G': 10, 'H': 26})
IR0, IR1 = 4, 4 + N_ITEM - 1
for i, it in enumerate(ITEMS):
    r = IR0 + i
    for c, v in enumerate(list(it) + ['启用', ''], 1):
        put(ws, f'{L(c)}{r}', v, font=F_IN, fill=FILL_IN,
            fmt=PCT if c == 5 else TXT, align=CL if c in (2, 3, 8) else C)
for r in range(IR0 + len(ITEMS), IR1 + 1):
    for c in range(1, 9):
        put(ws, f'{L(c)}{r}', None, font=F_IN, fill=FILL_IN,
            fmt=PCT if c == 5 else TXT, align=CL if c in (2, 3, 8) else C)
dv(ws, f'D{IR0}:D{IR1}', '=计量单位表')
dv(ws, f'E{IR0}:E{IR1}', '=税率表')
dv(ws, f'G{IR0}:G{IR1}', '"启用,停用"')
ws.auto_filter.ref = f'A{HDR}:H{IR1}'
ws.freeze_panes = 'B4'
page(ws, titles=f'{HDR}:{HDR}')
name('商品编码表', dynlist('商品档案', '商品编码', HDR, IR0, IR1, 'H'))

# ════════════════════════════════════════════════════════════
# ⑤ 店铺档案
# ════════════════════════════════════════════════════════════
ws = sheet('店铺档案')
title(ws, '店 铺 档 案', 'G',
      '★ 店铺挂在哪家公司、哪个平台，在这里定。《店铺利润分析》按这里的归属汇总。'
      '「平台佣金率」只是给你估算用的参考值，实际扣费以销售登记里填的为准。')
SC = ['店铺名称', '所属公司', '平台', '平台佣金率(参考)', '负责人', '状态', '备注']
headers(ws, HDR, SC)
widths(ws, {'A': 18, 'B': 14, 'C': 12, 'D': 16, 'E': 12, 'F': 10, 'G': 28})
SR0, SR1 = 4, 4 + N_SHOP - 1
for i, s in enumerate(SHOPS):
    r = SR0 + i
    for c, v in enumerate([s[0], s[1], s[2], s[3], s[4], '启用', s[5]], 1):
        put(ws, f'{L(c)}{r}', v, font=F_IN, fill=FILL_IN,
            fmt=PCT if c == 4 else TXT, align=CL if c in (1, 7) else C)
for r in range(SR0 + len(SHOPS), SR1 + 1):
    for c in range(1, 8):
        put(ws, f'{L(c)}{r}', None, font=F_IN, fill=FILL_IN,
            fmt=PCT if c == 4 else TXT, align=CL if c in (1, 7) else C)
dv(ws, f'B{SR0}:B{SR1}', '=公司表')
dv(ws, f'C{SR0}:C{SR1}', '=平台表')
dv(ws, f'F{SR0}:F{SR1}', '"启用,停用"')
ws.freeze_panes = 'B4'
page(ws, titles=f'{HDR}:{HDR}')
name('店铺表', dynlist('店铺档案', '店铺名称', HDR, SR0, SR1, 'G'))

# ════════════════════════════════════════════════════════════
# ⑥ 资金账户档案
# ════════════════════════════════════════════════════════════
ws = sheet('资金账户档案')
title(ws, '资 金 账 户 档 案', 'F',
      '★ 银行户、平台货款户（拼多多/抖音待结算）、现金都建在这里，各自挂一个会计科目。'
      '平台货款提现到银行＝《资金流水》里做一笔「内部转账-转出」＋一笔「内部转账-转入」，不影响损益。')
AC = ['资金账户', '所属公司', '账户类型', '对应会计科目', '状态', '备注']
headers(ws, HDR, AC)
widths(ws, {'A': 20, 'B': 14, 'C': 12, 'D': 24, 'E': 10, 'F': 24})
AR0, AR1 = 4, 4 + N_ACC - 1
for i, a in enumerate(ACCOUNTS):
    r = AR0 + i
    for c, v in enumerate([a[0], a[1], a[2], a[3], '启用', a[4]], 1):
        put(ws, f'{L(c)}{r}', v, font=F_IN, fill=FILL_IN, align=CL if c in (1, 4, 6) else C)
for r in range(AR0 + len(ACCOUNTS), AR1 + 1):
    for c in range(1, 7):
        put(ws, f'{L(c)}{r}', None, font=F_IN, fill=FILL_IN, align=CL if c in (1, 4, 6) else C)
dv(ws, f'B{AR0}:B{AR1}', '=公司表')
dv(ws, f'C{AR0}:C{AR1}', '=账户类型表')
dv(ws, f'E{AR0}:E{AR1}', '"启用,停用"')
ws.freeze_panes = 'B4'
page(ws, titles=f'{HDR}:{HDR}', landscape=False)
name('资金账户表', dynlist('资金账户档案', '资金账户', HDR, AR0, AR1, 'F'))

print('  ✓ 查询设置 / 基础资料 / 往来单位主档 / 商品档案 / 店铺档案 / 资金账户档案')

# ════════════════════════════════════════════════════════════
# ⑦ 科目表
# ════════════════════════════════════════════════════════════
ws = sheet('科目表')
title(ws, '会 计 科 目 表', 'F',
      '★ 小企业会计准则 + 电商常用明细。「报表项目」是资产负债表/利润表取数的钥匙，'
      '新增科目必须填，否则报表汇不上。「损益方向」只有损益类科目要填（收入/成本）。')
CC = ['科目编码', '科目名称', '科目类别', '余额方向', '报表项目', '损益方向']
headers(ws, HDR, CC)
widths(ws, {'A': 12, 'B': 32, 'C': 10, 'D': 10, 'E': 18, 'F': 10})
CR0, CR1 = 4, 4 + N_SUBJ - 1
for i, s in enumerate(SUBJECTS):
    r = CR0 + i
    for c, v in enumerate(s, 1):
        put(ws, f'{L(c)}{r}', v, font=F_IN, fill=FILL_IN,
            align=CL if c == 2 else C, fmt=TXT)
for r in range(CR0 + len(SUBJECTS), CR1 + 1):
    for c in range(1, 7):
        put(ws, f'{L(c)}{r}', None, font=F_IN, fill=FILL_IN, align=CL if c == 2 else C, fmt=TXT)
dv(ws, f'E{CR0}:E{CR1}', '"货币资金,应收账款,预付款项,其他应收款,存货,固定资产,无形资产,短期借款,应付账款,预收款项,应付职工薪酬,应交税费,其他应付款,实收资本,未分配利润,营业收入,营业成本,税金及附加,销售费用,管理费用,财务费用,营业外收入,营业外支出,所得税费用"',
   '★ 只能选这些 —— 报表就是按它取数的，填别的这个科目的钱会掉在报表外面', block=True)
dv(ws, f'C{CR0}:C{CR1}', '"资产,负债,权益,损益"')
dv(ws, f'D{CR0}:D{CR1}', '"借,贷"')
dv(ws, f'F{CR0}:F{CR1}', '"收入,成本"')
ws.auto_filter.ref = f'A{HDR}:F{CR1}'
ws.freeze_panes = 'C4'
page(ws, titles=f'{HDR}:{HDR}', landscape=False)
name('科目名称表', dynlist('科目表', '科目名称', HDR, CR0, CR1, 'F'))
SUBJ_N = f'科目表!$B${CR0}:$B${CR1}'
SUBJ_DIR = f'科目表!$D${CR0}:$D${CR1}'
SUBJ_ITEM = f'科目表!$E${CR0}:$E${CR1}'
SUBJ_CLS = f'科目表!$C${CR0}:$C${CR1}'

# ════════════════════════════════════════════════════════════
# ⑧ 资金规则（业务类型 → 对方科目 / 现金流类别）
# ════════════════════════════════════════════════════════════
ws = sheet('资金规则')
title(ws, '资 金 业 务 规 则（业务类型 → 对方科目 / 现金流类别）', 'E',
      '★《资金流水》选了业务类型，对方科目和现金流类别就按这张表自动判。'
      '要改口径就改右边两列；个别单据要例外，在资金流水最右边的「科目手工覆盖」里填。\n'
      '★「内部转账」两行的对方科目是按你填的【对方账户】自动取的，这里写什么都不影响。')
RC = ['业务类型', '方向', '对方科目', '现金流类别', '说明']
headers(ws, HDR, RC)
widths(ws, {'A': 20, 'B': 8, 'C': 26, 'D': 14, 'E': 44})
RR0, RR1 = 4, 4 + 60 - 1
for i, rr in enumerate(CASH_RULES):
    r = RR0 + i
    for c, v in enumerate(rr, 1):
        put(ws, f'{L(c)}{r}', v, font=F_IN, fill=FILL_IN, align=CL if c in (3, 5) else C, fmt=TXT)
for r in range(RR0 + len(CASH_RULES), RR1 + 1):
    for c in range(1, 6):
        put(ws, f'{L(c)}{r}', None, font=F_IN, fill=FILL_IN, align=CL if c in (3, 5) else C, fmt=TXT)
dv(ws, f'B{RR0}:B{RR1}', '"收,付"')
dv(ws, f'C{RR0}:C{RR1}', '=科目名称表')
dv(ws, f'D{RR0}:D{RR1}', '"经营活动,投资活动,筹资活动,不计入"')
page(ws, titles=f'{HDR}:{HDR}', landscape=False)
name('资金业务类型表', dynlist('资金规则', '业务类型', HDR, RR0, RR1, 'E'))
RULE_T = f'资金规则!$A${RR0}:$A${RR1}'
RULE_S = f'资金规则!$C${RR0}:$C${RR1}'
RULE_C = f'资金规则!$D${RR0}:$D${RR1}'

# ════════════════════════════════════════════════════════════
# ⑨ 期初余额
# ════════════════════════════════════════════════════════════
ws = sheet('期初余额')
title(ws, '期 初 余 额（年初数 · 一年一个账套）', 'N',
      '★ 只填浅黄色格子。主区是「公司 + 科目」的年初余额，资产/成本类填借方，负债/权益类填贷方。\n'
      '★ 右边三个小区是明细期初（往来单位 / 商品库存 / 资金账户），填了以后综合往来、库存、'
      '账户余额才接得上年初数；底部会自动跟主区核对，对不上会标红。\n'
      '★ 换年度：复制整个文件 → 把《查询设置》年度改成下一年 → 把上一年《资产负债表》的年末数抄到这里。')
OC = ['公司', '科目名称', '科目类别', '期初借方', '期初贷方', '备注']
headers(ws, HDR, OC)
widths(ws, {'A': 14, 'B': 30, 'C': 10, 'D': 15, 'E': 15, 'F': 22,
            'G': 3, 'H': 14, 'I': 20, 'J': 14, 'K': 14, 'L': 14, 'M': 14, 'N': 22})
OR0, OR1 = 4, 4 + N_OPEN - 1
OPENING = [
    ('甲公司', '银行存款',        120000.00, 0),
    ('甲公司', '其他货币资金—平台资金', 18000.00, 0),
    ('甲公司', '库存现金',          3000.00, 0),
    ('甲公司', '应收账款',         26000.00, 0),
    ('甲公司', '库存商品',         85000.00, 0),
    ('甲公司', '应付账款',              0, 62000.00),
    ('甲公司', '其他应收款—保证金',  5000.00, 0),
    ('甲公司', '实收资本',              0, 100000.00),
    ('甲公司', '利润分配—未分配利润',   0, 95000.00),
    ('乙公司', '银行存款',          60000.00, 0),
    ('乙公司', '库存商品',          30000.00, 0),
    ('乙公司', '应付账款',              0, 20000.00),
    ('乙公司', '实收资本',              0, 50000.00),
    ('乙公司', '利润分配—未分配利润',   0, 20000.00),
]
for i in range(N_OPEN):
    r = OR0 + i
    src = OPENING[i] if i < len(OPENING) else None
    put(ws, f'A{r}', src[0] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'B{r}', src[1] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'C{r}', f'=IF($B{r}="","",IFERROR(INDEX({SUBJ_CLS},MATCH($B{r},{SUBJ_N},0)),"★科目表里没有"))',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'D{r}', src[2] if src else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'E{r}', src[3] if src else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'F{r}', None, font=F_IN, fill=FILL_IN, align=CL)
dv(ws, f'A{OR0}:A{OR1}', '=公司表')
dv(ws, f'B{OR0}:B{OR1}', '=科目名称表')
OT = OR1 + 1
put(ws, f'A{OT}', '合  计', font=F_TOT, fill=FILL_TOT)
put(ws, f'B{OT}', None, font=F_TOT, fill=FILL_TOT)
put(ws, f'C{OT}', None, font=F_TOT, fill=FILL_TOT)
for c in 'DE':
    put(ws, f'{c}{OT}', f'=ROUND(SUM({c}{OR0}:{c}{OR1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
put(ws, f'F{OT}', f'=IF(ROUND($D{OT}-$E{OT},2)=0,"✔ 借贷平衡","✘ 差 "&TEXT($D{OT}-$E{OT},"#,##0.00"))',
    font=F_TOT, fill=FILL_TOT)

# 明细期初：往来单位
put(ws, 'H3', '往来单位期初（按公司分开）', font=F_HDR2, fill=FILL_HDR2)
ws.merge_cells('H3:L3')
headers(ws, 4, [None]*7 + ['公司', '往来单位', '应收账款', '应付账款', '应收返利', '借款余额'],
        fill=FILL_HDR2, font=F_HDR2, height=30)
OP2_0, OP2_1 = 5, 104
OPEN_PARTY = [
    ('甲公司', '优品电商服务部', 26000.00, 0, 0, 0),
    ('甲公司', '宏发供应链有限公司', 0, 40000.00, 0, 0),
    ('甲公司', '恒信百货批发行', 0, 22000.00, 0, 0),
    ('乙公司', '恒信百货批发行', 0, 20000.00, 0, 0),
]
for i in range(OP2_0, OP2_1 + 1):
    src = OPEN_PARTY[i - OP2_0] if i - OP2_0 < len(OPEN_PARTY) else None
    put(ws, f'H{i}', src[0] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'I{i}', src[1] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    for k, c in enumerate('JKLM'):
        put(ws, f'{c}{i}', src[2 + k] if src else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'N{i}', None, font=F_IN, fill=FILL_IN, align=CL)
dv(ws, f'H{OP2_0}:H{OP2_1}', '=公司表')
dv(ws, f'I{OP2_0}:I{OP2_1}', '=往来单位表')
PT = OP2_1 + 1
put(ws, f'I{PT}', '小计', font=F_TOT, fill=FILL_TOT)
for c in 'JKLM':
    put(ws, f'{c}{PT}', f'=ROUND(SUM({c}{OP2_0}:{c}{OP2_1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
CHK = PT + 1
put(ws, f'I{CHK}', '与主区核对', font=F_NOTE, fill=FILL_AUTO)
for c, subj in (('J', '应收账款'), ('K', '应付账款'), ('L', '其他应收款—应收返利')):
    side = 'D' if subj == '应收账款' or subj.startswith('其他应收') else 'E'
    put(ws, f'{c}{CHK}',
        f'=IF(ROUND({c}{PT}-SUMIF($B${OR0}:$B${OR1},"{subj}",${side}${OR0}:${side}${OR1}),2)=0,"✔","✘ 差 "'
        f'&TEXT({c}{PT}-SUMIF($B${OR0}:$B${OR1},"{subj}",${side}${OR0}:${side}${OR1}),"#,##0.00"))',
        font=F_TOT, fill=FILL_AUTO)
put(ws, f'M{CHK}', '（借款明细仅供对账）', font=F_NOTE, fill=FILL_AUTO)

# 明细期初：商品库存
put(ws, 'H107', '商品库存期初（决定期初单位成本）', font=F_HDR2, fill=FILL_HDR2)
ws.merge_cells('H107:L107')
headers(ws, 108, [None]*7 + ['公司', '商品编码', '期初数量', '期初金额', '期初单位成本', ''],
        fill=FILL_HDR2, font=F_HDR2, height=30)
OP3_0, OP3_1 = 109, 168
OPEN_ITEM = [
    ('甲公司', 'SP001', 500, 20000.00),
    ('甲公司', 'SP002', 800, 42000.00),
    ('甲公司', 'SP003', 1000, 23000.00),
    ('乙公司', 'SP001', 400, 16000.00),
    ('乙公司', 'SP005', 500, 14000.00),
]
for i in range(OP3_0, OP3_1 + 1):
    src = OPEN_ITEM[i - OP3_0] if i - OP3_0 < len(OPEN_ITEM) else None
    put(ws, f'H{i}', src[0] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'I{i}', src[1] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'J{i}', src[2] if src else None, font=F_IN, fill=FILL_IN, fmt=NUM)
    put(ws, f'K{i}', src[3] if src else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'L{i}', f'=IF(N($J{i})=0,"",ROUND(N($K{i})/$J{i},4))', font=F_AUTO, fill=FILL_AUTO, fmt=PRICE)
    put(ws, f'M{i}', None, font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'N{i}', None, font=F_IN, fill=FILL_IN, align=CL)
dv(ws, f'H{OP3_0}:H{OP3_1}', '=公司表')
dv(ws, f'I{OP3_0}:I{OP3_1}', '=商品编码表')
IT2 = OP3_1 + 1
put(ws, f'I{IT2}', '小计', font=F_TOT, fill=FILL_TOT)
for c in 'JK':
    put(ws, f'{c}{IT2}', f'=ROUND(SUM({c}{OP3_0}:{c}{OP3_1}),2)', font=F_TOT, fill=FILL_TOT,
        fmt=NUM if c == 'J' else MONEY)
put(ws, f'L{IT2}',
    f'=IF(ROUND($K{IT2}-SUMIF($B${OR0}:$B${OR1},"库存商品",$D${OR0}:$D${OR1}),2)=0,"✔ 与主区一致","✘ 差 "'
    f'&TEXT($K{IT2}-SUMIF($B${OR0}:$B${OR1},"库存商品",$D${OR0}:$D${OR1}),"#,##0.00"))',
    font=F_TOT, fill=FILL_TOT)

# 明细期初：资金账户
put(ws, 'H171', '资金账户期初', font=F_HDR2, fill=FILL_HDR2)
ws.merge_cells('H171:K171')
headers(ws, 172, [None]*7 + ['资金账户', '所属公司(自动)', '期初余额', '', '', ''],
        fill=FILL_HDR2, font=F_HDR2, height=30)
OP4_0, OP4_1 = 173, 212
OPEN_ACC = [('工行-甲基本户', 120000.00), ('其他', 0)]
ACC_N = f'资金账户档案!$A${AR0}:$A${AR1}'
ACC_CO = f'资金账户档案!$B${AR0}:$B${AR1}'
ACC_SUBJ = f'资金账户档案!$D${AR0}:$D${AR1}'
OPEN_ACC_MAP = {'工行-甲基本户': 120000.00, '拼多多货款户-甲': 18000.00,
                '甲现金': 3000.00, '建行-乙基本户': 60000.00}
for i in range(OP4_0, OP4_1 + 1):
    k = i - OP4_0
    put(ws, f'H{i}', f'=IF(INDEX({ACC_N},{k+1})="","",INDEX({ACC_N},{k+1}))',
        font=F_AUTO, fill=FILL_AUTO, align=CL)
    put(ws, f'I{i}', f'=IF($H{i}="","",IFERROR(INDEX({ACC_CO},MATCH($H{i},{ACC_N},0)),""))',
        font=F_AUTO, fill=FILL_AUTO)
    nm_ = ACCOUNTS[k][0] if k < len(ACCOUNTS) else None
    put(ws, f'J{i}', OPEN_ACC_MAP.get(nm_), font=F_IN, fill=FILL_IN, fmt=MONEY)
    for c in 'KLMN':
        put(ws, f'{c}{i}', None, font=F_AUTO, fill=FILL_AUTO)
AT2 = OP4_1 + 1
put(ws, f'I{AT2}', '小计', font=F_TOT, fill=FILL_TOT)
put(ws, f'J{AT2}', f'=ROUND(SUM(J{OP4_0}:J{OP4_1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
put(ws, f'K{AT2}',
    f'=IF(ROUND($J{AT2}-SUMIFS($D${OR0}:$D${OR1},$B${OR0}:$B${OR1},"库存现金")'
    f'-SUMIFS($D${OR0}:$D${OR1},$B${OR0}:$B${OR1},"银行存款")'
    f'-SUMIFS($D${OR0}:$D${OR1},$B${OR0}:$B${OR1},"其他货币资金—平台资金"),2)=0,'
    f'"✔ 与主区货币资金一致","✘ 差 "&TEXT($J{AT2}-SUMIFS($D${OR0}:$D${OR1},$B${OR0}:$B${OR1},"库存现金")'
    f'-SUMIFS($D${OR0}:$D${OR1},$B${OR0}:$B${OR1},"银行存款")'
    f'-SUMIFS($D${OR0}:$D${OR1},$B${OR0}:$B${OR1},"其他货币资金—平台资金"),"#,##0.00"))',
    font=F_TOT, fill=FILL_TOT)
ws.freeze_panes = 'A4'
page(ws)

OPEN_CO, OPEN_SUB = f'期初余额!$A${OR0}:$A${OR1}', f'期初余额!$B${OR0}:$B${OR1}'
OPEN_D, OPEN_C = f'期初余额!$D${OR0}:$D${OR1}', f'期初余额!$E${OR0}:$E${OR1}'
OPP_CO, OPP_PT = f'期初余额!$H${OP2_0}:$H${OP2_1}', f'期初余额!$I${OP2_0}:$I${OP2_1}'
OPP_AR, OPP_AP = f'期初余额!$J${OP2_0}:$J${OP2_1}', f'期初余额!$K${OP2_0}:$K${OP2_1}'
OPP_RB, OPP_LN = f'期初余额!$L${OP2_0}:$L${OP2_1}', f'期初余额!$M${OP2_0}:$M${OP2_1}'
OPI_CO, OPI_IT = f'期初余额!$H${OP3_0}:$H${OP3_1}', f'期初余额!$I${OP3_0}:$I${OP3_1}'
OPI_Q, OPI_A = f'期初余额!$J${OP3_0}:$J${OP3_1}', f'期初余额!$K${OP3_0}:$K${OP3_1}'
OPA_N, OPA_V = f'期初余额!$H${OP4_0}:$H${OP4_1}', f'期初余额!$J${OP4_0}:$J${OP4_1}'

print('  ✓ 科目表 / 资金规则 / 期初余额')

# ════════════════════════════════════════════════════════════
# 各表区间常量（公式里到处要用，先定好）
# ════════════════════════════════════════════════════════════
PN = {p[0]: p[1] for p in PARTIES}        # 编码 → 名称
def pn(code):
    return PN.get(code, code)

PT_N = f'往来单位主档!$B${PR0}:$B${PR1}'
PT_SUP = f'往来单位主档!$D${PR0}:$D${PR1}'
PT_CUS = f'往来单位主档!$E${PR0}:$E${PR1}'
PT_FIN = f'往来单位主档!$F${PR0}:$F${PR1}'
PT_PAR = f'往来单位主档!$G${PR0}:$G${PR1}'
PT_CRD = f'往来单位主档!$H${PR0}:$H${PR1}'
PT_ST  = f'往来单位主档!$M${PR0}:$M${PR1}'

IT_C = f'商品档案!$A${IR0}:$A${IR1}'
IT_N = f'商品档案!$B${IR0}:$B${IR1}'
IT_S = f'商品档案!$C${IR0}:$C${IR1}'
IT_U = f'商品档案!$D${IR0}:$D${IR1}'
IT_T = f'商品档案!$E${IR0}:$E${IR1}'

SP_N = f'店铺档案!$A${SR0}:$A${SR1}'
SP_CO = f'店铺档案!$B${SR0}:$B${SR1}'
SP_PF = f'店铺档案!$C${SR0}:$C${SR1}'

EX_I = f'基础资料!$M${BR0}:$M${BR1}'
EX_S = f'基础资料!$N${BR0}:$N${BR1}'
EX_C = f'基础资料!$O${BR0}:$O${BR1}'

PU = lambda c: f'采购登记!${c}${P0}:${c}${P1}'
SA = lambda c: f'销售登记!${c}${S0}:${S1 and S1}'
SA = lambda c: f'销售登记!${c}${S0}:${c}${S1}'
RB = lambda c: f'返利登记!${c}${B0}:${c}${B1}'
DP = lambda c: f'一件代发结算!${c}${F0}:${c}${F1}'
CA = lambda c: f'资金流水!${c}${K0}:${c}${K1}'
IV = lambda c: f'发票台账!${c}${V0}:${c}${V1}'
EP = lambda c: f'费用及其他!${c}${E0}:${c}${E1}'
MV = lambda c: f'手工凭证!${c}${M0}:${c}${M1}'

def ym(col, r):
    """日期 → 年月（202601）。日期录成文本时不炸，直接留空"""
    return f'IF(OR(${col}{r}="",NOT(ISNUMBER(${col}{r}))),"",--TEXT(${col}{r},"yyyymm"))'

def paid(ref_col, co_col, r):
    """按「关联单号＋公司」汇总真正付出去的钱。只认业务类型以「付-」开头的行；
       同一类型里如果填在收入栏（供应商退款），自动减回来。
       —— 收付必须分开算，否则一件代发那种「同一单号既收客户钱、又付代发商钱」会互相抵掉。"""
    base = f'{CA("G")},${ref_col}{r},{CA("C")},${co_col}{r},{CA("E")},"付-*"'
    return f'ROUND(SUMIFS({CA("I")},{base})-SUMIFS({CA("H")},{base}),2)'


def recv(ref_col, co_col, r):
    """按「关联单号＋公司」汇总真正收到的钱。只认业务类型以「收-」开头的行。"""
    base = f'{CA("G")},${ref_col}{r},{CA("C")},${co_col}{r},{CA("E")},"收-*"'
    return f'ROUND(SUMIFS({CA("H")},{base})-SUMIFS({CA("I")},{base}),2)'



def off_ap(ref_col, co_col, r):
    """这张采购单被《手工凭证》正式抵销掉的应付（借应付账款 − 贷应付账款）"""
    return (f'ROUND(SUMIFS({MV("I")},{MV("F")},"应付账款",{MV("H")},${ref_col}{r},{MV("C")},${co_col}{r})'
            f'-SUMIFS({MV("J")},{MV("F")},"应付账款",{MV("H")},${ref_col}{r},{MV("C")},${co_col}{r}),2)')


def off_ar(ref_col, co_col, r):
    """这张销售单被《手工凭证》正式抵销掉的应收（贷应收账款 − 借应收账款）"""
    return (f'ROUND(SUMIFS({MV("J")},{MV("F")},"应收账款",{MV("H")},${ref_col}{r},{MV("C")},${co_col}{r})'
            f'-SUMIFS({MV("I")},{MV("F")},"应收账款",{MV("H")},${ref_col}{r},{MV("C")},${co_col}{r}),2)')


def invoiced(ref_col, co_col, kind, r):
    return (f'ROUND(SUMIFS({IV("J")},{IV("K")},${ref_col}{r},{IV("C")},${co_col}{r},'
            f'{IV("D")},"{kind}"),2)')

def st_pay(total, unpaid, paid_, r):
    return (f'IF(ROUND({total},2)=0,"无需付款",IF(ROUND({unpaid},2)=0,"已付清",'
            f'IF(ROUND({paid_},2)=0,"未付款","部分付款")))')

def st_rec(total, unrec, rec_, r):
    return (f'IF(ROUND({total},2)=0,"无需回款",IF(ROUND({unrec},2)=0,"已回款",'
            f'IF(ROUND({rec_},2)=0,"未回款","部分回款")))')

def st_inv(total, done):
    return (f'IF(ROUND({total},2)=0,"无需开票",IF(ABS({done})+0.01>=ABS({total}),"已开票",'
            f'IF(ROUND({done},2)=0,"未开票","部分开票")))')

# ════════════════════════════════════════════════════════════
# ⑩ 采购登记
# ════════════════════════════════════════════════════════════
ws = sheet('采购登记')
title(ws, '采 购 登 记', 'AD',
      '★ 浅黄＝手工填，浅灰＝公式自动。「应收返利」优先用你填的固定金额，没填才按返利率算。\n'
      '★ 实际采购成本 = 不含税金额 − 应收返利 —— 入账的库存商品就是这个数，返利单独挂「其他应收款—应收返利」，'
      '所以返利不会在利润表里重复算一次。\n'
      '★ 退货就在「数量」里填负数，金额、成本、返利会一起变负。付款情况是按《资金流水》里填了同一个「采购单号」的行自动汇总的。')
PUC = ['序号', '日期', '采购单号', '公司', '店铺', '供应商', '商品编码', '商品名称', '规格',
       '数量', '含税单价', '含税金额', '税率', '不含税金额', '税额',
       '返利率', '返利金额(固定·填了就不按比例，填0＝没返利)', '应收返利', '实际采购成本', '单位实际成本',
       '已付/已结', '未付金额', '付款状态', '信用期(天)', '到期日', '逾期天数',
       '已开票金额', '票据状态', '备注', '年月']
headers(ws, HDR, PUC)
widths(ws, {'A': 6, 'B': 11, 'C': 14, 'D': 12, 'E': 14, 'F': 20, 'G': 11, 'H': 16, 'I': 14,
            'J': 9, 'K': 11, 'L': 13, 'M': 8, 'N': 13, 'O': 11,
            'P': 9, 'Q': 13, 'R': 12, 'S': 14, 'T': 13,
            'U': 13, 'V': 13, 'W': 11, 'X': 11, 'Y': 11, 'Z': 10,
            'AA': 13, 'AB': 11, 'AC': 36, 'AD': 9})
IN_COLS_P = 'BCDEFGJKMPQAC'
for i in range(N_PUR):
    r = P0 + i
    src = T.PURCHASE[i] if i < len(T.PURCHASE) else None
    put(ws, f'A{r}', f'=IF($B{r}="","",COUNT($B${P0}:$B{r}))', font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'B{r}', src[0] if src else None, font=F_IN, fill=FILL_IN, fmt=DATEF)
    put(ws, f'C{r}', src[1] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'D{r}', src[2] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'E{r}', src[3] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'F{r}', pn(src[4]) if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'G{r}', src[5] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'H{r}', f'=IF($G{r}="","",IFERROR(INDEX({IT_N},MATCH($G{r},{IT_C},0)),"★档案里没有"))',
        font=F_AUTO, fill=FILL_AUTO, align=CL)
    put(ws, f'I{r}', f'=IF($G{r}="","",IFERROR(INDEX({IT_S},MATCH($G{r},{IT_C},0)),""))',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'J{r}', src[6] if src else None, font=F_IN, fill=FILL_IN, fmt=NUM)
    put(ws, f'K{r}', src[7] if src else None, font=F_IN, fill=FILL_IN, fmt=PRICE)
    put(ws, f'L{r}', f'=IF(OR($J{r}="",$K{r}=""),"",ROUND($J{r}*$K{r},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'M{r}', src[8] if src else None, font=F_IN, fill=FILL_IN, fmt=PCT)
    put(ws, f'N{r}', f'=IF($L{r}="","",ROUND($L{r}/(1+N($M{r})),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'O{r}', f'=IF($L{r}="","",ROUND($L{r}-$N{r},2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'P{r}', src[9] if src and src[9] != '' else None, font=F_IN, fill=FILL_IN, fmt=PCT)
    put(ws, f'Q{r}', src[10] if src and src[10] != '' else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'R{r}', f'=IF($B{r}="","",IF($Q{r}<>"",ROUND($Q{r},2),ROUND(N($N{r})*N($P{r}),2)))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'S{r}', f'=IF($B{r}="","",ROUND(N($N{r})-N($R{r}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'T{r}', f'=IF(OR($B{r}="",N($J{r})=0),"",ROUND($S{r}/$J{r},4))',
        font=F_AUTO, fill=FILL_AUTO, fmt=PRICE)
    put(ws, f'U{r}', f'=IF($B{r}="","",ROUND({paid("C","D",r)}+{off_ap("C","D",r)},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'V{r}', f'=IF($B{r}="","",ROUND(N($L{r})-N($U{r}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'W{r}', f'=IF($B{r}="","",{st_pay(f"N($L{r})", f"N($V{r})", f"N($U{r})", r)})',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'X{r}', f'=IF($B{r}="","",IFERROR(INDEX({PT_CRD},MATCH($F{r},{PT_N},0)),0))',
        font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'Y{r}', f'=IF($B{r}="","",$B{r}+N($X{r}))', font=F_AUTO, fill=FILL_AUTO, fmt=DATEQ)
    put(ws, f'Z{r}', f'=IF(OR($B{r}="",ROUND(N($V{r}),2)=0),"",MAX(0,TODAY()-$Y{r}))',
        font=F_AUTO, fill=FILL_AUTO, fmt='0;;\\-')
    put(ws, f'AA{r}', f'=IF($B{r}="","",{invoiced("C","D","进项",r)})', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'AB{r}', f'=IF($B{r}="","",{st_inv(f"N($L{r})", f"N($AA{r})")})', font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'AC{r}', src[12] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'AD{r}', f'={ym("B", r)}', font=F_AUTO, fill=FILL_AUTO, fmt=YM)
PT_TOT = P1 + 1
put(ws, f'A{PT_TOT}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in list('BCDEFGHIKMPQTXYZ'):
    put(ws, f'{c}{PT_TOT}', None, font=F_TOT, fill=FILL_TOT)
for c in ['J', 'L', 'N', 'O', 'R', 'S', 'U', 'V', 'AA']:
    put(ws, f'{c}{PT_TOT}', f'=ROUND(SUM({c}{P0}:{c}{P1}),2)', font=F_TOT, fill=FILL_TOT,
        fmt=NUM if c == 'J' else MONEY)
for c in ['W', 'AB', 'AC', 'AD']:
    put(ws, f'{c}{PT_TOT}', None, font=F_TOT, fill=FILL_TOT)
dv(ws, f'D{P0}:D{P1}', '=公司表', '哪家公司买的')
dv(ws, f'E{P0}:E{P1}', '=店铺表', '专门给某个店铺备的货才填，一般留空')
dv(ws, f'F{P0}:F{P1}', '=供应商名单', '只列勾了「是否供应商＝是」的往来单位')
dv(ws, f'G{P0}:G{P1}', '=商品编码表')
dv(ws, f'M{P0}:M{P1}', '=税率表')
ws.auto_filter.ref = f'A{HDR}:AD{P1}'
ws.freeze_panes = 'D4'
page(ws, titles=f'{HDR}:{HDR}')

# ════════════════════════════════════════════════════════════
# ⑪ 销售登记
# ════════════════════════════════════════════════════════════
ws = sheet('销售登记')
title(ws, '销 售 登 记', 'AF',
      '★ 单位成本是「截至这一单当天的实际采购加权成本」（含期初库存），'
      '不是一卖就把整批采购全转成本；要手工指定成本就填「手工单位成本」那一列。\n'
      '★ 平台扣费（佣金/技术服务费）直接从货款里扣，所以它一边进销售费用、一边冲应收账款；'
      '「未回款 = 含税金额 − 平台扣费 − 已回/已结」，平台结算打过来的净额正好对上。\n'
      '★ 退款就把数量填负数（扣费也填负数），收入、成本、税一起冲回。')
SAC = ['序号', '日期', '销售单号', '公司', '店铺', '平台', '客户', '商品编码', '商品名称', '规格',
       '数量', '含税单价', '含税金额', '税率', '不含税收入', '销项税额', '平台扣费',
       '单位成本(自动)', '手工单位成本', '销售成本', '毛利', '毛利率',
       '已回/已结', '未回款', '回款状态', '信用期(天)', '到期日', '逾期天数',
       '已开票金额', '票据状态', '备注', '年月']
headers(ws, HDR, SAC)
widths(ws, {'A': 6, 'B': 11, 'C': 14, 'D': 12, 'E': 14, 'F': 10, 'G': 20, 'H': 11, 'I': 16, 'J': 14,
            'K': 9, 'L': 11, 'M': 13, 'N': 8, 'O': 13, 'P': 12, 'Q': 12,
            'R': 13, 'S': 14, 'T': 13, 'U': 12, 'V': 10,
            'W': 13, 'X': 13, 'Y': 11, 'Z': 11, 'AA': 11, 'AB': 10,
            'AC': 13, 'AD': 11, 'AE': 36, 'AF': 9})
COST_NUM = (lambda r: f'(SUMIFS({OPI_A},{OPI_CO},$D{r},{OPI_IT},$H{r})'
                      f'+SUMIFS({PU("S")},{PU("D")},$D{r},{PU("G")},$H{r},{PU("B")},"<="&$B{r}))')
COST_DEN = (lambda r: f'(SUMIFS({OPI_Q},{OPI_CO},$D{r},{OPI_IT},$H{r})'
                      f'+SUMIFS({PU("J")},{PU("D")},$D{r},{PU("G")},$H{r},{PU("B")},"<="&$B{r}))')
for i in range(N_SAL):
    r = S0 + i
    src = T.SALE[i] if i < len(T.SALE) else None
    put(ws, f'A{r}', f'=IF($B{r}="","",COUNT($B${S0}:$B{r}))', font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'B{r}', src[0] if src else None, font=F_IN, fill=FILL_IN, fmt=DATEF)
    put(ws, f'C{r}', src[1] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'D{r}', src[2] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'E{r}', src[3] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'F{r}', f'=IF($E{r}="","",IFERROR(INDEX({SP_PF},MATCH($E{r},{SP_N},0)),""))',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'G{r}', pn(src[4]) if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'H{r}', src[5] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'I{r}', f'=IF($H{r}="","",IFERROR(INDEX({IT_N},MATCH($H{r},{IT_C},0)),"★档案里没有"))',
        font=F_AUTO, fill=FILL_AUTO, align=CL)
    put(ws, f'J{r}', f'=IF($H{r}="","",IFERROR(INDEX({IT_S},MATCH($H{r},{IT_C},0)),""))',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'K{r}', src[6] if src else None, font=F_IN, fill=FILL_IN, fmt=NUM)
    put(ws, f'L{r}', src[7] if src else None, font=F_IN, fill=FILL_IN, fmt=PRICE)
    put(ws, f'M{r}', f'=IF(OR($K{r}="",$L{r}=""),"",ROUND($K{r}*$L{r},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'N{r}', src[8] if src else None, font=F_IN, fill=FILL_IN, fmt=PCT)
    put(ws, f'O{r}', f'=IF($M{r}="","",ROUND($M{r}/(1+N($N{r})),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'P{r}', f'=IF($M{r}="","",ROUND($M{r}-$O{r},2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'Q{r}', src[9] if src else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'R{r}', f'=IF($B{r}="","",IFERROR(ROUND({COST_NUM(r)}/{COST_DEN(r)},4),""))',
        font=F_AUTO, fill=FILL_AUTO, fmt=PRICE)
    put(ws, f'S{r}', None, font=F_IN, fill=FILL_IN, fmt=PRICE)
    put(ws, f'T{r}', f'=IF($B{r}="","",ROUND(IF($S{r}<>"",$S{r},N($R{r}))*N($K{r}),2))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'U{r}', f'=IF($B{r}="","",ROUND(N($O{r})-N($T{r})-N($Q{r}),2))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'V{r}', f'=IF(OR($B{r}="",ROUND(N($O{r}),2)=0),"",ROUND(N($U{r})/$O{r},4))',
        font=F_AUTO, fill=FILL_AUTO, fmt=PCT)
    put(ws, f'W{r}', f'=IF($B{r}="","",ROUND({recv("C","D",r)}+{off_ar("C","D",r)},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'X{r}', f'=IF($B{r}="","",ROUND(N($M{r})-N($Q{r})-N($W{r}),2))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'Y{r}', f'=IF($B{r}="","",{st_rec(f"N($M{r})-N($Q{r})", f"N($X{r})", f"N($W{r})", r)})',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'Z{r}', f'=IF($B{r}="","",IFERROR(INDEX({PT_CRD},MATCH($G{r},{PT_N},0)),0))',
        font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'AA{r}', f'=IF($B{r}="","",$B{r}+N($Z{r}))', font=F_AUTO, fill=FILL_AUTO, fmt=DATEQ)
    put(ws, f'AB{r}', f'=IF(OR($B{r}="",ROUND(N($X{r}),2)=0),"",MAX(0,TODAY()-$AA{r}))',
        font=F_AUTO, fill=FILL_AUTO, fmt='0;;\\-')
    put(ws, f'AC{r}', f'=IF($B{r}="","",{invoiced("C","D","销项",r)})', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'AD{r}', f'=IF($B{r}="","",{st_inv(f"N($M{r})", f"N($AC{r})")})', font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'AE{r}', src[11] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'AF{r}', f'={ym("B", r)}', font=F_AUTO, fill=FILL_AUTO, fmt=YM)
ST_TOT = S1 + 1
put(ws, f'A{ST_TOT}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in list('BCDEFGHIJLNRSVYZAAAB'):
    put(ws, f'{c}{ST_TOT}', None, font=F_TOT, fill=FILL_TOT)
for c in ['K', 'M', 'O', 'P', 'Q', 'T', 'U', 'W', 'X', 'AC']:
    put(ws, f'{c}{ST_TOT}', f'=ROUND(SUM({c}{S0}:{c}{S1}),2)', font=F_TOT, fill=FILL_TOT,
        fmt=NUM if c == 'K' else MONEY)
for c in ['AD', 'AE', 'AF']:
    put(ws, f'{c}{ST_TOT}', None, font=F_TOT, fill=FILL_TOT)
dv(ws, f'D{S0}:D{S1}', '=公司表')
dv(ws, f'E{S0}:E{S1}', '=店铺表', '哪个店卖的 —— 店铺利润分析按这一列汇总')
dv(ws, f'G{S0}:G{S1}', '=客户名单', '只列勾了「是否客户＝是」的往来单位')
dv(ws, f'H{S0}:H{S1}', '=商品编码表')
dv(ws, f'N{S0}:N{S1}', '=税率表')
ws.auto_filter.ref = f'A{HDR}:AF{S1}'
ws.freeze_panes = 'D4'
page(ws, titles=f'{HDR}:{HDR}')

print('  ✓ 采购登记 / 销售登记')

# ════════════════════════════════════════════════════════════
# ⑫ 返利登记
# ════════════════════════════════════════════════════════════
ws = sheet('返利登记')
title(ws, '返 利 登 记（采购返利的确认、到账与发票）', 'P',
      '★ 返利在《采购登记》按单已经预提过一次（挂「其他应收款—应收返利」，同时把库存成本降下来），'
      '所以这张表不再重复确认，只干三件事：跟供应商对账确认金额、跟踪到账、跟踪发票。\n'
      '★「账面应收返利」是采购登记里同公司同供应商同月份的合计，自动算的；'
      '你填的「确认返利金额」跟它的差额（差异调整）才会生成凭证去调成本 —— 这样返利在利润表里只算一次。\n'
      '★ 收到返利：去《资金流水》记一笔「收-返利」，关联单号填这一行的返利单号（FLxxxx）即可。\n'
      '★「确认返利金额」没填 ＝ 还没跟供应商对上账，这一行先不动账（差异调整留空）；'
      '真的一分返利都不给，就明明白白填 0，系统才会把账面预提的那笔冲掉。')
RBC = ['序号', '返利单号', '日期', '公司', '供应商', '返利期间(年月)', '约定返利率',
       '账面应收返利', '确认返利金额', '差异调整', '已收金额', '未收金额',
       '返利发票状态', '返利发票号', '状态', '备注']
headers(ws, HDR, RBC)
widths(ws, {'A': 6, 'B': 13, 'C': 11, 'D': 12, 'E': 22, 'F': 13, 'G': 11,
            'H': 14, 'I': 14, 'J': 12, 'K': 12, 'L': 12, 'M': 13, 'N': 14, 'O': 11, 'P': 40})
for i in range(N_RBT):
    r = B0 + i
    src = T.REBATE[i] if i < len(T.REBATE) else None
    put(ws, f'A{r}', f'=IF($C{r}="","",COUNT($C${B0}:$C{r}))', font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'B{r}', f'=IF($C{r}="","","FL{YEAR}"&TEXT(N($A{r}),"000"))', font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'C{r}', src[0] if src else None, font=F_IN, fill=FILL_IN, fmt=DATEF)
    put(ws, f'D{r}', src[1] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'E{r}', pn(src[2]) if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'F{r}', src[3] if src else None, font=F_IN, fill=FILL_IN, fmt=YM)
    put(ws, f'G{r}', src[4] if src and src[4] != 0 else None, font=F_IN, fill=FILL_IN, fmt=PCT)
    put(ws, f'H{r}', f'=IF($C{r}="","",ROUND(SUMIFS({PU("R")},{PU("D")},$D{r},{PU("F")},$E{r},'
                     f'{PU("AD")},$F{r}),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'I{r}', src[5] if src else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'J{r}', f'=IF(OR($C{r}="",$I{r}=""),"",ROUND(N($I{r})-N($H{r}),2))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'K{r}', f'=IF($C{r}="","",{recv("B","D",r)})', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'L{r}', f'=IF(OR($C{r}="",$I{r}=""),"",ROUND(N($I{r})-N($K{r}),2))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'M{r}', src[6] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'N{r}', src[7] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'O{r}', f'=IF($C{r}="","",IF($I{r}="","★ 待对账确认（确认金额没填，先不动账）",'
                     f'IF(ROUND(N($I{r}),2)=0,"无返利",'
                     f'IF(ROUND(N($L{r}),2)=0,"已收清",IF(ROUND(N($K{r}),2)=0,"未收","部分收")))))',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'P{r}', src[8] if src else None, font=F_IN, fill=FILL_IN, align=CL)
RT = B1 + 1
put(ws, f'A{RT}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in list('BCDEFGMNOP'):
    put(ws, f'{c}{RT}', None, font=F_TOT, fill=FILL_TOT)
for c in ['H', 'I', 'J', 'K', 'L']:
    put(ws, f'{c}{RT}', f'=ROUND(SUM({c}{B0}:{c}{B1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
dv(ws, f'D{B0}:D{B1}', '=公司表')
dv(ws, f'E{B0}:E{B1}', '=供应商名单')
dv(ws, f'M{B0}:M{B1}', '"已开票,未开票,无需开票"')
ws.auto_filter.ref = f'A{HDR}:P{B1}'
ws.freeze_panes = 'C4'
page(ws, titles=f'{HDR}:{HDR}')

# ════════════════════════════════════════════════════════════
# ⑬ 一件代发结算
# ════════════════════════════════════════════════════════════
ws = sheet('一件代发结算')
title(ws, '一 件 代 发 结 算', 'AC',
      '★ 代发的货不进自己仓库，所以不走库存商品：收入照记，成本直接进主营业务成本，'
      '欠代发商的钱挂应付账款 —— 跟正常采购＋销售分开，免得把库存搞乱。\n'
      '★ 回款用「结算单号」关联《资金流水》的「收-销售货款」；付代发商的钱用同一个单号记「付-代发货款」。\n'
      '★ 代发退货：数量填负数，售价成本扣费一起冲回。')
DPC = ['序号', '日期', '结算单号', '公司', '店铺', '代发供应商', '客户/平台', '商品编码', '商品名称',
       '数量', '含税售价', '含税销售额', '销售税率', '不含税收入', '销项税额',
       '代发含税单价', '代发含税成本', '代发税率', '代发不含税成本', '代发进项税',
       '平台扣费', '毛利', '已回款', '未回款', '已付代发款', '未付代发款', '结算状态', '备注', '年月']
headers(ws, HDR, DPC)
widths(ws, {'A': 6, 'B': 11, 'C': 14, 'D': 12, 'E': 14, 'F': 16, 'G': 16, 'H': 11, 'I': 16,
            'J': 9, 'K': 11, 'L': 13, 'M': 10, 'N': 13, 'O': 11,
            'P': 13, 'Q': 14, 'R': 10, 'S': 15, 'T': 12,
            'U': 11, 'V': 12, 'W': 12, 'X': 12, 'Y': 13, 'Z': 13, 'AA': 12, 'AB': 34, 'AC': 9})
for i in range(N_DRP):
    r = F0 + i
    src = T.DROPSHIP[i] if i < len(T.DROPSHIP) else None
    put(ws, f'A{r}', f'=IF($B{r}="","",COUNT($B${F0}:$B{r}))', font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'B{r}', src[0] if src else None, font=F_IN, fill=FILL_IN, fmt=DATEF)
    put(ws, f'C{r}', src[1] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'D{r}', src[2] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'E{r}', src[3] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'F{r}', pn(src[4]) if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'G{r}', pn(src[5]) if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'H{r}', src[6] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'I{r}', f'=IF($H{r}="","",IFERROR(INDEX({IT_N},MATCH($H{r},{IT_C},0)),"★档案里没有"))',
        font=F_AUTO, fill=FILL_AUTO, align=CL)
    put(ws, f'J{r}', src[7] if src else None, font=F_IN, fill=FILL_IN, fmt=NUM)
    put(ws, f'K{r}', src[8] if src else None, font=F_IN, fill=FILL_IN, fmt=PRICE)
    put(ws, f'L{r}', f'=IF(OR($J{r}="",$K{r}=""),"",ROUND($J{r}*$K{r},2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'M{r}', src[9] if src else None, font=F_IN, fill=FILL_IN, fmt=PCT)
    put(ws, f'N{r}', f'=IF($L{r}="","",ROUND($L{r}/(1+N($M{r})),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'O{r}', f'=IF($L{r}="","",ROUND($L{r}-$N{r},2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'P{r}', src[10] if src else None, font=F_IN, fill=FILL_IN, fmt=PRICE)
    put(ws, f'Q{r}', f'=IF(OR($J{r}="",$P{r}=""),"",ROUND($J{r}*$P{r},2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'R{r}', src[11] if src else None, font=F_IN, fill=FILL_IN, fmt=PCT)
    put(ws, f'S{r}', f'=IF($Q{r}="","",ROUND($Q{r}/(1+N($R{r})),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'T{r}', f'=IF($Q{r}="","",ROUND($Q{r}-$S{r},2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'U{r}', src[12] if src else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'V{r}', f'=IF($B{r}="","",ROUND(N($N{r})-N($S{r})-N($U{r}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'W{r}', f'=IF($B{r}="","",ROUND({recv("C","D",r)}+{off_ar("C","D",r)},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'X{r}', f'=IF($B{r}="","",ROUND(N($L{r})-N($U{r})-N($W{r}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'Y{r}', f'=IF($B{r}="","",ROUND({paid("C","D",r)}+{off_ap("C","D",r)},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'Z{r}', f'=IF($B{r}="","",ROUND(N($Q{r})-N($Y{r}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'AA{r}', f'=IF($B{r}="","",IF(AND(ROUND(N($X{r}),2)=0,ROUND(N($Z{r}),2)=0),"已结清",'
                      f'IF(AND(ROUND(N($X{r}),2)<=0,ROUND(N($Z{r}),2)<=0),"退货待冲回",'
                      f'IF(ROUND(N($X{r}),2)>0,IF(ROUND(N($Z{r}),2)>0,"收付都欠","欠客户回款"),"欠代发商货款"))))',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'AB{r}', src[13] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'AC{r}', f'={ym("B", r)}', font=F_AUTO, fill=FILL_AUTO, fmt=YM)
DT = F1 + 1
put(ws, f'A{DT}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in list('BCDEFGHIKMPRAAABAC'):
    put(ws, f'{c}{DT}', None, font=F_TOT, fill=FILL_TOT)
for c in ['J', 'L', 'N', 'O', 'Q', 'S', 'T', 'U', 'V', 'W', 'X', 'Y', 'Z']:
    put(ws, f'{c}{DT}', f'=ROUND(SUM({c}{F0}:{c}{F1}),2)', font=F_TOT, fill=FILL_TOT,
        fmt=NUM if c == 'J' else MONEY)
dv(ws, f'D{F0}:D{F1}', '=公司表')
dv(ws, f'E{F0}:E{F1}', '=店铺表')
dv(ws, f'F{F0}:F{F1}', '=供应商名单')
dv(ws, f'G{F0}:G{F1}', '=客户名单')
dv(ws, f'H{F0}:H{F1}', '=商品编码表')
dv(ws, f'M{F0}:M{F1}', '=税率表')
dv(ws, f'R{F0}:R{F1}', '=税率表')
ws.auto_filter.ref = f'A{HDR}:AC{F1}'
ws.freeze_panes = 'D4'
page(ws, titles=f'{HDR}:{HDR}')

# ════════════════════════════════════════════════════════════
# ⑭ 资金流水
# ════════════════════════════════════════════════════════════
ws = sheet('资金流水')
title(ws, '资 金 流 水', 'R',
      '★ 所有进出钱的动作都记在这里。「业务类型」决定对方科目和现金流类别（见《资金规则》）；'
      '要例外就在「科目手工覆盖」里填一个科目名。\n'
      '★ 关联单号很重要：填了采购单号/销售单号/返利单号/代发结算单号/费用单号，'
      '那边的「已付/已回」才会自动跟着动，《收付款核销中心》也才对得上。\n'
      '★ 「已回款」只认「收-」开头的业务类型，「已付款」只认「付-」开头的 —— '
      '所以一件代发那种同一个单号既收客户钱、又付代发商钱的，两边各算各的，不会互相抵掉。\n'
      '★ 平台货款提现到银行：记两行 —— 平台户「内部转账-转出」＋银行户「内部转账-转入」，'
      '两行都要填「对方账户」，这样不会被当成收入或费用。')
CAC = ['序号', '日期', '公司', '资金账户', '业务类型', '往来单位', '关联单号', '收入金额', '支出金额',
       '账户余额', '摘要', '对方账户(仅内部转账)', '科目手工覆盖',
       '账户对应科目', '对方科目', '现金流类别', '备注', '年月']
headers(ws, HDR, CAC)
widths(ws, {'A': 6, 'B': 11, 'C': 12, 'D': 18, 'E': 16, 'F': 20, 'G': 14, 'H': 13, 'I': 13,
            'J': 14, 'K': 22, 'L': 18, 'M': 16, 'N': 20, 'O': 22, 'P': 12, 'Q': 26, 'R': 9})
for i in range(N_CASH):
    r = K0 + i
    src = T.CASH[i] if i < len(T.CASH) else None
    put(ws, f'A{r}', f'=IF($B{r}="","",COUNT($B${K0}:$B{r}))', font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'B{r}', src[0] if src else None, font=F_IN, fill=FILL_IN, fmt=DATEF)
    put(ws, f'C{r}', src[1] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'D{r}', src[2] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'E{r}', src[3] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'F{r}', pn(src[4]) if src and src[4] else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'G{r}', src[5] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'H{r}', src[6] if src and src[6] else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'I{r}', src[7] if src and src[7] else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'J{r}', f'=IF($D{r}="","",ROUND(IFERROR(INDEX({OPA_V},MATCH($D{r},{OPA_N},0)),0)'
                     f'+SUMIFS($H${K0}:$H{r},$D${K0}:$D{r},$D{r})-SUMIFS($I${K0}:$I{r},$D${K0}:$D{r},$D{r}),2))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'K{r}', src[8] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'L{r}', src[9] if src and src[9] else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'M{r}', src[10] if src and src[10] else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'N{r}', f'=IF($D{r}="","",IFERROR(INDEX({ACC_SUBJ},MATCH($D{r},{ACC_N},0)),"★账户档案里没有"))',
        font=F_AUTO, fill=FILL_AUTO, align=CL)
    put(ws, f'O{r}', f'=IF($B{r}="","",IF($M{r}<>"",$M{r},'
                     f'IF(LEFT($E{r},4)="内部转账",IFERROR(INDEX({ACC_SUBJ},MATCH($L{r},{ACC_N},0)),"★请填对方账户"),'
                     f'IFERROR(INDEX({RULE_S},MATCH($E{r},{RULE_T},0)),"★资金规则里没有这个业务类型"))))',
        font=F_AUTO, fill=FILL_AUTO, align=CL)
    put(ws, f'P{r}', f'=IF($B{r}="","",IF(LEFT($E{r},4)="内部转账","不计入",'
                     f'IFERROR(INDEX({RULE_C},MATCH($E{r},{RULE_T},0)),"经营活动")))',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'Q{r}', src[11] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'R{r}', f'={ym("B", r)}', font=F_AUTO, fill=FILL_AUTO, fmt=YM)
KT = K1 + 1
put(ws, f'A{KT}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in list('BCDEFGJKLMNOPQR'):
    put(ws, f'{c}{KT}', None, font=F_TOT, fill=FILL_TOT)
for c in ['H', 'I']:
    put(ws, f'{c}{KT}', f'=ROUND(SUM({c}{K0}:{c}{K1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
dv(ws, f'C{K0}:C{K1}', '=公司表')
dv(ws, f'D{K0}:D{K1}', '=资金账户表')
dv(ws, f'E{K0}:E{K1}', '=资金业务类型表', '决定对方科目和现金流类别')
dv(ws, f'F{K0}:F{K1}', '=往来单位表')
dv(ws, f'L{K0}:L{K1}', '=资金账户表', '只有内部转账要填')
dv(ws, f'M{K0}:M{K1}', '=科目名称表', '留空＝按资金规则自动判')
ws.auto_filter.ref = f'A{HDR}:R{K1}'
ws.freeze_panes = 'D4'
page(ws, titles=f'{HDR}:{HDR}')

print('  ✓ 返利登记 / 一件代发结算 / 资金流水')

# ════════════════════════════════════════════════════════════
# ⑮ 发票台账
# ════════════════════════════════════════════════════════════
ws = sheet('发票台账')
title(ws, '发 票 台 账（进项 / 销项）', 'N',
      '★ 这张表只管「票」，不生成凭证 —— 税额在采购、销售、费用登记时就已经入账了。\n'
      '★ 填了「关联单号」，采购登记/销售登记的「票据状态」才会从「未开票」变成「已开票」，'
      '《异常预警中心》的缺票提醒也才消得掉。')
IVC = ['序号', '日期', '公司', '票据类型', '往来单位', '发票号码', '不含税金额', '税率', '税额',
       '价税合计', '关联单号', '状态', '备注', '年月']
headers(ws, HDR, IVC)
widths(ws, {'A': 6, 'B': 11, 'C': 12, 'D': 10, 'E': 22, 'F': 15, 'G': 14, 'H': 8, 'I': 12,
            'J': 14, 'K': 14, 'L': 11, 'M': 30, 'N': 9})
for i in range(N_INV):
    r = V0 + i
    src = T.INVOICE[i] if i < len(T.INVOICE) else None
    put(ws, f'A{r}', f'=IF($B{r}="","",COUNT($B${V0}:$B{r}))', font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'B{r}', src[0] if src else None, font=F_IN, fill=FILL_IN, fmt=DATEF)
    put(ws, f'C{r}', src[1] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'D{r}', src[2] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'E{r}', pn(src[3]) if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'F{r}', src[4] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'G{r}', src[5] if src else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'H{r}', src[6] if src else None, font=F_IN, fill=FILL_IN, fmt=PCT)
    put(ws, f'I{r}', f'=IF($G{r}="","",ROUND($G{r}*N($H{r}),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'J{r}', f'=IF($G{r}="","",ROUND($G{r}+N($I{r}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'K{r}', src[7] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'L{r}', src[8] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'M{r}', src[9] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'N{r}', f'={ym("B", r)}', font=F_AUTO, fill=FILL_AUTO, fmt=YM)
VT = V1 + 1
put(ws, f'A{VT}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in list('BCDEFHKLMN'):
    put(ws, f'{c}{VT}', None, font=F_TOT, fill=FILL_TOT)
for c in ['G', 'I', 'J']:
    put(ws, f'{c}{VT}', f'=ROUND(SUM({c}{V0}:{c}{V1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
dv(ws, f'C{V0}:C{V1}', '=公司表')
dv(ws, f'D{V0}:D{V1}', '=票据类型表')
dv(ws, f'E{V0}:E{V1}', '=往来单位表')
dv(ws, f'H{V0}:H{V1}', '=税率表')
dv(ws, f'L{V0}:L{V1}', '=发票状态表')
ws.auto_filter.ref = f'A{HDR}:N{V1}'
ws.freeze_panes = 'C4'
page(ws, titles=f'{HDR}:{HDR}')

# ════════════════════════════════════════════════════════════
# ⑯ 费用及其他
# ════════════════════════════════════════════════════════════
ws = sheet('费用及其他')
title(ws, '费 用 及 其 他', 'S',
      '★ 一律按权责发生制记：先在这里登记（借 费用科目 / 贷 应付账款），付钱的时候再到《资金流水》'
      '记「付-费用」并填上这里的费用单号，应付就冲掉了 —— 跟采购是同一套路。\n'
      '★ 填了「店铺」的费用会进《店铺利润分析》；不填就算公司共同费用。\n'
      '★ 想把某笔支出记成资产而不是费用（比如买设备），把费用项目选成「购置固定资产」即可，'
      '对应科目会自动变成「固定资产」。')
EPC = ['序号', '日期', '公司', '店铺', '费用项目', '往来单位', '费用单号', '摘要',
       '不含税金额', '税率', '税额', '价税合计', '借方科目(费用/资产)', '贷方科目(挂账)',
       '已付/已结', '未付金额', '付款状态', '备注', '年月']
headers(ws, HDR, EPC)
widths(ws, {'A': 6, 'B': 11, 'C': 12, 'D': 14, 'E': 18, 'F': 20, 'G': 14, 'H': 22,
            'I': 14, 'J': 8, 'K': 12, 'L': 14, 'M': 20, 'N': 20,
            'O': 13, 'P': 13, 'Q': 11, 'R': 26, 'S': 9})
for i in range(N_EXP):
    r = E0 + i
    src = T.EXPENSE[i] if i < len(T.EXPENSE) else None
    put(ws, f'A{r}', f'=IF($B{r}="","",COUNT($B${E0}:$B{r}))', font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'B{r}', src[0] if src else None, font=F_IN, fill=FILL_IN, fmt=DATEF)
    put(ws, f'C{r}', src[1] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'D{r}', src[2] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'E{r}', src[3] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'F{r}', pn(src[4]) if src and src[4] else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'G{r}', src[5] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'H{r}', src[6] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'I{r}', src[7] if src else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'J{r}', src[8] if src else None, font=F_IN, fill=FILL_IN, fmt=PCT)
    put(ws, f'K{r}', f'=IF($I{r}="","",ROUND($I{r}*N($J{r}),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'L{r}', f'=IF($I{r}="","",ROUND($I{r}+N($K{r}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'M{r}', f'=IF($E{r}="","",IFERROR(INDEX({EX_S},MATCH($E{r},{EX_I},0)),"★基础资料里没配借方科目"))',
        font=F_AUTO, fill=FILL_AUTO, align=CL)
    put(ws, f'N{r}', f'=IF($E{r}="","",IFERROR(INDEX({EX_C},MATCH($E{r},{EX_I},0)),"应付账款"))',
        font=F_AUTO, fill=FILL_AUTO, align=CL)
    put(ws, f'O{r}', f'=IF($B{r}="","",ROUND({paid("G","C",r)}+{off_ap("G","C",r)},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'P{r}', f'=IF($B{r}="","",ROUND(N($L{r})-N($O{r}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'Q{r}', f'=IF($B{r}="","",{st_pay(f"N($L{r})", f"N($P{r})", f"N($O{r})", r)})',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'R{r}', src[9] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'S{r}', f'={ym("B", r)}', font=F_AUTO, fill=FILL_AUTO, fmt=YM)
ET = E1 + 1
put(ws, f'A{ET}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in list('BCDEFGHJMNQRS'):
    put(ws, f'{c}{ET}', None, font=F_TOT, fill=FILL_TOT)
for c in ['I', 'K', 'L', 'O', 'P']:
    put(ws, f'{c}{ET}', f'=ROUND(SUM({c}{E0}:{c}{E1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
dv(ws, f'C{E0}:C{E1}', '=公司表')
dv(ws, f'D{E0}:D{E1}', '=店铺表', '店铺专属费用才填')
dv(ws, f'E{E0}:E{E1}', '=费用项目表')
dv(ws, f'F{E0}:F{E1}', '=往来单位表')
dv(ws, f'J{E0}:J{E1}', '=税率表')
dv(ws, f'M{E0}:M{E1}', '=科目名称表')
dv(ws, f'N{E0}:N{E1}', '=科目名称表')
ws.auto_filter.ref = f'A{HDR}:S{E1}'
ws.freeze_panes = 'C4'
page(ws, titles=f'{HDR}:{HDR}')

# ════════════════════════════════════════════════════════════
# ⑰ 手工凭证
# ════════════════════════════════════════════════════════════
ws = sheet('手工凭证')
title(ws, '手 工 凭 证（只做业务登记做不出来的那几件事）', 'N',
      '★ 一行＝一个分录行，同一个「凭证字号」的若干行凑成一张凭证。右边会自动校验这张凭证平不平。\n'
      '★ 该在这里做的：折旧摊销、税费计提、正式往来抵销、固定资产处置、存货报废、期末调整。\n'
      '★ 不该在这里做的：采购、销售、返利、收付款、费用 —— 那些在各自的登记页记，系统自动出凭证，'
      '在这儿再记一遍就重复了。\n'
      '★ 往来抵销一定要填「关联单号」：填了采购单号，那张采购单的「已付/已结」才认；'
      '填了销售单号，那张销售单的「已回/已结」才认。')
MVC = ['序号', '日期', '公司', '凭证字号', '摘要', '科目', '往来单位', '关联单号',
       '借方金额', '贷方金额', '用途标记', '本张凭证', '备注', '年月']
headers(ws, HDR, MVC)
widths(ws, {'A': 6, 'B': 11, 'C': 12, 'D': 15, 'E': 34, 'F': 28, 'G': 20, 'H': 14,
            'I': 13, 'J': 13, 'K': 12, 'L': 14, 'M': 40, 'N': 9})
for i in range(N_MAN):
    r = M0 + i
    src = T.MANUAL[i] if i < len(T.MANUAL) else None
    put(ws, f'A{r}', f'=IF($B{r}="","",COUNT($B${M0}:$B{r}))', font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'B{r}', src[0] if src else None, font=F_IN, fill=FILL_IN, fmt=DATEF)
    put(ws, f'C{r}', src[1] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'D{r}', src[2] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'E{r}', src[3] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'F{r}', src[4] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'G{r}', pn(src[5]) if src and src[5] else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'H{r}', src[6] if src and src[6] else None, font=F_IN, fill=FILL_IN)
    put(ws, f'I{r}', src[7] if src and src[7] else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'J{r}', src[8] if src and src[8] else None, font=F_IN, fill=FILL_IN, fmt=MONEY)
    put(ws, f'K{r}', src[9] if src else None, font=F_IN, fill=FILL_IN)
    put(ws, f'L{r}', f'=IF($D{r}="","",IF(ROUND(SUMIF($D${M0}:$D${M1},$D{r},$I${M0}:$I${M1})'
                     f'-SUMIF($D${M0}:$D${M1},$D{r},$J${M0}:$J${M1}),2)=0,"✔ 平",'
                     f'"✘ 差 "&TEXT(SUMIF($D${M0}:$D${M1},$D{r},$I${M0}:$I${M1})'
                     f'-SUMIF($D${M0}:$D${M1},$D{r},$J${M0}:$J${M1}),"#,##0.00")))',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'M{r}', src[10] if src else None, font=F_IN, fill=FILL_IN, align=CL)
    put(ws, f'N{r}', f'={ym("B", r)}', font=F_AUTO, fill=FILL_AUTO, fmt=YM)
MT = M1 + 1
put(ws, f'A{MT}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in list('BCDEFGHKLMN'):
    put(ws, f'{c}{MT}', None, font=F_TOT, fill=FILL_TOT)
for c in ['I', 'J']:
    put(ws, f'{c}{MT}', f'=ROUND(SUM({c}{M0}:{c}{M1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
put(ws, f'K{MT}', f'=IF(ROUND($I{MT}-$J{MT},2)=0,"✔ 全部平","✘ 合计差 "&TEXT($I{MT}-$J{MT},"#,##0.00"))',
    font=F_TOT, fill=FILL_TOT)
dv(ws, f'C{M0}:C{M1}', '=公司表')
dv(ws, f'F{M0}:F{M1}', '=科目名称表')
dv(ws, f'G{M0}:G{M1}', '=往来单位表')
dv(ws, f'K{M0}:K{M1}', '=手工凭证用途表')
ws.auto_filter.ref = f'A{HDR}:N{M1}'
ws.freeze_panes = 'C4'
page(ws, titles=f'{HDR}:{HDR}')

print('  ✓ 发票台账 / 费用及其他 / 手工凭证')

# ════════════════════════════════════════════════════════════
# ⑱ 自动凭证 —— 把业务登记按固定规则展开成分录行
# ════════════════════════════════════════════════════════════
AV_LINES = {'采购': 4, '销售': 7, '返利': 2, '代发': 8, '资金': 2, '费用': 3}
AV_ROWS = (N_PUR * 4 + N_SAL * 7 + N_RBT * 2 + N_DRP * 8 + N_CASH * 2 + N_EXP * 3)
A0 = 4
A1 = A0 + AV_ROWS - 1

ws = sheet('自动凭证')
title(ws, '自 动 凭 证（由业务登记自动展开 · 请勿手工修改）', 'L',
      '★ 采购、销售、返利、一件代发、资金、费用每记一行，这里就自动展开成对应的分录行。'
      '金额是 0 的分录行自动隐去，所以这张表上看到的都是真有内容的。\n'
      '★ 它和《手工凭证》一起，是《科目余额表》和三张报表唯一的取数来源 —— '
      '所以业务登记页改一个数，报表立刻跟着变，不用再录一遍总账。\n'
      '★ 展开规则写在《使用说明》里；想改口径就改那边的规则，别在这张表上手工改数。')
AVC = ['序号', '日期', '公司', '凭证字号', '摘要', '会计科目', '往来单位', '关联单号',
       '借方金额', '贷方金额', '来源', '年月']
headers(ws, HDR, AVC)
widths(ws, {'A': 7, 'B': 11, 'C': 12, 'D': 18, 'E': 30, 'F': 30, 'G': 20, 'H': 14,
            'I': 14, 'J': 14, 'K': 12, 'L': 9})

def Q(s):
    return '"' + s.replace('"', '""') + '"'

def emit(ws, r, active, date_f, co_f, vno_f, sum_f, subj_f, party_f, ref_f, dr_f, cr_f, tag):
    """写一条自动凭证行。active 为假（源行空 / 本行金额为 0）时整行留空"""
    g = lambda x: f'=IF({active},"",{x})'
    put(ws, f'A{r}', f'=IF($B{r}="","",COUNT($B${A0}:$B{r}))', font=F_AUTO, fill=FILL_AUTO, fmt='0', border=None)
    put(ws, f'B{r}', g(date_f), font=F_AUTO, fill=FILL_AUTO, fmt=DATEQ, border=None)
    put(ws, f'C{r}', g(co_f), font=F_AUTO, fill=FILL_AUTO, border=None)
    put(ws, f'D{r}', g(vno_f), font=F_AUTO, fill=FILL_AUTO, border=None)
    put(ws, f'E{r}', g(sum_f), font=F_AUTO, fill=FILL_AUTO, align=CL, border=None)
    put(ws, f'F{r}', g(subj_f), font=F_AUTO, fill=FILL_AUTO, align=CL, border=None)
    put(ws, f'G{r}', g(party_f), font=F_AUTO, fill=FILL_AUTO, align=CL, border=None)
    put(ws, f'H{r}', g(ref_f), font=F_AUTO, fill=FILL_AUTO, border=None)
    put(ws, f'I{r}', g(f'ROUND({dr_f},2)'), font=F_AUTO, fill=FILL_AUTO, fmt=MONEY, border=None)
    put(ws, f'J{r}', g(f'ROUND({cr_f},2)'), font=F_AUTO, fill=FILL_AUTO, fmt=MONEY, border=None)
    put(ws, f'K{r}', g(Q(tag)), font=F_AUTO, fill=FILL_AUTO, border=None)
    put(ws, f'L{r}', f'=IF(OR($B{r}="",NOT(ISNUMBER($B{r}))),"",--TEXT($B{r},"yyyymm"))',
        font=F_AUTO, fill=FILL_AUTO, fmt=YM, border=None)

r = A0
# —— 采购：库存商品 / 进项税 / 应收返利 ← 应付账款 ——
for i in range(N_PUR):
    sr = P0 + i
    S_ = f'N(采购登记!$S{sr})'; O_ = f'N(采购登记!$O{sr})'
    R_ = f'N(采购登记!$R{sr})'; L_ = f'N(采购登记!$L{sr})'
    date_f = f'采购登记!$B{sr}'; co_f = f'采购登记!$D{sr}'
    vno_f = f'"采-"&采购登记!$C{sr}'; ref_f = f'采购登记!$C{sr}'
    party_f = f'采购登记!$F{sr}'
    dead = f'采购登记!$B{sr}=""'
    lines = [
        (Q('库存商品'), S_, '0', f'"采购入库 "&采购登记!$H{sr}'),
        (Q('应交税费—应交增值税(进项税额)'), O_, '0', '"采购进项税额"'),
        (Q('其他应收款—应收返利'), R_, '0', '"按约定预提采购返利"'),
        (Q('应付账款'), '0', L_, f'"采购应付 "&采购登记!$F{sr}'),
    ]
    for subj, dr, cr, smy in lines:
        emit(ws, r, f'OR({dead},ROUND(ABS({dr})+ABS({cr}),2)=0)', date_f, co_f, vno_f, smy,
             subj, party_f, ref_f, dr, cr, '采购登记')
        r += 1
# —— 销售：应收 / 收入 / 销项 / 成本 / 库存 / 平台扣费 ——
for i in range(N_SAL):
    sr = S0 + i
    M_ = f'N(销售登记!$M{sr})'; O_ = f'N(销售登记!$O{sr})'
    P_ = f'N(销售登记!$P{sr})'; T_ = f'N(销售登记!$T{sr})'; Qq = f'N(销售登记!$Q{sr})'
    date_f = f'销售登记!$B{sr}'; co_f = f'销售登记!$D{sr}'
    vno_f = f'"销-"&销售登记!$C{sr}'; ref_f = f'销售登记!$C{sr}'
    party_f = f'销售登记!$G{sr}'
    dead = f'销售登记!$B{sr}=""'
    lines = [
        (Q('应收账款'), M_, '0', f'"销售 "&销售登记!$I{sr}'),
        (Q('主营业务收入'), '0', O_, f'"确认收入 "&销售登记!$E{sr}'),
        (Q('应交税费—应交增值税(销项税额)'), '0', P_, '"销项税额"'),
        (Q('主营业务成本'), T_, '0', f'"结转销售成本 "&销售登记!$I{sr}'),
        (Q('库存商品'), '0', T_, '"结转销售成本"'),
        (Q('销售费用'), Qq, '0', '"平台扣费（佣金/技术服务费）"'),
        (Q('应收账款'), '0', Qq, '"平台扣费直接从货款里扣"'),
    ]
    for subj, dr, cr, smy in lines:
        emit(ws, r, f'OR({dead},ROUND(ABS({dr})+ABS({cr}),2)=0)', date_f, co_f, vno_f, smy,
             subj, party_f, ref_f, dr, cr, '销售登记')
        r += 1
# —— 返利：只做「确认金额 − 账面预提」的差异调整 ——
for i in range(N_RBT):
    sr = B0 + i
    J_ = f'N(返利登记!$J{sr})'
    date_f = f'返利登记!$C{sr}'; co_f = f'返利登记!$D{sr}'
    vno_f = f'"返-"&返利登记!$B{sr}'; ref_f = f'返利登记!$B{sr}'
    party_f = f'返利登记!$E{sr}'
    dead = f'返利登记!$C{sr}=""'
    lines = [
        (Q('其他应收款—应收返利'), f'MAX(0,{J_})', f'MAX(0,-{J_})', '"返利差异调整（对账确认数 − 账面预提数）"'),
        (Q('主营业务成本'), f'MAX(0,-{J_})', f'MAX(0,{J_})', '"返利差异调整冲成本"'),
    ]
    for subj, dr, cr, smy in lines:
        emit(ws, r, f'OR({dead},ROUND(ABS({dr})+ABS({cr}),2)=0)', date_f, co_f, vno_f, smy,
             subj, party_f, ref_f, dr, cr, '返利登记')
        r += 1
# —— 一件代发：收入/成本两头都走，但不碰库存商品 ——
for i in range(N_DRP):
    sr = F0 + i
    L_ = f'N(一件代发结算!$L{sr})'; N_ = f'N(一件代发结算!$N{sr})'
    O_ = f'N(一件代发结算!$O{sr})'; S_ = f'N(一件代发结算!$S{sr})'
    T_ = f'N(一件代发结算!$T{sr})'; Qq = f'N(一件代发结算!$Q{sr})'; U_ = f'N(一件代发结算!$U{sr})'
    date_f = f'一件代发结算!$B{sr}'; co_f = f'一件代发结算!$D{sr}'
    vno_f = f'"代-"&一件代发结算!$C{sr}'; ref_f = f'一件代发结算!$C{sr}'
    cus = f'一件代发结算!$G{sr}'; sup = f'一件代发结算!$F{sr}'
    dead = f'一件代发结算!$B{sr}=""'
    lines = [
        (Q('应收账款'), L_, '0', f'"一件代发销售 "&一件代发结算!$I{sr}', cus),
        (Q('主营业务收入'), '0', N_, f'"确认代发收入 "&一件代发结算!$E{sr}', cus),
        (Q('应交税费—应交增值税(销项税额)'), '0', O_, '"代发销项税额"', cus),
        (Q('主营业务成本'), S_, '0', f'"代发成本 "&一件代发结算!$F{sr}', sup),
        (Q('应交税费—应交增值税(进项税额)'), T_, '0', '"代发进项税额"', sup),
        (Q('应付账款'), '0', Qq, f'"应付代发商 "&一件代发结算!$F{sr}', sup),
        (Q('销售费用'), U_, '0', '"平台扣费（代发）"', cus),
        (Q('应收账款'), '0', U_, '"平台扣费直接从货款里扣"', cus),
    ]
    for subj, dr, cr, smy, pty in lines:
        emit(ws, r, f'OR({dead},ROUND(ABS({dr})+ABS({cr}),2)=0)', date_f, co_f, vno_f, smy,
             subj, pty, ref_f, dr, cr, '一件代发')
        r += 1
# —— 资金流水：账户科目 ←→ 对方科目 ——
for i in range(N_CASH):
    sr = K0 + i
    H_ = f'N(资金流水!$H{sr})'; I_ = f'N(资金流水!$I{sr})'
    date_f = f'资金流水!$B{sr}'; co_f = f'资金流水!$C{sr}'
    vno_f = f'"资-"&TEXT(N(资金流水!$A{sr}),"0000")'; ref_f = f'资金流水!$G{sr}'
    party_f = f'资金流水!$F{sr}'
    smy = f'IF(资金流水!$K{sr}="",资金流水!$E{sr},资金流水!$K{sr})'
    # 内部转账要记两行（一出一入）才能让两个账户的余额都动起来，
    # 但凭证只由「转出」那一行生成（借 对方账户科目 / 贷 本账户科目），已经是一张完整凭证；
    # 「转入」行再生成一次就把这笔钱记了两遍，所以这里直接跳过。
    dead = (f'OR(资金流水!$B{sr}="",AND(LEFT(资金流水!$E{sr},4)="内部转账",'
            f'RIGHT(资金流水!$E{sr},2)="转入"))')
    lines = [
        (f'资金流水!$N{sr}', H_, I_),
        (f'资金流水!$O{sr}', I_, H_),
    ]
    for subj, dr, cr in lines:
        emit(ws, r, f'OR({dead},ROUND(ABS({dr})+ABS({cr}),2)=0)', date_f, co_f, vno_f, smy,
             subj, party_f, ref_f, dr, cr, '资金流水')
        r += 1
# ↑ dead 里已经含了「转入行不出凭证」的判断
# —— 费用及其他：费用科目 + 进项税 ← 应付账款 ——
for i in range(N_EXP):
    sr = E0 + i
    I_ = f'N(费用及其他!$I{sr})'; K_ = f'N(费用及其他!$K{sr})'; L_ = f'N(费用及其他!$L{sr})'
    date_f = f'费用及其他!$B{sr}'; co_f = f'费用及其他!$C{sr}'
    vno_f = f'"费-"&费用及其他!$G{sr}'; ref_f = f'费用及其他!$G{sr}'
    party_f = f'费用及其他!$F{sr}'
    smy = f'费用及其他!$E{sr}&" "&费用及其他!$H{sr}'
    dead = f'费用及其他!$B{sr}=""'
    lines = [
        (f'费用及其他!$M{sr}', I_, '0', smy),
        (Q('应交税费—应交增值税(进项税额)'), K_, '0', '"费用进项税额"'),
        (f'费用及其他!$N{sr}', '0', L_, f'"挂账 "&费用及其他!$F{sr}'),
    ]
    for subj, dr, cr, sm in lines:
        emit(ws, r, f'OR({dead},ROUND(ABS({dr})+ABS({cr}),2)=0)', date_f, co_f, vno_f, sm,
             subj, party_f, ref_f, dr, cr, '费用及其他')
        r += 1
assert r - 1 == A1, (r, A1)
AT = A1 + 1
put(ws, f'A{AT}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in list('BCDEFGHKL'):
    put(ws, f'{c}{AT}', None, font=F_TOT, fill=FILL_TOT)
for c in ['I', 'J']:
    put(ws, f'{c}{AT}', f'=ROUND(SUM({c}{A0}:{c}{A1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
put(ws, f'K{AT}', f'=IF(ROUND($I{AT}-$J{AT},2)=0,"✔ 借贷平衡","✘ 差 "&TEXT($I{AT}-$J{AT},"#,##0.00"))',
    font=F_TOT, fill=FILL_TOT)
ws.auto_filter.ref = f'A{HDR}:L{A1}'
ws.freeze_panes = 'C4'
page(ws, titles=f'{HDR}:{HDR}')

# 凭证池常量
AV = lambda c: f'自动凭证!${c}${A0}:${c}${A1}'
MVv = lambda c: f'手工凭证!${c}${M0}:${c}${M1}'

def vamt(side, subj, ym_lo, ym_hi, co='公司条件'):
    """凭证池（自动＋手工）里某科目、某公司、某月份区间的借方或贷方合计"""
    a_amt = AV('I') if side == 'D' else AV('J')
    m_amt = MVv('I') if side == 'D' else MVv('J')
    return (f'(SUMIFS({a_amt},{AV("F")},{subj},{AV("C")},{co},{AV("L")},">="&{ym_lo},{AV("L")},"<="&{ym_hi})'
            f'+SUMIFS({m_amt},{MVv("F")},{subj},{MVv("C")},{co},{MVv("N")},">="&{ym_lo},{MVv("N")},"<="&{ym_hi}))')

print(f'  ✓ 自动凭证（{AV_ROWS} 行展开位）')

# ════════════════════════════════════════════════════════════
# ⑲ 科目余额表
# ════════════════════════════════════════════════════════════
ws = sheet('科目余额表')
title(ws, '科 目 余 额 表', 'O', None)
put(ws, 'A2', '=本期标题&"　｜　年初＝账套年初数，期初＝滚到查询月上月末，本期＝查询月发生额"',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells('A2:O2')
SBC = ['科目编码', '科目名称', '科目类别', '余额方向', '报表项目',
       '年初借方', '年初贷方', '期初借方', '期初贷方', '本期借方', '本期贷方',
       '本年累计借方', '本年累计贷方', '期末借方', '期末贷方']
headers(ws, HDR, SBC)
widths(ws, {'A': 11, 'B': 30, 'C': 9, 'D': 9, 'E': 15,
            'F': 14, 'G': 14, 'H': 14, 'I': 14, 'J': 14, 'K': 14,
            'L': 15, 'M': 15, 'N': 14, 'O': 14, 'P': 13, 'Q': 13})
SB0, SB1 = 4, 4 + N_SUBJ - 1
for i in range(N_SUBJ):
    r = SB0 + i
    k = i + 1
    put(ws, f'A{r}', f'=IF(INDEX(科目表!$B${CR0}:$B${CR1},{k})="","",INDEX(科目表!$A${CR0}:$A${CR1},{k}))',
        font=F_AUTO, fill=FILL_AUTO, fmt=TXT)
    put(ws, f'B{r}', f'=IF(INDEX(科目表!$B${CR0}:$B${CR1},{k})="","",INDEX(科目表!$B${CR0}:$B${CR1},{k}))',
        font=F_TOT, fill=FILL_AUTO, align=CL)
    for c, src in (('C', 'C'), ('D', 'D'), ('E', 'E')):
        put(ws, f'{c}{r}', f'=IF($B{r}="","",INDEX(科目表!${src}${CR0}:${src}${CR1},{k}))',
            font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'F{r}', f'=IF($B{r}="","",ROUND(SUMIFS({OPEN_D},{OPEN_SUB},$B{r},{OPEN_CO},公司条件),2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'G{r}', f'=IF($B{r}="","",ROUND(SUMIFS({OPEN_C},{OPEN_SUB},$B{r},{OPEN_CO},公司条件),2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'P{r}', f'=IF($B{r}="","",ROUND({vamt("D", f"$B{r}", "年初年月", "起始年月-1")},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY, border=None)
    put(ws, f'Q{r}', f'=IF($B{r}="","",ROUND({vamt("C", f"$B{r}", "年初年月", "起始年月-1")},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY, border=None)
    net_open = f'(N($F{r})-N($G{r})+N($P{r})-N($Q{r}))'
    put(ws, f'H{r}', f'=IF($B{r}="","",ROUND(MAX(0,{net_open}),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'I{r}', f'=IF($B{r}="","",ROUND(MAX(0,-{net_open}),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'J{r}', f'=IF($B{r}="","",ROUND({vamt("D", f"$B{r}", "起始年月", "截止年月")},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'K{r}', f'=IF($B{r}="","",ROUND({vamt("C", f"$B{r}", "起始年月", "截止年月")},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'L{r}', f'=IF($B{r}="","",ROUND({vamt("D", f"$B{r}", "年初年月", "截止年月")},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'M{r}', f'=IF($B{r}="","",ROUND({vamt("C", f"$B{r}", "年初年月", "截止年月")},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    net_end = f'(N($F{r})-N($G{r})+N($L{r})-N($M{r}))'
    put(ws, f'N{r}', f'=IF($B{r}="","",ROUND(MAX(0,{net_end}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'O{r}', f'=IF($B{r}="","",ROUND(MAX(0,-{net_end}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
ws.column_dimensions['P'].hidden = True
ws.column_dimensions['Q'].hidden = True
SBT = SB1 + 1
put(ws, f'A{SBT}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in 'BCDE':
    put(ws, f'{c}{SBT}', None, font=F_TOT, fill=FILL_TOT)
for c in 'FGHIJKLMNO':
    put(ws, f'{c}{SBT}', f'=ROUND(SUM({c}{SB0}:{c}{SB1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
put(ws, f'B{SBT}', f'=IF(AND(ROUND($J{SBT}-$K{SBT},2)=0,ROUND($N{SBT}-$O{SBT},2)=0),'
                   f'"✔ 本期借贷平、期末借贷平","✘ 本期差 "&TEXT($J{SBT}-$K{SBT},"#,##0.00")'
                   f'&"；期末差 "&TEXT($N{SBT}-$O{SBT},"#,##0.00"))',
    font=F_TOT, fill=FILL_TOT, align=CL)
ws.auto_filter.ref = f'A{HDR}:O{SB1}'
ws.freeze_panes = 'C4'
page(ws, titles=f'{HDR}:{HDR}')

SB = lambda c: f'科目余额表!${c}${SB0}:${c}${SB1}'
def rep(item, mode, per):
    """按报表项目汇总。mode: 'DR'＝借减贷（资产/成本），'CR'＝贷减借（负债/权益/收入）
       per: 'cur' 本期 / 'ytd' 本年累计 / 'end' 期末 / 'beg' 年初"""
    cols = {'cur': ('J', 'K'), 'ytd': ('L', 'M'), 'end': ('N', 'O'), 'beg': ('F', 'G')}[per]
    d, c = cols
    if mode == 'DR':
        return f'ROUND(SUMIF({SB("E")},{item},{SB(d)})-SUMIF({SB("E")},{item},{SB(c)}),2)'
    return f'ROUND(SUMIF({SB("E")},{item},{SB(c)})-SUMIF({SB("E")},{item},{SB(d)}),2)'

# ════════════════════════════════════════════════════════════
# ⑳ 利润表
# ════════════════════════════════════════════════════════════
ws = sheet('利润表')
title(ws, '利 润 表', 'D', None)
put(ws, 'A2', '=本期标题&"　｜　单位：元"', font=F_NOTE, align=CL, border=None)
ws.merge_cells('A2:D2')
headers(ws, HDR, ['项　目', '行次', '本期金额', '本年累计'])
widths(ws, {'A': 40, 'B': 8, 'C': 18, 'D': 18})
PL = [
    ('一、营业收入', 1, 'CR', '"营业收入"', False),
    ('　　减：营业成本', 2, 'DR', '"营业成本"', False),
    ('　　　　税金及附加', 3, 'DR', '"税金及附加"', False),
    ('　　　　销售费用', 4, 'DR', '"销售费用"', False),
    ('　　　　管理费用', 5, 'DR', '"管理费用"', False),
    ('　　　　财务费用', 6, 'DR', '"财务费用"', False),
    ('二、营业利润（亏损以“－”号填列）', 7, None, 'C4-C5-C6-C7-C8-C9', True),
    ('　　加：营业外收入', 8, 'CR', '"营业外收入"', False),
    ('　　减：营业外支出', 9, 'DR', '"营业外支出"', False),
    ('三、利润总额（亏损总额以“－”号填列）', 10, None, 'C10+C11-C12', True),
    ('　　减：所得税费用', 11, 'DR', '"所得税费用"', False),
    ('四、净利润（净亏损以“－”号填列）', 12, None, 'C13-C14', True),
]
r = 4
for lbl, no, mode, expr, is_sum in PL:
    bold = lbl.startswith(('一', '二', '三', '四'))
    put(ws, f'A{r}', lbl, font=F_TOT if bold else F_TXT, align=CL,
        fill=FILL_TOT if is_sum else None)
    put(ws, f'B{r}', no, font=F_NOTE, fill=FILL_TOT if is_sum else None, fmt='0')
    if is_sum:
        put(ws, f'C{r}', f'=ROUND({expr},2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
        put(ws, f'D{r}', f'=ROUND({expr.replace("C", "D")},2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
    else:
        put(ws, f'C{r}', f'={rep(expr, mode, "cur")}', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
        put(ws, f'D{r}', f'={rep(expr, mode, "ytd")}', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    r += 1
PL_NET = r - 1        # 净利润所在行
put(ws, f'A{r+1}', '注：收入、成本、费用全部取自《科目余额表》的损益类科目；'
                   '科目余额表又只认《自动凭证》＋《手工凭证》—— 所以业务登记页改一个数，这里立刻跟着变。\n'
                   '　　「本期」＝《查询设置》里选的那个月（选「全年」就是 1-12 月）；「本年累计」＝从 1 月到查询月。\n'
                   '　　返利已经在采购入库时冲减了库存成本，销售结转成本时自然带走，所以不会在这里重复算一次。',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells(f'A{r+1}:D{r+3}')
ws.row_dimensions[r + 1].height = 20
page(ws, landscape=False)
name('净利润本期', f'利润表!$C${PL_NET}')
name('净利润累计', f'利润表!$D${PL_NET}')

# ════════════════════════════════════════════════════════════
# ㉑ 资产负债表
# ════════════════════════════════════════════════════════════
ws = sheet('资产负债表')
title(ws, '资 产 负 债 表', 'G', None)
put(ws, 'A2', '=本期标题&"　｜　单位：元"', font=F_NOTE, align=CL, border=None)
ws.merge_cells('A2:G2')
headers(ws, HDR, ['资　产', '期末余额', '年初余额', '', '负债和所有者权益', '期末余额', '年初余额'])
widths(ws, {'A': 26, 'B': 17, 'C': 17, 'D': 3, 'E': 26, 'F': 17, 'G': 17})
ASSETS = ['货币资金', '应收账款', '预付款项', '其他应收款', '存货', '固定资产', '无形资产']
LIABS = ['短期借款', '应付账款', '预收款项', '应付职工薪酬', '应交税费', '其他应付款']
EQUITY = ['实收资本', '未分配利润']
r = 4
for a in ASSETS:
    put(ws, f'A{r}', '　' + a, font=F_TXT, align=CL)
    put(ws, f'B{r}', f'={rep(Q(a), "DR", "end")}', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'C{r}', f'={rep(Q(a), "DR", "beg")}', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    r += 1
A_END = r - 1
put(ws, f'A{r}', '资 产 总 计', font=F_TOT, fill=FILL_TOT, align=CL)
put(ws, f'B{r}', f'=ROUND(SUM(B4:B{A_END}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
put(ws, f'C{r}', f'=ROUND(SUM(C4:C{A_END}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
BS_A_TOT = r
r2 = 4
for x in LIABS:
    put(ws, f'E{r2}', '　' + x, font=F_TXT, align=CL)
    put(ws, f'F{r2}', f'={rep(Q(x), "CR", "end")}', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'G{r2}', f'={rep(Q(x), "CR", "beg")}', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    r2 += 1
L_END = r2 - 1
put(ws, f'E{r2}', '负 债 合 计', font=F_TOT, fill=FILL_TOT, align=CL)
put(ws, f'F{r2}', f'=ROUND(SUM(F4:F{L_END}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
put(ws, f'G{r2}', f'=ROUND(SUM(G4:G{L_END}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
L_TOT = r2
r2 += 1
put(ws, f'E{r2}', '　实收资本', font=F_TXT, align=CL)
put(ws, f'F{r2}', f'={rep(Q("实收资本"), "CR", "end")}', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
put(ws, f'G{r2}', f'={rep(Q("实收资本"), "CR", "beg")}', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
r2 += 1
put(ws, f'E{r2}', '　未分配利润', font=F_TXT, align=CL)
put(ws, f'F{r2}', f'=ROUND({rep(Q("未分配利润"), "CR", "end")}+N(净利润累计),2)',
    font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
put(ws, f'G{r2}', f'={rep(Q("未分配利润"), "CR", "beg")}', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
EQ_END = r2
r2 += 1
put(ws, f'E{r2}', '所有者权益合计', font=F_TOT, fill=FILL_TOT, align=CL)
put(ws, f'F{r2}', f'=ROUND(SUM(F{L_TOT+1}:F{EQ_END}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
put(ws, f'G{r2}', f'=ROUND(SUM(G{L_TOT+1}:G{EQ_END}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
EQ_TOT = r2
r2 += 1
put(ws, f'E{r2}', '负债和所有者权益总计', font=F_TOT, fill=FILL_TOT, align=CL)
put(ws, f'F{r2}', f'=ROUND($F{L_TOT}+$F{EQ_TOT},2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
put(ws, f'G{r2}', f'=ROUND($G{L_TOT}+$G{EQ_TOT},2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
BS_L_TOT = r2
CHK_R = max(BS_A_TOT, BS_L_TOT) + 2
put(ws, f'A{CHK_R}', '平 衡 检 查', font=F_TOT, fill=FILL_SEC, align=CR)
put(ws, f'B{CHK_R}', f'=IF(ROUND($B${BS_A_TOT}-$F${BS_L_TOT},2)=0,"✔ 期末资产＝负债＋权益",'
                     f'"✘ 期末差 "&TEXT($B${BS_A_TOT}-$F${BS_L_TOT},"#,##0.00"))',
    font=F_TOT, fill=FILL_TOT, align=CL)
ws.merge_cells(f'B{CHK_R}:C{CHK_R}')
put(ws, f'E{CHK_R}', f'=IF(ROUND($C${BS_A_TOT}-$G${BS_L_TOT},2)=0,"✔ 年初资产＝负债＋权益",'
                     f'"✘ 年初差 "&TEXT($C${BS_A_TOT}-$G${BS_L_TOT},"#,##0.00"))',
    font=F_TOT, fill=FILL_TOT, align=CL)
ws.merge_cells(f'E{CHK_R}:G{CHK_R}')
put(ws, f'A{CHK_R+2}',
    '注：① 本表不做「结转本年利润」凭证 —— 未分配利润＝年初未分配利润 ＋ 利润表的本年累计净利润，'
    '所以你不用每月手工结转损益。\n'
    '　　② 固定资产是净额（原值 − 累计折旧）；应交税费是进项、销项、已交税金轧差后的净额，'
    '留抵进项多的时候会显示成负数，这是正常的。\n'
    '　　③ 平衡检查两个都要是 ✔。不平优先查《报表勾稽检查》。',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells(f'A{CHK_R+2}:G{CHK_R+4}')
name('资产总计', f'资产负债表!$B${BS_A_TOT}')
name('货币资金期末', '资产负债表!$B$4')
page(ws, landscape=False)

print('  ✓ 科目余额表 / 利润表 / 资产负债表')

# ════════════════════════════════════════════════════════════
# _自动清单 的区间（先定好，管理表要用）
# ════════════════════════════════════════════════════════════
AX = '_自动清单'
# 「这一行的公司符不符合《查询设置》选的公司」—— 公司条件是 * 或具体公司名
CO = lambda cell: f'COUNTIF({cell},公司条件)>0'
X_PT0, X_PT1 = 4, 4 + N_PARTY - 1                      # 往来单位角色名单
X_PR0, X_PR1 = 4, 4 + N_CO_MAX * N_PARTY - 1           # 公司×往来单位 配对池
X_IT0, X_IT1 = 4, 4 + N_CO_MAX * N_ITEM - 1            # 公司×商品 配对池
X_AR0, X_AR1 = 4, 4 + N_SAL + N_DRP - 1                # 应收核销池
X_AP0, X_AP1 = 4, 4 + N_PUR + N_DRP + N_EXP - 1        # 应付核销池
X_EX0, X_EX1 = 4, 4 + N_PUR + N_SAL + N_RBT + N_DRP + N_CASH + N_EXP - 1   # 异常池
XR = lambda c, a, b: f'{AX}!${c}${a}:${c}${b}'

# ════════════════════════════════════════════════════════════
# ㉒ 现金流简表
# ════════════════════════════════════════════════════════════
ws = sheet('现金流简表')
title(ws, '现 金 流 简 表', 'D', None)
put(ws, 'A2', '=本期标题&"　｜　直接法 · 取自《资金流水》，按《资金规则》里的现金流类别分类"',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells('A2:D2')
headers(ws, HDR, ['项　目', '行次', '本期金额', '本年累计'])
widths(ws, {'A': 46, 'B': 8, 'C': 18, 'D': 18})

def cf(kind, side, per):
    lo, hi = ('起始年月', '截止年月') if per == 'cur' else ('年初年月', '截止年月')
    amt = CA('H') if side == 'in' else CA('I')
    return (f'ROUND(SUMIFS({amt},{CA("P")},{kind},{CA("C")},公司条件,'
            f'{CA("R")},">="&{lo},{CA("R")},"<="&{hi}),2)')

CFL = []
for k, (nm, key) in enumerate([('经营活动', '"经营活动"'), ('投资活动', '"投资活动"'), ('筹资活动', '"筹资活动"')]):
    CFL.append((f'{"一二三"[k]}、{nm}产生的现金流量', None, None, True))
    CFL.append((f'　　现金流入小计', 'in', key, False))
    CFL.append((f'　　现金流出小计', 'out', key, False))
    CFL.append((f'　　{nm}产生的现金流量净额', 'net', key, True))
r = 4
ROWMAP = {}
for lbl, side, key, bold in CFL:
    put(ws, f'A{r}', lbl, font=F_TOT if bold else F_TXT, align=CL, fill=FILL_TOT if side == 'net' else None)
    put(ws, f'B{r}', r - 3, font=F_NOTE, fmt='0', fill=FILL_TOT if side == 'net' else None)
    if side is None:
        put(ws, f'C{r}', None, font=F_TOT, fill=FILL_SEC)
        put(ws, f'D{r}', None, font=F_TOT, fill=FILL_SEC)
    elif side == 'net':
        put(ws, f'C{r}', f'=ROUND(C{r-2}-C{r-1},2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
        put(ws, f'D{r}', f'=ROUND(D{r-2}-D{r-1},2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
        ROWMAP[key] = r
    else:
        put(ws, f'C{r}', f'={cf(key, side, "cur")}', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
        put(ws, f'D{r}', f'={cf(key, side, "ytd")}', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    r += 1
NET_ROWS = list(ROWMAP.values())
put(ws, f'A{r}', '四、内部账户调拨（平台提现等 · 一出一入，轧差应为 0）', font=F_TXT, align=CL)
put(ws, f'B{r}', r - 3, font=F_NOTE, fmt='0')
NOCF = Q('不计入')
put(ws, f'C{r}', '=ROUND(' + cf(NOCF, 'in', 'cur') + '-' + cf(NOCF, 'out', 'cur') + ',2)',
    font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
put(ws, f'D{r}', '=ROUND(' + cf(NOCF, 'in', 'ytd') + '-' + cf(NOCF, 'out', 'ytd') + ',2)',
    font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
INNER = r
r += 1
put(ws, f'A{r}', '五、现金及现金等价物净增加额', font=F_TOT, fill=FILL_TOT, align=CL)
put(ws, f'B{r}', r - 3, font=F_NOTE, fill=FILL_TOT, fmt='0')
for c in 'CD':
    put(ws, f'{c}{r}', f'=ROUND({"+".join(f"{c}{x}" for x in NET_ROWS)}+{c}{INNER},2)',
        font=F_TOT, fill=FILL_TOT, fmt=MONEY)
NETINC = r
r += 1
OPEN_CASH_BASE = f'ROUND(SUMIFS({OPA_V},期初余额!$I${OP4_0}:$I${OP4_1},公司条件),2)'
PRIOR_FLOW = (f'ROUND(SUMIFS({CA("H")},{CA("C")},公司条件,{CA("R")},">="&年初年月,{CA("R")},"<="&(起始年月-1))'
              f'-SUMIFS({CA("I")},{CA("C")},公司条件,{CA("R")},">="&年初年月,{CA("R")},"<="&(起始年月-1)),2)')
put(ws, f'A{r}', '　　加：期初现金及现金等价物余额', font=F_TXT, align=CL)
put(ws, f'B{r}', r - 3, font=F_NOTE, fmt='0')
put(ws, f'C{r}', f'=ROUND({OPEN_CASH_BASE}+{PRIOR_FLOW},2)', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
put(ws, f'D{r}', f'={OPEN_CASH_BASE}', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
BEGCASH = r
r += 1
put(ws, f'A{r}', '六、期末现金及现金等价物余额', font=F_TOT, fill=FILL_TOT, align=CL)
put(ws, f'B{r}', r - 3, font=F_NOTE, fill=FILL_TOT, fmt='0')
for c in 'CD':
    put(ws, f'{c}{r}', f'=ROUND({c}{NETINC}+{c}{BEGCASH},2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
ENDCASH = r
r += 2
put(ws, f'A{r}', '与《资产负债表》货币资金核对', font=F_TOT, fill=FILL_SEC, align=CR)
put(ws, f'C{r}', f'=IF(ROUND($C${ENDCASH}-N(货币资金期末),2)=0,"✔ 一致",'
                 f'"✘ 差 "&TEXT($C${ENDCASH}-N(货币资金期末),"#,##0.00"))',
    font=F_TOT, fill=FILL_TOT)
ws.merge_cells(f'C{r}:D{r}')
name('期末现金', f'现金流简表!$C${ENDCASH}')
put(ws, f'A{r+2}',
    '注：① 现金流类别由《资金规则》里每个业务类型对应的那一列决定，改口径就改那里。\n'
    '　　② 平台货款提现记成「内部转账」，既不算经营流入也不算流出，只在第四行轧差显示，应当永远是 0。\n'
    '　　③ 期初现金＝《期初余额》里各资金账户的期初 ＋ 本年到上月末的净流量；选「全年」时就是年初数。',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells(f'A{r+2}:D{r+4}')
page(ws, landscape=False)

# ════════════════════════════════════════════════════════════
# ㉓ 商品库存与成本
# ════════════════════════════════════════════════════════════
ws = sheet('商品库存与成本')
title(ws, '商 品 库 存 与 成 本（按公司分开算）', 'P', None)
put(ws, 'A2', '=本期标题&"　｜　结存＝年初 ＋ 本年到查询月末的采购 － 销售；跟着《查询设置》的公司走；一件代发不进库存，不在这张表里"',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells('A2:P2')
ICC = ['序号', '公司', '商品编码', '商品名称', '单位',
       '年初数量', '年初金额', '年初单位成本',
       '本年采购数量', '本年采购成本(实际)', '本年销售数量', '本年销售成本',
       '结存数量', '结存金额', '当前单位成本', '提醒']
headers(ws, HDR, ICC)
widths(ws, {'A': 6, 'B': 13, 'C': 12, 'D': 18, 'E': 8,
            'F': 12, 'G': 14, 'H': 14, 'I': 13, 'J': 16, 'K': 13, 'L': 14,
            'M': 12, 'N': 14, 'O': 14, 'P': 30})
IC0, IC1 = 4, 303
for i in range(IC0, IC1 + 1):
    k = i - IC0 + 1
    put(ws, f'A{i}', f'=IF($B{i}="","",{k})', font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'B{i}', f'=IFERROR(INDEX({XR("S", X_IT0, X_IT1)},MATCH({k},{XR("V", X_IT0, X_IT1)},0)),"")',
        font=F_TOT, fill=FILL_AUTO)
    put(ws, f'C{i}', f'=IFERROR(INDEX({XR("T", X_IT0, X_IT1)},MATCH({k},{XR("V", X_IT0, X_IT1)},0)),"")',
        font=F_TOT, fill=FILL_AUTO)
    put(ws, f'D{i}', f'=IF($C{i}="","",IFERROR(INDEX({IT_N},MATCH($C{i},{IT_C},0)),""))',
        font=F_AUTO, fill=FILL_AUTO, align=CL)
    put(ws, f'E{i}', f'=IF($C{i}="","",IFERROR(INDEX({IT_U},MATCH($C{i},{IT_C},0)),""))',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'F{i}', f'=IF($B{i}="","",ROUND(SUMIFS({OPI_Q},{OPI_CO},$B{i},{OPI_IT},$C{i}),4))',
        font=F_AUTO, fill=FILL_AUTO, fmt=NUM)
    put(ws, f'G{i}', f'=IF($B{i}="","",ROUND(SUMIFS({OPI_A},{OPI_CO},$B{i},{OPI_IT},$C{i}),2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'H{i}', f'=IF(OR($B{i}="",N($F{i})=0),"",ROUND($G{i}/$F{i},4))',
        font=F_AUTO, fill=FILL_AUTO, fmt=PRICE)
    put(ws, f'I{i}', f'=IF($B{i}="","",ROUND(SUMIFS({PU("J")},{PU("D")},$B{i},{PU("G")},$C{i},'
                     f'{PU("AD")},">="&年初年月,{PU("AD")},"<="&截止年月),4))',
        font=F_AUTO, fill=FILL_AUTO, fmt=NUM)
    put(ws, f'J{i}', f'=IF($B{i}="","",ROUND(SUMIFS({PU("S")},{PU("D")},$B{i},{PU("G")},$C{i},'
                     f'{PU("AD")},">="&年初年月,{PU("AD")},"<="&截止年月),2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'K{i}', f'=IF($B{i}="","",ROUND(SUMIFS({SA("K")},{SA("D")},$B{i},{SA("H")},$C{i},'
                     f'{SA("AF")},">="&年初年月,{SA("AF")},"<="&截止年月),4))',
        font=F_AUTO, fill=FILL_AUTO, fmt=NUM)
    put(ws, f'L{i}', f'=IF($B{i}="","",ROUND(SUMIFS({SA("T")},{SA("D")},$B{i},{SA("H")},$C{i},'
                     f'{SA("AF")},">="&年初年月,{SA("AF")},"<="&截止年月),2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'M{i}', f'=IF($B{i}="","",ROUND(N($F{i})+N($I{i})-N($K{i}),4))', font=F_TOT, fill=FILL_AUTO, fmt=NUM)
    put(ws, f'N{i}', f'=IF($B{i}="","",ROUND(N($G{i})+N($J{i})-N($L{i}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'O{i}', f'=IF(OR($B{i}="",ROUND(N($M{i}),4)=0),"",ROUND($N{i}/$M{i},4))',
        font=F_AUTO, fill=FILL_AUTO, fmt=PRICE)
    put(ws, f'P{i}', f'=IF($B{i}="","",IF(N($M{i})<0,"★ 负库存：卖得比进得多，查采购是不是漏登了",'
                     f'IF(AND(N($M{i})=0,ROUND(N($N{i}),2)<>0),"★ 数量为 0 但还有金额，成本结转有问题",'
                     f'IF(N($N{i})<0,"★ 负金额：单位成本可能填错了",""))))',
        font=F_WARN, fill=FILL_AUTO, align=CL)
ICT = IC1 + 1
put(ws, f'A{ICT}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in 'BCDEHOP':
    put(ws, f'{c}{ICT}', None, font=F_TOT, fill=FILL_TOT)
for c in ['F', 'G', 'I', 'J', 'K', 'L', 'M', 'N']:
    put(ws, f'{c}{ICT}', f'=ROUND(SUM({c}{IC0}:{c}{IC1}),2)', font=F_TOT, fill=FILL_TOT,
        fmt=NUM if c in 'FIKM' else MONEY)
put(ws, f'P{ICT}', '按商品逐个算出来的小计', font=F_NOTE, fill=FILL_TOT, align=CL)
ADJ_R = ICT + 1
MANU_INV = (f'ROUND(SUMIFS({MVv("I")},{MVv("F")},"库存商品",{MVv("C")},公司条件,'
            f'{MVv("N")},"<="&截止年月)-SUMIFS({MVv("J")},{MVv("F")},"库存商品",{MVv("C")},公司条件,'
            f'{MVv("N")},"<="&截止年月),2)')
put(ws, f'A{ADJ_R}', '调  整', font=F_TOT, fill=FILL_WARN)
put(ws, f'B{ADJ_R}', '手工凭证对「库存商品」的调整', font=F_TOT, fill=FILL_WARN, align=CL)
for c in 'CDEFGHIJKLMO':
    put(ws, f'{c}{ADJ_R}', None, font=F_TOT, fill=FILL_WARN)
put(ws, f'N{ADJ_R}', f'={MANU_INV}', font=F_TOT, fill=FILL_WARN, fmt=MONEY)
put(ws, f'P{ADJ_R}', '报废、盘盈盘亏这些不按商品算，只在总额上调', font=F_NOTE, fill=FILL_WARN, align=CL)
INV_TOT = ADJ_R + 1
put(ws, f'A{INV_TOT}', '合  计', font=F_TOT, fill=FILL_TOT)
put(ws, f'B{INV_TOT}', '（与资产负债表「存货」对）', font=F_TOT, fill=FILL_TOT, align=CL)
for c in 'CDEFGHIJKLMO':
    put(ws, f'{c}{INV_TOT}', None, font=F_TOT, fill=FILL_TOT)
put(ws, f'N{INV_TOT}', f'=ROUND(N($N{ICT})+N($N{ADJ_R}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
put(ws, f'P{INV_TOT}', f'=IF(ROUND($N{INV_TOT}-{rep(Q("存货"), "DR", "end")},2)=0,"✔ 与资产负债表存货一致",'
                       f'"✘ 差 "&TEXT($N{INV_TOT}-{rep(Q("存货"), "DR", "end")},"#,##0.00"))',
    font=F_TOT, fill=FILL_TOT, align=CL)
ws.auto_filter.ref = f'A{HDR}:P{IC1}'
ws.freeze_panes = 'D4'
page(ws, titles=f'{HDR}:{HDR}')

print('  ✓ 现金流简表 / 商品库存与成本')

# ════════════════════════════════════════════════════════════
# ㉔ 综合往来对账（按「公司 + 往来单位」，不跨公司抵销）
# ════════════════════════════════════════════════════════════
def vpt(side, subj, party, co, lo='年初年月', hi='截止年月'):
    """凭证池里某科目、某公司、某往来单位的借方或贷方累计"""
    a_amt = AV('I') if side == 'D' else AV('J')
    m_amt = MVv('I') if side == 'D' else MVv('J')
    return (f'(SUMIFS({a_amt},{AV("F")},{subj},{AV("C")},{co},{AV("G")},{party},'
            f'{AV("L")},">="&{lo},{AV("L")},"<="&{hi})'
            f'+SUMIFS({m_amt},{MVv("F")},{subj},{MVv("C")},{co},{MVv("G")},{party},'
            f'{MVv("N")},">="&{lo},{MVv("N")},"<="&{hi}))')

ws = sheet('综合往来对账')
title(ws, '综 合 往 来 对 账（按「公司 + 往来单位」汇总）', 'S', None)
put(ws, 'A2', '=本期标题&"　｜　★ 同一家单位在不同公司的账是分开的，甲公司欠它的钱不会自动拿乙公司的应收去冲；'
              '应付货款、应收销售款、应收返利、借款也各算各的，想抵销必须在《手工凭证》做正式抵销凭证"',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells('A2:S2')
PTC = ['序号', '公司', '往来单位', '身份',
       '应付发生', '已付/已结', '应付余额',
       '应收发生', '已收/已结', '应收余额',
       '返利应收发生', '返利已收', '返利余额',
       '借款余额', '经营净往来', '净资金敞口', '欠款方向', '跨公司提示', '备注']
headers(ws, HDR, PTC)
widths(ws, {'A': 6, 'B': 13, 'C': 24, 'D': 22,
            'E': 14, 'F': 14, 'G': 14, 'H': 14, 'I': 14, 'J': 14,
            'K': 14, 'L': 13, 'M': 13, 'N': 14, 'O': 15, 'P': 15,
            'Q': 12, 'R': 26, 'S': 20})
PTD0, PTD1 = 4, 203
AP_S, AR_S = Q('应付账款'), Q('应收账款')
RB_S = Q('其他应收款—应收返利')
LN1, LN2 = Q('短期借款'), Q('其他应付款—股东借款')
for i in range(PTD0, PTD1 + 1):
    k = i - PTD0 + 1
    put(ws, f'A{i}', f'=IF($B{i}="","",{k})', font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'B{i}', f'=IFERROR(INDEX({XR("O", X_PR0, X_PR1)},MATCH({k},{XR("R", X_PR0, X_PR1)},0)),"")',
        font=F_TOT, fill=FILL_AUTO)
    put(ws, f'C{i}', f'=IFERROR(INDEX({XR("P", X_PR0, X_PR1)},MATCH({k},{XR("R", X_PR0, X_PR1)},0)),"")',
        font=F_TOT, fill=FILL_AUTO, align=CL)
    put(ws, f'D{i}', f'=IF($C{i}="","",IFERROR(SUBSTITUTE(TRIM('
                     f'IF(INDEX({PT_SUP},MATCH($C{i},{PT_N},0))="是","供应商 ","")&'
                     f'IF(INDEX({PT_CUS},MATCH($C{i},{PT_N},0))="是","客户 ","")&'
                     f'IF(INDEX({PT_FIN},MATCH($C{i},{PT_N},0))="是","资方 ","")&'
                     f'IF(INDEX({PT_PAR},MATCH($C{i},{PT_N},0))="是","店铺合作方","")),"  "," "),""))',
        font=F_AUTO, fill=FILL_AUTO)
    op_ap = f'SUMIFS({OPP_AP},{OPP_CO},$B{i},{OPP_PT},$C{i})'
    op_ar = f'SUMIFS({OPP_AR},{OPP_CO},$B{i},{OPP_PT},$C{i})'
    op_rb = f'SUMIFS({OPP_RB},{OPP_CO},$B{i},{OPP_PT},$C{i})'
    op_ln = f'SUMIFS({OPP_LN},{OPP_CO},$B{i},{OPP_PT},$C{i})'
    put(ws, f'E{i}', f'=IF($B{i}="","",ROUND({vpt("C", AP_S, f"$C{i}", f"$B{i}")},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'F{i}', f'=IF($B{i}="","",ROUND({vpt("D", AP_S, f"$C{i}", f"$B{i}")},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'G{i}', f'=IF($B{i}="","",ROUND({op_ap}+N($E{i})-N($F{i}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'H{i}', f'=IF($B{i}="","",ROUND({vpt("D", AR_S, f"$C{i}", f"$B{i}")},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'I{i}', f'=IF($B{i}="","",ROUND({vpt("C", AR_S, f"$C{i}", f"$B{i}")},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'J{i}', f'=IF($B{i}="","",ROUND({op_ar}+N($H{i})-N($I{i}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'K{i}', f'=IF($B{i}="","",ROUND({vpt("D", RB_S, f"$C{i}", f"$B{i}")},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'L{i}', f'=IF($B{i}="","",ROUND({vpt("C", RB_S, f"$C{i}", f"$B{i}")},2))',
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'M{i}', f'=IF($B{i}="","",ROUND({op_rb}+N($K{i})-N($L{i}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    ln = (f'{vpt("C", LN1, f"$C{i}", f"$B{i}")}-{vpt("D", LN1, f"$C{i}", f"$B{i}")}'
          f'+{vpt("C", LN2, f"$C{i}", f"$B{i}")}-{vpt("D", LN2, f"$C{i}", f"$B{i}")}')
    put(ws, f'N{i}', f'=IF($B{i}="","",ROUND({op_ln}+{ln},2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'O{i}', f'=IF($B{i}="","",ROUND(N($J{i})-N($G{i}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'P{i}', f'=IF($B{i}="","",ROUND(N($J{i})+N($M{i})-N($G{i})-N($N{i}),2))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'Q{i}', f'=IF($B{i}="","",IF(ROUND(N($P{i}),2)=0,"两清",IF(N($P{i})>0,"对方欠我","我欠对方")))',
        font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'T{i}', f'=IF($B{i}="",0,IF(ROUND(ABS(N($G{i}))+ABS(N($J{i}))+ABS(N($M{i}))'
                     f'+ABS(N($N{i})),2)>0,1,0))', font=F_AUTO, fill=FILL_AUTO, fmt='0', border=None)
    put(ws, f'U{i}', f'=IF($T{i}=0,0,IF(SUMPRODUCT(($C${PTD0}:$C${PTD1}=$C{i})*'
                     f'($T${PTD0}:$T${PTD1}))>1,1,0))', font=F_AUTO, fill=FILL_AUTO, fmt='0', border=None)
    put(ws, f'R{i}', f'=IF($B{i}="","",IF($U{i}=1,"⚠ 这家在 "&SUMPRODUCT(($C${PTD0}:$C${PTD1}=$C{i})'
                     f'*($T${PTD0}:$T${PTD1}))&" 家公司都有余额，不能互相抵销",""))',
        font=F_WARN, fill=FILL_AUTO, align=CL)
    put(ws, f'S{i}', None, font=F_IN, fill=FILL_IN, align=CL)
# 工资、税费这些没有往来对象的，也要有个去处，否则合计跟资产负债表对不上
UNS = PTD1 + 1
put(ws, f'A{UNS}', '—', font=F_TOT, fill=FILL_WARN)
put(ws, f'B{UNS}', '（全部公司）', font=F_TOT, fill=FILL_WARN)
put(ws, f'C{UNS}', '未指定往来单位', font=F_TOT, fill=FILL_WARN, align=CL)
put(ws, f'D{UNS}', '工资/税费/手续费等没有往来对象的挂账', font=F_NOTE, fill=FILL_WARN, align=CL)
TOTALS = {
    'E': f'{vamt("C", AP_S, "年初年月", "截止年月")}',
    'F': f'{vamt("D", AP_S, "年初年月", "截止年月")}',
    'G': f'SUMIFS({OPEN_C},{OPEN_SUB},"应付账款",{OPEN_CO},公司条件)'
         f'+{vamt("C", AP_S, "年初年月", "截止年月")}-{vamt("D", AP_S, "年初年月", "截止年月")}',
    'H': f'{vamt("D", AR_S, "年初年月", "截止年月")}',
    'I': f'{vamt("C", AR_S, "年初年月", "截止年月")}',
    'J': f'SUMIFS({OPEN_D},{OPEN_SUB},"应收账款",{OPEN_CO},公司条件)'
         f'+{vamt("D", AR_S, "年初年月", "截止年月")}-{vamt("C", AR_S, "年初年月", "截止年月")}',
    'K': f'{vamt("D", RB_S, "年初年月", "截止年月")}',
    'L': f'{vamt("C", RB_S, "年初年月", "截止年月")}',
    'M': f'SUMIFS({OPEN_D},{OPEN_SUB},"其他应收款—应收返利",{OPEN_CO},公司条件)'
         f'+{vamt("D", RB_S, "年初年月", "截止年月")}-{vamt("C", RB_S, "年初年月", "截止年月")}',
    'N': f'SUMIFS({OPEN_C},{OPEN_SUB},"短期借款",{OPEN_CO},公司条件)'
         f'+SUMIFS({OPEN_C},{OPEN_SUB},"其他应付款—股东借款",{OPEN_CO},公司条件)'
         f'+{vamt("C", LN1, "年初年月", "截止年月")}-{vamt("D", LN1, "年初年月", "截止年月")}'
         f'+{vamt("C", LN2, "年初年月", "截止年月")}-{vamt("D", LN2, "年初年月", "截止年月")}',
}
for c, f in TOTALS.items():
    put(ws, f'{c}{UNS}', f'=ROUND(({f})-SUM({c}{PTD0}:{c}{PTD1}),2)',
        font=F_TOT, fill=FILL_WARN, fmt=MONEY)
put(ws, f'O{UNS}', f'=ROUND(N($J{UNS})-N($G{UNS}),2)', font=F_TOT, fill=FILL_WARN, fmt=MONEY)
put(ws, f'P{UNS}', f'=ROUND(N($J{UNS})+N($M{UNS})-N($G{UNS})-N($N{UNS}),2)',
    font=F_TOT, fill=FILL_WARN, fmt=MONEY)
for c in ['Q', 'R', 'S']:
    put(ws, f'{c}{UNS}', None, font=F_TOT, fill=FILL_WARN)
PTT = PTD1 + 2
put(ws, f'A{PTT}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in ['B', 'C', 'D', 'Q', 'R', 'S']:
    put(ws, f'{c}{PTT}', None, font=F_TOT, fill=FILL_TOT)
for c in list('EFGHIJKLMNOP'):
    put(ws, f'{c}{PTT}', f'=ROUND(SUM({c}{PTD0}:{c}{UNS}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
ws.column_dimensions['T'].hidden = True
ws.column_dimensions['U'].hidden = True
ws.auto_filter.ref = f'A{HDR}:S{PTD1}'
ws.freeze_panes = 'D4'
page(ws, titles=f'{HDR}:{HDR}')

# ════════════════════════════════════════════════════════════
# ㉕ 收付款核销中心
# ════════════════════════════════════════════════════════════
ws = sheet('收付款核销中心')
title(ws, '收 付 款 核 销 中 心（只列还没结清的单）', 'L', None)
put(ws, 'A2', '=本期标题&"　｜　跟着《查询设置》选的公司走，但不分月份 —— 没结清就是没结清。上半张＝客户/平台还欠我们的，下半张＝我们还欠供应商的。结清的单不在这里出现"',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells('A2:L2')
WOC = ['序号', '来源', '单　号', '日期', '公司', '往来单位', '应收/应付金额',
       '已收/已付', '未结金额', '状态', '逾期天数', '账龄段']
def wo_block(ws, hdr_row, r0, r1, pool_cols, pool_a, pool_b, label, kind):
    put(ws, f'A{hdr_row-1}', label, font=F_SEC, fill=FILL_SEC, align=CL)
    ws.merge_cells(f'A{hdr_row-1}:L{hdr_row-1}')
    headers(ws, hdr_row, WOC, fill=FILL_HDR2, font=F_HDR2, height=30)
    src, num, dt, co, pt, amt, done, left = pool_cols
    for i in range(r0, r1 + 1):
        k = i - r0 + 1
        idx = f'MATCH({k},{XR(pool_b, pool_a[0], pool_a[1])},0)'
        g = lambda c: f'=IFERROR(INDEX({XR(c, pool_a[0], pool_a[1])},{idx}),"")'
        put(ws, f'A{i}', f'=IF($C{i}="","",{k})', font=F_AUTO, fill=FILL_AUTO, fmt='0')
        put(ws, f'B{i}', g(src), font=F_AUTO, fill=FILL_AUTO)
        put(ws, f'C{i}', g(num), font=F_TOT, fill=FILL_AUTO)
        put(ws, f'D{i}', g(dt), font=F_AUTO, fill=FILL_AUTO, fmt=DATEQ)
        put(ws, f'E{i}', g(co), font=F_AUTO, fill=FILL_AUTO)
        put(ws, f'F{i}', g(pt), font=F_AUTO, fill=FILL_AUTO, align=CL)
        put(ws, f'G{i}', g(amt), font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
        put(ws, f'H{i}', g(done), font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
        put(ws, f'I{i}', g(left), font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
        if kind == 'AR':
            put(ws, f'J{i}', f'=IF($C{i}="","",IF(ROUND(N($H{i}),2)=0,"一分未回","部分回款"))',
                font=F_AUTO, fill=FILL_AUTO)
        else:
            put(ws, f'J{i}', f'=IF($C{i}="","",IF(ROUND(N($H{i}),2)=0,"一分未付","部分付款"))',
                font=F_AUTO, fill=FILL_AUTO)
        put(ws, f'K{i}', f'=IF($C{i}="","",MAX(0,TODAY()-N($D{i})-IFERROR(INDEX({PT_CRD},'
                         f'MATCH($F{i},{PT_N},0)),0)))', font=F_AUTO, fill=FILL_AUTO, fmt='0;;\\-')
        put(ws, f'L{i}', f'=IF($C{i}="","",IF(N($K{i})=0,"未到期",IF(N($K{i})<=30,"逾期30天内",'
                         f'IF(N($K{i})<=60,"逾期31-60天",IF(N($K{i})<=90,"逾期61-90天","逾期90天以上")))))',
            font=F_AUTO, fill=FILL_AUTO)
    tr = r1 + 1
    put(ws, f'A{tr}', '小  计', font=F_TOT, fill=FILL_TOT)
    for c in 'BCDEFJKL':
        put(ws, f'{c}{tr}', None, font=F_TOT, fill=FILL_TOT)
    for c in 'GHI':
        put(ws, f'{c}{tr}', f'=ROUND(SUM({c}{r0}:{c}{r1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
    return tr

widths(ws, {'A': 6, 'B': 14, 'C': 16, 'D': 11, 'E': 13, 'F': 24, 'G': 16,
            'H': 15, 'I': 15, 'J': 12, 'K': 11, 'L': 14})
AR_T = wo_block(ws, 4, 5, 204, ('X', 'Y', 'Z', 'AA', 'AB', 'AC', 'AD', 'AE'),
                (X_AR0, X_AR1), 'AG', '一、应收未结（销售 / 一件代发）—— 客户和平台还欠我们的', 'AR')
AP_T = wo_block(ws, AR_T + 3, AR_T + 4, AR_T + 203, ('AI', 'AJ', 'AK', 'AL', 'AM', 'AN', 'AO', 'AP'),
                (X_AP0, X_AP1), 'AR', '二、应付未结（采购 / 一件代发 / 费用）—— 我们还欠供应商的', 'AP')
put(ws, f'A{AP_T+2}',
    '注：① 一张单出现在这里，就说明它还没结清。结清的判断＝「资金流水里按单号收/付的钱」＋'
    '「手工凭证里按单号做的正式抵销」≥ 单据金额。\n'
    '　　② 逾期天数＝今天 − 单据日期 − 这家往来单位的信用期（信用期在《往来单位主档》里设）。\n'
    '　　③ 想让某笔钱认到某张单上，《资金流水》的「关联单号」一定要填对，否则这里永远挂着。',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells(f'A{AP_T+2}:L{AP_T+4}')
ws.freeze_panes = 'C5'
page(ws)

# ════════════════════════════════════════════════════════════
# ㉖ 店铺利润分析
# ════════════════════════════════════════════════════════════
ws = sheet('店铺利润分析')
title(ws, '店 铺 利 润 分 析', 'Q', None)
put(ws, 'A2', '=本期标题&"　｜　自营（走库存）和一件代发分开列，最后合成店铺净利；'
              '公司层面的共同费用（没填店铺的）不摊到店铺上，单独在最后一行显示"',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells('A2:Q2')
SPC = ['序号', '店铺', '所属公司', '平台',
       '自营收入', '自营成本', '自营平台扣费', '自营毛利',
       '代发收入', '代发成本', '代发平台扣费', '代发毛利',
       '店铺直接费用', '店铺净利', '毛利率', '净利率', '订单笔数']
headers(ws, HDR, SPC)
widths(ws, {'A': 6, 'B': 18, 'C': 13, 'D': 10,
            'E': 14, 'F': 14, 'G': 14, 'H': 14,
            'I': 14, 'J': 14, 'K': 14, 'L': 14,
            'M': 14, 'N': 14, 'O': 10, 'P': 10, 'Q': 10})
SPD0, SPD1 = 4, 4 + N_SHOP - 1
YMR = lambda rng: f'{rng},">="&起始年月,{rng},"<="&截止年月'
for i in range(SPD0, SPD1 + 1):
    k = i - SPD0 + 1
    put(ws, f'A{i}', f'=IF($B{i}="","",{k})', font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'B{i}', f'=IF(INDEX({SP_N},{k})="","",INDEX({SP_N},{k}))', font=F_TOT, fill=FILL_AUTO, align=CL)
    put(ws, f'C{i}', f'=IF($B{i}="","",INDEX({SP_CO},{k}))', font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'D{i}', f'=IF($B{i}="","",INDEX({SP_PF},{k}))', font=F_AUTO, fill=FILL_AUTO)
    base_s = f'{SA("E")},$B{i},{SA("D")},公司条件,{YMR(SA("AF"))}'
    base_d = f'{DP("E")},$B{i},{DP("D")},公司条件,{YMR(DP("AC"))}'
    base_e = f'{EP("D")},$B{i},{EP("C")},公司条件,{YMR(EP("S"))}'
    put(ws, f'E{i}', f'=IF($B{i}="","",ROUND(SUMIFS({SA("O")},{base_s}),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'F{i}', f'=IF($B{i}="","",ROUND(SUMIFS({SA("T")},{base_s}),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'G{i}', f'=IF($B{i}="","",ROUND(SUMIFS({SA("Q")},{base_s}),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'H{i}', f'=IF($B{i}="","",ROUND(N($E{i})-N($F{i})-N($G{i}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'I{i}', f'=IF($B{i}="","",ROUND(SUMIFS({DP("N")},{base_d}),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'J{i}', f'=IF($B{i}="","",ROUND(SUMIFS({DP("S")},{base_d}),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'K{i}', f'=IF($B{i}="","",ROUND(SUMIFS({DP("U")},{base_d}),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'L{i}', f'=IF($B{i}="","",ROUND(N($I{i})-N($J{i})-N($K{i}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'M{i}', f'=IF($B{i}="","",ROUND(SUMIFS({EP("I")},{base_e}),2))', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'N{i}', f'=IF($B{i}="","",ROUND(N($H{i})+N($L{i})-N($M{i}),2))', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'O{i}', f'=IF(OR($B{i}="",ROUND(N($E{i})+N($I{i}),2)=0),"",'
                     f'ROUND((N($H{i})+N($L{i}))/(N($E{i})+N($I{i})),4))', font=F_AUTO, fill=FILL_AUTO, fmt=PCT)
    put(ws, f'P{i}', f'=IF(OR($B{i}="",ROUND(N($E{i})+N($I{i}),2)=0),"",'
                     f'ROUND(N($N{i})/(N($E{i})+N($I{i})),4))', font=F_AUTO, fill=FILL_AUTO, fmt=PCT)
    put(ws, f'Q{i}', f'=IF($B{i}="","",COUNTIFS({base_s})+COUNTIFS({base_d}))',
        font=F_AUTO, fill=FILL_AUTO, fmt='0')
SPT = SPD1 + 1
put(ws, f'A{SPT}', '小  计', font=F_TOT, fill=FILL_TOT)
for c in 'BCDOP':
    put(ws, f'{c}{SPT}', None, font=F_TOT, fill=FILL_TOT)
for c in list('EFGHIJKLMN'):
    put(ws, f'{c}{SPT}', f'=ROUND(SUM({c}{SPD0}:{c}{SPD1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
put(ws, f'Q{SPT}', f'=SUM(Q{SPD0}:Q{SPD1})', font=F_TOT, fill=FILL_TOT, fmt='0')
CM = SPT + 1
put(ws, f'A{CM}', '公司共同费用', font=F_TOT, fill=FILL_WARN)
put(ws, f'B{CM}', '（没填店铺的那些费用，不摊到店铺）', font=F_NOTE, fill=FILL_WARN, align=CL)
for c in 'CDEFGHIJKLOPQ':
    put(ws, f'{c}{CM}', None, font=F_TOT, fill=FILL_WARN)
put(ws, f'M{CM}', f'=ROUND(SUMIFS({EP("I")},{EP("D")},"",{EP("C")},公司条件,{YMR(EP("S"))}),2)',
    font=F_TOT, fill=FILL_WARN, fmt=MONEY)
put(ws, f'N{CM}', f'=ROUND(-N($M{CM}),2)', font=F_TOT, fill=FILL_WARN, fmt=MONEY)
GT = CM + 1
put(ws, f'A{GT}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in 'BCDOPQ':
    put(ws, f'{c}{GT}', None, font=F_TOT, fill=FILL_TOT)
for c in list('EFGHIJKLMN'):
    put(ws, f'{c}{GT}', f'=ROUND({c}{SPT}+N({c}{CM}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
DIFF_R = GT + 1
put(ws, f'A{DIFF_R}', '与利润表对账', font=F_TOT, fill=FILL_SEC, align=CR)
put(ws, f'B{DIFF_R}', '利润表「本期净利润」', font=F_TXT, fill=FILL_SEC, align=CL)
put(ws, f'C{DIFF_R}', f'=ROUND(利润表!$C${PL_NET},2)', font=F_TOT, fill=FILL_SEC, fmt=MONEY)
put(ws, f'D{DIFF_R}', '差额', font=F_TXT, fill=FILL_SEC)
put(ws, f'E{DIFF_R}', f'=ROUND($N{GT}-$C{DIFF_R},2)', font=F_TOT, fill=FILL_SEC, fmt=MONEY)
put(ws, f'F{DIFF_R}', '＝折旧、税金及附加、所得税、营业外收支等「不摊到店铺」的项目',
    font=F_NOTE, fill=FILL_SEC, align=CL)
for c in 'GHIJKLMNOPQ':
    put(ws, f'{c}{DIFF_R}', None, font=F_TOT, fill=FILL_SEC)
ws.merge_cells(f'F{DIFF_R}:Q{DIFF_R}')
put(ws, f'A{DIFF_R+2}',
    '注：① 这里的「净利」是管理口径 —— 只扣了直接挂在店铺上的费用（《费用及其他》里填了店铺的那些），'
    '折旧、税金及附加、所得税、营业外收支这些不摊到店铺，所以上面那一行的差额是正常的、不是错。\n'
    '　　② 想让一笔费用进某个店铺，在《费用及其他》把「店铺」那一列填上即可；不填就算公司共同费用。\n'
    '　　③ 自营毛利已经扣掉了平台扣费；代发毛利同理。返利已经在采购成本里冲过，不在这里单列。\n'
    '　　④ 本表跟着《查询设置》的公司和月份走。',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells(f'A{DIFF_R+2}:Q{DIFF_R+5}')
ws.freeze_panes = 'C4'
page(ws, titles=f'{HDR}:{HDR}')

print('  ✓ 综合往来对账 / 收付款核销中心 / 店铺利润分析')

# ════════════════════════════════════════════════════════════
# ㉗ 异常预警中心
# ════════════════════════════════════════════════════════════
name('报表项目清单', '{"货币资金","应收账款","预付款项","其他应收款","存货","固定资产","无形资产","短期借款","应付账款","预收款项","应付职工薪酬","应交税费","其他应付款","实收资本","未分配利润","营业收入","营业成本","税金及附加","销售费用","管理费用","财务费用","营业外收入","营业外支出","所得税费用"}')

ws = sheet('异常预警中心')
title(ws, '异 常 预 警 中 心', 'H', None)
put(ws, 'A2', '"★ 预警永远看全量数据（不分公司、不分月份）—— 免得切了查询条件就漏掉老问题。上半张是规则汇总（看有没有红的），下半张是逐笔明细（照着去改）"',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells('A2:H2')
headers(ws, HDR, ['编号', '预警规则', '触发笔数', '涉及金额', '严重度', '怎么处理', '', ''])
widths(ws, {'A': 7, 'B': 34, 'C': 11, 'D': 16, 'E': 10, 'F': 60, 'G': 14, 'H': 14})
NOTBLANK = '"<>"'
LINKED_R = ('(LEFT(' + CA("E") + ',4)="付-采购")+(LEFT(' + CA("E") + ',4)="付-代发")'
            '+(LEFT(' + CA("E") + ',4)="付-费用")+(LEFT(' + CA("E") + ',4)="收-销售")'
            '+(LEFT(' + CA("E") + ',4)="收-平台")+(LEFT(' + CA("E") + ',4)="收-返利")'
            '+(LEFT(' + CA("E") + ',4)="收-退货")')
RULES = [
    ('W01', '采购已过信用期还没付清',
     f'COUNTIFS({PU("Z")},">0")', f'SUMIFS({PU("V")},{PU("Z")},">0")', '高',
     '去《收付款核销中心》下半张看逐笔；该付就付，付了记《资金流水》并填采购单号'),
    ('W02', '采购缺票（没收到进项发票）',
     f'COUNTIFS({PU("AB")},"未开票")+COUNTIFS({PU("AB")},"部分开票")',
     f'SUMIFS({PU("L")},{PU("AB")},"未开票")+SUMIFS({PU("L")},{PU("AB")},"部分开票")', '高',
     '催供应商开票；收到票在《发票台账》登记并填上采购单号，这条就消了'),
    ('W03', '采购付超了（已付大于应付）',
     f'COUNTIFS({PU("V")},"<-0.01",{PU("L")},">0")',
     f'-SUMIFS({PU("V")},{PU("V")},"<-0.01",{PU("L")},">0")', '高',
     '多半是《资金流水》关联单号填错，或者一笔钱记了两遍'),
    ('W04', '销售已过账期还没回款',
     f'COUNTIFS({SA("AB")},">0")', f'SUMIFS({SA("X")},{SA("AB")},">0")', '高',
     '平台结算一般 7 天，超了要查是不是被冻结或漏记结算单'),
    ('W05', '销售缺票（该开没开）',
     f'COUNTIFS({SA("AD")},"未开票")+COUNTIFS({SA("AD")},"部分开票")',
     f'SUMIFS({SA("M")},{SA("AD")},"未开票")+SUMIFS({SA("M")},{SA("AD")},"部分开票")', '中',
     'C 端零售不要票是正常的；批发客户要票的记得开并登记'),
    ('W06', '销售收超了（已收大于应收）',
     f'COUNTIFS({SA("X")},"<-0.01",{SA("M")},">0")',
     f'-SUMIFS({SA("X")},{SA("X")},"<-0.01",{SA("M")},">0")', '高',
     '查关联单号；客户先打款后发货请用「收-预收货款」（挂预收账款），不要用「收-其他」'),
    ('W07', '销售取不到成本（这个商品之前没进过货）',
     f'SUMPRODUCT(({SA("B")}<>"")*({SA("K")}<>0)*({SA("R")}=""))',
     f'SUMPRODUCT(({SA("B")}<>"")*({SA("K")}<>0)*({SA("R")}="")*N({SA("O")}))', '高',
     '要么补登采购，要么在《期初余额》填这个商品的期初库存，要么直接填「手工单位成本」'),
    ('W08', '返利确认了但还没到账',
     f'COUNTIFS({RB("O")},"未收")+COUNTIFS({RB("O")},"部分收")',
     f'SUMIFS({RB("L")},{RB("O")},"未收")+SUMIFS({RB("L")},{RB("O")},"部分收")', '中',
     '催供应商打款；到账了记《资金流水》「收-返利」并填返利单号'),
    ('W09', '返利确认数和账面预提差得多',
     f'SUMPRODUCT(({RB("C")}<>"")*(ABS(N({RB("J")}))>0.01))',
     f'SUMPRODUCT(({RB("C")}<>"")*ABS(N({RB("J")})))', '中',
     '差异会自动冲/加主营业务成本；差太多先跟供应商核对采购基数对不对'),
    ('W10', '返利没拿到发票',
     f'COUNTIFS({RB("M")},"未开票")',
     f'SUMIFS({RB("I")},{RB("M")},"未开票")', '中',
     '返利要不要票看当地税务口径；要票就催，红字票收到后做一张进项转出的手工凭证'),
    ('W11', '一件代发还没结清（欠客户回款 或 欠代发商货款）',
     f'COUNTIFS({DP("AA")},"<>已结清")-COUNTIFS({DP("AA")},"")',
     f'SUMIFS({DP("X")},{DP("AA")},"<>已结清")+SUMIFS({DP("Z")},{DP("AA")},"<>已结清")', '中',
     '回款填「收-销售货款」＋结算单号；付代发商填「付-代发货款」＋同一个结算单号'),
    ('W12', '资金流水填了关联单号，但找不到对应的业务单',
     f'SUMPRODUCT(({CA("G")}<>"")*({LINKED_R})*(COUNTIF({PU("C")},{CA("G")})+COUNTIF({SA("C")},{CA("G")})'
     f'+COUNTIF({RB("B")},{CA("G")})+COUNTIF({DP("C")},{CA("G")})+COUNTIF({EP("G")},{CA("G")})=0))',
     f'SUMPRODUCT(({CA("G")}<>"")*({LINKED_R})*(COUNTIF({PU("C")},{CA("G")})+COUNTIF({SA("C")},{CA("G")})'
     f'+COUNTIF({RB("B")},{CA("G")})+COUNTIF({DP("C")},{CA("G")})+COUNTIF({EP("G")},{CA("G")})=0)'
     f'*(N({CA("H")})+N({CA("I")})))', '高',
     '单号打错了，或者业务单还没登记 —— 这笔钱就核销不到任何一张单上。'
     '借款、股东投入、税费、保证金这些本来就没有业务单，不在这条规则里'),
    ('W13', '资金流水的对方科目取不到',
     f'COUNTIF({CA("O")},"★*")', f'SUMIFS({CA("H")},{CA("O")},"★*")+SUMIFS({CA("I")},{CA("O")},"★*")', '高',
     '业务类型不在《资金规则》里，或者内部转账没填「对方账户」'),
    ('W14', '费用的对应科目取不到',
     f'COUNTIF({EP("M")},"★*")', f'SUMIFS({EP("L")},{EP("M")},"★*")', '高',
     '《基础资料》里这个费用项目右边没配「费用对应科目」'),
    ('W15', '手工凭证有借贷不平的',
     f'SUMPRODUCT(({MV("D")}<>"")*(LEFT({MV("L")},1)="✘")/MAX(1,COUNTIF({MV("D")},{MV("D")}&"")))',
     f'ROUND(SUMIF({MV("D")},"<>",{MV("I")})-SUMIF({MV("D")},"<>",{MV("J")}),2)', '高',
     '同一个凭证字号的借方合计必须等于贷方合计，去《手工凭证》L 列看是哪一张'),
    ('W16', '有商品是负库存 / 负金额',
     f'COUNTIF(商品库存与成本!$P${IC0}:$P${IC1},"★*")',
     f'SUMIFS(商品库存与成本!$N${IC0}:$N${IC1},商品库存与成本!$P${IC0}:$P${IC1},"★*")', '高',
     '一般是采购漏登、期初没填，或者退货把数量填成正数了'),
    ('W17', '凭证里出现了科目表里没有的科目（自动＋手工都查）',
     f'SUMPRODUCT(({AV("F")}<>"")*(COUNTIF({SUBJ_N},{AV("F")})=0))'
     f'+SUMPRODUCT(({MVv("F")}<>"")*(COUNTIF({SUBJ_N},{MVv("F")})=0))',
     f'SUMPRODUCT(({AV("F")}<>"")*(COUNTIF({SUBJ_N},{AV("F")})=0)*(N({AV("I")})+N({AV("J")})))'
     f'+SUMPRODUCT(({MVv("F")}<>"")*(COUNTIF({SUBJ_N},{MVv("F")})=0)*(N({MVv("I")})+N({MVv("J")})))',
     '高',
     '★ 科目名称写错或《科目表》里缺这个科目 —— 这笔钱会掉在报表外面，而且借贷还是平的，很难发现。'
     '手工凭证尤其容易，一定要从下拉里选'),
    ('W21', '登记表里有金额被填成了文本（那一单会被整单剔出报表）',
     f'SUMPRODUCT(({PU("B")}<>"")*(ISTEXT({PU("J")})+ISTEXT({PU("K")})+ISTEXT({PU("M")})))'
     f'+SUMPRODUCT(({SA("B")}<>"")*(ISTEXT({SA("K")})+ISTEXT({SA("L")})+ISTEXT({SA("N")})+ISTEXT({SA("Q")})))'
     f'+SUMPRODUCT(({CA("B")}<>"")*(ISTEXT({CA("H")})+ISTEXT({CA("I")})))'
     f'+SUMPRODUCT(({EP("B")}<>"")*(ISTEXT({EP("I")})+ISTEXT({EP("J")})))'
     f'+SUMPRODUCT(({MV("B")}<>"")*(ISTEXT({MV("I")})+ISTEXT({MV("J")})))', '0', '高',
     '★ 从别的表复制粘贴最容易出这个。文本金额参与不了计算，那一单凭证直接不生成，'
     '报表照样是平的 —— 选中那一列 →【数据】→【分列】→ 直接完成，就能转回数字'),
    ('W22', '内部转账选了别家公司的账户',
     f'SUMPRODUCT((LEFT({CA("E")},4)="内部转账")*({CA("L")}<>"")'
     f'*(COUNTIFS({ACC_N},{CA("L")},{ACC_CO},{CA("C")})=0))', '0', '高',
     '内部转账只能在同一家公司的账户之间调钱；跨公司要走往来，不能用内部转账'),
    ('W23', '科目表的「报表项目」填了非法值',
     f'SUMPRODUCT(({SUBJ_N}<>"")*(COUNTIF(报表项目清单,{SUBJ_ITEM})=0))', '0', '高',
     '报表按「报表项目」取数，填了清单外的值，这个科目的钱就进不了报表'),
    ('W18', '业务单上的往来单位没在主档里建过',
     f'SUMPRODUCT(({PU("F")}<>"")*(COUNTIF({PT_N},{PU("F")})=0))'
     f'+SUMPRODUCT(({SA("G")}<>"")*(COUNTIF({PT_N},{SA("G")})=0))'
     f'+SUMPRODUCT(({CA("F")}<>"")*(COUNTIF({PT_N},{CA("F")})=0))', '0', '中',
     '去《往来单位主档》补一条，并勾上它的身份'),
    ('W19', '同一家往来单位在两家以上公司都有余额',
     f'SUM(综合往来对账!$U${PTD0}:$U${PTD1})', '0', '提示',
     '这不是错 —— 只是提醒你：跨公司的应收应付不能互相抵销，想抵要走内部交易或分别结算'),
    ('W20', '内部转账两边对不上（一出一入应当轧差为 0）',
     f'IF(ROUND(SUMIFS({CA("H")},{CA("P")},"不计入")-SUMIFS({CA("I")},{CA("P")},"不计入"),2)=0,0,1)',
     f'ROUND(SUMIFS({CA("H")},{CA("P")},"不计入")-SUMIFS({CA("I")},{CA("P")},"不计入"),2)', '高',
     '平台提现要记两行：平台户「内部转账-转出」＋银行户「内部转账-转入」，金额一样'),
]
r = 4
for no, nm, cnt, amt, sev, how in RULES:
    put(ws, f'A{r}', no, font=F_TOT, fill=FILL_AUTO)
    put(ws, f'B{r}', nm, font=F_TXT, align=CL)
    put(ws, f'C{r}', f'=IFERROR(ROUND({cnt},0),0)', font=F_TOT, fill=FILL_AUTO, fmt='#,##0;;\\-')
    put(ws, f'D{r}', f'=IFERROR(ROUND({amt},2),0)' if amt != '0' else '', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'E{r}', f'=IF(N($C{r})=0,"✔ 正常","{sev}")', font=F_WARN, fill=FILL_AUTO)
    put(ws, f'F{r}', how, font=F_NOTE, align=CL)
    r += 1
RT2 = r
put(ws, f'A{RT2}', '汇  总', font=F_TOT, fill=FILL_TOT)
put(ws, f'B{RT2}', f'=IF(SUM(C4:C{RT2-1})=0,"✔ 一条预警都没有，可以放心出报表",'
                   f'"✘ 一共 "&SUM(C4:C{RT2-1})&" 条待处理（其中高风险 "'
                   f'&SUMPRODUCT((E4:E{RT2-1}="高")*1)&" 类）")',
    font=F_TOT, fill=FILL_TOT, align=CL)
put(ws, f'C{RT2}', f'=SUM(C4:C{RT2-1})', font=F_TOT, fill=FILL_TOT, fmt='#,##0')
for c in 'DEF':
    put(ws, f'{c}{RT2}', None, font=F_TOT, fill=FILL_TOT)

DT0 = RT2 + 2
put(ws, f'A{DT0}', '逐 笔 明 细（只列真有问题的行，最多 400 条）', font=F_SEC, fill=FILL_SEC, align=CL)
ws.merge_cells(f'A{DT0}:H{DT0}')
headers(ws, DT0 + 1, ['序号', '来源表', '单　号', '日期', '公司', '往来单位', '异常类型', '涉及金额'],
        fill=FILL_HDR2, font=F_HDR2, height=30)
ED0, ED1 = DT0 + 2, DT0 + 401
for i in range(ED0, ED1 + 1):
    k = i - ED0 + 1
    idx = f'MATCH({k},{XR("BB", X_EX0, X_EX1)},0)'
    g = lambda c: f'=IFERROR(INDEX({XR(c, X_EX0, X_EX1)},{idx}),"")'
    put(ws, f'A{i}', f'=IF($C{i}="","",{k})', font=F_AUTO, fill=FILL_AUTO, fmt='0')
    put(ws, f'B{i}', g('AT'), font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'C{i}', g('AU'), font=F_TOT, fill=FILL_AUTO)
    put(ws, f'D{i}', g('AV'), font=F_AUTO, fill=FILL_AUTO, fmt=DATEQ)
    put(ws, f'E{i}', g('AW'), font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'F{i}', g('AX'), font=F_AUTO, fill=FILL_AUTO, align=CL)
    put(ws, f'G{i}', g('AY'), font=F_WARN, fill=FILL_AUTO, align=CL)
    put(ws, f'H{i}', g('AZ'), font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
put(ws, f'A{ED1+1}', '小  计', font=F_TOT, fill=FILL_TOT)
for c in 'BCDEFG':
    put(ws, f'{c}{ED1+1}', None, font=F_TOT, fill=FILL_TOT)
put(ws, f'H{ED1+1}', f'=ROUND(SUM(H{ED0}:H{ED1}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
ws.freeze_panes = 'A4'
page(ws)

# ════════════════════════════════════════════════════════════
# ㉘ 报表勾稽检查
# ════════════════════════════════════════════════════════════
ws = sheet('报表勾稽检查')
title(ws, '报 表 勾 稽 检 查', 'F', None)
put(ws, 'A2', '=本期标题&"　｜　出报表之前先看这一页，全是 ✔ 再往外发"', font=F_NOTE, align=CL, border=None)
ws.merge_cells('A2:F2')
headers(ws, HDR, ['编号', '勾稽关系', '应该等于', '实际是', '差异', '结论 / 不平了怎么查'])
widths(ws, {'A': 7, 'B': 36, 'C': 18, 'D': 18, 'E': 15, 'F': 52})
AVT = f'自动凭证!$I${AT}'
AVTC = f'自动凭证!$J${AT}'
MVT = f'手工凭证!$I${MT}'
MVTC = f'手工凭证!$J${MT}'
CHECKS = [
    ('J01', '自动凭证：借方合计 = 贷方合计', AVT, AVTC,
     '不平说明展开规则被改坏了，或者某张业务单的金额列有非数字'),
    ('J02', '手工凭证：借方合计 = 贷方合计', MVT, MVTC,
     '去《手工凭证》L 列找 ✘ 的那张凭证字号'),
    ('J03', '科目余额表：本期借方合计 = 本期贷方合计',
     f'科目余额表!$J${SBT}', f'科目余额表!$K${SBT}',
     '通常是凭证里用了《科目表》里没有的科目名 —— 看预警 W17'),
    ('J04', '科目余额表：期末借方合计 = 期末贷方合计',
     f'科目余额表!$N${SBT}', f'科目余额表!$O${SBT}',
     '先确认《期初余额》本身是平的（见 J05）'),
    ('J05', '期初余额：借方合计 = 贷方合计', f'期初余额!$D${OT}', f'期初余额!$E${OT}',
     '年初数没录平，报表一定不平。资产/成本填借方，负债/权益填贷方'),
    ('J06', '资产负债表：期末 资产 = 负债 + 所有者权益',
     f'资产负债表!$B${BS_A_TOT}', f'资产负债表!$F${BS_L_TOT}',
     '先看 J03/J04/J05；都平还不平就是有科目的「报表项目」没填或填错了'),
    ('J07', '资产负债表：年初 资产 = 负债 + 所有者权益',
     f'资产负债表!$C${BS_A_TOT}', f'资产负债表!$G${BS_L_TOT}',
     '只跟《期初余额》有关，跟本年业务无关'),
    ('J08', '现金流简表期末现金 = 资产负债表货币资金',
     f'现金流简表!$C${ENDCASH}', '资产负债表!$B$4',
     '常见原因：有一笔钱只做了手工凭证没走《资金流水》，或者资金账户没挂对科目'),
    ('J09', '商品库存（含手工调整）= 资产负债表「存货」',
     f'商品库存与成本!$N${INV_TOT}', f'{rep(Q("存货"), "DR", "end")}',
     '按商品算的结存 ＋ 手工凭证对库存商品的调整，应当正好等于存货'),
    ('J10', '综合往来：应付余额合计 = 资产负债表「应付账款」',
     f'综合往来对账!$G${PTT}', f'{rep(Q("应付账款"), "CR", "end")}',
     '差额多半是某笔应付没填往来单位，或者往来单位名字写得不一样'),
    ('J11', '综合往来：应收余额合计 = 资产负债表「应收账款」',
     f'综合往来对账!$J${PTT}', rep(Q('应收账款'), 'DR', 'end'),
     '同上；C 端散户记得统一用「散户/个人消费者」这一个往来单位'),
    ('J12', '综合往来：返利余额合计 = 科目「其他应收款—应收返利」期末',
     f'综合往来对账!$M${PTT}',
     f'ROUND(SUMIF({SB("B")},"其他应收款—应收返利",{SB("N")})-SUMIF({SB("B")},"其他应收款—应收返利",{SB("O")}),2)',
     '返利没填供应商，或者返利单号对不上'),
    ('J13', '利润表净利润 = 科目余额表全部损益类科目轧差',
     f'利润表!$D${PL_NET}',
     f'ROUND(SUMPRODUCT(({SB("C")}="损益")*(N({SB("M")})-N({SB("L")}))),2)',
     '不等＝有损益类科目的「报表项目」填错或没填，那笔钱进了总账却没进利润表'),
    ('J16', '利润表净利润 = 各报表项目分项加总（分项自检）',
     f'利润表!$D${PL_NET}',
     f'ROUND({rep(Q("营业收入"), "CR", "ytd")}+{rep(Q("营业外收入"), "CR", "ytd")}'
     f'-{rep(Q("营业成本"), "DR", "ytd")}-{rep(Q("税金及附加"), "DR", "ytd")}'
     f'-{rep(Q("销售费用"), "DR", "ytd")}-{rep(Q("管理费用"), "DR", "ytd")}'
     f'-{rep(Q("财务费用"), "DR", "ytd")}-{rep(Q("营业外支出"), "DR", "ytd")}'
     f'-{rep(Q("所得税费用"), "DR", "ytd")},2)',
     '不等说明有损益类科目的「报表项目」没填'),
    ('J14', '内部转账轧差 = 0（平台提现一出一入）',
     '0', f'ROUND(SUMIFS({CA("H")},{CA("P")},"不计入")-SUMIFS({CA("I")},{CA("P")},"不计入"),2)',
     '少记了配对的那一行，或者两行金额填得不一样'),
    ('J15', '期初往来明细合计 = 期初主区（应收/应付）',
     f'ROUND(期初余额!$J${PT}+期初余额!$K${PT},2)',
     f'ROUND(SUMIF({OPEN_SUB},"应收账款",{OPEN_D})+SUMIF({OPEN_SUB},"应付账款",{OPEN_C}),2)',
     '右边三个明细小区是给综合往来接年初数用的，要跟左边主区对上'),
]
r = 4
for no, nm, exp, act, how in CHECKS:
    put(ws, f'A{r}', no, font=F_TOT, fill=FILL_AUTO)
    put(ws, f'B{r}', nm, font=F_TXT, align=CL)
    put(ws, f'C{r}', f'=ROUND({exp},2)' if not exp.startswith('=') else exp,
        font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'D{r}', f'=ROUND({act},2)', font=F_AUTO, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'E{r}', f'=ROUND(N($C{r})-N($D{r}),2)', font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'F{r}', f'=IF(ABS(N($E{r}))<0.01,"✔ 对得上","✘ 不平 —— {how}")',
        font=F_TXT, fill=FILL_AUTO, align=CL)
    r += 1
JT = r
put(ws, f'A{JT}', '结  论', font=F_TOT, fill=FILL_TOT)
put(ws, f'B{JT}', f'=IF(SUMPRODUCT((ABS(N(E4:E{JT-1}))>=0.01)*1)=0,'
                  f'"✔ {len(CHECKS)} 项勾稽全部通过，报表可以对外出了",'
                  f'"✘ 有 "&SUMPRODUCT((ABS(N(E4:E{JT-1}))>=0.01)*1)&" 项不平，先按右边的提示查")',
    font=F_TOT, fill=FILL_TOT, align=CL)
ws.merge_cells(f'B{JT}:F{JT}')
page(ws)

print('  ✓ 异常预警中心 / 报表勾稽检查')

# ════════════════════════════════════════════════════════════
# ㉙ 老板经营看板
# ════════════════════════════════════════════════════════════
ws = sheet('老板经营看板')
title(ws, '老 板 经 营 看 板', 'L', None)
put(ws, 'A2', '=本期标题&"　｜　只看结果，不看过程。数字全部来自各张登记表，改了那边这里立刻变"',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells('A2:L2')
widths(ws, {'A': 3, 'B': 20, 'C': 18, 'D': 18, 'E': 3, 'F': 22, 'G': 18, 'H': 18,
            'I': 3, 'J': 24, 'K': 16, 'L': 16})
PLr = {'收入': 4, '成本': 5, '税金': 6, '销售费用': 7, '管理费用': 8, '财务费用': 9,
       '营业利润': 10, '营业外收入': 11, '营业外支出': 12, '利润总额': 13, '所得税': 14, '净利润': 15}
KPI = [
    ('营业收入', f'利润表!$C${PLr["收入"]}', f'利润表!$D${PLr["收入"]}'),
    ('毛　利', f'ROUND(利润表!$C${PLr["收入"]}-利润表!$C${PLr["成本"]},2)',
     f'ROUND(利润表!$D${PLr["收入"]}-利润表!$D${PLr["成本"]},2)'),
    ('净利润', f'利润表!$C${PLr["净利润"]}', f'利润表!$D${PLr["净利润"]}'),
    ('期末现金', f'现金流简表!$C${ENDCASH}', f'现金流简表!$D${ENDCASH}'),
]
put(ws, 'B4', '关 键 指 标', font=F_SEC, fill=FILL_SEC, align=CL)
ws.merge_cells('B4:D4')
headers(ws, 5, [None, '指　标', '本期', '本年累计'], fill=FILL_HDR2, font=F_HDR2, height=26)
r = 6
for nm, cur, ytd in KPI:
    put(ws, f'B{r}', nm, font=F_TOT, fill=FILL_KPI, align=CL)
    put(ws, f'C{r}', f'=ROUND({cur},2)', font=F_BIG, fill=FILL_KPI, fmt=MONEY0)
    put(ws, f'D{r}', f'=ROUND({ytd},2)', font=F_BIG, fill=FILL_KPI, fmt=MONEY0)
    ws.row_dimensions[r].height = 30
    r += 1
put(ws, f'B{r}', '毛利率', font=F_TOT, fill=FILL_KPI, align=CL)
put(ws, f'C{r}', f'=IF(ROUND(利润表!$C${PLr["收入"]},2)=0,"",ROUND($C7/利润表!$C${PLr["收入"]},4))',
    font=F_TOT, fill=FILL_KPI, fmt=PCT)
put(ws, f'D{r}', f'=IF(ROUND(利润表!$D${PLr["收入"]},2)=0,"",ROUND($D7/利润表!$D${PLr["收入"]},4))',
    font=F_TOT, fill=FILL_KPI, fmt=PCT)
r += 1
put(ws, f'B{r}', '净利率', font=F_TOT, fill=FILL_KPI, align=CL)
put(ws, f'C{r}', f'=IF(ROUND(利润表!$C${PLr["收入"]},2)=0,"",ROUND($C8/利润表!$C${PLr["收入"]},4))',
    font=F_TOT, fill=FILL_KPI, fmt=PCT)
put(ws, f'D{r}', f'=IF(ROUND(利润表!$D${PLr["收入"]},2)=0,"",ROUND($D8/利润表!$D${PLr["收入"]},4))',
    font=F_TOT, fill=FILL_KPI, fmt=PCT)

put(ws, 'F4', '钱 和 往 来', font=F_SEC, fill=FILL_SEC, align=CL)
ws.merge_cells('F4:H4')
headers(ws, 5, [None]*5 + ['项　目', '金　额', '说　明'], fill=FILL_HDR2, font=F_HDR2, height=26)
MONEYS = [
    ('期末货币资金', '资产负债表!$B$4', '银行＋平台货款户＋现金'),
    ('别人欠我（应收账款）', f'{rep(Q("应收账款"), "DR", "end")}', '客户和平台还没结算的货款'),
    ('应收返利', f'ROUND(SUMIF({SB("B")},"其他应收款—应收返利",{SB("N")})'
                 f'-SUMIF({SB("B")},"其他应收款—应收返利",{SB("O")}),2)', '供应商确认了还没打的返利'),
    ('我欠别人（应付账款）', f'{rep(Q("应付账款"), "CR", "end")}', '欠供应商、代发商、费用的钱'),
    ('借款余额', f'ROUND({rep(Q("短期借款"), "CR", "end")}+SUMIF({SB("B")},"其他应付款—股东借款",{SB("O")})'
                 f'-SUMIF({SB("B")},"其他应付款—股东借款",{SB("N")}),2)', '资方和股东借给公司的钱'),
    ('库存金额', f'{rep(Q("存货"), "DR", "end")}', '还压在仓库里的钱'),
]
r = 6
for nm, f, note in MONEYS:
    put(ws, f'F{r}', nm, font=F_TXT, fill=FILL_KPI, align=CL)
    put(ws, f'G{r}', f'=ROUND({f},2)', font=F_TOT, fill=FILL_KPI, fmt=MONEY0)
    put(ws, f'H{r}', note, font=F_NOTE, fill=FILL_KPI, align=CL)
    r += 1
put(ws, f'F{r}', '净 资 金 敞 口', font=F_TOT, fill=FILL_TOT, align=CL)
put(ws, f'G{r}', '=ROUND($G6+$G7+$G8-$G9-$G10,2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY0)
put(ws, f'H{r}', '现金＋应收＋返利 −应付 −借款；为负说明账上的钱不够还', font=F_NOTE, fill=FILL_TOT, align=CL)

put(ws, 'J4', '红 灯 与 勾 稽', font=F_SEC, fill=FILL_SEC, align=CL)
ws.merge_cells('J4:L4')
headers(ws, 5, [None]*9 + ['检　查', '结　果', ''], fill=FILL_HDR2, font=F_HDR2, height=26)
put(ws, 'J6', '异常预警条数', font=F_TXT, fill=FILL_KPI, align=CL)
put(ws, 'K6', f'=N(异常预警中心!$C${RT2})', font=F_TOT, fill=FILL_KPI, fmt='#,##0')
put(ws, 'L6', f'=IF(N($K6)=0,"✔ 干净","去《异常预警中心》")', font=F_WARN, fill=FILL_KPI)
put(ws, 'J7', '报表勾稽', font=F_TXT, fill=FILL_KPI, align=CL)
put(ws, 'K7', f'=SUMPRODUCT((ABS(N(报表勾稽检查!$E$4:$E${JT-1}))>=0.01)*1)', font=F_TOT, fill=FILL_KPI, fmt='#,##0')
put(ws, 'L7', f'=IF(N($K7)=0,"✔ 全部对得上","✘ 有不平项")', font=F_WARN, fill=FILL_KPI)
put(ws, 'J8', '逾期未收（笔）', font=F_TXT, fill=FILL_KPI, align=CL)
put(ws, 'K8', f'=COUNTIFS({SA("AB")},">0")', font=F_TOT, fill=FILL_KPI, fmt='#,##0')
put(ws, 'L8', '销售过了账期还没回款', font=F_NOTE, fill=FILL_KPI, align=CL)
put(ws, 'J9', '逾期未付（笔）', font=F_TXT, fill=FILL_KPI, align=CL)
put(ws, 'K9', f'=COUNTIFS({PU("Z")},">0")', font=F_TOT, fill=FILL_KPI, fmt='#,##0')
put(ws, 'L9', '采购过了信用期还没付', font=F_NOTE, fill=FILL_KPI, align=CL)
put(ws, 'J10', '缺进项票金额', font=F_TXT, fill=FILL_KPI, align=CL)
put(ws, 'K10', f'=ROUND(SUMIFS({PU("L")},{PU("AB")},"未开票")+SUMIFS({PU("L")},{PU("AB")},"部分开票"),2)',
    font=F_TOT, fill=FILL_KPI, fmt=MONEY0)
put(ws, 'L10', '影响进项抵扣和所得税', font=F_NOTE, fill=FILL_KPI, align=CL)
put(ws, 'J11', '返利未收金额', font=F_TXT, fill=FILL_KPI, align=CL)
put(ws, 'K11', f'=ROUND(SUMIFS({RB("L")},{RB("O")},"未收")+SUMIFS({RB("L")},{RB("O")},"部分收"),2)',
    font=F_TOT, fill=FILL_KPI, fmt=MONEY0)
put(ws, 'L11', '确认了还没到账的返利', font=F_NOTE, fill=FILL_KPI, align=CL)

TOPS_R = 15
def top5(col0, title_, src_name, src_val, src_key, fmt=MONEY0, big=True, co_range=None):
    put(ws, f'{col0}{TOPS_R}', title_, font=F_SEC, fill=FILL_SEC, align=CL)
    c0 = CI(col0)
    ws.merge_cells(f'{col0}{TOPS_R}:{L(c0+1)}{TOPS_R}')
    put(ws, f'{col0}{TOPS_R+1}', '名　称', font=F_HDR2, fill=FILL_HDR2)
    put(ws, f'{L(c0+1)}{TOPS_R+1}', '金　额', font=F_HDR2, fill=FILL_HDR2)
    for k in range(1, 6):
        rr = TOPS_R + 1 + k
        big_f = f'LARGE({src_val},{k})' if big else f'SMALL({src_val},{k})'
        nm_f = (f'INDEX({src_name},MATCH({big_f},{src_val},0))' if not co_range else
                f'INDEX({co_range},MATCH({big_f},{src_val},0))&" / "&'
                f'INDEX({src_name},MATCH({big_f},{src_val},0))')
        put(ws, f'{col0}{rr}', f'=IFERROR({nm_f},"")', font=F_TXT, fill=FILL_AUTO, align=CL)
        put(ws, f'{L(c0+1)}{rr}', f'=IFERROR({big_f},"")', font=F_TOT, fill=FILL_AUTO, fmt=fmt)

top5('B', 'TOP5 店铺净利', f'店铺利润分析!$B${SPD0}:$B${SPD1}',
     f'店铺利润分析!$N${SPD0}:$N${SPD1}', None)
top5('F', 'TOP5 商品毛利', XR('BD', 4, 103), XR('BF', 4, 103), None)
top5('J', 'TOP5 欠我钱的', f'综合往来对账!$C${PTD0}:$C${PTD1}',
     f'综合往来对账!$J${PTD0}:$J${PTD1}', None, co_range=f'综合往来对账!$B${PTD0}:$B${PTD1}')
put(ws, f'B{TOPS_R+8}', 'TOP5 我欠钱的', font=F_SEC, fill=FILL_SEC, align=CL)
ws.merge_cells(f'B{TOPS_R+8}:C{TOPS_R+8}')
put(ws, f'B{TOPS_R+9}', '名　称', font=F_HDR2, fill=FILL_HDR2)
put(ws, f'C{TOPS_R+9}', '金　额', font=F_HDR2, fill=FILL_HDR2)
for k in range(1, 6):
    rr = TOPS_R + 9 + k
    put(ws, f'B{rr}', f'=IFERROR(INDEX(综合往来对账!$B${PTD0}:$B${PTD1},'
                      f'MATCH(LARGE(综合往来对账!$G${PTD0}:$G${PTD1},{k}),'
                      f'综合往来对账!$G${PTD0}:$G${PTD1},0))&" / "&'
                      f'INDEX(综合往来对账!$C${PTD0}:$C${PTD1},'
                      f'MATCH(LARGE(综合往来对账!$G${PTD0}:$G${PTD1},{k}),'
                      f'综合往来对账!$G${PTD0}:$G${PTD1},0)),"")', font=F_TXT, fill=FILL_AUTO, align=CL)
    put(ws, f'C{rr}', f'=IFERROR(LARGE(综合往来对账!$G${PTD0}:$G${PTD1},{k}),"")',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY0)
put(ws, f'F{TOPS_R+8}', '各 资 金 账 户 余 额', font=F_SEC, fill=FILL_SEC, align=CL)
ws.merge_cells(f'F{TOPS_R+8}:H{TOPS_R+8}')
put(ws, f'F{TOPS_R+9}', '账　户', font=F_HDR2, fill=FILL_HDR2)
put(ws, f'G{TOPS_R+9}', '余　额', font=F_HDR2, fill=FILL_HDR2)
put(ws, f'H{TOPS_R+9}', '所属公司', font=F_HDR2, fill=FILL_HDR2)
for k in range(1, 11):
    rr = TOPS_R + 9 + k
    put(ws, f'F{rr}', f'=IF(INDEX({ACC_N},{k})="","",IF(COUNTIF(INDEX({ACC_CO},{k}),公司条件)>0,'
                      f'INDEX({ACC_N},{k}),""))', font=F_TXT, fill=FILL_AUTO, align=CL)
    put(ws, f'G{rr}', f'=IF($F{rr}="","",ROUND(IFERROR(INDEX({OPA_V},MATCH($F{rr},{OPA_N},0)),0)'
                      f'+SUMIFS({CA("H")},{CA("D")},$F{rr},{CA("R")},"<="&截止年月)'
                      f'-SUMIFS({CA("I")},{CA("D")},$F{rr},{CA("R")},"<="&截止年月),2))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY0)
    put(ws, f'H{rr}', f'=IF($F{rr}="","",IFERROR(INDEX({ACC_CO},MATCH($F{rr},{ACC_N},0)),""))',
        font=F_AUTO, fill=FILL_AUTO)
END_R = TOPS_R + 21
put(ws, f'B{END_R}',
    '看板怎么用：① 先看右上角「红灯与勾稽」两个数是不是 0 —— 不是 0 就先去修，修完数才可信。'
    '② 再看左上角四个大数；③ TOP5 是拿来定动作的：谁欠钱最多就先催谁，哪个店最赚就往哪个店投钱。\n'
    '　 换公司/换月份：去《查询设置》改，这一页会跟着变。TOP5 里如果有并列相同金额，只会显示排在前面的那一个。',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells(f'B{END_R}:L{END_R+2}')
page(ws)

# ════════════════════════════════════════════════════════════
# ㉚ 测试说明
# ════════════════════════════════════════════════════════════
ws = sheet('测试说明')
title(ws, '测 试 说 明（TEST 案例清单 · 验收照着这个走）', 'E',
      '★ 每张登记表里带「TESTxx」备注的都是给你验收用的样例数据。'
      '看懂了、对过账了，就可以整片删掉换成你们自己的真数据 —— 删数据不会破坏公式。\n'
      '★ 删的时候：选中浅黄色那些格子按 Delete 即可，不要整行删除（整行删会把下面行的公式一起带走）。')
headers(ws, HDR, ['所在表', '案例编号', '测试点', '预期看到的结果', '在哪验证'])
widths(ws, {'A': 16, 'B': 10, 'C': 30, 'D': 52, 'E': 26})
CASES = [
    ('采购登记', 'TEST01', '按比例返利', '应收返利＝不含税×2%；实际采购成本＝不含税−返利', '采购登记 R/S 列'),
    ('采购登记', 'TEST03', '部分付款', '应付 49,600 只付了 30,000 → 付款状态「部分付款」', '收付款核销中心 下半张'),
    ('采购登记', 'TEST05', '采购缺票', '票据状态「未开票」', '异常预警 W02'),
    ('采购登记', 'TEST10', '采购退货', '数量 −50，金额/成本/返利一起变负', '采购登记 L/S 列'),
    ('采购登记', 'TEST11', '固定金额返利', '填了返利金额 1500 就不再按比例算', '采购登记 Q/R 列'),
    ('采购登记', 'TEST13', '逾期未付', '过了 30 天信用期一分没付 → 逾期天数 > 0', '异常预警 W01'),
    ('销售登记', 'TEST01', '平台扣费', '扣费一边进销售费用、一边冲应收；到账＝含税−扣费', '自动凭证 销-XS2026001'),
    ('销售登记', 'TEST04', '抖音高佣金', '同样的货，抖音扣 5% 比拼多多 0.6% 毛利低很多', '店铺利润分析'),
    ('销售登记', 'TEST07', '批发部分回款', '应收 23,800 只回 15,000 → 未回 8,800', '收付款核销中心 上半张'),
    ('销售登记', 'TEST10', '退款', '数量 −20、扣费也退 → 收入成本税一起冲回', '利润表'),
    ('销售登记', 'TEST12', '卖给供应商', '宏发既是供应商又是客户，应收应付分开列', '综合往来对账 宏发那一行'),
    ('销售登记', 'TEST16', '逾期未回款', '过了 7 天账期还没结算 → 逾期天数 > 0', '异常预警 W04'),
    ('返利登记', 'TEST02', '返利未收', '确认了 658.41 但一直没到账', '异常预警 W08'),
    ('返利登记', 'TEST03', '返利部分收', '确认 594.69 只收到 300', '返利登记 K/L 列'),
    ('返利登记', 'TEST05', '返利上调差异', '确认 700 > 账面 647.79 → 差异 +52.21 自动冲成本', '自动凭证 返-FL2026005'),
    ('返利登记', 'TEST09', '返利核减差异', '确认 150 < 账面 180.53 → 差异 −30.53 自动加回成本', '自动凭证 返-FL2026009'),
    ('返利登记', 'TEST10', '账面没有的临时返利', '账面 0、确认 500 → 全额都是差异调整', '返利登记 J 列'),
    ('一件代发', 'TEST01', '代发不进库存', '成本直接进主营业务成本，库存商品一分不动', '商品库存与成本（查不到这几笔）'),
    ('一件代发', 'TEST06', '欠代发商货款', '一分没付 → 结算状态「欠代发商货款」', '收付款核销中心 下半张'),
    ('一件代发', 'TEST10', '代发退货', '数量 −10，售价成本扣费一起冲回', '一件代发结算 L/S/U 列'),
    ('资金流水', 'TEST04/05', '平台提现（内部转账）', '一出一入两行，损益不受影响，现金流第四行轧差为 0', '现金流简表 第四行'),
    ('资金流水', 'TEST12', '资方借款', '宏发借 10 万走短期借款，不混进应付货款', '综合往来对账 宏发「借款余额」'),
    ('资金流水', 'TEST13', '股东借款', '张伟借 5 万走其他应付款—股东借款，不是收入', '科目余额表'),
    ('资金流水', 'TEST23', '交保证金', '保证金是资产（其他应收款—保证金），不是费用', '资产负债表 其他应收款'),
    ('资金流水', 'TEST27', '买固定资产', '走投资活动，不进经营活动现金流', '现金流简表 二、投资活动'),
    ('资金流水', 'TEST35', '科目手工覆盖', '填了「营业外支出」就不按资金规则自动判了', '资金流水 O 列'),
    ('发票台账', 'TEST05', '返利发票', '返利收到红字票，配套一张进项转出的手工凭证', '手工凭证 记-2026-006'),
    ('费用及其他', 'TEST09', '固定资产', '费用项目选「购置固定资产」→ 对应科目自动变成固定资产', '费用及其他 M 列'),
    ('费用及其他', 'TEST01', '店铺费用', '填了店铺的费用会摊到那个店；没填的算公司共同费用', '店铺利润分析'),
    ('手工凭证', 'TEST01', '折旧', '业务登记做不出来，只能手工计提', '利润表 管理费用'),
    ('手工凭证', 'TEST03', '正式往来抵销', '填了关联单号，采购 CG2026009 和销售 XS2026012 双边自动认',
     '采购登记 U 列 / 销售登记 W 列'),
    ('手工凭证', 'TEST08', '没票也要计提', '权责发生制：6 月房租没拿到票也进当期费用', '利润表'),
    ('手工凭证', 'TEST10', '故意做不平', '现在借贷都是 0＝平的；填个数试试勾稽会不会抓到', '报表勾稽检查 J02'),
    ('期初余额', '—', '一年一个账套', '年初数在这里录，明细三小区要跟主区对上', '报表勾稽检查 J05/J15'),
    ('查询设置', '—', '月份＋公司统一切换', '把月份改成 3，所有报表和看板一起跳到 3 月', '利润表 / 老板看板'),
]
r = 4
for a, b, c, d, e in CASES:
    put(ws, f'A{r}', a, font=F_TOT, fill=FILL_AUTO)
    put(ws, f'B{r}', b, font=F_AUTO, fill=FILL_AUTO)
    put(ws, f'C{r}', c, font=F_TXT, align=CL)
    put(ws, f'D{r}', d, font=F_NOTE, align=CL)
    put(ws, f'E{r}', e, font=F_LINK, align=CL)
    r += 1
put(ws, f'A{r+1}',
    f'一共 {len(CASES)} 个测试点，覆盖了：部分付款、部分回款、返利未收、返利差异（上调/核减/凭空）、采购缺票、'
    '返利发票、平台扣费、退货退款、一件代发、资方兼客户兼供应商、融资借款、股东借款、正式往来抵销、'
    '固定资产、保证金、内部转账、店铺合作方分成、跨公司不抵销、科目手工覆盖。',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells(f'A{r+1}:E{r+2}')
ws.freeze_panes = 'A4'
page(ws, titles=f'{HDR}:{HDR}')

print('  ✓ 老板经营看板 / 测试说明')

# ════════════════════════════════════════════════════════════
# ㉛ _自动清单（隐藏 · 各种去重名单、配对池、核销池、异常池）
# ════════════════════════════════════════════════════════════
ws = sheet('_自动清单', hidden=True)
put(ws, 'A1', '本表全部由公式维护，供其他表取数用。请勿手工修改，也不要删行删列。',
    font=F_WARN, align=CL, border=None)
HD = {'A': '往来单位', 'B': '是供应商', 'C': '累计', 'D': '供应商名单',
      'E': '是客户', 'F': '累计', 'G': '客户名单',
      'H': '是资方', 'I': '累计', 'J': '资方名单',
      'K': '是合作方', 'L': '累计', 'M': '合作方名单',
      'O': '公司', 'P': '往来单位', 'Q': '有效', 'R': '累计',
      'S': '公司', 'T': '商品编码', 'U': '有效', 'V': '累计',
      'X': '来源', 'Y': '单号', 'Z': '日期', 'AA': '公司', 'AB': '往来单位',
      'AC': '应收', 'AD': '已收', 'AE': '未收', 'AF': '有效', 'AG': '累计',
      'AI': '来源', 'AJ': '单号', 'AK': '日期', 'AL': '公司', 'AM': '往来单位',
      'AN': '应付', 'AO': '已付', 'AP': '未付', 'AQ': '有效', 'AR': '累计',
      'AT': '来源', 'AU': '单号', 'AV': '日期', 'AW': '公司', 'AX': '往来单位',
      'AY': '异常类型', 'AZ': '金额', 'BA': '有效', 'BB': '累计',
      'BD': '商品编码', 'BE': '收入', 'BF': '毛利'}
for c, t in HD.items():
    put(ws, f'{c}{HDR}', t, font=F_HDR2, fill=FILL_HDR2)

# —— ① 往来单位四种身份的紧凑名单 ——
ROLE = [('B', 'C', 'D', PT_SUP), ('E', 'F', 'G', PT_CUS), ('H', 'I', 'J', PT_FIN), ('K', 'L', 'M', PT_PAR)]
for i in range(N_PARTY):
    r = X_PT0 + i
    sr = PR0 + i
    put(ws, f'A{r}', f'=IF(往来单位主档!$B{sr}="","",往来单位主档!$B{sr})', font=F_AUTO, border=None)
    for fc, cc, nc, _ in ROLE:
        role_col = {'B': 'D', 'E': 'E', 'H': 'F', 'K': 'G'}[fc]
        put(ws, f'{fc}{r}', f'=IF(AND($A{r}<>"",往来单位主档!${role_col}{sr}="是",'
                            f'往来单位主档!$M{sr}<>"停用"),1,0)', font=F_AUTO, border=None, fmt='0')
        put(ws, f'{cc}{r}', f'=N({cc}{r-1})+${fc}{r}', font=F_AUTO, border=None, fmt='0')
        put(ws, f'{nc}{r}', f'=IFERROR(INDEX($A${X_PT0}:$A${X_PT1},'
                            f'MATCH(ROW()-{X_PT0-1},${cc}${X_PT0}:${cc}${X_PT1},0)),"")',
            font=F_AUTO, border=None)
for nm, cc, nc in (('供应商名单', 'C', 'D'), ('客户名单', 'F', 'G'),
                   ('资方名单', 'I', 'J'), ('合作方名单', 'L', 'M')):
    name(nm, f'OFFSET({AX}!${nc}${X_PT0},0,0,MAX(1,MAX({AX}!${cc}${X_PT0}:${cc}${X_PT1})),1)')

# —— ② 公司 × 往来单位 配对池 ——
for p in range(N_CO_MAX * N_PARTY):
    r = X_PR0 + p
    ci, pi = p // N_PARTY, p % N_PARTY
    co_cell = f'基础资料!$A{BR0 + ci}'
    pt_cell = f'往来单位主档!$B{PR0 + pi}'
    put(ws, f'O{r}', f'=IF(OR({co_cell}="",{pt_cell}=""),"",{co_cell})', font=F_AUTO, border=None)
    put(ws, f'P{r}', f'=IF($O{r}="","",{pt_cell})', font=F_AUTO, border=None)
    put(ws, f'Q{r}', f'=IF($O{r}="",0,IF({CO(f"$O{r}")},1,0))', font=F_AUTO, border=None, fmt='0')
    put(ws, f'R{r}', f'=N(R{r-1})+$Q{r}', font=F_AUTO, border=None, fmt='0')

# —— ③ 公司 × 商品 配对池 ——
for p in range(N_CO_MAX * N_ITEM):
    r = X_IT0 + p
    ci, ii = p // N_ITEM, p % N_ITEM
    co_cell = f'基础资料!$A{BR0 + ci}'
    it_cell = f'商品档案!$A{IR0 + ii}'
    put(ws, f'S{r}', f'=IF(OR({co_cell}="",{it_cell}=""),"",{co_cell})', font=F_AUTO, border=None)
    put(ws, f'T{r}', f'=IF($S{r}="","",{it_cell})', font=F_AUTO, border=None)
    put(ws, f'U{r}', f'=IF($S{r}="",0,IF({CO(f"$S{r}")},1,0))', font=F_AUTO, border=None, fmt='0')
    put(ws, f'V{r}', f'=N(V{r-1})+$U{r}', font=F_AUTO, border=None, fmt='0')

# —— ④ 应收核销池（销售 + 一件代发）——
for p in range(N_SAL + N_DRP):
    r = X_AR0 + p
    if p < N_SAL:
        sr = S0 + p
        cols = ('"销售"', f'销售登记!$C{sr}', f'销售登记!$B{sr}', f'销售登记!$D{sr}', f'销售登记!$G{sr}',
                f'ROUND(N(销售登记!$M{sr})-N(销售登记!$Q{sr}),2)', f'N(销售登记!$W{sr})', f'N(销售登记!$X{sr})')
        key = f'销售登记!$C{sr}'
    else:
        sr = F0 + (p - N_SAL)
        cols = ('"一件代发"', f'一件代发结算!$C{sr}', f'一件代发结算!$B{sr}', f'一件代发结算!$D{sr}',
                f'一件代发结算!$G{sr}', f'ROUND(N(一件代发结算!$L{sr})-N(一件代发结算!$U{sr}),2)',
                f'N(一件代发结算!$W{sr})', f'N(一件代发结算!$X{sr})')
        key = f'一件代发结算!$C{sr}'
    for c, f in zip(('X', 'Y', 'Z', 'AA', 'AB', 'AC', 'AD', 'AE'), cols):
        put(ws, f'{c}{r}', f'=IF({key}="","",{f})', font=F_AUTO, border=None,
            fmt=MONEY if c in ('AC', 'AD', 'AE') else (DATEQ if c == 'Z' else None))
    put(ws, f'AF{r}', f'=IF(AND($Y{r}<>"",ROUND(N($AE{r}),2)<>0,{CO(f"$AA{r}")}),1,0)',
        font=F_AUTO, border=None, fmt='0')
    put(ws, f'AG{r}', f'=N(AG{r-1})+$AF{r}', font=F_AUTO, border=None, fmt='0')

# —— ⑤ 应付核销池（采购 + 一件代发 + 费用）——
for p in range(N_PUR + N_DRP + N_EXP):
    r = X_AP0 + p
    if p < N_PUR:
        sr = P0 + p
        key = f'采购登记!$C{sr}'
        cols = ('"采购"', f'采购登记!$C{sr}', f'采购登记!$B{sr}', f'采购登记!$D{sr}', f'采购登记!$F{sr}',
                f'N(采购登记!$L{sr})', f'N(采购登记!$U{sr})', f'N(采购登记!$V{sr})')
    elif p < N_PUR + N_DRP:
        sr = F0 + (p - N_PUR)
        key = f'一件代发结算!$C{sr}'
        cols = ('"一件代发"', f'一件代发结算!$C{sr}', f'一件代发结算!$B{sr}', f'一件代发结算!$D{sr}',
                f'一件代发结算!$F{sr}', f'N(一件代发结算!$Q{sr})', f'N(一件代发结算!$Y{sr})',
                f'N(一件代发结算!$Z{sr})')
    else:
        sr = E0 + (p - N_PUR - N_DRP)
        key = f'费用及其他!$G{sr}'
        cols = ('"费用"', f'费用及其他!$G{sr}', f'费用及其他!$B{sr}', f'费用及其他!$C{sr}',
                f'费用及其他!$F{sr}', f'N(费用及其他!$L{sr})', f'N(费用及其他!$O{sr})', f'N(费用及其他!$P{sr})')
    for c, f in zip(('AI', 'AJ', 'AK', 'AL', 'AM', 'AN', 'AO', 'AP'), cols):
        put(ws, f'{c}{r}', f'=IF({key}="","",{f})', font=F_AUTO, border=None,
            fmt=MONEY if c in ('AN', 'AO', 'AP') else (DATEQ if c == 'AK' else None))
    put(ws, f'AQ{r}', f'=IF(AND($AJ{r}<>"",ROUND(N($AP{r}),2)<>0,{CO(f"$AL{r}")}),1,0)',
        font=F_AUTO, border=None, fmt='0')
    put(ws, f'AR{r}', f'=N(AR{r-1})+$AQ{r}', font=F_AUTO, border=None, fmt='0')

# —— ⑥ 异常池 ——
NOPARTY = lambda cell: f'AND({cell}<>"",COUNTIF({PT_N},{cell})=0)'
EX_SRC = []
for i in range(N_PUR):
    sr = P0 + i
    EX_SRC.append(('"采购登记"', f'采购登记!$C{sr}', f'采购登记!$B{sr}', f'采购登记!$D{sr}', f'采购登记!$F{sr}',
        f'IF(N(采购登记!$Z{sr})>0,"采购逾期未付",'
        f'IF(AND(N(采购登记!$V{sr})<-0.01,N(采购登记!$L{sr})>0),"采购付超（已付大于应付）",'
        f'IF(OR(采购登记!$AB{sr}="未开票",采购登记!$AB{sr}="部分开票"),"采购缺票",'
        f'IF({NOPARTY(f"采购登记!$F{sr}")},"供应商没在主档里建过",""))))',
        f'IF(N(采购登记!$Z{sr})>0,N(采购登记!$V{sr}),'
        f'IF(N(采购登记!$V{sr})<-0.01,-N(采购登记!$V{sr}),N(采购登记!$L{sr})))'))
for i in range(N_SAL):
    sr = S0 + i
    EX_SRC.append(('"销售登记"', f'销售登记!$C{sr}', f'销售登记!$B{sr}', f'销售登记!$D{sr}', f'销售登记!$G{sr}',
        f'IF(AND(N(销售登记!$K{sr})<>0,销售登记!$R{sr}=""),"销售取不到成本（这个商品没进过货）",'
        f'IF(N(销售登记!$AB{sr})>0,"销售逾期未回款",'
        f'IF(AND(N(销售登记!$X{sr})<-0.01,N(销售登记!$M{sr})>0),"销售收超（已收大于应收）",'
        f'IF(OR(销售登记!$AD{sr}="未开票",销售登记!$AD{sr}="部分开票"),"销售缺票",'
        f'IF({NOPARTY(f"销售登记!$G{sr}")},"客户没在主档里建过","")))))',
        f'IF(N(销售登记!$AB{sr})>0,N(销售登记!$X{sr}),N(销售登记!$M{sr}))'))
for i in range(N_RBT):
    sr = B0 + i
    EX_SRC.append(('"返利登记"', f'返利登记!$B{sr}', f'返利登记!$C{sr}', f'返利登记!$D{sr}', f'返利登记!$E{sr}',
        f'IF(返利登记!$O{sr}="未收","返利确认了还没到账",'
        f'IF(返利登记!$O{sr}="部分收","返利只收到一部分",'
        f'IF(ABS(N(返利登记!$J{sr}))>0.01,"返利确认数与账面预提有差异",'
        f'IF(返利登记!$M{sr}="未开票","返利还没拿到发票",""))))',
        f'IF(ABS(N(返利登记!$J{sr}))>0.01,ABS(N(返利登记!$J{sr})),N(返利登记!$L{sr}))'))
for i in range(N_DRP):
    sr = F0 + i
    EX_SRC.append(('"一件代发"', f'一件代发结算!$C{sr}', f'一件代发结算!$B{sr}', f'一件代发结算!$D{sr}',
        f'一件代发结算!$F{sr}',
        f'IF(OR(一件代发结算!$AA{sr}="",一件代发结算!$AA{sr}="已结清"),"","代发"&一件代发结算!$AA{sr})',
        f'N(一件代发结算!$X{sr})+N(一件代发结算!$Z{sr})'))
for i in range(N_CASH):
    sr = K0 + i
    known = (f'COUNTIF({PU("C")},资金流水!$G{sr})+COUNTIF({SA("C")},资金流水!$G{sr})'
             f'+COUNTIF({RB("B")},资金流水!$G{sr})+COUNTIF({DP("C")},资金流水!$G{sr})'
             f'+COUNTIF({EP("G")},资金流水!$G{sr})')
    EX_SRC.append(('"资金流水"', f'IF(资金流水!$G{sr}="","(无单号)"&TEXT(N(资金流水!$A{sr}),"0000"),资金流水!$G{sr})',
        f'资金流水!$B{sr}', f'资金流水!$C{sr}', f'资金流水!$F{sr}',
        f'IF(LEFT(资金流水!$O{sr},1)="★",资金流水!$O{sr},'
        f'IF(AND(LEFT(资金流水!$E{sr},4)="内部转账",资金流水!$L{sr}=""),"内部转账没填对方账户",'
        f'IF(AND(LEFT(资金流水!$E{sr},4)="内部转账",资金流水!$L{sr}<>"",'
        f'COUNTIFS({ACC_N},资金流水!$L{sr},{ACC_CO},资金流水!$C{sr})=0),'
        f'"内部转账选了别家公司的账户",'
        f'IF(AND(资金流水!$G{sr}<>"",{known}=0,OR(LEFT(资金流水!$E{sr},4)="付-采购",'
        f'LEFT(资金流水!$E{sr},4)="付-代发",LEFT(资金流水!$E{sr},4)="付-费用",'
        f'LEFT(资金流水!$E{sr},4)="收-销售",LEFT(资金流水!$E{sr},4)="收-平台",'
        f'LEFT(资金流水!$E{sr},4)="收-返利",LEFT(资金流水!$E{sr},4)="收-退货")),'
        f'"关联单号找不到对应的业务单",'
        f'IF({NOPARTY(f"资金流水!$F{sr}")},"往来单位没在主档里建过","")))))',
        f'N(资金流水!$H{sr})+N(资金流水!$I{sr})'))
for i in range(N_EXP):
    sr = E0 + i
    EX_SRC.append(('"费用及其他"', f'费用及其他!$G{sr}', f'费用及其他!$B{sr}', f'费用及其他!$C{sr}',
        f'费用及其他!$F{sr}',
        f'IF(LEFT(费用及其他!$M{sr},1)="★","费用项目没配借方科目",'
        f'IF(N(费用及其他!$P{sr})<-0.01,"费用付超",""))',
        f'N(费用及其他!$L{sr})'))
assert len(EX_SRC) == X_EX1 - X_EX0 + 1, (len(EX_SRC), X_EX1 - X_EX0 + 1)
for p, (src, num, dt, co, pt, typ, amt) in enumerate(EX_SRC):
    r = X_EX0 + p
    put(ws, f'AT{r}', f'=IF($AY{r}="","",{src})', font=F_AUTO, border=None)
    put(ws, f'AU{r}', f'=IF($AY{r}="","",{num})', font=F_AUTO, border=None)
    put(ws, f'AV{r}', f'=IF($AY{r}="","",{dt})', font=F_AUTO, border=None, fmt=DATEQ)
    put(ws, f'AW{r}', f'=IF($AY{r}="","",{co})', font=F_AUTO, border=None)
    put(ws, f'AX{r}', f'=IF($AY{r}="","",{pt})', font=F_AUTO, border=None)
    put(ws, f'AY{r}', f'={typ}', font=F_AUTO, border=None)
    put(ws, f'AZ{r}', f'=IF($AY{r}="","",ROUND({amt},2))', font=F_AUTO, border=None, fmt=MONEY)
    put(ws, f'BA{r}', f'=IF($AY{r}="",0,1)', font=F_AUTO, border=None, fmt='0')
    put(ws, f'BB{r}', f'=N(BB{r-1})+$BA{r}', font=F_AUTO, border=None, fmt='0')

# —— ⑦ 商品收入/毛利汇总（给看板 TOP5 用）——
for i in range(N_ITEM):
    r = 4 + i
    it = f'商品档案!$A{IR0 + i}'
    put(ws, f'BD{r}', f'=IF({it}="","",{it})', font=F_AUTO, border=None)
    put(ws, f'BE{r}', f'=IF($BD{r}="","",ROUND(SUMIFS({SA("O")},{SA("H")},$BD{r},{SA("D")},公司条件,'
                      f'{YMR(SA("AF"))})+SUMIFS({DP("N")},{DP("H")},$BD{r},{DP("D")},公司条件,'
                      f'{YMR(DP("AC"))}),2))', font=F_AUTO, border=None, fmt=MONEY)
    put(ws, f'BF{r}', f'=IF($BD{r}="","",ROUND(SUMIFS({SA("U")},{SA("H")},$BD{r},{SA("D")},公司条件,'
                      f'{YMR(SA("AF"))})+SUMIFS({DP("V")},{DP("H")},$BD{r},{DP("D")},公司条件,'
                      f'{YMR(DP("AC"))}),2))', font=F_AUTO, border=None, fmt=MONEY)
ws.sheet_view.showGridLines = False
print('  ✓ _自动清单（名单 / 配对池 / 核销池 / 异常池）')

# ════════════════════════════════════════════════════════════
# ㉜ 主页
# ════════════════════════════════════════════════════════════
ws = sheet('主页')
title(ws, '电 商 一 体 化 账 务 模 板', 'H', None)
put(ws, 'A2', f'={YEAR}&" 年账套　｜　业务登记 → 自动凭证 → 报表，一条线打通，不用再手工录总账"',
    font=F_NOTE, align=CL, border=None)
ws.merge_cells('A2:H2')
widths(ws, {'A': 4, 'B': 24, 'C': 46, 'D': 4, 'E': 24, 'F': 20, 'G': 20, 'H': 20})
NAV = [
    ('① 先设置', [
        ('查询设置', '改年度 / 公司 / 月份 —— 所有报表跟着这里走'),
        ('基础资料', '各种下拉选项；加了这里，登记页的下拉就多出来'),
        ('往来单位主档', '★ 一家单位只建一次，同时勾供应商/客户/资方/店铺合作方'),
        ('商品档案', '商品编码、规格、默认税率'),
        ('店铺档案', '店铺挂哪家公司、哪个平台'),
        ('资金账户档案', '银行户 / 平台货款户 / 现金，各挂一个会计科目'),
        ('科目表', '会计科目 + 报表项目（报表靠它取数）'),
        ('资金规则', '资金业务类型 → 对方科目 / 现金流类别'),
        ('期初余额', '★ 一年一个账套，年初数在这里录'),
    ]),
    ('② 天天录', [
        ('采购登记', '进货。含返利率 → 实际采购成本 = 不含税 − 应收返利'),
        ('销售登记', '卖货。平台扣费单列，成本按实际采购加权成本结转'),
        ('返利登记', '跟供应商对账确认返利、跟踪到账和发票'),
        ('一件代发结算', '不进库存的代发订单'),
        ('资金流水', '所有进出钱。★ 关联单号一定要填对'),
        ('发票台账', '进项 / 销项票，只管票不生成凭证'),
        ('费用及其他', '推广、运费、工资、房租、买设备…'),
        ('手工凭证', '折旧、计提、正式往来抵销这些做不出来的'),
    ]),
    ('③ 看结果', [
        ('老板经营看板', '★ 老板只看这一页'),
        ('报表勾稽检查', '★ 出报表前先看这一页，全 ✔ 再发'),
        ('异常预警中心', '哪些单有问题，逐笔列给你'),
        ('综合往来对账', '按「公司＋往来单位」算净往来，不跨公司抵销'),
        ('收付款核销中心', '还没结清的单，带账龄'),
        ('店铺利润分析', '哪个店赚钱、哪个店亏'),
        ('商品库存与成本', '每个商品还剩多少、值多少钱'),
        ('利润表', '本期 / 本年累计'),
        ('资产负债表', '期末 / 年初'),
        ('现金流简表', '经营 / 投资 / 筹资'),
        ('科目余额表', '所有科目的年初、本期、期末'),
        ('自动凭证', '业务登记自动展开的分录（查账用）'),
        ('测试说明', 'TEST 样例数据说明，验收照着走'),
        ('使用说明', '完整说明书：口径、凭证规则、常见问题'),
    ]),
]
r = 4
for grp, items in NAV:
    put(ws, f'B{r}', grp, font=F_SEC, fill=FILL_SEC, align=CL)
    put(ws, f'C{r}', None, font=F_SEC, fill=FILL_SEC)
    ws.row_dimensions[r].height = 24
    r += 1
    for nm, note in items:
        c = put(ws, f'B{r}', nm, font=F_LINK, fill=FILL_AUTO, align=CL)
        c.hyperlink = f"#'{nm}'!A1"
        put(ws, f'C{r}', note, font=F_NOTE, align=CL)
        r += 1
    r += 1
STAT = 4
put(ws, f'E{STAT}', '当 前 状 态', font=F_SEC, fill=FILL_SEC, align=CL)
ws.merge_cells(f'E{STAT}:H{STAT}')
STATS = [
    ('查询范围', '=本期标题', TXT),
    ('营业收入(本年)', f'=ROUND(利润表!$D${PLr["收入"]},2)', MONEY0),
    ('净利润(本年)', f'=ROUND(利润表!$D${PLr["净利润"]},2)', MONEY0),
    ('期末货币资金', '=ROUND(资产负债表!$B$4,2)', MONEY0),
    ('应收余额', f'=ROUND({rep(Q("应收账款"), "DR", "end")},2)', MONEY0),
    ('应付余额', f'=ROUND({rep(Q("应付账款"), "CR", "end")},2)', MONEY0),
    ('库存金额', f'=ROUND({rep(Q("存货"), "DR", "end")},2)', MONEY0),
    ('异常预警条数', f'=N(异常预警中心!$C${RT2})', '#,##0'),
    ('勾稽不平项', f'=SUMPRODUCT((ABS(N(报表勾稽检查!$E$4:$E${JT-1}))>=0.01)*1)', '#,##0'),
]
r = STAT + 1
for nm, f, fmt in STATS:
    put(ws, f'E{r}', nm, font=F_TXT, fill=FILL_KPI, align=CL)
    put(ws, f'F{r}', f, font=F_TOT, fill=FILL_KPI, fmt=fmt, align=CL if fmt == TXT else C)
    if fmt == TXT:
        ws.merge_cells(f'F{r}:H{r}')
    r += 1
put(ws, f'E{r}', '整 体 结 论', font=F_TOT, fill=FILL_TOT, align=CL)
put(ws, f'F{r}', f'=IF(AND(N($F{STAT+8})=0,N($F{STAT+9})=0),"✔ 账是干净的，报表可以对外出",'
                 f'IF(N($F{STAT+9})>0,"✘ 报表不平，先去《报表勾稽检查》",'
                 f'"⚠ 报表是平的，但有 "&N($F{STAT+8})&" 条业务预警待处理"))',
    font=F_TOT, fill=FILL_TOT, align=CL)
ws.merge_cells(f'F{r}:H{r}')
put(ws, f'E{r+2}',
    '三句话上手：\n'
    '① 先把《往来单位主档》《商品档案》《店铺档案》《资金账户档案》《期初余额》填好；\n'
    '② 之后天天只用管中间那八张登记表 —— 凭证和报表会自己长出来；\n'
    '③ 每月末看《报表勾稽检查》和《异常预警中心》，全绿了再把《老板经营看板》发出去。\n\n'
    '模板里带的 TEST 数据是样例，对过账之后整片删掉换成你们的真数据即可（删值不删行）。',
    font=F_NOTE, align=CT, fill=FILL_KPI)
ws.merge_cells(f'E{r+2}:H{r+8}')
page(ws, landscape=False)

# ════════════════════════════════════════════════════════════
# ㉝ 使用说明
# ════════════════════════════════════════════════════════════
ws = sheet('使用说明')
title(ws, '使 用 说 明', 'C', None)
widths(ws, {'A': 4, 'B': 30, 'C': 108})
DOC = [
    ('H', '一、这套表是怎么转起来的', ''),
    ('T', '一句话', '采购 / 销售 / 返利 / 代发 / 资金 / 发票 / 费用 七张登记表 → 《自动凭证》按固定规则展开成分录 '
          '→ 加上《手工凭证》→ 《科目余额表》→ 利润表 / 资产负债表 / 现金流简表。'
          '你只录业务，不用再录一遍总账。'),
    ('T', '什么要手工做', '只有四类：折旧摊销、税费计提、正式往来抵销、期末调整（存货报废、重分类等）。'
          '这些在《手工凭证》里做，其余一律别手工记 —— 记了就是重复。'),
    ('T', '颜色约定', '浅黄＝手工填；浅灰＝公式自动算的，别往里面打字；浅绿＝合计或关键结果；浅红＝预警。'),
    ('H', '二、同一家单位既是供应商又是客户怎么办', ''),
    ('T', '一家只建一次', '《往来单位主档》里一家公司只建一个编码，右边四个「是否」按实际情况勾。'
          '勾了供应商，采购登记的下拉里才有它；勾了客户，销售登记的下拉里才有它。'),
    ('T', '账还是分开的', '采购走「应付账款」、销售走「应收账款」、返利走「其他应收款—应收返利」、'
          '融资走「短期借款」或「其他应付款—股东借款」。四条线各算各的，《综合往来对账》里分四组列出来，'
          '最后才给一个「经营净往来」和「净资金敞口」给老板看。'),
    ('T', '不会误抵销', '汇总的钥匙是「公司 + 往来单位」。甲公司欠宏发的钱，不会拿乙公司应收宏发的钱去冲。'
          '同一家在多家公司都有往来时，《综合往来对账》最右边会亮一条提示。'),
    ('T', '真要抵销怎么做', '必须有书面抵销协议，然后到《手工凭证》做一张：借 应付账款 / 贷 应收账款，'
          '用途标记选「往来抵销」，并且把「关联单号」填上（填采购单号，那张采购单的「已付/已结」就认这笔；'
          '填销售单号，那张销售单的「已回/已结」就认这笔）。系统不会自动帮你抵。'),
    ('H', '三、采购返利的口径', ''),
    ('T', '核心公式', '实际采购成本 = 不含税采购成本 − 应收返利。返利在《采购登记》按单预提，'
          '入库的「库存商品」就是净额，返利单独挂「其他应收款—应收返利」。'),
    ('T', '为什么不会重复算', '成本已经在入库时降下来了，卖的时候按净成本结转 —— '
          '所以返利不会再在利润表里当成一笔收入算第二次。'),
    ('T', '《返利登记》干什么', '三件事：跟供应商对账确认金额、跟踪到账、跟踪发票。'
          '你填的「确认返利金额」跟系统算的「账面应收返利」之间的差额（差异调整）才会生成凭证去调成本。'),
    ('T', '「确认返利金额」留空 vs 填 0', '★ 留空 ＝ 还没跟供应商对上账，这一行先不动账，状态显示「待对账确认」；'
          '填 0 ＝ 对完账了、这个月一分返利都没有，系统会把采购时预提的那笔冲掉。两者完全不同，别混。'),
    ('T', '收到返利', '去《资金流水》记「收-返利」，关联单号填返利单号（FLxxxx），'
          '《返利登记》的「已收金额」就自动跟上。'),
    ('H', '四、销售成本怎么结转', ''),
    ('T', '不是一卖就全转', '单位成本 = 截至这一单当天的加权平均实际采购成本（含期初库存），'
          '销售成本 = 单位成本 × 本单数量。进了 1000 个只卖 300 个，就只转 300 个的成本。'),
    ('T', '要手工指定', '《销售登记》里填「手工单位成本」那一列，填了就以它为准。'),
    ('T', '取不到成本', '说明这个商品在这家公司名下没进过货、期初也没填 —— 《异常预警中心》W07 会点名。'),
    ('H', '五、平台扣费和平台提现', ''),
    ('T', '平台扣费', '佣金/技术服务费是平台从货款里直接扣的，所以一边进「销售费用」、一边冲「应收账款」。'
          '「未回款 = 含税金额 − 平台扣费 − 已回/已结」，平台打过来的净额正好对得上。'),
    ('T', '平台提现', '平台货款户和银行户都是我们自己的账户，提现只是左口袋到右口袋。'
          '在《资金流水》记两行：平台户「内部转账-转出」＋银行户「内部转账-转入」，两行都填「对方账户」。'
          '《现金流简表》第四行会显示它们的轧差，正常永远是 0。'),
    ('T', '内部转账的凭证只出一张', '★「转出」那一行已经生成了一张完整凭证（借 对方账户科目／贷 本账户科目），'
          '「转入」那一行不再出凭证 —— 它的作用只是让收款账户的余额动起来。'
          '所以在《自动凭证》里只能查到一张，这是对的，不是漏了。'
          '另外：对方账户必须是同一家公司的账户，跨公司调钱要走往来，不能用内部转账（预警 W22 会抓）。'),
    ('H', '六、自动凭证的展开规则（想改口径就照这个改）', ''),
    ('T', '采购登记（4 行）', '借 库存商品＝实际采购成本；借 应交税费—进项税额＝税额；'
          '借 其他应收款—应收返利＝应收返利；贷 应付账款＝含税金额。'),
    ('T', '销售登记（7 行）', '借 应收账款＝含税金额；贷 主营业务收入＝不含税；贷 销项税额；'
          '借 主营业务成本／贷 库存商品＝销售成本；借 销售费用／贷 应收账款＝平台扣费。'),
    ('T', '返利登记（2 行）', '只对「差异调整」做：差额为正 → 借 其他应收款—应收返利／贷 主营业务成本；为负则反过来。'),
    ('T', '一件代发（8 行）', '销售那三行同上；成本这边 借 主营业务成本＋借 进项税／贷 应付账款；'
          '另加平台扣费两行。★ 不碰库存商品。'),
    ('T', '资金流水（2 行）', '收钱：借 账户科目／贷 对方科目；付钱反过来。'
          '对方科目按《资金规则》判，内部转账按「对方账户」判，填了「科目手工覆盖」就以它为准。'),
    ('T', '费用及其他（3 行）', '借 对应科目＝不含税；借 进项税＝税额；贷 应付账款＝价税合计。'
          '付款时在《资金流水》记「付-费用」并填费用单号，应付就冲掉了。'),
    ('T', '金额是 0 的行', '自动隐去，所以《自动凭证》上看到的都是真有内容的分录。'),
    ('H', '七、一年一个账套 / 换年度', ''),
    ('T', '步骤', '① 复制整个文件，改名成下一年；② 《查询设置》里把年度改掉；'
          '③ 把上一年《资产负债表》的年末数，按「公司＋科目」抄进《期初余额》主区；'
          '④ 右边三个明细小区（往来单位 / 商品库存 / 资金账户）也要跟着更新；'
          '⑤ 把七张登记表里上一年的数据整片删掉（删值不删行）。'),
    ('T', '为什么明细小区也要填', '《综合往来对账》的年初往来、《商品库存与成本》的年初库存、'
          '《现金流简表》的期初现金，都是从那三个小区取的。不填的话本年发生额对，但余额从 0 起算。'),
    ('H', '八、每月的固定动作', ''),
    ('T', '月中', '有业务就记：进货记采购、卖货记销售、收付钱记资金流水、收到票记发票台账。'),
    ('T', '月末', '① 跟供应商对返利 → 《返利登记》；② 计提折旧、税费 → 《手工凭证》；'
          '③ 看《异常预警中心》把红的处理掉；④ 看《报表勾稽检查》全 ✔；'
          '⑤ 《查询设置》选上这个月，把《老板经营看板》发出去。'),
    ('H', '九、几个最容易踩的坑（都有预警盯着）', ''),
    ('T', '金额被粘成文本', '★ 从别的表复制粘贴最容易出。文本金额参与不了计算，那一单**凭证直接不生成**，'
          '而报表照样是平的，肉眼完全看不出来。预警 W21 专门抓这个；'
          '修法：选中那一列 →【数据】→【分列】→ 直接点完成。'),
    ('T', '科目名写错', '★ 手工凭证里科目名打错一个字，这笔钱就掉在报表外面，借贷却还是平的。'
          '所以科目一定要从下拉里选。预警 W17 会扫自动凭证和手工凭证两边。'),
    ('T', '新增科目忘了填报表项目', '★ 报表是按「报表项目」取数的。《科目表》E 列有下拉，'
          '只能选那 24 个合法值；填别的，这个科目的钱进不了报表。预警 W23 会抓。'),
    ('T', '查询参数填错', '★《查询设置》的年度和月份一定要从下拉里选。'
          '手打的文本「3」以前会让所有报表静默归零 —— 现在已经能认了，'
          '而且那一页的「参数自检」格会直接告诉你对不对。'),
    ('T', '本模板不做损益结转', '★ 月末**不要**做「结转本年利润」的凭证。'
          '未分配利润 ＝ 年初未分配利润 ＋ 利润表本年累计净利润，报表自己算好了；'
          '再手工结转一次就等于把利润减了两遍。所以手工凭证的用途下拉里没有这个选项。'),
    ('H', '十、常见问题', ''),
    ('T', '报表全是 0', '《查询设置》里的「公司」选了一家没业务的，或者「月份」选错了。'),
    ('T', '资产负债表不平', '按顺序查《报表勾稽检查》的 J05（期初平不平）→ J01/J02（凭证平不平）'
          '→ J03/J04（科目余额表平不平）。多半是某个科目的「报表项目」没填。'),
    ('T', '某笔钱核销不到单上', '《资金流水》的「关联单号」没填或填错。预警 W12 会列出来。'),
    ('T', '下拉里没有我要的选项', '去对应的档案表加一行就行 —— 下拉区域是按表头名字动态找列的，'
          '加了内容自动进下拉，不用改公式。'),
    ('T', '往来单位改名了', '先在《往来单位主档》改，然后用「查找替换」把各登记表里的旧名字一起换掉 —— '
          '往来是按名字匹配的。'),
    ('T', '表变卡', '这套表公式多，属正常。实在慢就把《自动凭证》《_自动清单》的空白行往上删一些'
          '（保留比实际业务量多一倍的余量即可）。'),
    ('T', '不要另存为别的格式', 'xlsx 存成 xls 或用别的软件转一道，公式和数据校验可能变形。正常 Ctrl+S 保存没问题。'),
]
r = 4
for kind, a, b in DOC:
    if kind == 'H':
        put(ws, f'B{r}', a, font=F_SEC, fill=FILL_SEC, align=CL)
        put(ws, f'C{r}', None, font=F_SEC, fill=FILL_SEC)
        ws.row_dimensions[r].height = 24
    else:
        put(ws, f'B{r}', a, font=F_TOT, fill=FILL_AUTO, align=CL)
        put(ws, f'C{r}', b, font=F_TXT, align=CL)
        ws.row_dimensions[r].height = max(18, 15 * (1 + len(b) // 54))
    r += 1
page(ws, titles=None, landscape=False)

# ════════════════════════════════════════════════════════════
# 收尾：排序、页签颜色、只选中主页
# ════════════════════════════════════════════════════════════
ORDER = ['主页', '使用说明', '查询设置', '基础资料', '往来单位主档', '商品档案', '店铺档案',
         '资金账户档案', '科目表', '资金规则', '期初余额',
         '采购登记', '销售登记', '返利登记', '一件代发结算', '资金流水', '发票台账',
         '费用及其他', '手工凭证',
         '自动凭证', '科目余额表', '利润表', '资产负债表', '现金流简表',
         '综合往来对账', '收付款核销中心', '店铺利润分析', '商品库存与成本',
         '异常预警中心', '报表勾稽检查', '老板经营看板', '测试说明', '_自动清单']
assert set(ORDER) == set(wb.sheetnames), set(ORDER) ^ set(wb.sheetnames)
wb._sheets = [wb[n] for n in ORDER]
TAB = {'主页': '1F3864', '使用说明': '1F3864',
       '查询设置': '2E75B6', '基础资料': '2E75B6', '往来单位主档': '2E75B6', '商品档案': '2E75B6',
       '店铺档案': '2E75B6', '资金账户档案': '2E75B6', '科目表': '2E75B6', '资金规则': '2E75B6',
       '期初余额': '2E75B6',
       '采购登记': 'ED7D31', '销售登记': 'ED7D31', '返利登记': 'ED7D31', '一件代发结算': 'ED7D31',
       '资金流水': 'ED7D31', '发票台账': 'ED7D31', '费用及其他': 'ED7D31', '手工凭证': 'ED7D31',
       '自动凭证': '7F7F7F', '科目余额表': '70AD47', '利润表': '70AD47', '资产负债表': '70AD47',
       '现金流简表': '70AD47',
       '综合往来对账': 'FFC000', '收付款核销中心': 'FFC000', '店铺利润分析': 'FFC000',
       '商品库存与成本': 'FFC000',
       '异常预警中心': 'C00000', '报表勾稽检查': 'C00000', '老板经营看板': '7030A0',
       '测试说明': '7F7F7F'}
for n, c in TAB.items():
    wb[n].sheet_properties.tabColor = c
for n in wb.sheetnames:
    wb[n].sheet_view.tabSelected = (n == '主页')
wb.active = 0
wb.save(OUT)
print(f'已写 {OUT}')
