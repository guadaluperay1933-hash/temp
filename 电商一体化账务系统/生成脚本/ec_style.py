# -*- coding: utf-8 -*-
"""电商一体化账务模板 —— 样式常量与放格子的小工具"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

FN = '微软雅黑'
F_TITLE = Font(name=FN, size=15, bold=True, color='1F3864')
F_NOTE  = Font(name=FN, size=9,  color='7F7F7F')
F_HDR   = Font(name=FN, size=10, bold=True, color='FFFFFF')
F_HDR2  = Font(name=FN, size=10, bold=True, color='1F3864')
F_TXT   = Font(name=FN, size=10)
F_IN    = Font(name=FN, size=10, color='0000C0')          # 手工录入＝蓝字
F_AUTO  = Font(name=FN, size=10, color='404040')          # 公式自动＝深灰
F_TOT   = Font(name=FN, size=10, bold=True)
F_LINK  = Font(name=FN, size=10, color='1F4E79', underline='single')
F_WARN  = Font(name=FN, size=10, bold=True, color='C00000')
F_OK    = Font(name=FN, size=10, bold=True, color='2E7D32')
F_BIG   = Font(name=FN, size=18, bold=True, color='1F3864')
F_SEC   = Font(name=FN, size=11, bold=True, color='1F3864')

FILL_HDR   = PatternFill('solid', fgColor='2F5597')       # 深蓝表头
FILL_HDR2  = PatternFill('solid', fgColor='D6E4F0')       # 浅蓝副表头
FILL_IN    = PatternFill('solid', fgColor='FFF2CC')       # 淡黄＝手工填
FILL_AUTO  = PatternFill('solid', fgColor='F2F2F2')       # 浅灰＝公式
FILL_TOT   = PatternFill('solid', fgColor='E2EFDA')       # 浅绿＝合计
FILL_WARN  = PatternFill('solid', fgColor='FCE4E4')       # 浅红＝预警
FILL_KPI   = PatternFill('solid', fgColor='EAF1FA')
FILL_SEC   = PatternFill('solid', fgColor='DDEBF7')

C  = Alignment(horizontal='center', vertical='center', wrap_text=True)
CL = Alignment(horizontal='left',   vertical='center', wrap_text=True)
CR = Alignment(horizontal='right',  vertical='center')
CT = Alignment(horizontal='left',   vertical='top',    wrap_text=True)

_s = Side(style='thin', color='BFBFBF')
BOX = Border(left=_s, right=_s, top=_s, bottom=_s)
_m = Side(style='medium', color='2F5597')
BOXM = Border(left=_m, right=_m, top=_m, bottom=_m)

MONEY = '#,##0.00;[Red]\\-#,##0.00;\\-'
MONEY0 = '#,##0;[Red]\\-#,##0;\\-'
NUM   = '#,##0.####;[Red]\\-#,##0.####;\\-'
PRICE = '#,##0.0000;[Red]\\-#,##0.0000;\\-'
PCT   = '0.00%;[Red]\\-0.00%;\\-'
DATEF = 'yyyy-mm-dd'
DATEQ = 'yyyy-mm-dd;;;@'          # 自动区日期：0 不显示
YM    = '000000'
TXT   = '@'


def put(ws, coord, value=None, font=F_TXT, fill=None, align=C, fmt=None, border=BOX):
    c = ws[coord]
    if value is not None:
        c.value = value
    c.font = font
    if fill is not None:
        c.fill = fill
    if align is not None:
        c.alignment = align
    if fmt is not None:
        c.number_format = fmt
    if border is not None:
        c.border = border
    return c


def title(ws, text, last_col, sub=None):
    """第1行大标题，第2行灰色说明"""
    ws.merge_cells(f'A1:{last_col}1')
    put(ws, 'A1', text, font=F_TITLE, align=CL, border=None)
    ws.row_dimensions[1].height = 30
    if sub is not None:
        ws.merge_cells(f'A2:{last_col}2')
        put(ws, 'A2', sub, font=F_NOTE, align=CL, border=None)
        ws.row_dimensions[2].height = 30


def headers(ws, row, names, fill=FILL_HDR, font=F_HDR, height=32):
    for i, n in enumerate(names, 1):
        if n is None:
            continue
        put(ws, f'{L(i)}{row}', n, font=font, fill=fill)
    ws.row_dimensions[row].height = height


def widths(ws, m):
    for k, v in m.items():
        ws.column_dimensions[k].width = v


def page(ws, titles=None, landscape=True, fit=1):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = fit
    ws.page_setup.fitToHeight = 0
    if titles:
        ws.print_title_rows = titles


def band(ws, r0, r1, c0, c1, fill=None, font=F_AUTO, fmt=None, align=C):
    """整片区域铺样式"""
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            put(ws, f'{L(c)}{r}', None, font=font, fill=fill, align=align, fmt=fmt)
