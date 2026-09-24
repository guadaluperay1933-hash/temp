# -*- coding: utf-8 -*-
"""A050 第八轮 · 《01 水果进销存台账》

 ① 基础资料新增「购买方」列后，所有下拉改成「按表头名字找列」的动态区域 ——
    以后再插列、再往下加内容，下拉都跟着走，不会再错位，也不用回来改公式。
 ② 成品出库明细 F/G/H 三列原来共用一个下拉（都指向 C 货主/客户）：
    现在 F/G 取「货主/客户」、H 取「购买方」。
 ③ 金额口径按客户确认的来：销售金额(应收)=数量×销售单价、采购金额(应付)=数量×采购单价，
    运费不并进去；另开「应收合计」「应付合计」两列把运费加回来。
 ④ 果然鲜销售汇总的「已收金额」原来比对的是「已收清」，
    但录入下拉里写的是「已收讫」—— 一字之差，已收永远算成 0，全都显示未收。

跑法：python3 fix08_01.py <入> <出>
"""
import sys, os, copy
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from style01 import *
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)

B0, B1, HDR, LAST = 4, 103, 3, 'P'     # 基础资料：数据行 / 表头行 / 下拉选项区最后一列

# ══ ① 基础资料：把已经在用的「购买方」种进去，下拉一打开就有东西选 ══
ws = wb['基础资料']
names = {ws.cell(HDR, c).value: c for c in range(1, 17) if ws.cell(HDR, c).value}
assert names.get('购买方') == 4, names

det = wb['成品出库明细']
seen, buyers = set(), []
for r in range(4, 3004):
    v = det.cell(r, 8).value
    if isinstance(v, str) and v.strip() and v.strip() not in seen:
        seen.add(v.strip()); buyers.append(v.strip())
have = {ws.cell(r, 4).value for r in range(B0, B1 + 1)}
row, added = B0, 0
for b in buyers:
    if b in have: continue
    while row <= B1 and ws.cell(row, 4).value not in (None, ''): row += 1
    if row > B1: break
    put(ws, f'D{row}', b, font=F_IN, fill=FILL_IN, align=CL)
    row += 1; added += 1
print(f'  ✓ 基础资料「购买方」补了 {added} 个已用到的名字：{"、".join(buyers) or "（无）"}')

# ══ ② 定义名称：按表头名字定位列，长度随内容走 ══
def dyn(t):
    m = f'MATCH("{t}",基础资料!$A${HDR}:${LAST}${HDR},0)'
    n = f'COUNTA(INDEX(基础资料!$A${B0}:${LAST}${B1},0,{m}))'
    return f'OFFSET(基础资料!$A${B0},0,{m}-1,MAX(1,{n}),1)'

DN = {'品种表': '品种', '等级表': '等级', '货主表': '货主/客户', '购买方表': '购买方',
      '供应商表': '供应商', '货权属性表': '货权属性', '单位规格表': '单位/规格',
      '筐子类型表': '筐子类型', '出库原因表': '出库原因', '销售类型表': '销售类型',
      '装卸方式表': '装卸方式', '核查结果表': '核查结果', '结算方式表': '结算方式',
      '收款状态表': '收款状态', '出入方向表': '出入方向', '业务类别表': '业务类别'}
for nm, t in DN.items():
    assert t in names, f'基础资料缺表头「{t}」'
    if nm in wb.defined_names: del wb.defined_names[nm]
    wb.defined_names.add(DefinedName(nm, attr_text=dyn(t)))
print(f'  ✓ 建了 {len(DN)} 个动态下拉区（按表头名字找列，新增内容自动进下拉）')

# ══ ③ 各子表下拉改指向定义名称 ══
BIND = {
    '计费规则':     {'B11:B79': '货主表'},
    '原料入库明细': {'D4:D3003': '品种表', 'E4:E3003': '等级表', 'F4:F3003': '货主表',
                     'G4:G3003': '货权属性表', 'I4:I3003': '单位规格表', 'L4:L3003': '筐子类型表',
                     'O4:O3003': '装卸方式表', 'P4:P3003': '核查结果表'},
    '原料出库明细': {'D4:D3003': '品种表', 'E4:E3003': '等级表', 'F4:F3003': '货主表',
                     'G4:G3003': '货权属性表', 'I4:I3003': '单位规格表', 'L4:L3003': '筐子类型表',
                     'N4:N3003': '出库原因表', 'P4:P3003': '装卸方式表', 'S4:S3003': '核查结果表'},
    '成品出库明细': {'D4:D3003': '品种表', 'E4:E3003': '等级表',
                     'F4:F3003': '货主表', 'G4:G3003': '货主表', 'H4:H3003': '购买方表',
                     'I4:I3003': '销售类型表', 'K4:K3003': '单位规格表', 'O4:O3003': '筐子类型表',
                     'S4:S3003': '装卸方式表', 'W4:W3003': '结算方式表', 'AC4:AC3003': '结算方式表',
                     'AD4:AD3003': '收款状态表', 'AE4:AE3003': '核查结果表'},
    '其他代存明细': {'D4:D803': '出入方向表', 'E4:E803': '品种表', 'F4:F803': '等级表',
                     'G4:G803': '货主表', 'I4:I803': '单位规格表', 'M4:M803': '核查结果表'},
}
KEEP = {'成品出库明细': ['X4:X3003']}          # 付款状态是写死的四个词，留着
for sn, mp in BIND.items():
    s = wb[sn]
    s.data_validations.dataValidation = [
        dv for dv in s.data_validations.dataValidation if str(dv.sqref) in KEEP.get(sn, [])]
    for sq, nm in mp.items():
        dv = DataValidation(type='list', formula1=f'={nm}', allow_blank=True, showDropDown=False,
                            showInputMessage=True)
        dv.errorTitle, dv.error = '不在基础资料里', '请从下拉里选；没有的先去【基础资料】对应列加一行'
        dv.promptTitle, dv.prompt = '从基础资料取', '点右边小箭头选。新加的内容会自动出现在这里'
        s.add_data_validation(dv); dv.add(sq)
    print(f'  ✓ {sn}：{len(mp)} 个下拉改成动态区')

# ══ ④ 成品出库明细：金额口径 + 应收/应付合计 ══
ws = wb['成品出库明细']
for c, txt in (('AI3', '应收合计\n(自动)'), ('AJ3', '应付合计\n(自动)')):
    ws[c]._style = copy.copy(ws['AG3']._style)
    ws[c].value = txt
ws.column_dimensions['AI'].width = 13
ws.column_dimensions['AJ'].width = 13
for r in range(4, 3004):
    put(ws, f'V{r}',  f'=IF($B{r}="","",IF(N($T{r})=0,"",ROUND(N($AH{r}),2)))',
        font=F_TXT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'AB{r}', f'=IF($B{r}="","",IF(N($Y{r})=0,"",ROUND(N($AG{r}),2)))',
        font=F_TXT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'AI{r}', f'=IF($B{r}="","",IF(AND(N($AB{r})=0,N($Z{r})=0,N($AA{r})=0),"",'
                      f'ROUND(N($AB{r})+N($Z{r})+N($AA{r}),2)))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'AJ{r}', f'=IF($B{r}="","",IF(AND(N($V{r})=0,N($U{r})=0),"",'
                      f'ROUND(N($V{r})+N($U{r}),2)))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
ws['A2'].value = ('★「销售金额(应收)」＝数量×销售单价，「采购金额(应付)」＝数量×采购单价，都只算货款；'
                  '运费在「自结运费 / 代发运费 / 采购代发运费」三列单独填。'
                  '最右边的「应收合计」＝销售金额＋自结运费＋代发运费，「应付合计」＝采购金额＋采购代发运费，'
                  '这两列才是真正要收、要付的钱，对账单和财务册都按它取数。')
print('  ✓ 成品出库明细：销售/采购金额只算货款，另加「应收合计」「应付合计」两列')

# ══ ⑤ 下游：凡是「实际要收/要付多少钱」的都改指 AI / AJ ══
def swap(ws, rows, cols, pairs):
    n = 0
    for r in rows:
        for cl in cols:
            v = ws[f'{cl}{r}'].value
            if not isinstance(v, str) or not v.startswith('='): continue
            new = v
            for a, b in pairs: new = new.replace(a, b)
            if new != v: ws[f'{cl}{r}'].value = new; n += 1
    return n

AB2AI = [('成品出库明细!$AB$4:$AB$3003', '成品出库明细!$AI$4:$AI$3003')]
V2AJ  = [('成品出库明细!$V$4:$V$3003',   '成品出库明细!$AJ$4:$AJ$3003')]

ws = wb['收入结算汇总']
n = 0                       # F「运费及其他」原来还加了一个空的 AJ 列，AJ 现在有含义了，得摘掉
for r in range(6, 15):
    v = ws[f'F{r}'].value
    if isinstance(v, str) and '+SUMIFS(成品出库明细!$AJ$' in v:
        i = v.find('+SUMIFS(成品出库明细!$AJ$')
        j, depth = i + 1, 0
        while j < len(v):
            if v[j] == '(': depth += 1
            elif v[j] == ')':
                depth -= 1
                if depth == 0: break
            j += 1
        ws[f'F{r}'].value = v[:i] + v[j + 1:]; n += 1
print(f'  ✓ 收入结算汇总 运费列摘掉失效的 AJ 项：{n} 格；'
      f'已收金额改按应收合计：{swap(ws, range(6, 15), ["H"], AB2AI)} 格')

ws = wb['财务取数接口']
print(f'  ✓ 财务取数接口 月度已收金额：{swap(ws, range(76, 100), ["H"], AB2AI)} 格；'
      f'往来款应收/应付：{swap(ws, range(207, 506), ["D", "E", "H"], AB2AI + V2AJ)} 格')

# ══ ⑥ 对账明细接口：金额＝应收合计，另补两列给 03 / 04 取数 ══
ws = wb['对账明细接口']
put(ws, 'AE3', '销售货款', font=F_TOT, fill=FILL_HDR2)
put(ws, 'AF3', '应付合计', font=F_TOT, fill=FILL_HDR2)
ws.column_dimensions['AE'].width = 12
ws.column_dimensions['AF'].width = 12
for r in range(4, 1204):
    v = ws[f'K{r}'].value
    if isinstance(v, str):
        ws[f'K{r}'].value = v.replace('成品出库明细!$AB$4:$AB$3003', '成品出库明细!$AI$4:$AI$3003')
    for col, src in (('AE', 'AB'), ('AF', 'AJ')):
        put(ws, f'{col}{r}',
            f'=IF(OR($Z{r}="",$Z{r}<=6000,$Z{r}>9000),"",'
            f'IF(N(INDEX(成品出库明细!${src}$4:${src}$3003,$Z{r}-6000))=0,"",'
            f'N(INDEX(成品出库明细!${src}$4:${src}$3003,$Z{r}-6000))))',
            font=F_NOTE, border=None, fmt=MONEY)
ws['K3'].value = '金额(应收合计)'
print('  ✓ 对账明细接口：金额改成应收合计，补「销售货款」「应付合计」两列')

# ══ ⑦ 果然鲜销售明细 / 汇总 / 采购明细 ══
ws = wb['果然鲜销售明细']
ws['U3']._style = copy.copy(ws['O3']._style); ws['U3'].value = '应收合计'
ws.column_dimensions['U'].width = 13
for r in range(4, 404):
    put(ws, f'U{r}', f'=IF($W{r}="","",IF(N(INDEX(成品出库明细!$AI$4:$AI$3003,$W{r}))=0,"",'
                     f'N(INDEX(成品出库明细!$AI$4:$AI$3003,$W{r}))))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
put(ws, 'U404', '=ROUND(SUM(U$4:U$403),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
ws['J407']._style = copy.copy(ws['I407']._style); ws['J407'].value = '应收合计'
ws['K407']._style = copy.copy(ws['J407']._style); ws['K407'].value = '备注'
for r in range(408, 412):
    put(ws, f'J{r}', f'=ROUND(SUMIF($D$4:$D$403,$B{r},$U$4:$U$403),2)',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
put(ws, 'J412', '=ROUND(SUM(J$408:J$411),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
print('  ✓ 果然鲜销售明细：补「应收合计」列，按销售类型汇总也跟着加一列')

# 购买方名单只列真有果然鲜销售的人，不再把几百个普通出库客户也排进来
ws = wb['_自动清单']
for r in range(4, 3004):
    put(ws, f'CD{r}', f'=IF($CA{r}<>1,"",IF(成品出库明细!$H{r}<>"",成品出库明细!$H{r},'
                      f'成品出库明细!$G{r}))', font=F_NOTE, border=None)

ws = wb['果然鲜销售汇总']
for r in range(4, 124):
    put(ws, f'I{r}', f'=IF($B{r}="","",ROUND(N($F{r})+N($G{r})+N($H{r}),2))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    # 已收金额要按「应收合计」算 —— 收款状态写了已收讫，运费也是一起收回来的
    put(ws, f'J{r}', f'=IF($B{r}="","",ROUND(SUMIFS(果然鲜销售明细!$U$4:$U$403,'
                     f'果然鲜销售明细!$F$4:$F$403,$B{r},果然鲜销售明细!$S$4:$S$403,"已收讫"),2))',
        font=F_TXT, fill=FILL_AUTO, fmt=MONEY)
ws['A126'].value = ('注：「销售金额(应收)」只算货款（数量×销售单价），后面两列是运费；'
                    '「应收合计」＝销售金额＋自结运费＋代发运费，才是客户实际要付的钱。'
                    '「已收金额」按【成品出库明细】收款状态选了「已收讫」的行统计。')
print('  ✓ 果然鲜销售汇总：应收合计＝货款＋运费；已收金额改比对「已收讫」')

ws = wb['果然鲜采购明细']
ws['O3']._style = copy.copy(ws['K3']._style); ws['O3'].value = '应付合计'
ws.column_dimensions['O'].width = 13
for r in range(4, 404):
    put(ws, f'O{r}', f'=IF($P{r}="","",IF(N(INDEX(成品出库明细!$AJ$4:$AJ$3003,$P{r}))=0,"",'
                     f'N(INDEX(成品出库明细!$AJ$4:$AJ$3003,$P{r}))))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
put(ws, 'O404', '=ROUND(SUM(O4:O403),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
# 按货主汇总：E 采购代发运费 / F 采购金额(应付) / G 应付合计 / H 已付金额 / I 未付金额
for col, txt in (('G', '应付合计'), ('H', '已付金额'), ('I', '未付金额')):
    ws[f'{col}407']._style = copy.copy(ws['F407']._style); ws[f'{col}407'].value = txt
for r in range(408, 523):
    put(ws, f'G{r}', f'=IF($B{r}="","",ROUND(N($E{r})+N($F{r}),2))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'H{r}', f'=IF($B{r}="","",ROUND(SUMIFS($O$4:$O$403,$D$4:$D$403,$B{r},'
                     f'$L$4:$L$403,"已付清"),2))', font=F_TXT, fill=FILL_AUTO, fmt=MONEY)
    put(ws, f'I{r}', f'=IF($B{r}="","",ROUND(N($G{r})-N($H{r}),2))',
        font=F_TOT, fill=FILL_AUTO, fmt=MONEY)
print('  ✓ 果然鲜采购明细：补「应付合计」列，按货主汇总改成 应付合计/已付/未付')

wb.save(OUT)
print('已写', OUT)
