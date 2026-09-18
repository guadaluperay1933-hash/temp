# -*- coding: utf-8 -*-
"""A050 第八轮 · 《04 综合查询对账单》

 ① 【对接源_01水果】最后 6 列（购买方 / 销售单价 / 自结运费 / 代发运费 / 采购代发运费 / 销售类型）
    整整错了一行 —— 第 4 行取的是源表第 5 行。购买方一错位，
    「果然鲜销售对账单」按购买方筛出来的明细自然全是错的。
 ② 【_自动清单】里「购买方」那几列上一轮直接盖在了 AG/AH/AI 上，
    而 AG/AI 本来是【物料与筐子销售对账单】的取数序号 —— 两边互相踩。
    现在购买方挪到 AN~AQ，AG/AH/AI 还原。
 ③ 【公司购买果品对账单】的取数序号 AK 是「累计的累计」，纯乱码；
    改成「本期真买下的那几笔」——只列填了采购单价/采购金额的行。
 ④ 两张对账单按新的金额口径重排：货款、运费、合计分开列；「已收清」改「已收讫」。

跑法：python3 fix08_04.py <入> <出>
"""
import sys, os, re, copy, zipfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from style01 import *
from openpyxl.worksheet.datavalidation import DataValidation

SRC, OUT = sys.argv[1], sys.argv[2]

z = zipfile.ZipFile(SRC)
wbxml = z.read('xl/workbook.xml').decode('utf8')
ids = re.findall(r'r:id="(rId\d+)"',
                 re.search(r'<externalReferences>(.*?)</externalReferences>', wbxml, re.S).group(1))
rels = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"',
                       z.read('xl/_rels/workbook.xml.rels').decode('utf8')))
IDX = {}
for k, rid in enumerate(ids, 1):
    nm = rels[rid].replace('../', '').split('/')[-1]
    tgt = re.search(rb'Target="([^"]+)"', z.read('xl/externalLinks/_rels/' + nm + '.rels')).group(1).decode()
    IDX[tgt] = k
z.close()
E01 = IDX['01_水果进销存台账模板.xlsx']
print(f'  《01》在本册里是 [{E01}]')

wb = openpyxl.load_workbook(SRC)
S0, S1 = 4, 1203            # 对接源两张表的数据行
A0, A1 = 4, 2403            # _自动清单的数据行
FRUIT = 1200                # _自动清单前 1200 行对应 01，后 1200 行对应 02

# ══ ① 对接源_01水果：行号回正，再补两列 ══
ws = wb['对接源_01水果']
FIX = [('AC', 'X'), ('AD', 'Y'), ('AE', 'AA'), ('AF', 'AB'), ('AG', 'AC'), ('AH', 'AD')]
for r in range(S0, S1 + 1):
    for col, sc in FIX:
        ws[f'{col}{r}'].value = (f'=IFERROR(IF([{E01}]对账明细接口!${sc}{r}=0,"",'
                                 f'[{E01}]对账明细接口!${sc}{r}),"")')
print(f'  ✓ 对接源_01水果：6 列错位的行号回正（{(S1 - S0 + 1) * 6} 格）')
for col, w, name, sc in (('AI', 12, '销售货款', 'AE'), ('AJ', 12, '应付合计', 'AF')):
    ws.column_dimensions[col].width = w
    put(ws, f'{col}3', name, font=F_TOT, fill=FILL_HDR2)
    for r in range(S0, S1 + 1):
        put(ws, f'{col}{r}', f'=IFERROR(IF([{E01}]对账明细接口!${sc}{r}=0,"",'
                             f'[{E01}]对账明细接口!${sc}{r}),"")', font=F_NOTE, border=None)
ws['K3'].value = '金额(应收合计)'
print('  ✓ 对接源_01水果：补「销售货款」「应付合计」两列')

# ══ ② _自动清单：AG/AH/AI 还原，购买方挪到 AN~AQ，公司购买另起 AL/AM ══
ws = wb['_自动清单']
st = copy.copy(ws['AF4']._style)
sh = copy.copy(ws['AF3']._style)

def cell(coord, v, style=None):
    c = ws[coord]; c.value = v
    if style is not None: c._style = copy.copy(style)
    return c

# AG 累计（物料销售）、AH 周转筐销售、AI 累计 —— 前 1200 行被上一轮盖掉了，补回来
ws['AG3'].value = '累计'; ws['AG3']._style = copy.copy(sh)
ws['AH3'].value = '周转筐销售'; ws['AH3']._style = copy.copy(sh)
ws['AI3'].value = '累计'; ws['AI3']._style = copy.copy(sh)
for r in range(A0, A0 + FRUIT):
    src = f'对接源_01水果!$D{r}', f'对接源_01水果!$K{r}'
    cell(f'AG{r}', f'=N(AG{r - 1})+$AF{r}', st)
    cell(f'AH{r}', f'=IF(AND($N{r}=1,AND(LEFT({src[0]},3)="周转筐",N({src[1]})<>0)),1,0)', st)
    cell(f'AI{r}', f'=N(AI{r - 1})+$AH{r}', st)
print(f'  ✓ _自动清单：AG/AH/AI 三列前 {FRUIT} 行还原（物料与筐子销售对账单的取数序号）')

for r in range(1, A1 + 1):                       # AJ/AK 整列废弃
    ws[f'AJ{r}'].value = None
    ws[f'AK{r}'].value = None
ws['AJ3'].value = '（已弃用）'; ws['AJ3']._style = copy.copy(sh)
ws['AK3'].value = '（已弃用）'; ws['AK3']._style = copy.copy(sh)

# AL 公司购买命中 / AM 累计
cell('AL1', '=公司购买果品对账单!$B$2', sh)
cell('AL2', '=IF(公司购买果品对账单!$F$2="",0,N(公司购买果品对账单!$F$2))', sh)
cell('AL3', '=IF(公司购买果品对账单!$H$2="",99999999,N(公司购买果品对账单!$H$2))', sh)
cell('AM1', '=MAX(AM$4:AM$2403)', sh)
cell('AM3', '累计', sh)
for r in range(A0, A1 + 1):
    if r < A0 + FRUIT:
        cell('AL%d' % r,
             f'=IF(OR($AL$1="",对接源_01水果!$Y{r}="",对接源_01水果!$Y{r}<>$AL$1,'
             f'对接源_01水果!$B{r}="",N(对接源_01水果!$AJ{r})<=0),0,'
             f'IF(OR(N(对接源_01水果!$B{r})<$AL$2,N(对接源_01水果!$B{r})>$AL$3),0,1))', st)
    else:
        cell(f'AL{r}', 0, st)
    cell(f'AM{r}', f'=N(AM{r - 1})+$AL{r}', st)
print('  ✓ _自动清单：新建 AL/AM —— 公司真买下的那几笔（填了采购金额的行）')

# AN 购买方 / AO 首次 / AP 累计 / AQ 紧凑
for coord, v in (('AN3', '购买方'), ('AO3', '首次'), ('AP3', '累计'), ('AQ3', '购买方(紧凑)')):
    cell(coord, v, sh)
cell('AQ1', '=MAX(AP$4:AP$1203)', sh)
for r in range(A0, A1 + 1):
    if r < A0 + FRUIT:
        cell(f'AN{r}', f'=IF(对接源_01水果!$AC{r}="","",对接源_01水果!$AC{r})', st)
        cell(f'AO{r}', f'=IF(AN{r}="",0,IF(MATCH(AN{r},$AN$4:$AN$1203,0)=ROW()-3,1,0))', st)
    else:
        cell(f'AN{r}', '=""', st)
        cell(f'AO{r}', 0, st)
    cell(f'AP{r}', f'=N(AP{r - 1})+$AO{r}', st)
for r in range(A0, A0 + 200):
    cell(f'AQ{r}', '=IFERROR(INDEX($AN$4:$AN$1203,MATCH(ROW()-3,$AP$4:$AP$1203,0)),"")', st)
print('  ✓ _自动清单：购买方名单挪到 AN~AQ，不再跟别人抢列')

# ══ ③ 果然鲜销售对账单 ══
ws = wb['果然鲜销售对账单']
R0, R1 = 6, 205
TOT = R1 + 1
for m in [str(x) for x in ws.merged_cells.ranges]:
    if m.startswith('A3:'): ws.unmerge_cells(m)
ws.merge_cells('A3:M3')
stH = copy.copy(ws['J5']._style)
stMoney = copy.copy(ws['J6']._style)
stTxt = copy.copy(ws['K6']._style)
stNote = copy.copy(ws['L6']._style)
stTot = copy.copy(ws['J206']._style)
stLbl = copy.copy(ws['N3']._style)
stVal = copy.copy(ws['P3']._style)

HDRS = ['日　期', '票据单号', '品　名', '等　级', '数　量', '单位', '销售单价',
        '销售金额（应收）', '自结运费', '代发运费', '应收合计', '收款状态', '备　注']
for i, t in enumerate(HDRS):
    c = ws.cell(5, i + 1); c.value = t; c._style = copy.copy(stH)
for k, w in zip('ABCDEFGHIJKLM', [12, 13, 10, 12, 9, 7, 11, 15, 11, 11, 15, 11, 16]):
    ws.column_dimensions[k].width = w

SRC_COL = {'A': 'B', 'B': 'E', 'C': 'F', 'D': 'G', 'E': 'H', 'F': 'I',
           'G': 'AD', 'H': 'AI', 'I': 'AE', 'J': 'AF', 'K': 'K', 'L': 'L', 'M': 'M'}
def pick(col, r):
    return (f'IF(INDEX(对接源_01水果!${col}$4:${col}$1203,$S{r})=0,"",'
            f'INDEX(对接源_01水果!${col}$4:${col}$1203,$S{r}))')
for r in range(R0, R1 + 1):
    for out, src in SRC_COL.items():
        style = stMoney if out in 'GHIJK' else (stNote if out == 'M' else stTxt)
        c = ws[f'{out}{r}']; c._style = copy.copy(style)
        c.value = f'=IF($S{r}="","",{pick(src, r)})'
    ws[f'S{r}'].value = (
        f'=IFERROR(_xlfn.AGGREGATE(15,6,(ROW(对接源_01水果!$B$4:$B$1203)-3)/'
        f'((对接源_01水果!$AC$4:$AC$1203=$B$2)*(对接源_01水果!$B$4:$B$1203>=$T$1)*'
        f'(对接源_01水果!$B$4:$B$1203<=$T$2)*(N(对接源_01水果!$K$4:$K$1203)>0)),ROW()-{R0 - 1}),"")')
for c in 'ABCDEFGHIJKLM':
    cell_ = ws[f'{c}{TOT}']; cell_._style = copy.copy(stTot)
    cell_.value = (f'=ROUND(SUM({c}{R0}:{c}{R1}),2)' if c in 'EHIJK'
                   else ('合  计' if c == 'A' else None))

COND = (f'对接源_01水果!$AC$4:$AC$1203,$B$2,对接源_01水果!$B$4:$B$1203,">="&$T$1,'
        f'对接源_01水果!$B$4:$B$1203,"<="&$T$2,对接源_01水果!$K$4:$K$1203,">0"')
SUM_ = [('销售总数量（箱）', f'=ROUND(SUMIFS(对接源_01水果!$H$4:$H$1203,{COND}),0)', NUM),
        ('销售金额（应收）', f'=ROUND(SUMIFS(对接源_01水果!$AI$4:$AI$1203,{COND}),2)', MONEY),
        ('其中·自结运费',   f'=ROUND(SUMIFS(对接源_01水果!$AE$4:$AE$1203,{COND}),2)', MONEY),
        ('其中·代发运费',   f'=ROUND(SUMIFS(对接源_01水果!$AF$4:$AF$1203,{COND}),2)', MONEY),
        ('应收合计',        f'=ROUND(SUMIFS(对接源_01水果!$K$4:$K$1203,{COND}),2)', MONEY),
        ('已收金额',        f'=ROUND(SUMIFS(对接源_01水果!$K$4:$K$1203,{COND},'
                            f'对接源_01水果!$L$4:$L$1203,"已收讫"),2)', MONEY),
        ('未收金额（应收）', '=ROUND(N($Q$7)-N($Q$8),2)', MONEY)]
for col in ('N', 'P'):
    for r in range(2, 10): ws[f'{col}{r}'].value = None
ws['O2'].value = '销 售 汇 总'; ws['O2']._style = copy.copy(ws['N2']._style); ws['N2'].value = None
for i, (lbl, f, fmt) in enumerate(SUM_):
    r = 3 + i
    c = ws[f'O{r}']; c.value = lbl; c._style = copy.copy(stLbl)
    v = ws[f'Q{r}']; v.value = f; v._style = copy.copy(stVal); v.number_format = fmt
for k, w in zip('NOPQ', [3, 18, 3, 14]): ws.column_dimensions[k].width = w
ws['E4'].value = '=ROUND(N($Q$7)-N($Q$8),2)'
ws['A4'].value = '客户欠我方（应收）总金额'
ws[f'F{TOT + 2}'].value = f'=ROUND($K${TOT},2)'
ws[f'F{TOT + 3}'].value = '=ROUND(N($Q$8),2)'
ws[f'F{TOT + 4}'].value = f'=ROUND(N($F${TOT + 2})-N($F${TOT + 3}),2)'
ws['A3'].value = ('★ 只填 B2 购买方与起止日期。这里列的是【01 成品出库明细】里填了「销售单价」的行 —— '
                  '公司（果然鲜）卖给这位客户的果品。「销售金额（应收）」只算货款（数量×销售单价），'
                  '运费分「自结运费」「代发运费」两列单列，三者相加＝「应收合计」，才是他要付给我们的钱。')
ws.data_validations.dataValidation = []
dv = DataValidation(type='list', formula1='=_自动清单!$AQ$4:$AQ$203', allow_blank=True, showDropDown=False,
                            showInputMessage=True)
dv.promptTitle, dv.prompt = '选购买方', '名单由【01 成品出库明细】的「购买方」列自动去重生成'
ws.add_data_validation(dv); dv.add('B2')
ws.print_title_rows = '$5:$5'
print('  ✓ 果然鲜销售对账单：明细改成 货款/自结运费/代发运费/应收合计 四列，已收改比对「已收讫」')

# ══ ④ 公司购买果品对账单 ══
ws = wb['公司购买果品对账单']
R0, R1 = 6, 155
for m in [str(x) for x in ws.merged_cells.ranges]:
    if m.startswith(('A3:', 'L2:', 'L3:', 'L4:', 'L5:', 'L6:')): ws.unmerge_cells(m)
for col in ('L', 'M', 'N', 'O', 'P'):          # 老的汇总块在 L/N，先腾干净再重排
    for r in range(1, 10): ws[f'{col}{r}'].value = None
ws.merge_cells('A3:L3')
stH = copy.copy(ws['H5']._style)
stMoney = copy.copy(ws['H6']._style)
stTxt = copy.copy(ws['I6']._style)
stNote = copy.copy(ws['J6']._style)
stLbl = copy.copy(ws['L3']._style)
stVal = copy.copy(ws['N3']._style)

HDRS = ['日　期', '票据单号', '品　名', '等　级', '数　量', '单位', '采购单价',
        '采购金额（应付）', '采购代发运费', '应付合计', '付款状态', '备　注']
for i, t in enumerate(HDRS):
    c = ws.cell(5, i + 1); c.value = t; c._style = copy.copy(stH)
for k, w in zip('ABCDEFGHIJKL', [13, 16, 11, 13, 10, 7, 11, 15, 13, 15, 11, 18]):
    ws.column_dimensions[k].width = w

SRC_COL = {'A': 'B', 'B': 'E', 'C': 'F', 'D': 'G', 'E': 'H', 'F': 'I',
           'G': 'Z', 'H': 'AA', 'I': 'AG', 'J': 'AJ', 'K': 'AB', 'L': 'M'}
for r in range(R0, R1 + 1):
    ws[f'AE{r}'].value = f'=IFERROR(MATCH({r - R0 + 1},_自动清单!$AM$4:$AM$2403,0),"")'
    for out, src in SRC_COL.items():
        style = stMoney if out in 'GHIJ' else (stNote if out == 'L' else stTxt)
        c = ws[f'{out}{r}']; c._style = copy.copy(style)
        c.value = (f'=IF($AE{r}="","",IF(INDEX(对接源_01水果!${src}$4:${src}$1203,$AE{r})=0,"",'
                   f'INDEX(对接源_01水果!${src}$4:${src}$1203,$AE{r})))')

COND = (f'对接源_01水果!$Y$4:$Y$1203,$B$2,对接源_01水果!$B$4:$B$1203,">="&$R$1,'
        f'对接源_01水果!$B$4:$B$1203,"<="&$R$2,对接源_01水果!$AJ$4:$AJ$1203,">0"')
SUM_ = [('采购总数量', f'=ROUND(SUMIFS(对接源_01水果!$H$4:$H$1203,{COND}),0)', NUM),
        ('采购金额（应付）', f'=ROUND(SUMIFS(对接源_01水果!$AA$4:$AA$1203,{COND}),2)', MONEY),
        ('其中·采购代发运费', f'=ROUND(SUMIFS(对接源_01水果!$AG$4:$AG$1203,{COND}),2)', MONEY),
        ('应付合计', f'=ROUND(SUMIFS(对接源_01水果!$AJ$4:$AJ$1203,{COND}),2)', MONEY),
        ('已付金额', f'=ROUND(SUMIFS(对接源_01水果!$AJ$4:$AJ$1203,{COND},'
                     f'对接源_01水果!$AB$4:$AB$1203,"已付清"),2)', MONEY),
        ('未付金额（应付）', '=ROUND(N($P$6)-N($P$7),2)', MONEY)]
ws['R1'].value = '=IF($F$2="",DATE(1900,1,1),$F$2)'
ws['R2'].value = '=IF($H$2="",DATE(2999,12,31),$H$2)'
ws['N2'].value = '采 购 汇 总'; ws['N2']._style = copy.copy(stLbl)
ws.merge_cells('N2:P2')
for i, (lbl, f, fmt) in enumerate(SUM_):
    r = 3 + i
    c = ws[f'N{r}']; c.value = lbl; c._style = copy.copy(stLbl)
    ws.merge_cells(f'N{r}:O{r}')
    v = ws[f'P{r}']; v.value = f; v._style = copy.copy(stVal); v.number_format = fmt
for k, w in zip('MNOPQR', [2, 16, 12, 14, 2, 12]): ws.column_dimensions[k].width = w
ws['H4'].value = '=ROUND(N($P$6),2)'
ws['A3'].value = ('★ 只填 B2 货主/客户与起止日期。这里只列公司（果然鲜）真买下来的那几笔 —— '
                  '也就是【01 成品出库明细】里填了「采购单价 / 采购金额」的行；'
                  '别的行是客户自己的货放在我们这里，不是我们买的，不在这里出现。'
                  '「采购金额（应付）」只算货款，运费单列，两者相加＝「应付合计」，才是我们要付给他的钱。')
print('  ✓ 公司购买果品对账单：只列真买下的那几笔，明细拆成 货款/运费/应付合计')

# 【应收应付汇总】顶上那句「往来单位 0/250」读的是 _自动清单!$A$1，可那格放的是说明文字，
# N(文字)=0，所以永远显示 0 —— 改成读真正的去重计数
ws = wb['应收应付汇总']
v = ws['A2'].value
if isinstance(v, str) and '_自动清单!$A$1' in v:
    ws['A2'].value = v.replace('_自动清单!$A$1', 'MAX(_自动清单!$C$4:$C$2403)')
    print('  ✓ 应收应付汇总：顶部「往来单位 n/250」的计数改成真的去重数')

wb.save(OUT)
print('已写', OUT)
