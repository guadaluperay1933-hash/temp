# -*- coding: utf-8 -*-
"""A050 第九轮 · 《02 物料与周转物台账》

① 《对账明细接口》的「发出筐数」原来是 出库数量 ＋ 销售数量 两列相加，
结果卖断给客户的筐子也被当成「借出去要还的往来筐」传到《04》，
在《筐子对账单》里跟领用出库的筐混在一起，还进了「未回库筐子（欠筐）」。

卖断的筐子已经是客户的了，不存在还不还 —— 这里把「发出筐数」改成只认「出库数量」，
销售数量走金额线（《04 物料与筐子销售对账单》）。

02 表自己的《筐子往来汇总》本来就是分开算的（发出用 J 列、销售单列在 U:Z），不受影响。

② 《_自动清单》里读《周转筐出入库明细》和《客户自备加工筐明细》的那几列，
   从某一行起整体错位了一行（指到了下一行去）：
       周转筐相关（Q / W / CS / CP 的周转筐段）—— 第 43 行起
       自备筐相关（AC / AI / CP 的自备筐段）    —— 第 5 行起
   后果：这两张明细表的**最后一行永远传不到《04》**。
   在明细表末尾新录一笔（比如卖掉一批筐子），04 的对账单就是看不到它，
   直到下面再录一行才会冒出来 —— 这次就是这么撞上的。
   这里按「一行对一行」把这些列整列重建。

跑法：python3 fix09_02.py <入> <出>
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)
ws = wb['对账明细接口']

# ══ 发出筐数：去掉「＋销售数量」 ══
fixed = 0
for r in range(4, 1204):
    c = ws[f'O{r}']
    v = c.value
    if not isinstance(v, str):
        continue
    dup = f'+N(INDEX(周转筐出入库明细!$K$4:$K$2004,$Z{r}-2000))'
    if dup in v:
        c.value = v.replace(dup, '')
        fixed += 1
assert fixed == 1200, f'只改到 {fixed} 行，公式结构可能变了'

# A2 是一整条 IF 公式，说明文字要塞进最后那段字符串里面，不能直接往后面接
note = ws['A2'].value
tail = ('　★「发出筐数」只取《周转筐出入库明细》的「出库数量」，不含「销售数量」——'
        '卖断的筐子不是往来筐，它走金额线，在《04 物料与筐子销售对账单》里对账。')
assert note.rstrip().endswith('")'), 'A2 结构变了，说明文字插不进去'
if tail not in note:
    body = note.rstrip()
    ws['A2'].value = body[:-2] + tail + '")'

# ══ _自动清单：把错位一行的取数列整列重建 ══
import re
al = wb['_自动清单']
#  列, 起行, 止行, 行偏移（本表行 − 明细行）, 明细表名
REBUILD = [
    ('Q',     4, 2003,    0, '周转筐出入库明细'),
    ('W',     4, 2003,    0, '周转筐出入库明细'),
    ('CS',    4, 2003,    0, '周转筐出入库明细'),
    ('AC',    4, 1203,    0, '客户自备加工筐明细'),
    ('AI',    4, 1203,    0, '客户自备加工筐明细'),
    ('CP', 2004, 4003, 2000, '周转筐出入库明细'),
    ('CP', 4004, 5203, 4000, '客户自备加工筐明细'),
]
shift = 0
for col, r0, r1, off, sheet in REBUILD:
    pat = re.compile(re.escape(sheet) + r'!\$([A-Z]{1,2})(\d+)')
    for r in range(r0, r1 + 1):
        c = al[f'{col}{r}']
        if not isinstance(c.value, str):
            continue
        want = r - off
        new = pat.sub(lambda m: f'{sheet}!${m.group(1)}{want}', c.value)
        if new != c.value:
            c.value = new
            shift += 1

wb.save(OUT)
print(f'  ✓ 02 表：对账明细接口「发出筐数」剔除销售数量，改了 {fixed} 行')
print(f'  ✓ 02 表：_自动清单 取数列错位修正，改了 {shift} 格')
