# -*- coding: utf-8 -*-
"""⑤ 《04 综合查询对账单》：
   · 【对接源_01水果】补 6 列（购买方 / 销售单价 / 自结运费 / 代发运费 / 采购代发运费 / 销售类型）
   · 【水果对账单】右上角汇总补上单位：入库出库是「筐」，成品出库是「箱」，不再看着像能相减
   · 新增【果然鲜销售对账单】：按购买方出，卖了什么、多少钱、运费多少、收了多少、还欠多少
   跑法：python3 fix05_04表.py <入> <出>"""
import sys, os, re, copy, zipfile, openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from style01 import *
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

SRC, OUT = sys.argv[1], sys.argv[2]

# 先弄清 [n] 各指向哪个文件
z = zipfile.ZipFile(SRC)
wbxml = z.read('xl/workbook.xml').decode('utf8')
ids = re.findall(r'r:id="(rId\d+)"', re.search(r'<externalReferences>(.*?)</externalReferences>', wbxml, re.S).group(1))
rels = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', z.read('xl/_rels/workbook.xml.rels').decode('utf8')))
IDX = {}
for k, rid in enumerate(ids, 1):
    nm = rels[rid].replace('../', '').split('/')[-1]
    tgt = re.search(rb'Target="([^"]+)"', z.read('xl/externalLinks/_rels/' + nm + '.rels')).group(1).decode()
    IDX[tgt] = k
z.close()
E01 = IDX['01_水果进销存台账模板.xlsx']
print(f'  《01》在本册里是 [{E01}]')

wb = openpyxl.load_workbook(SRC)
S0, S1 = 4, 1203

# ---------------- ① 对接源_01水果 补列 ----------------
ws = wb['对接源_01水果']
ADD = [('AC', 12, '购买方', 'X'), ('AD', 11, '销售单价', 'Y'), ('AE', 11, '自结运费', 'AA'),
       ('AF', 11, '代发运费', 'AB'), ('AG', 13, '采购代发运费', 'AC'), ('AH', 11, '销售类型', 'AD')]
for col, w, name, sc in ADD:
    ws.column_dimensions[col].width = w
    put(ws, f'{col}3', name, font=F_TOT, fill=FILL_HDR2)
    for r in range(S0, S1 + 1):
        put(ws, f'{col}{r}',
            f'=IFERROR(IF([{E01}]对账明细接口!${sc}{r + 1}=0,"",[{E01}]对账明细接口!${sc}{r + 1}),"")',
            font=F_NOTE, border=None)
print('  ✓ 对接源_01水果 补 6 列')

# ---------------- ② 水果对账单：汇总加单位 ----------------
ws = wb['水果对账单']
LBL = {3: '入库总数量（筐）', 4: '入库总毛重KG', 5: '出库总数量（筐）', 6: '结存数量（筐）',
       7: '成品出库数量（箱）', 8: '成品出库金额'}
for r, t in LBL.items():
    if ws.cell(row=r, column=12).value: ws.cell(row=r, column=12).value = t
# 成品出库拆成「箱」和「筐等其它单位」两行，不再把两种单位混成一个数
f7 = ws['N7'].value
if isinstance(f7, str) and f7.startswith('='):
    base = f7.rstrip(')')
    # 原式形如 =ROUND(SUMIFS(...)+SUMIFS(...),0)
    box = f7.replace('对接源_01水果!$D$4:$D$1203,"成品出库*"',
                     '对接源_01水果!$D$4:$D$1203,"成品出库*",对接源_01水果!$I$4:$I$1203,"箱"')
    box = box.replace('对接源_02物料!$D$4:$D$1203,"成品出库*"',
                      '对接源_02物料!$D$4:$D$1203,"成品出库*",对接源_02物料!$I$4:$I$1203,"箱"')
    oth = f7.replace('对接源_01水果!$D$4:$D$1203,"成品出库*"',
                     '对接源_01水果!$D$4:$D$1203,"成品出库*",对接源_01水果!$I$4:$I$1203,"<>箱"')
    oth = oth.replace('对接源_02物料!$D$4:$D$1203,"成品出库*"',
                      '对接源_02物料!$D$4:$D$1203,"成品出库*",对接源_02物料!$I$4:$I$1203,"<>箱"')
    ws['N7'] = box
    ws['L9'] = '成品出库数量（筐等）'
    ws['L9']._style = copy.copy(ws['L7']._style)
    ws['N9'] = oth
    ws['N9']._style = copy.copy(ws['N7']._style)
    print('  ✓ 水果对账单：成品出库数量拆成「箱」和「筐等」两行')

a3 = ws['A3']
if isinstance(a3.value, str) and '箱' not in a3.value:
    a3.value = (a3.value.rstrip('　 ') +
                '　★ 注意单位：原料入库/出库/结存是「筐」，成品出库是「箱」，两边不能直接相减；'
                '结存数量＝入库筐数−出库筐数，跟成品出库箱数没有加减关系。')
print('  ✓ 水果对账单：汇总项补上单位（筐 / 箱）')

# ---------------- ③ 新增 果然鲜销售对账单 ----------------
NM = '果然鲜销售对账单'
if NM in wb.sheetnames: del wb[NM]
ws = wb.create_sheet(NM, wb.sheetnames.index('公司购买果品对账单') + 1)
N = 200
R0, R1 = 6, 5 + N
title(ws, '新疆果然鲜仓储有限公司', 'G', sub_last='A',
      sub='★ 只填 B2 购买方与起止日期。这里列的是【01 成品出库明细】里填了「销售单价 / 销售金额」的行 —— '
          '也就是公司（果然鲜）卖给这位客户的果品，金额是他要付给我们的钱（应收）。'
          '运费分「自结运费」和「代发运费」两列，都已经含在销售金额里。')
put(ws, 'I1', '对账单', font=F_TITLE, align=C, border=None)
widths(ws, {'A': 12, 'B': 13, 'C': 10, 'D': 12, 'E': 9, 'F': 7, 'G': 11, 'H': 11, 'I': 11,
            'J': 15, 'K': 11, 'L': 16, 'M': 3, 'N': 18, 'O': 3, 'P': 14, 'Q': 3, 'R': 3, 'S': 6,
            'T': 12})
put(ws, 'A2', '购 买 方:', font=F_TOT, fill=FILL_HDR2, align=CR)
put(ws, 'B2', '王照雪', font=F_IN, fill=FILL_IN)
put(ws, 'E2', '开始日期:', font=F_TOT, fill=FILL_HDR2, align=CR)
put(ws, 'F2', None, font=F_IN, fill=FILL_IN, fmt=DATE)
put(ws, 'G2', '结束日期:', font=F_TOT, fill=FILL_HDR2, align=CR)
put(ws, 'H2', None, font=F_IN, fill=FILL_IN, fmt=DATE)
put(ws, 'T1', '=IF($F$2="",DATE(1900,1,1),$F$2)', font=F_NOTE, fmt=DATE, border=None)
put(ws, 'T2', '=IF($H$2="",DATE(2999,12,31),$H$2)', font=F_NOTE, fmt=DATE, border=None)
ws.column_dimensions['S'].hidden = True
ws.column_dimensions['T'].hidden = True
dv = DataValidation(type='list', formula1='=_自动清单!$AH$4:$AH$203', allow_blank=True,
                    showErrorMessage=False)
ws.add_data_validation(dv); dv.add('B2')

SRC01 = '对接源_01水果'
RNG = lambda c: f'{SRC01}!${c}$4:${c}$1203'
COND = (f'{RNG("AC")},$B$2,{RNG("B")},">="&$T$1,{RNG("B")},"<="&$T$2,{RNG("K")},">0"')
put(ws, 'N2', '销 售 汇 总', font=F_TOT, fill=FILL_HDR2)
SUMS = [('销售总数量（箱）', f'=ROUND(SUMIFS({RNG("H")},{COND}),0)', NUM),
        ('销售金额（应收）', f'=ROUND(SUMIFS({RNG("K")},{COND}),2)', MONEY),
        ('其中·自结运费', f'=ROUND(SUMIFS({RNG("AE")},{COND}),2)', MONEY),
        ('其中·代发运费', f'=ROUND(SUMIFS({RNG("AF")},{COND}),2)', MONEY),
        ('已收金额', f'=ROUND(SUMIFS({RNG("K")},{COND},{RNG("L")},"已收清"),2)', MONEY),
        ('未收金额（应收）', '=ROUND(N($P$4)-N($P$7),2)', MONEY)]
for i, (lab, f, fmt) in enumerate(SUMS):
    r = 3 + i
    put(ws, f'N{r}', lab, font=F_TOT, fill=FILL_HDR2, align=CL)
    put(ws, f'P{r}', f, font=F_TOT, fill=FILL_AUTO, fmt=fmt)
put(ws, 'A4', '客户欠我方（应收）总金额', font=F_TOT, fill=FILL_HDR2, align=CL)
ws.merge_cells('A4:D4')
put(ws, 'E4', '=ROUND(N($P$4)-N($P$7),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
ws.merge_cells('E4:H4')
headers(ws, 5, ['日　期', '票据单号', '品　名', '等　级', '数　量', '单位', '销售单价',
                '自结运费', '代发运费', '销售金额（应收）', '收款状态', '备　注'])
def pick(col, r):
    return (f'IF(INDEX({RNG(col)},$S{r})=0,"",INDEX({RNG(col)},$S{r}))')
for i in range(N):
    r = R0 + i
    # AGGREGATE(15,6,…) ＝ 免 Ctrl+Shift+Enter 的 SMALL+IF，Excel 2010+ 和 WPS 都认
    put(ws, f'S{r}',
        f'=IFERROR(AGGREGATE(15,6,(ROW({RNG("B")})-3)/'
        f'(({RNG("AC")}=$B$2)*({RNG("B")}>=$T$1)*({RNG("B")}<=$T$2)'
        f'*(N({RNG("K")})>0)),ROW()-{R0 - 1}),"")',
        font=F_NOTE, border=None)
    put(ws, f'A{r}', f'=IF($S{r}="","",{pick("B", r)})', font=F_LINK, fmt=DATE)
    put(ws, f'B{r}', f'=IF($S{r}="","",{pick("E", r)})', font=F_LINK)
    put(ws, f'C{r}', f'=IF($S{r}="","",{pick("F", r)})', font=F_LINK)
    put(ws, f'D{r}', f'=IF($S{r}="","",{pick("G", r)})', font=F_LINK)
    put(ws, f'E{r}', f'=IF($S{r}="","",{pick("H", r)})', font=F_LINK, fmt=NUM)
    put(ws, f'F{r}', f'=IF($S{r}="","",{pick("I", r)})', font=F_LINK)
    put(ws, f'G{r}', f'=IF($S{r}="","",{pick("AD", r)})', font=F_LINK, fmt=MONEY)
    put(ws, f'H{r}', f'=IF($S{r}="","",{pick("AE", r)})', font=F_LINK, fmt=MONEY)
    put(ws, f'I{r}', f'=IF($S{r}="","",{pick("AF", r)})', font=F_LINK, fmt=MONEY)
    put(ws, f'J{r}', f'=IF($S{r}="","",{pick("K", r)})', font=F_TOT, fmt=MONEY)
    put(ws, f'K{r}', f'=IF($S{r}="","",{pick("L", r)})', font=F_LINK)
    put(ws, f'L{r}', f'=IF($S{r}="","",{pick("M", r)})', font=F_NOTE, align=CL)
    ws.row_dimensions[r].height = 16
T = R1 + 1
put(ws, f'A{T}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in 'BCDFKL': put(ws, f'{c}{T}', None, font=F_TOT, fill=FILL_TOT)
for c in 'EGHIJ':
    put(ws, f'{c}{T}', f'=ROUND(SUM({c}{R0}:{c}{R1}),2)', font=F_TOT, fill=FILL_TOT,
        fmt=NUM if c == 'E' else MONEY)
put(ws, f'A{T + 2}', '=("对账日期：　"&TEXT(TODAY(),"yyyy 年 m 月 d 日"))', font=F_TXT, align=CL, border=None)
put(ws, f'E{T + 2}', '合计金额:', font=F_TOT, align=CR)
put(ws, f'F{T + 2}', f'=ROUND($J${T},2)', font=F_TOT, fmt=MONEY)
put(ws, f'E{T + 3}', '已 收 款:', font=F_TOT, align=CR)
put(ws, f'F{T + 3}', '=ROUND(N($P$7),2)', font=F_TOT, fmt=MONEY)
put(ws, f'E{T + 4}', '应 收 款:', font=F_TOT, align=CR)
put(ws, f'F{T + 4}', f'=ROUND(N($F${T + 2})-N($F${T + 3}),2)', font=F_TOT, fill=FILL_TOT, fmt=MONEY)
ws.freeze_panes = 'C6'
page(ws, titles='5:5')
print('  ✓ 新表【果然鲜销售对账单】（按购买方）')

# 04 的 _自动清单 里补一段「购买方(紧凑)」给上面的下拉用
ws = wb['_自动清单']
put(ws, 'AG3', '购买方', font=F_TOT, fill=FILL_HDR2)
put(ws, 'AH3', '购买方(紧凑)', font=F_TOT, fill=FILL_HDR2)
put(ws, 'AI3', '首次', font=F_TOT, fill=FILL_HDR2)
put(ws, 'AJ3', '累计', font=F_TOT, fill=FILL_HDR2)
for r in range(4, 1204):
    put(ws, f'AG{r}', f'=IF({SRC01}!$AC{r}="","",{SRC01}!$AC{r})', font=F_NOTE, border=None)
    put(ws, f'AI{r}', f'=IF(AG{r}="",0,IF(MATCH(AG{r},$AG$4:$AG$1203,0)=ROW()-3,1,0))',
        font=F_NOTE, border=None)
    put(ws, f'AJ{r}', f'=N(AJ{r - 1})+AI{r}', font=F_NOTE, border=None)
for r in range(4, 204):
    put(ws, f'AH{r}', f'=IFERROR(INDEX($AG$4:$AG$1203,MATCH(ROW()-3,$AJ$4:$AJ$1203,0)),"")',
        font=F_NOTE, border=None)
print('  ✓ _自动清单 补「购买方(紧凑)」给对账单下拉用')

wb.save(OUT)
print('已写', OUT)
