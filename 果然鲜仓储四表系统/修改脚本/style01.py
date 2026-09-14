# -*- coding: utf-8 -*-
"""照《01》原有的样子做新表用的样式常量"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
F_TITLE = Font(name='微软雅黑', size=14, bold=True, color='1F3864')
F_NOTE  = Font(name='微软雅黑', size=9,  color='7F7F7F')
F_HDR   = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
F_TXT   = Font(name='微软雅黑', size=10)
F_LINK  = Font(name='微软雅黑', size=10, color='1F4E79')
F_TOT   = Font(name='微软雅黑', size=10, bold=True)
F_IN    = Font(name='微软雅黑', size=10, color='0000C0')
FILL_HDR  = PatternFill('solid', fgColor='2F5597')
FILL_HDR2 = PatternFill('solid', fgColor='D6E4F0')
FILL_AUTO = PatternFill('solid', fgColor='F2F2F2')
FILL_IN   = PatternFill('solid', fgColor='FFF2CC')
FILL_TOT  = PatternFill('solid', fgColor='E2EFDA')
FILL_WARN = PatternFill('solid', fgColor='FCE4E4')
C  = Alignment(horizontal='center', vertical='center', wrap_text=True)
CL = Alignment(horizontal='left',   vertical='center', wrap_text=True)
CR = Alignment(horizontal='right',  vertical='center')
_s = Side(style='thin', color='BFBFBF')
BOX = Border(left=_s, right=_s, top=_s, bottom=_s)
MONEY = '#,##0.00'; NUM = '#,##0'; DATE = 'yyyy-mm-dd'; PCT = '0.00%'

def put(ws, coord, value=None, font=F_TXT, fill=None, align=C, fmt=None, border=BOX):
    c = ws[coord]
    if value is not None: c.value = value
    c.font = font
    if fill is not None: c.fill = fill
    if align is not None: c.alignment = align
    if fmt is not None: c.number_format = fmt
    if border is not None: c.border = border
    return c

def title(ws, text, last_col, sub=None, sub_last=None):
    ws.merge_cells(f'A1:{last_col}1')
    put(ws, 'A1', text, font=F_TITLE, align=CL, border=None)
    ws.row_dimensions[1].height = 26
    if sub:
        if sub_last != 'A':
            ws.merge_cells(f'A2:{sub_last or last_col}2')
            put(ws, 'A2', sub, font=F_NOTE, align=CL, border=None)
            ws.row_dimensions[2].height = 16
        else:
            ws.merge_cells(f'A3:{last_col}3')      # 说明挪到第 3 行，第 2 行留给筛选条件
            put(ws, 'A3', sub, font=F_NOTE, align=CL, border=None)
            ws.row_dimensions[3].height = 26

def headers(ws, row, names, fill=FILL_HDR, font=F_HDR):
    for i, n in enumerate(names, 1):
        put(ws, f'{ws.cell(row=row, column=i).column_letter}{row}', n, font=font, fill=fill)
    ws.row_dimensions[row].height = 30

def widths(ws, m):
    for k, v in m.items(): ws.column_dimensions[k].width = v

def page(ws, titles=None, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if titles: ws.print_title_rows = titles
