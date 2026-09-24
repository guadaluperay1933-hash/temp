# -*- coding: utf-8 -*-
"""A050 第八轮 · 《02 物料与周转物台账》

 ① 次果筐汇总：退回的四列（大托盘 / 公司一次性筐 / 公司杂周转筐 / 公司新周转筐）
    原来显示成正数，跟旁边「退回-公司红周转筐」「退回合计」的红色负数不一致 ——
    统一成红色负数显示（数值不动，只改显示格式，合计照样对）。
 ② 下拉全部改成「按表头名字找列」的动态区域，基础资料里新增的内容自动进下拉。

跑法：python3 fix08_02.py <入> <出>
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from style01 import *
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)

B0, B1, HDR, LAST = 4, 103, 3, 'P'

# ══ ① 次果筐汇总：退回四列改红色负数 ══
NEG = '[Red]\\-#,##0;[Red]\\-#,##0;\\-'
ws = wb['次果筐汇总']
n = 0
for col in ('I', 'J', 'K', 'L'):
    for r in range(4, 204):
        ws[f'{col}{r}'].number_format = NEG; n += 1
ws['A2'].value = ('★ 全自动。出库＝发给客户的次果筐/托盘，退回＝客户还回来的（红色负数显示），'
                  '未回合计＝出库－退回。押金同理：收押金正数、退押金红色负数。'
                  '数值本身都是正的，只是显示成负数，方便一眼看出方向；合计不受影响。')
print(f'  ✓ 次果筐汇总：退回 I~L 四列改成红色负数显示（{n} 格）')

# ══ ② 动态下拉 ══
def dyn(t):
    m = f'MATCH("{t}",基础资料!$A${HDR}:${LAST}${HDR},0)'
    c = f'COUNTA(INDEX(基础资料!$A${B0}:${LAST}${B1},0,{m}))'
    return f'OFFSET(基础资料!$A${B0},0,{m}-1,MAX(1,{c}),1)'

ws = wb['基础资料']
names = {ws.cell(HDR, c).value: c for c in range(1, 17) if ws.cell(HDR, c).value}
DN = {'物料种类表': '物料种类', '筐子类型表': '筐子类型', '托盘类型表': '托盘类型',
      '客户表': '客户/领取人', '供应商表': '供应商', '物料业务类型表': '物料业务类型',
      '筐子业务类型表': '筐子业务类型', '装卸方式表': '装卸方式', '核查结果表': '核查结果',
      '结算方式表': '结算方式', '收付款状态表': '收/付款状态', '筐子用途表': '筐子用途',
      '物料名称表': '物料名称', '规格表': '规格', '计量单位表': '计量单位'}
for nm, t in DN.items():
    assert t in names, f'基础资料缺表头「{t}」'
    if nm in wb.defined_names: del wb.defined_names[nm]
    wb.defined_names.add(DefinedName(nm, attr_text=dyn(t)))
print(f'  ✓ 建了 {len(DN)} 个动态下拉区')

BIND = {
    '周转筐出入库明细': {'D4:D2004': '筐子类型表', 'F4:F2004': '筐子业务类型表',
                         'N4:N2004': '装卸方式表', 'Q4:Q2004': '收付款状态表',
                         'R4:R2004': '结算方式表', 'V4:V2004': '核查结果表'},
    '包装物料出入库明细': {'E4:E2003': '物料种类表', 'H4:H2003': '物料业务类型表',
                           'P4:P2003': '收付款状态表', 'Q4:Q2003': '核查结果表'},
    '客户自备加工筐明细': {'C4:C1204': '筐子类型表', 'I4:I1204': '装卸方式表',
                           'J4:J1204': '筐子用途表', 'K4:K1204': '核查结果表'},
    '次果筐+托盘明细': {'D4:D603': '筐子类型表', 'F4:F603': '筐子业务类型表',
                        'I4:I603': '装卸方式表', 'M4:M603': '收付款状态表',
                        'N4:N603': '筐子用途表', 'O4:O603': '核查结果表'},
}
# 客户/领取人那几列取的是「_自动清单」里自动去重的紧凑名单，比基础资料更全，保持原样
KEEP = {'周转筐出入库明细': ['C4:C2004', 'U4:U2004'], '包装物料出入库明细': ['D4:D2003'],
        '客户自备加工筐明细': ['E4:E1204'], '次果筐+托盘明细': ['C4:C603']}
for sn, mp in BIND.items():
    s = wb[sn]
    s.data_validations.dataValidation = [
        dv for dv in s.data_validations.dataValidation if str(dv.sqref) in KEEP.get(sn, [])]
    for sq, nm in mp.items():
        dv = DataValidation(type='list', formula1=f'={nm}', allow_blank=True, showDropDown=False,
                            showInputMessage=True)
        dv.errorTitle, dv.error = '不在基础资料里', '请从下拉里选；没有的先去【基础资料】对应列加一行'
        dv.promptTitle, dv.prompt = '从基础资料取', '点右边小箭头选。新加的内容会自动出现在这里'
        s.add_data_validation(dv); dv.add(sq)
    print(f'  ✓ {sn}：{len(mp)} 个下拉改成动态区')

wb.save(OUT)
print('已写', OUT)
