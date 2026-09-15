# -*- coding: utf-8 -*-
"""⑥ 《03 财务账套与报表》新增两张表：
   · 【对接源_01库存】 跨文件取《01》的【库存等级接口】—— 每个货主每个等级还剩多少
   · 【库存价值与欠款比对】上半部分按客户给出「库存价值 vs 应收款 vs 差额」，
     下半部分按客户＋等级列出结余数量，单价那一列你自己填，库存价值自动算。
     差额为正＝货够抵账，可以先不收；为负＝要催收。
   跑法：python3 fix06_03新表.py <入> <出>"""
import sys, os, re, zipfile, openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from style01 import *
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

SRC, OUT = sys.argv[1], sys.argv[2]
z = zipfile.ZipFile(SRC)
wbxml = z.read('xl/workbook.xml').decode('utf8')
ids = re.findall(r'r:id="(rId\d+)"', re.search(r'<externalReferences>(.*?)</externalReferences>', wbxml, re.S).group(1))
rels = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', z.read('xl/_rels/workbook.xml.rels').decode('utf8')))
IDX = {}
for k, rid in enumerate(ids, 1):
    nm = rels[rid].replace('../', '').split('/')[-1]
    tgt = re.search(rb'Target="([^"]+)"', z.read('xl/externalLinks/_rels/' + nm + '.rels')).group(1).decode()
    IDX[tgt] = k
z.close()
E01 = IDX['01_水果进销存台账模板.xlsx']
print(f'  《01》在本册里是 [{E01}]')

wb = openpyxl.load_workbook(SRC)
N_INV = 400

# ---------------- ① 对接源_01库存 ----------------
NM = '对接源_01库存'
if NM in wb.sheetnames: del wb[NM]
ws = wb.create_sheet(NM, wb.sheetnames.index('对接源_01入库吨位') + 1)
title(ws, '对 接 源 · 01 库 存 等 级（跨文件自动取数 · 请勿改动结构）', 'J',
      sub='★ 1:1 对应《01 水果进销存台账》的【库存等级接口】。链接断了会显示空白，'
          '不会报错；把四个文件放同一个文件夹、打开时选「更新链接」就能取到最新的。')
widths(ws, {'A': 7, 'B': 14, 'C': 11, 'D': 13, 'E': 7, 'F': 12, 'G': 13, 'H': 14, 'I': 7, 'J': 7})
headers(ws, 3, ['源行', '货主/客户', '品种', '等级', '单位', '结余数量', '结余总重KG', '备注', '有效', '累计'])
SRC_COL = {'B': 'B', 'C': 'C', 'D': 'D', 'E': 'E', 'F': 'F', 'G': 'G', 'H': 'H'}
for i in range(N_INV):
    r = 4 + i
    put(ws, f'A{r}', str(i + 1), font=F_NOTE); ws[f'A{r}'] = i + 1
    for tgt, sc in SRC_COL.items():
        put(ws, f'{tgt}{r}',
            f'=IFERROR(IF([{E01}]库存等级接口!${sc}${r}=0,"",[{E01}]库存等级接口!${sc}${r}),"")',
            font=F_LINK, align=CL if tgt in 'BH' else C,
            fmt=NUM if tgt == 'F' else (MONEY if tgt == 'G' else None))
    put(ws, f'I{r}', f'=IF($B{r}="",0,1)', font=F_NOTE)
    put(ws, f'J{r}', f'=N(J{r - 1})+$I{r}', font=F_NOTE)
    ws.row_dimensions[r].height = 15
put(ws, 'L1', '有效行数', font=F_TOT, fill=FILL_HDR2)
put(ws, 'M1', f'=MAX($J$4:$J${3 + N_INV})', font=F_TOT, fill=FILL_AUTO, fmt=NUM)
ws.freeze_panes = 'C4'
page(ws, titles='3:3', landscape=False)
print('  ✓ 新表【对接源_01库存】')

# ---------------- ② 库存价值与欠款比对 ----------------
NM2 = '库存价值与欠款比对'
if NM2 in wb.sheetnames: del wb[NM2]
ws = wb.create_sheet(NM2, wb.sheetnames.index('应收应付汇总表') + 1)
title(ws, '库 存 价 值 与 欠 款 比 对', 'J',
      sub='★ 下半部分「单价」那一列是淡黄色的手工格，你按等级填个心里价，库存价值自动算。'
          '上半部分按客户把库存价值和应收款摆在一起：差额＝库存价值−应收款。'
          '差额为正＝他压在我们库里的货够抵账，可以先不催；为负＝货不够抵，该收钱了。')
widths(ws, {'A': 6, 'B': 15, 'C': 11, 'D': 13, 'E': 7, 'F': 12, 'G': 11, 'H': 15, 'I': 15,
            'J': 15, 'K': 14, 'L': 18})
S_R0 = 4
S_N = 60
ws.merge_cells(f'A{S_R0}:L{S_R0}')
put(ws, f'A{S_R0}', '一、按 客 户 比 对（库存价值 ↔ 应收款 ↔ 差额）', font=F_TOT, fill=FILL_HDR2, align=CL)
H1 = S_R0 + 1
headers(ws, H1, ['序号', '客户', '库存笔数', '结余数量', '结余总重KG', '', '',
                 '库存价值', '应收款余额', '差额（库存价值−应收）', '建议', '说明'])
D0 = H1 + 1
D1 = D0 + S_N - 1
INV = '对接源_01库存'
DET_0, DET_1 = D1 + 4, D1 + 3 + N_INV
for i in range(S_N):
    r = D0 + i
    put(ws, f'A{r}', f'=IF($B{r}="","",ROW()-{D0 - 1})', font=F_LINK)
    put(ws, f'B{r}', f'=IFERROR(INDEX(_自动清单!$E$4:$E$203,ROW()-{D0 - 1}),"")', font=F_LINK, align=CL)
    put(ws, f'C{r}', f'=IF($B{r}="","",COUNTIF({INV}!$B$4:$B${3 + N_INV},$B{r}))',
        font=F_LINK, fmt=NUM)
    put(ws, f'D{r}', f'=IF($B{r}="","",SUMIF({INV}!$B$4:$B${3 + N_INV},$B{r},{INV}!$F$4:$F${3 + N_INV}))',
        font=F_LINK, fmt=NUM)
    put(ws, f'E{r}', f'=IF($B{r}="","",ROUND(SUMIF({INV}!$B$4:$B${3 + N_INV},$B{r},'
                     f'{INV}!$G$4:$G${3 + N_INV}),2))', font=F_LINK, fmt=MONEY)
    put(ws, f'F{r}', None, border=None); put(ws, f'G{r}', None, border=None)
    put(ws, f'H{r}', f'=IF($B{r}="","",ROUND(SUMIF($B${DET_0}:$B${DET_1},$B{r},'
                     f'$H${DET_0}:$H${DET_1}),2))', font=F_TOT, fmt=MONEY)
    put(ws, f'I{r}', f'=IF($B{r}="","",ROUND(IFERROR(INDEX(应收应付汇总表!$F$5:$F$204,'
                     f'MATCH($B{r},应收应付汇总表!$B$5:$B$204,0)),0),2))', font=F_LINK, fmt=MONEY)
    put(ws, f'J{r}', f'=IF($B{r}="","",ROUND(N($H{r})-N($I{r}),2))', font=F_TOT, fmt=MONEY)
    put(ws, f'K{r}', f'=IF($B{r}="","",IF(N($I{r})<=0,"没有欠款",'
                     f'IF(N($J{r})>=0,"货够抵账 · 可以先不收","货不够抵 · 该收钱了")))',
        font=F_TXT, fill=FILL_CHK if False else FILL_AUTO)
    put(ws, f'L{r}', f'=IF($B{r}="","",$B{r}&" 库里还压着 "&TEXT(N($H{r}),"#,##0.00")'
                     f'&" 元的货，欠我们 "&TEXT(N($I{r}),"#,##0.00")&" 元，差额 "'
                     f'&TEXT(N($J{r}),"#,##0.00")&" 元")', font=F_NOTE, align=CL)
    ws.row_dimensions[r].height = 16
T1 = D1 + 1
put(ws, f'A{T1}', '合  计', font=F_TOT, fill=FILL_TOT)
put(ws, f'B{T1}', None, font=F_TOT, fill=FILL_TOT)
for c in 'CDEHIJ':
    put(ws, f'{c}{T1}', f'=ROUND(SUM({c}{D0}:{c}{D1}),2)', font=F_TOT, fill=FILL_TOT,
        fmt=NUM if c in 'CD' else MONEY)
for c in 'FGKL': put(ws, f'{c}{T1}', None, font=F_TOT, fill=FILL_TOT)
ws.conditional_formatting.add(f'J{D0}:J{D1}',
    FormulaRule(formula=[f'AND($B{D0}<>"",$J{D0}<0)'], fill=FILL_WARN,
                font=Font(color='9C0006', bold=True)))
ws.conditional_formatting.add(f'J{D0}:J{D1}',
    FormulaRule(formula=[f'AND($B{D0}<>"",$J{D0}>=0,N($I{D0})>0)'], fill=FILL_TOT))

ws.merge_cells(f'A{D1 + 2}:L{D1 + 2}')
put(ws, f'A{D1 + 2}', '二、按 客 户 ＋ 等 级 明 细（「单价」是淡黄色手工格，填了库存价值就自动出）',
    font=F_TOT, fill=FILL_HDR2, align=CL)
H2 = D1 + 3
headers(ws, H2, ['序号', '客户', '品种', '等级', '单位', '结余数量', '结余总重KG',
                 '库存价值', '单价（手工填）', '库存状态', '', '备注'])
for i in range(N_INV):
    r = DET_0 + i
    sr = 4 + i
    put(ws, f'A{r}', f'=IF($B{r}="","",ROW()-{DET_0 - 1})', font=F_LINK)
    put(ws, f'B{r}', f'=IF({INV}!$B{sr}="","",{INV}!$B{sr})', font=F_LINK, align=CL)
    put(ws, f'C{r}', f'=IF($B{r}="","",{INV}!$C{sr})', font=F_LINK)
    put(ws, f'D{r}', f'=IF($B{r}="","",{INV}!$D{sr})', font=F_LINK)
    put(ws, f'E{r}', f'=IF($B{r}="","",{INV}!$E{sr})', font=F_LINK)
    put(ws, f'F{r}', f'=IF($B{r}="","",N({INV}!$F{sr}))', font=F_LINK, fmt=NUM)
    put(ws, f'G{r}', f'=IF($B{r}="","",N({INV}!$G{sr}))', font=F_LINK, fmt=MONEY)
    put(ws, f'I{r}', None, font=F_IN, fill=FILL_IN, fmt=MONEY)      # 单价手工填
    put(ws, f'H{r}', f'=IF($B{r}="","",ROUND(N($F{r})*N($I{r}),2))', font=F_TOT, fmt=MONEY)
    put(ws, f'J{r}', f'=IF($B{r}="","",{INV}!$H{sr})', font=F_TXT, fill=FILL_AUTO)
    put(ws, f'K{r}', None, border=None)
    put(ws, f'L{r}', f'=IF($B{r}="","",IF(N($I{r})=0,"← 这一行还没填单价",""))',
        font=F_NOTE, align=CL)
    ws.row_dimensions[r].height = 15
T2 = DET_1 + 1
put(ws, f'A{T2}', '合  计', font=F_TOT, fill=FILL_TOT)
for c in 'BCDEJKL': put(ws, f'{c}{T2}', None, font=F_TOT, fill=FILL_TOT)
for c in 'FGH':
    put(ws, f'{c}{T2}', f'=ROUND(SUM({c}{DET_0}:{c}{DET_1}),2)', font=F_TOT, fill=FILL_TOT,
        fmt=NUM if c == 'F' else MONEY)
put(ws, f'I{T2}', None, font=F_TOT, fill=FILL_TOT)
ws.freeze_panes = 'C6'
page(ws, titles=None)
print('  ✓ 新表【库存价值与欠款比对】')

wb.save(OUT)
print('已写', OUT)
