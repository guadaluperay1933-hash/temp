# -*- coding: utf-8 -*-
"""① 在《01 水果进销存台账》的【成品出库明细】里，「客户/领用方」后面插一列「购买方」。
   插列会让 H 以后的所有列往右挪一格，所以本表自己的公式、全册引用这张表的公式、
   数据校验、合并区、筛选区、列宽全部跟着平移。
   跑法：python3 fix01_购买方.py <入> <出>"""
import sys, os, openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xlsx_util import insert_column
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

SRC = sys.argv[1]; OUT = sys.argv[2]
AT = 8                      # 插在第 8 列（H）＝「客户/领用方」G 的后面
SHEET = '成品出库明细'
MAXR, MAXC = 3003, 40

wb = openpyxl.load_workbook(SRC)
before = sum(1 for w in wb.worksheets for row in w.iter_rows() for c in row
             if isinstance(c.value, str) and c.value.startswith('='))
insert_column(wb, SHEET, AT, MAXR, MAXC, width=12)

ws = wb[SHEET]
hdr = ws.cell(row=3, column=7)          # 照抄「客户/领用方」表头的样子
cell = ws.cell(row=3, column=AT)
cell.value = '购买方'
if hdr.has_style:
    import copy as _c
    cell._style = _c.copy(hdr._style)
for r in range(4, MAXR + 1):
    src = ws.cell(row=r, column=7)
    dst = ws.cell(row=r, column=AT)
    if src.has_style:
        import copy as _c
        dst._style = _c.copy(src._style)
# 购买方用和「客户/领用方」同一份客户名单
dv = DataValidation(type='list', formula1='基础资料!$C$4:$C$103', allow_blank=True, showErrorMessage=False)
ws.add_data_validation(dv)
dv.add(f'H4:H{MAXR}')

note = ws.cell(row=2, column=1)
if isinstance(note.value, str) and '购买方' not in note.value:
    note.value = (note.value.rstrip('　 ') +
                  '　★ 新增「购买方」列：公司（果然鲜）买下货主的果品再卖出去时，'
                  'F 列填货主（卖给公司的那一方）、H 列填购买方（公司卖给谁）。'
                  '只是代加工出库、公司没买断的，购买方留空即可。')

wb.save(OUT)
after = sum(1 for w in wb.worksheets for row in w.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith('='))
print(f'公式数 {before} → {after}（应当不变）')
print('已写', OUT)
