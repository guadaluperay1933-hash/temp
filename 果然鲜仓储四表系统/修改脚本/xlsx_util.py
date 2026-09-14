# -*- coding: utf-8 -*-
"""给四本联动台账做外科手术用的小工具：整列插入、公式列号平移、跨文件链接改指向。"""
import re, copy
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

CELL = r'\$?[A-Z]{1,3}\$?\d+'
RANGE = rf'{CELL}(?::{CELL})?'

def shift_ref(ref, at, delta=1):
    """把一个 A1 / $A$1:$B$9 形式的引用里，列号 >= at 的部分整体平移 delta 列"""
    def one(m):
        d1, col, d2, row = m.group(1), m.group(2), m.group(3), m.group(4)
        i = CI(col)
        return f'{d1}{L(i + delta) if i >= at else col}{d2}{row}'
    return re.sub(r'(\$?)([A-Z]{1,3})(\$?)(\d+)', one, ref)

def shift_formula_for_sheet(f, sheet, at, delta=1):
    """只平移「指名道姓引用 sheet」的那些引用（含带引号的表名）"""
    if not isinstance(f, str) or not f.startswith('='):
        return f
    pat = re.compile(rf"((?:'{re.escape(sheet)}'|{re.escape(sheet)})!)({RANGE})")
    return pat.sub(lambda m: m.group(1) + shift_ref(m.group(2), at, delta), f)

def shift_formula_self(f, at, delta=1):
    """本表内部的相对引用全部平移（用于被插列的那张表自己的公式）"""
    if not isinstance(f, str) or not f.startswith('='):
        return f
    out, i = [], 0
    # 跳过字符串常量与带表名前缀的引用，其余的单元格引用才平移
    tok = re.compile(rf'"(?:[^"]|"")*"|(?:\'[^\']+\'|[A-Za-z_一-鿿][\w一-鿿.]*)!{RANGE}|{RANGE}')
    for m in tok.finditer(f):
        s = m.group(0)
        out.append(f[i:m.start()])
        prev = f[m.start() - 1] if m.start() else ''
        nxt = f[m.end()] if m.end() < len(f) else ''
        if s.startswith('"') or '!' in s or prev in '.0123456789' or nxt == '(':
            out.append(s)            # 字符串 / 别表引用 / 科学计数法尾巴 / 函数名，都不是单元格
        else:
            out.append(shift_ref(s, at, delta))
        i = m.end()
    out.append(f[i:])
    return ''.join(out)

def insert_column(wb, sheet, at, max_row, max_col, header=None, width=None):
    """在 sheet 的第 at 列前插入一整列：单元格、列宽、数据校验、合并区、筛选区、公式全部跟着走。
       返回被插入列的字母。"""
    ws = wb[sheet]
    # ⓪ 先把合并区拆了，不然合并区里的单元格是只读的，搬不动
    olds = [str(m) for m in ws.merged_cells.ranges]
    for m in olds: ws.unmerge_cells(m)
    # ① 单元格从右往左搬
    for c in range(max_col, at - 1, -1):
        for r in range(1, max_row + 1):
            src = ws.cell(row=r, column=c)
            dst = ws.cell(row=r, column=c + 1)
            dst.value = src.value
            if src.has_style:
                dst._style = copy.copy(src._style)
            src.value = None
    # ② 列宽
    dims = {k: (v.width, v.hidden) for k, v in ws.column_dimensions.items()}
    for k in list(ws.column_dimensions):
        del ws.column_dimensions[k]
    for k, (w, h) in dims.items():
        i = CI(k)
        nk = L(i + 1) if i >= at else k
        if w: ws.column_dimensions[nk].width = w
        if h: ws.column_dimensions[nk].hidden = h
    if width: ws.column_dimensions[L(at)].width = width
    # ③ 合并区按平移后的范围重新合上
    # ④ 数据校验
    for dv in list(ws.data_validations.dataValidation):
        dv.sqref = type(dv.sqref)(' '.join(shift_ref(str(x), at) for x in str(dv.sqref).split()))
    # ⑤ 筛选区 / 打印区
    if ws.auto_filter.ref: ws.auto_filter.ref = shift_ref(ws.auto_filter.ref, at)
    if ws.print_area:
        ws.print_area = [shift_ref(str(a).split('!')[-1], at) for a in ws.print_area]
    # ⑥ 本表公式
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col + 1):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                c.value = shift_formula_self(c.value, at)
    # ⑦ 全工作簿里引用这张表的公式
    for w in wb.worksheets:
        if w.title == sheet: continue
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    nv = shift_formula_for_sheet(c.value, sheet, at)
                    if nv != c.value: c.value = nv
        for dv in list(w.data_validations.dataValidation):
            if dv.formula1 and sheet in str(dv.formula1):
                dv.formula1 = shift_formula_for_sheet('=' + str(dv.formula1).lstrip('='), sheet, at)[1:]
    for m in olds: ws.merge_cells(shift_ref(m, at))
    return L(at)

def each_formula(wb):
    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    yield w, c
