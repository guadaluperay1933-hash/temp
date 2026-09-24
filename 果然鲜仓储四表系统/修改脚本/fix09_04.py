# -*- coding: utf-8 -*-
"""A050 第九轮 · 《04 综合查询对账单》

① 《物料与筐子销售对账单》原来是上下两块：上面 100 行只列包装物料，
   筐子销售明细被放在第 107 行以下 —— 屏幕外，翻不到，所以看着像「没有显示」。
   现在合并成一张连续明细（200 行），物料和筐子按业务顺序混排，
   「规格 / 业务类别」那一列：物料行显示规格，筐子行显示「周转筐－销售出库」。
   顶部的总金额也改成 物料 ＋ 筐子 的合计。

② 筐子销售的判定从「周转筐*」（会把公司向筐厂采购的筐也算进来）
   收紧成「周转筐－销售出库」，只认真正卖断给客户的那几笔。

③ 《筐子对账单》补一句说明：卖断的筐子不在这张表里。

跑法：python3 fix09_04.py <入> <出>
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from copy import copy

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)

AL0, AL1 = 4, 2403          # _自动清单 的标记行区间（前 1200 行＝01 水果，后 1200 行＝02 物料）
R0, R1 = 6, 205             # 合并后的明细区
OLD1 = 208                  # 原来第二块的末行

# ══════════════════════════════════════════════
# ① _自动清单：筐子销售判定收紧 + 新增「物料或筐子销售」合并标记
# ══════════════════════════════════════════════
ws = wb['_自动清单']
tight = 0
for r in range(AL0, AL1 + 1):
    c = ws[f'AH{r}']
    v = c.value
    if not isinstance(v, str):
        continue
    k = r if r <= 1203 else r - 1200             # 对应的对接源行号
    src = '对接源_01水果' if r <= 1203 else '对接源_02物料'
    old = f'LEFT({src}!$D{k},3)="周转筐"'
    new = f'{src}!$D{k}="周转筐－销售出库"'
    if old in v:
        c.value = v.replace(old, new)
        tight += 1
assert tight == AL1 - AL0 + 1, f'筐子销售判定只改到 {tight} 行'

ws['AR3'] = '物料+筐子销售'
ws['AS3'] = '累计'
for r in range(AL0, AL1 + 1):
    ws[f'AR{r}'] = f'=IF(OR($AF{r}=1,$AH{r}=1),1,0)'
    ws[f'AS{r}'] = f'=N(AS{r-1})+$AR{r}'
for col in ('AR', 'AS'):
    ws.column_dimensions[col].hidden = True

# ══════════════════════════════════════════════
# ② 物料与筐子销售对账单：两块合并成一张
# ══════════════════════════════════════════════
ws = wb['物料与筐子销售对账单']

for rng in ('A107:G107', 'H107:J107'):
    if rng in [str(x) for x in ws.merged_cells.ranges]:
        ws.unmerge_cells(rng)

# 明细行公式模板 —— idx 是那一行的「源行号」单元格
def pick(col, idx):
    return (f'IF({idx}<=1200,IF(INDEX(对接源_01水果!${col}$4:${col}$1203,{idx})=0,"",'
            f'INDEX(对接源_01水果!${col}$4:${col}$1203,{idx})),'
            f'IF(INDEX(对接源_02物料!${col}$4:${col}$1203,{idx}-1200)=0,"",'
            f'INDEX(对接源_02物料!${col}$4:${col}$1203,{idx}-1200)))')

def num(col, idx):
    return (f'IF({idx}<=1200,N(INDEX(对接源_01水果!${col}$4:${col}$1203,{idx})),'
            f'N(INDEX(对接源_02物料!${col}$4:${col}$1203,{idx}-1200)))')

STY = {c: copy(ws[f'{c}6']._style) for c in 'ABCDEFGHIJ'}
STY_AE = copy(ws['AE6']._style)

for r in range(R0, R1 + 1):
    idx = f'$AE{r}'
    g = lambda f: f'=IF({idx}="","",{f})'
    ws[f'A{r}'] = g(pick('B', idx))
    ws[f'B{r}'] = g(pick('E', idx))
    ws[f'C{r}'] = g(pick('F', idx))
    # 规格为空（筐子行）就显示业务类别，这样一眼看得出是物料还是卖筐
    ws[f'D{r}'] = g(f'IF({pick("G", idx)}="",{pick("D", idx)},{pick("G", idx)})')
    ws[f'E{r}'] = g(pick('H', idx))
    ws[f'F{r}'] = g(pick('I', idx))
    ws[f'G{r}'] = g(f'IF({num("H", idx)}=0,"",ROUND({num("K", idx)}/{num("H", idx)},2))')
    ws[f'H{r}'] = g(pick('K', idx))
    ws[f'I{r}'] = g(pick('L', idx))
    ws[f'J{r}'] = g(pick('M', idx))
    ws[f'AE{r}'] = f'=IFERROR(MATCH({r - R0 + 1},_自动清单!$AS$4:$AS$2403,0),"")'
    for c in 'ABCDEFGHIJ':
        ws[f'{c}{r}']._style = copy(STY[c])
    ws[f'AE{r}']._style = copy(STY_AE)

# 原第二块残留（旧标题行、旧表头、旧驱动列）整片清掉
for r in range(R1 + 1, OLD1 + 1):
    for c in list('ABCDEFGHIJ') + ['AE', 'AF']:
        cell = ws[f'{c}{r}']
        cell.value = None
        cell.border = openpyxl.styles.Border()
        cell.fill = openpyxl.styles.PatternFill()
        cell.number_format = 'General'
for r in range(R0, R1 + 1):
    ws[f'AF{r}'] = None

ws['D5'] = '规格 / 业务类别'
ws['A3'] = ('★ 只填 B2 客户与起止日期。这里是「要收钱」的那部分：包装物料的销售/领用 ＋ '
            '卖断给客户的筐子，两类按业务顺序混排在同一张明细里。'
            '筐子行的「规格 / 业务类别」显示「周转筐－销售出库」，一眼认得出。'
            '单价＝金额÷数量，自动算。★ 客户借走要还的周转筐不在这里，去《筐子对账单》看。')
ws['A4'] = '物料 ＋ 筐子销售总金额'

SALE = '"周转筐－销售出库"'
for addr in ('N4', 'N5'):
    v = ws[addr].value
    assert '"周转筐*"' in v, addr
    ws[addr] = v.replace('"周转筐*"', SALE)
ws['H4'] = '=ROUND(N($N$3)+N($N$4),2)'

if ws.auto_filter and ws.auto_filter.ref:
    ws.auto_filter.ref = f'A5:J{R1}'
ws.print_area = f'A1:J{R1}'

# ══════════════════════════════════════════════
# ③ 筐子对账单：说明里点一句
# ══════════════════════════════════════════════
k = wb['筐子对账单']
tail = '★ 卖断给客户的筐子（业务类型＝销售出库）不算往来，不在这张表里，请去《物料与筐子销售对账单》对金额。'
if tail not in (k['A3'].value or ''):
    k['A3'] = (k['A3'].value or '') + tail

wb.save(OUT)
print(f'  ✓ 04 表：销售明细合并成一张（{R0}~{R1} 行），筐子销售判定收紧为「周转筐－销售出库」')
